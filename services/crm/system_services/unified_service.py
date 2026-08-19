from __future__ import annotations
import re
import sqlite3
import uuid
import json
import os
import hashlib
import threading
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from openpyxl import load_workbook

from .config import SystemServiceSettings, load_system_settings, resolve_system_db_path
from .field_mapping import (
    SHEET_TABLE_MAPPINGS,
    TRIPS_DATA_START_ROW,
    TRIPS_SHEET_COLUMNS,
    TRIPS_SHEET_NAME,
)
from .phone_normalization import normalize_phone_input
from .payment_rules import validate_payment_transition
from .revenue_rules import booking_revenue
from .trip_program import build_trip_program
from .trip_pricing import parse_room_prices


BLOCKED_STATUSES = {"blacklisted", "blacklist"}
ARCHIVED_TRAVELER_STATUSES = {"inactive", "archived", "blocked"}
REVIEW_STATUSES = {"payment risk", "high maintenance"}
INACTIVE_TRIP_STATUSES = {"cancelled", "closed", "archived"}
QUALIFIED_LEAD_STAGES = {"Qualified", "VIP Priority", "Repeat Priority"}
# Every lead_stage spelling (current + legacy aliases, see leads.py's
# PIPELINE_GROUPS['Booking Draft']) that means "this lead reached Booking
# Draft" -- used to trigger auto_create_booking_from_lead_if_ready.
BOOKING_DRAFT_LEAD_STAGES = {"Booking Draft", "Booked", "Booking Draft Created", "VIP Booking Draft", "Repeat Booking Draft"}
GROUP_DISCOUNT_THRESHOLD = 4
ROOM_HOLD_COLUMNS = {
    "Single": "draft_holds_single",
    "Double": "draft_holds_double",
    "Triple": "draft_holds_triple",
}
ROOM_REMAINING_COLUMNS = {
    "Single": "single_remaining",
    "Double": "double_remaining",
    "Triple": "triple_remaining",
}
BOOKING_LIFECYCLE_STATUSES = [
    "Draft",
    "Waiting Customer",
    "Pending Confirmation",
    "Confirmed",
    "Payment Pending",
    "Paid",
    "Completed",
    "Cancelled",
]
BOOKING_ACTIVE_STATUSES = {
    "Draft",
    "Waiting Customer",
    "Pending Confirmation",
    "Confirmed",
    "Payment Pending",
    "Paid",
}
BOOKING_TRANSITIONS = {
    "Draft": {"Waiting Customer", "Pending Confirmation", "Cancelled"},
    "Waiting Customer": {"Pending Confirmation", "Confirmed", "Cancelled"},
    "Pending Confirmation": {"Confirmed", "Cancelled"},
    "Confirmed": {"Payment Pending", "Cancelled"},
    "Payment Pending": {"Paid", "Cancelled"},
    "Paid": {"Completed", "Cancelled"},
    "Completed": set(),
    "Cancelled": set(),
}
# Automated writes use BOOKING_TRANSITIONS. Employee CRM corrections use this
# complete set only when a reason is recorded in the status history.
EMPLOYEE_BOOKING_STATUSES = set(BOOKING_LIFECYCLE_STATUSES)
_SHEET_WRITE_LOCK = threading.RLock()
_SCHEMA_READY_LOCK = threading.RLock()


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


@dataclass
class IdentityResolution:
    match_status: str
    handoff_required: bool
    handoff_reason: str
    traveler: dict[str, Any] | None
    lookup_phone: dict[str, str]
    name_match_status: str
    actions: list[str]


class UnifiedCRMService:
    """The SQLite-backed half of this project's dual CRM backend.

    Owns raw sqlite3/SQL access plus this class's own idempotent lazy schema
    migrations (ensure_operational_schema() and the _migrate_* methods) for
    columns added to models after their Alembic migration -- see
    database/README.md's "Schema drift" section. When the same operations
    run against Postgres instead, apps/api/app/services/agent_crm_bridge.py's
    PostgresAgentBridgeService is the other half; see services/crm/README.md
    for how the two relate and how a caller ends up on one or the other.
    """

    _schema_ready_paths: set[str] = set()

    def __init__(self, settings: SystemServiceSettings | None = None) -> None:
        self.settings = settings or load_system_settings()

    @contextmanager
    def connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.settings.db_path)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
        finally:
            connection.close()

    @staticmethod
    def parse_trip_sales_notes(raw_notes: str | None, *, default_group_threshold: int = GROUP_DISCOUNT_THRESHOLD) -> dict[str, Any]:
        notes = str(raw_notes or "").replace("\r\n", "\n")
        vip_offer = ""
        group_offer = ""
        group_threshold = default_group_threshold
        freeform_lines: list[str] = []

        for raw_line in notes.splitlines():
            line = raw_line.strip()
            if not line:
                freeform_lines.append(raw_line)
                continue
            upper = line.upper()
            if upper.startswith("VIP_DISCOUNT:"):
                vip_offer = line.split(":", 1)[1].strip()
                continue
            if upper.startswith("GROUP_DISCOUNT:"):
                group_offer = line.split(":", 1)[1].strip()
                continue
            if upper.startswith("GROUP_THRESHOLD:"):
                try:
                    group_threshold = max(int(line.split(":", 1)[1].strip() or default_group_threshold), 2)
                except (TypeError, ValueError):
                    group_threshold = default_group_threshold
                continue
            freeform_lines.append(raw_line)

        return {
            "vip_offer": vip_offer,
            "group_offer": group_offer,
            "group_threshold": group_threshold,
            "notes_body": "\n".join(freeform_lines).strip(),
        }

    @classmethod
    def compose_trip_sales_notes(
        cls,
        *,
        existing_notes: str | None = None,
        vip_offer: str = "",
        group_offer: str = "",
        group_threshold: int = GROUP_DISCOUNT_THRESHOLD,
        notes_body: str = "",
    ) -> str:
        parsed = cls.parse_trip_sales_notes(existing_notes, default_group_threshold=group_threshold)
        final_notes_body = str(notes_body).strip() if str(notes_body).strip() else parsed["notes_body"]
        lines: list[str] = []
        if str(vip_offer).strip():
            lines.append(f"VIP_DISCOUNT: {str(vip_offer).strip()}")
        if str(group_offer).strip():
            lines.append(f"GROUP_DISCOUNT: {str(group_offer).strip()}")
        if group_threshold and int(group_threshold) != GROUP_DISCOUNT_THRESHOLD:
            lines.append(f"GROUP_THRESHOLD: {int(group_threshold)}")
        if final_notes_body:
            lines.append(final_notes_body)
        return "\n".join(lines).strip()

    @classmethod
    def load_trip_commercial_config(cls, trip_id: str) -> dict[str, Any]:
        if not str(trip_id or "").strip():
            return cls.parse_trip_sales_notes("")
        try:
            db_path = resolve_system_db_path()
            if not db_path.exists():
                return cls.parse_trip_sales_notes("")
            connection = sqlite3.connect(db_path)
            try:
                row = connection.execute(
                    "SELECT sales_notes FROM trips WHERE trip_id = ?",
                    (str(trip_id).strip(),),
                ).fetchone()
            finally:
                connection.close()
            return cls.parse_trip_sales_notes(row[0] if row else "")
        except Exception:
            return cls.parse_trip_sales_notes("")

    @staticmethod
    def normalize_phone(raw_phone: str, country_code: str = "20") -> dict[str, str]:
        detected = normalize_phone_input(
            raw_phone,
            country_code or "",
            default_country_is_explicit=bool(country_code),
        )
        lookup_key = ""
        if detected.country_code and detected.local_number:
            lookup_key = f"{detected.country_code}:{detected.local_number}"
        elif detected.normalized_e164:
            lookup_key = re.sub(r"\D", "", detected.normalized_e164)
        legacy_lookup_keys = list(detected.phone_variants_for_lookup)
        return {
            "raw_input": detected.raw_input,
            "raw_phone": detected.raw_input,
            "phone": detected.raw_input,
            "country_code": detected.country_code or country_code or "",
            "local_number": detected.local_number,
            "normalized_whatsapp": detected.normalized_e164,
            "normalized_e164": detected.normalized_e164,
            "lookup_key": lookup_key,
            "phone_lookup_key": lookup_key,
            "phone_variants_for_lookup": detected.phone_variants_for_lookup,
            "legacy_lookup_keys": legacy_lookup_keys,
            "confidence": detected.confidence,
            "requires_country_confirmation": detected.requires_country_confirmation,
            "inferred_country": detected.inferred_country,
            "inferred_nationality": detected.inferred_nationality,
        }

    @staticmethod
    def lookup_key_variants(phone: dict[str, str] | None) -> list[str]:
        """Return lookup-key forms that may exist in older workbook rows."""
        if not phone:
            return []
        candidates: list[str] = []
        candidates.extend(str(v).strip() for v in phone.get("phone_variants_for_lookup", []) or [])
        candidates.extend(str(v).strip() for v in phone.get("legacy_lookup_keys", []) or [])
        raw_phone = str(phone.get("raw_phone") or phone.get("raw_input") or "").strip()
        lookup_key = str(phone.get("lookup_key") or "").strip()
        normalized = str(phone.get("normalized_whatsapp") or "").strip()
        country_code = str(phone.get("country_code") or "").strip()
        local_number = str(phone.get("local_number") or "").strip()

        if raw_phone:
            candidates.append(raw_phone)
            candidates.append(re.sub(r"\D", "", raw_phone))
        if lookup_key:
            candidates.append(lookup_key)
        if normalized:
            candidates.append(normalized)
            candidates.append(re.sub(r"\D", "", normalized))
        if country_code and local_number:
            candidates.append(f"{country_code}:{local_number}")
            candidates.append(f"{country_code}{local_number}")
            candidates.append(local_number)
        if country_code and normalized:
            candidates.append(normalized.replace("+", ""))
        return list(dict.fromkeys(candidate for candidate in candidates if candidate))

    @staticmethod
    def _stable_json_hash(value: Any) -> str:
        payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]

    @classmethod
    def booking_idempotency_key(
        cls,
        *,
        traveler_id: str,
        trip_id: str,
        room_requirements: Any,
        session_id: str = "",
    ) -> str:
        requirements_hash = cls._stable_json_hash(room_requirements or [])
        return f"booking:{str(traveler_id or '').strip()}:{str(trip_id or '').strip()}:{requirements_hash}:{str(session_id or '').strip()}"

    @staticmethod
    def lead_idempotency_key(
        *,
        normalized_phone: str,
        trip_interest_or_general: str = "general",
        session_id: str = "",
    ) -> str:
        return (
            f"lead:{str(normalized_phone or '').strip()}:"
            f"{str(trip_interest_or_general or 'general').strip()}:"
            f"{str(session_id or '').strip()}"
        )

    @staticmethod
    def handoff_idempotency_key(
        *,
        traveler_id_or_phone: str,
        reason_code: str,
        session_id_or_open_lead_id: str = "",
    ) -> str:
        return (
            f"handoff:{str(traveler_id_or_phone or '').strip()}:"
            f"{str(reason_code or 'manual_handoff').strip()}:"
            f"{str(session_id_or_open_lead_id or '').strip()}"
        )

    @staticmethod
    def document_idempotency_key(*, traveler_id: str, file_hash: str) -> str:
        return f"document:{str(traveler_id or '').strip()}:{str(file_hash or '').strip()}"

    @staticmethod
    def write_result_contract(
        *,
        status: str,
        executed: bool,
        reused: bool,
        record_type: str,
        record_id: str = "",
        idempotency_key: str = "",
        customer_confirmation_allowed: bool = False,
        error_code: str = "",
        safe_customer_message_key: str = "",
        audit: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        return {
            "status": status,
            "executed": bool(executed),
            "reused": bool(reused),
            "record_type": record_type,
            "record_id": str(record_id or ""),
            "idempotency_key": str(idempotency_key or ""),
            "customer_confirmation_allowed": bool(customer_confirmation_allowed),
            "error_code": str(error_code or ""),
            "safe_customer_message_key": str(safe_customer_message_key or ""),
            "audit": audit or {},
        }

    @staticmethod
    def normalize_trip_type(raw_type: str | None) -> str | None:
        lowered = (raw_type or "").strip().lower()
        if lowered in {"local", "inside egypt", "egypt", "egypt trips"}:
            return "Local"
        if lowered in {"international", "abroad", "outside egypt"}:
            return "International"
        return None

    @staticmethod
    def normalize_name(value: str) -> str:
        lowered = (value or "").casefold().strip()
        lowered = re.sub(r"[^0-9a-z\u0600-\u06ff\s]+", " ", lowered)
        return " ".join(lowered.split())

    @classmethod
    def names_are_compatible(cls, submitted_name: str, existing_name: str) -> bool:
        submitted = cls.normalize_name(submitted_name)
        existing = cls.normalize_name(existing_name)
        if not existing:
            return False
        if not submitted:
            # Phone-first CRM flows must be allowed to identify an existing traveler
            # before the agent asks for the traveler's name.
            return True
        if submitted == existing:
            return True
        submitted_tokens = submitted.split()
        existing_tokens = existing.split()
        if len(submitted_tokens) < 2 or len(existing_tokens) < 2:
            return False
        if submitted_tokens[0] == existing_tokens[0] and submitted_tokens[-1] == existing_tokens[-1]:
            return True
        shared = set(submitted_tokens) & set(existing_tokens)
        shorter = min(len(set(submitted_tokens)), len(set(existing_tokens)))
        return bool(shorter and len(shared) >= 2 and (len(shared) / shorter) >= 0.80)

    def resolve_identity(self, full_name: str, raw_phone: str, country_code: str = "20") -> IdentityResolution:
        phone = self.normalize_phone(raw_phone, country_code)
        lookup_keys = self.lookup_key_variants(phone)
        with self.connect() as connection:
            query = """
                SELECT traveler_id, status, full_name, birthday, gender, nationality, preferred_currency,
                       phone_code, whatsapp_raw, integrated_whatsapp, normalized_whatsapp,
                       phone_lookup_key, local_trips_count, international_trips_count, total_trips,
                       lead_source, created_at, last_contacted_at, agent_notes, data_audit
                FROM travelers
                WHERE integrated_whatsapp = ?
                   OR normalized_whatsapp = ?
                   OR whatsapp_raw = ?
                   OR REPLACE(integrated_whatsapp, '+', '') = ?
                   OR REPLACE(normalized_whatsapp, '+', '') = ?
            """
            params: list[Any] = [
                phone["normalized_whatsapp"],
                phone["normalized_whatsapp"],
                phone["raw_phone"],
                re.sub(r"\D", "", phone["normalized_whatsapp"]),
                re.sub(r"\D", "", phone["normalized_whatsapp"]),
            ]
            if lookup_keys:
                placeholders = ", ".join("?" for _ in lookup_keys)
                query = query.replace(
                    "WHERE integrated_whatsapp = ?",
                    f"WHERE phone_lookup_key IN ({placeholders}) OR integrated_whatsapp = ?",
                )
                params = [*lookup_keys, *params]
            matches = connection.execute(
                query,
                params,
            ).fetchall()

        if not matches:
            return IdentityResolution(
                match_status="not_found",
                handoff_required=False,
                handoff_reason="",
                traveler=None,
                lookup_phone=phone,
                name_match_status="",
                actions=["collect_new_traveler_data"],
            )

        if len(matches) > 1:
            return IdentityResolution(
                match_status="multiple_matches",
                handoff_required=True,
                handoff_reason="duplicate_phone_match",
                traveler=None,
                lookup_phone=phone,
                name_match_status="",
                actions=["human_review_duplicate_phone"],
            )

        traveler = dict(matches[0])
        status = (traveler.get("status") or "").strip().lower()
        if status in BLOCKED_STATUSES:
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason="blacklisted_customer",
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="blocked_status",
                actions=["block_sales_flow"],
            )
        if status in ARCHIVED_TRAVELER_STATUSES:
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason="archived_traveler",
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="archived_status",
                actions=["human_review_archived_traveler"],
            )
        if not self.names_are_compatible(full_name, traveler.get("full_name") or ""):
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason="phone_name_conflict",
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="conflict",
                actions=["human_review_identity_conflict"],
            )
        if status in REVIEW_STATUSES:
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason=status.replace(" ", "_"),
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="matched",
                actions=["allow_conversation_but_require_human_review"],
            )
        return IdentityResolution(
            match_status="single_match",
            handoff_required=False,
            handoff_reason="",
            traveler=traveler,
            lookup_phone=phone,
            name_match_status="matched",
            actions=["continue_sales_flow"],
        )

    def recalculate_traveler_stats(self, traveler_id: str) -> None:
        if not traveler_id:
            return
        with self.connect() as connection:
            try:
                # 1. Count local trips
                local_trips = connection.execute(
                    """
                    SELECT COUNT(*) FROM trip_bookings b
                    JOIN trips t ON TRIM(b.trip_id) = TRIM(t.trip_id)
                    WHERE TRIM(b.traveler_id) = ?
                      AND IFNULL(b.booking_status, '') != 'Cancelled'
                      AND TRIM(t.type) = 'Local'
                    """,
                    (traveler_id.strip(),),
                ).fetchone()[0] or 0

                # 2. Count international trips
                intl_trips = connection.execute(
                    """
                    SELECT COUNT(*) FROM trip_bookings b
                    JOIN trips t ON TRIM(b.trip_id) = TRIM(t.trip_id)
                    WHERE TRIM(b.traveler_id) = ?
                      AND IFNULL(b.booking_status, '') != 'Cancelled'
                      AND TRIM(t.type) = 'International'
                    """,
                    (traveler_id.strip(),),
                ).fetchone()[0] or 0

                # 3. Count community events
                comm_events = connection.execute(
                    """
                    SELECT COUNT(*) FROM ce_bookings
                    WHERE TRIM(traveler_id) = ? AND IFNULL(status, '') != 'Cancelled'
                    """,
                    (traveler_id.strip(),),
                ).fetchone()[0] or 0

                revenue_rows = connection.execute(
                    """
                    SELECT b.booking_status, b.payment_status, b.currency, b.room_type, b.group_size,
                           b.refund_amount,
                           t.trip_id AS trip_trip_id, t.public_price AS trip_public_price,
                           t.room_prices_json AS trip_room_prices_json
                    FROM trip_bookings b
                    LEFT JOIN trips t ON TRIM(b.trip_id) = TRIM(t.trip_id)
                    WHERE TRIM(b.traveler_id) = ?
                    """,
                    (traveler_id.strip(),),
                ).fetchall()
                # booking_revenue() is the same rule Revenue Analytics and the
                # live per-traveler revenue summary use (room/currency-specific
                # pricing, requires a recognized currency) -- not a second,
                # currency-blind calculation that only ever looked at
                # Trip.public_price regardless of what currency the booking
                # was actually recorded in. It returns revenue net of refunds,
                # so refund_amount has to be selected above and carried here:
                # omit it and this path quietly reports gross while every
                # other surface reports net.
                lifetime_revenue = 0.0
                for row in revenue_rows:
                    booking_like = SimpleNamespace(
                        booking_status=row["booking_status"],
                        payment_status=row["payment_status"],
                        currency=row["currency"],
                        room_type=row["room_type"],
                        group_size=row["group_size"],
                        refund_amount=row["refund_amount"],
                    )
                    trip_like = None
                    if row["trip_trip_id"]:
                        trip_dict = {
                            "public_price": row["trip_public_price"],
                            "room_prices_json": row["trip_room_prices_json"],
                        }
                        trip_like = SimpleNamespace(to_dict=lambda d=trip_dict: d)
                    result = booking_revenue(booking_like, trip_like)
                    if result is not None:
                        _currency, amount = result
                        lifetime_revenue += amount

                # Update the database
                connection.execute(
                    """
                    UPDATE travelers
                    SET local_trips_count = ?,
                        international_trips_count = ?,
                        total_trips = ?,
                        community_events_count = ?,
                        lifetime_revenue = ?
                    WHERE traveler_id = ?
                    """,
                    (local_trips, intl_trips, local_trips + intl_trips, comm_events, lifetime_revenue, traveler_id),
                )
                connection.commit()
            except sqlite3.OperationalError:
                pass

    def refresh_traveler_sheet_stats(self, traveler_id: str) -> bool:
        """Pull profile counters from the Travelers sheet for one traveler."""
        if not traveler_id:
            return False
        return bool(self.refresh_travelers_sheet_stats([traveler_id.strip()]))

    def refresh_travelers_sheet_stats(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        """Pull profile counters from the Travelers sheet for multiple travelers."""
        if self.settings.data_authority == "crm":
            return {}
        clean_ids = sorted({str(traveler_id or "").strip() for traveler_id in traveler_ids if str(traveler_id or "").strip()})
        if not clean_ids:
            return {}

        stats_by_id = self._read_traveler_stats_by_id_from_sheet(clean_ids)
        if not stats_by_id:
            return {}

        with self.connect() as connection:
            for traveler_id, values in stats_by_id.items():
                if not values:
                    continue
                assignments = []
                params: list[Any] = []
                for field_name, value in values.items():
                    assignments.append(f"{field_name} = ?")
                    params.append(value)
                params.append(traveler_id)
                connection.execute(
                    f"UPDATE travelers SET {', '.join(assignments)} WHERE traveler_id = ?",
                    params,
                )
            connection.commit()
        return stats_by_id

    def _read_traveler_stats_from_sheet(self, traveler_id: str) -> dict[str, Any]:
        stats_by_id = self._read_traveler_stats_by_id_from_sheet([traveler_id])
        return stats_by_id.get(traveler_id, {})

    def _read_traveler_stats_by_id_from_sheet(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        if self.settings.sheet_backend == "google_sheets":
            return self._read_traveler_stats_by_id_from_google_sheet(traveler_ids)
        return self._read_traveler_stats_by_id_from_excel(traveler_ids)

    def _read_traveler_stats_by_id_from_excel(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        wanted = {str(traveler_id or "").strip() for traveler_id in traveler_ids if str(traveler_id or "").strip()}
        if not wanted:
            return {}

        workbook_path = self.settings.excel_runtime_workbook or self.settings.excel_source_workbook
        if not workbook_path or not workbook_path.exists():
            return {}

        stats_by_id: dict[str, dict[str, Any]] = {}
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if "Travelers" not in wb.sheetnames:
                return {}
            ws = wb["Travelers"]
            headers = {
                str(ws.cell(1, col_idx).value or "").strip().lower(): col_idx
                for col_idx in range(1, ws.max_column + 1)
                if ws.cell(1, col_idx).value
            }
            id_col = headers.get("traveler id")
            if not id_col:
                return {}
            for row_idx in range(2, ws.max_row + 1):
                current = str(ws.cell(row_idx, id_col).value or "").strip()
                if current not in wanted:
                    continue
                stats_by_id[current] = self._extract_traveler_stats_from_cells(
                    lambda header, row_idx=row_idx: ws.cell(row_idx, headers[header]).value if header in headers else None
                )
                if len(stats_by_id) == len(wanted):
                    break
        finally:
            wb.close()
        return stats_by_id

    def _read_traveler_stats_by_id_from_google_sheet(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        wanted = {str(traveler_id or "").strip() for traveler_id in traveler_ids if str(traveler_id or "").strip()}
        if not wanted or not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            return {}

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets.readonly",
            "https://www.googleapis.com/auth/drive.readonly",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        ws = client.open_by_key(self.settings.google_sheet_id).worksheet("Travelers")
        rows = ws.get_all_values()
        if not rows:
            return {}

        headers = {
            str(header or "").strip().lower(): col_idx
            for col_idx, header in enumerate(rows[0])
            if str(header or "").strip()
        }
        id_idx = headers.get("traveler id")
        if id_idx is None:
            return {}

        stats_by_id: dict[str, dict[str, Any]] = {}
        for row in rows[1:]:
            current = row[id_idx] if len(row) > id_idx else ""
            current = str(current).strip()
            if current not in wanted:
                continue
            stats_by_id[current] = self._extract_traveler_stats_from_cells(
                lambda header, row=row: row[headers[header]] if header in headers and len(row) > headers[header] else None
            )
            if len(stats_by_id) == len(wanted):
                break
        return stats_by_id

    def _read_traveler_stats_from_excel(self, traveler_id: str) -> dict[str, Any]:
        stats_by_id = self._read_traveler_stats_by_id_from_excel([traveler_id])
        return stats_by_id.get(traveler_id, {})

    def _read_traveler_stats_from_google_sheet(self, traveler_id: str) -> dict[str, Any]:
        stats_by_id = self._read_traveler_stats_by_id_from_google_sheet([traveler_id])
        return stats_by_id.get(traveler_id, {})

    def _extract_traveler_stats_from_cells(self, value_for_header) -> dict[str, Any]:
        field_headers = {
            "local_trips_count": ("loc. trips", "local trips count"),
            "international_trips_count": ("int. trips", "international trips count"),
            "total_trips": ("total trips",),
            "community_events_count": ("comm. events", "community events count"),
            "lifetime_revenue": ("lifetime revenue",),
        }
        values: dict[str, Any] = {}
        for field_name, headers in field_headers.items():
            raw_value = None
            for header in headers:
                raw_value = value_for_header(header)
                if raw_value not in (None, ""):
                    break
            if raw_value in (None, ""):
                continue
            if field_name == "lifetime_revenue":
                values[field_name] = self._parse_sheet_float(raw_value)
            else:
                values[field_name] = self._parse_sheet_int(raw_value)
        return values

    @staticmethod
    def _parse_sheet_int(value: Any) -> int:
        try:
            return int(float(str(value).replace(",", "").strip()))
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _parse_sheet_float(value: Any) -> float:
        try:
            return float(str(value).replace(",", "").replace("$", "").strip())
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _extract_numeric_traveler_id(traveler_id: str) -> int:
        match = re.fullmatch(r"TR(\d+)", str(traveler_id or "").strip())
        return int(match.group(1)) if match else 0

    def _max_traveler_id_from_workbook(self, workbook_path: Path | None) -> int:
        if not workbook_path or not workbook_path.exists():
            return 0
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if "Travelers" not in wb.sheetnames:
                return 0
            ws = wb["Travelers"]
            headers = {
                str(ws.cell(1, col_idx).value or "").strip().lower(): col_idx
                for col_idx in range(1, ws.max_column + 1)
                if ws.cell(1, col_idx).value
            }
            id_col = headers.get("traveler id") or 2
            max_number = 0
            for row_idx in range(2, ws.max_row + 1):
                traveler_id = str(ws.cell(row_idx, id_col).value or "").strip()
                current_number = self._extract_numeric_traveler_id(traveler_id)
                if current_number > max_number:
                    max_number = current_number
            return max_number
        finally:
            wb.close()

    def _max_traveler_id_from_db(self) -> int:
        with self.connect() as connection:
            row = connection.execute(
                """
                SELECT traveler_id
                FROM travelers
                WHERE traveler_id LIKE 'TR%'
                ORDER BY CAST(SUBSTR(traveler_id, 3) AS INTEGER) DESC
                LIMIT 1
                """
            ).fetchone()

        if row and row["traveler_id"]:
            return self._extract_numeric_traveler_id(str(row["traveler_id"]).strip())
        return 0

    def next_traveler_id(self) -> str:
        max_number = self._max_traveler_id_from_db()
        if max_number <= 0:
            max_number = max(
                self._max_traveler_id_from_workbook(self.settings.excel_runtime_workbook),
                self._max_traveler_id_from_workbook(self.settings.excel_source_workbook),
            )

        next_number = max_number + 1
        return f"TR{next_number:05d}"

    def create_traveler(
        self,
        *,
        full_name: str,
        raw_phone: str,
        birthday: str = "",
        gender: str = "",
        nationality: str = "",
        preferred_currency: str = "",
        lead_source: str = "",
        agent_notes: str = "",
        country_code: str = "20",
    ) -> dict[str, Any]:
        phone = self.normalize_phone(raw_phone, country_code)
        first_name, last_name = self._split_name(full_name)
        now = _utc_now().replace(microsecond=0).isoformat()

        max_retries = 3
        traveler_id = self.next_traveler_id()
        for attempt in range(max_retries):
            try:
                with self.connect() as connection:
                    connection.execute(
                        """
                        INSERT INTO travelers (
                            traveler_id, status, full_name, first_name, last_name, birthday, gender,
                            nationality, preferred_currency, phone_code, whatsapp_raw, integrated_whatsapp, normalized_whatsapp,
                            phone_lookup_key, lead_source, created_at, last_contacted_at, agent_notes
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            traveler_id,
                            "Active",
                            full_name.strip(),
                            first_name,
                            last_name,
                            birthday or None,
                            gender or None,
                            nationality or None,
                            preferred_currency or None,
                            phone["country_code"] or None,
                            phone["local_number"] or None,
                            phone["normalized_whatsapp"] or None,
                            phone["normalized_whatsapp"] or None,
                            phone["lookup_key"] or None,
                            lead_source or None,
                            now,
                            now,
                            agent_notes or None,
                        ),
                    )
                    connection.commit()
                break
            except sqlite3.IntegrityError as exc:
                error_text = str(exc)
                if "travelers.traveler_id" not in error_text and "UNIQUE constraint failed" not in error_text:
                    raise
                if attempt == max_retries - 1:
                    raise
                traveler_id = f"TR{self._extract_numeric_traveler_id(traveler_id) + 1:05d}"
        return {
            "traveler_id": traveler_id,
            "full_name": full_name.strip(),
            "birthday": birthday or None,
            "gender": gender or None,
            "nationality": nationality or None,
            "preferred_currency": preferred_currency or None,
            "phone_lookup_key": phone["lookup_key"],
            "integrated_whatsapp": phone["normalized_whatsapp"],
        }

    def upsert_lead(
        self,
        *,
        customer_name: str,
        raw_phone: str,
        traveler_id: str | None,
        lead_stage: str,
        lead_source: str,
        channel: str,
        preferred_trip_type: str = "",
        interested_trip_ids: str = "",
        suggested_trip_ids: str = "",
        priority: str = "Medium",
        follow_up_status: str = "",
        follow_up_due_date: str = "",
        notes: str = "",
        booking_id: str = "",
        traveler_status: str = "",
        customer_tier: str = "",
        match_status: str = "",
        last_interaction_id: str = "",
        flow_key: str = "",
        current_step: str = "",
        handoff_required: bool = False,
        handoff_reason: str = "",
        language: str = "",
        country_code: str = "20",
        group_size: int | str = 1,
        force_create_new: bool = False,
        session_id: str = "",
        idempotency_key: str = "",
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        phone = self.normalize_phone(raw_phone, country_code)
        lookup_keys = self.lookup_key_variants(phone)
        interest_key = interested_trip_ids or preferred_trip_type or "general"
        resolved_idempotency_key = str(idempotency_key or "").strip() or self.lead_idempotency_key(
            normalized_phone=phone.get("lookup_key") or phone.get("normalized_whatsapp") or raw_phone,
            trip_interest_or_general=interest_key,
            session_id=session_id or flow_key,
        )
        now = _utc_now().replace(microsecond=0).isoformat()
        created = False
        with self.connect() as connection:
            existing = None
            if not force_create_new and resolved_idempotency_key:
                existing = connection.execute(
                    """
                    SELECT lead_id, interaction_count FROM leads
                    WHERE idempotency_key = ?
                    ORDER BY created_at DESC
                    LIMIT 1
                    """,
                    (resolved_idempotency_key,),
                ).fetchone()
            if not existing and not force_create_new:
                query = """
                    SELECT lead_id, interaction_count FROM leads
                    WHERE (? <> '' AND traveler_id = ?)
                    ORDER BY created_at DESC
                    LIMIT 1
                """
                params: list[Any] = [traveler_id or "", traveler_id or ""]
                if lookup_keys:
                    placeholders = ", ".join("?" for _ in lookup_keys)
                    query = query.replace(
                        "WHERE (? <> '' AND traveler_id = ?)",
                        f"WHERE (? <> '' AND phone_lookup_key IN ({placeholders})) OR (? <> '' AND traveler_id = ?)",
                    )
                    params = [phone["lookup_key"] or "", *lookup_keys, *params]
                existing = connection.execute(
                    query,
                    params,
                ).fetchone()
            if existing:
                lead_id = str(existing["lead_id"])
                interaction_count = self._as_int(existing["interaction_count"], default=0) or 0
                connection.execute(
                    """
                    UPDATE leads
                    SET updated_at = ?, customer_name = ?, raw_phone = ?, integrated_whatsapp = ?,
                        phone_lookup_key = ?, traveler_id = ?, traveler_status = ?, customer_tier = ?,
                        match_status = ?, lead_stage = ?, lead_source = ?, channel = ?,
                        preferred_trip_type = ?, group_size = ?, interested_trip_ids = ?, suggested_trip_ids = ?,
                        priority = ?, follow_up_status = ?, follow_up_due_date = ?, last_interaction_id = ?,
                        interaction_count = ?, notes = ?, booking_id = ?, flow_key = ?, current_step = ?,
                        handoff_required = ?, handoff_reason = ?, language = ?,
                        idempotency_key = COALESCE(idempotency_key, ?)
                    WHERE lead_id = ?
                    """,
                    (
                        now,
                        customer_name,
                        phone["local_number"] or raw_phone,
                        phone["normalized_whatsapp"],
                        phone["lookup_key"],
                        traveler_id,
                        traveler_status or None,
                        customer_tier or None,
                        match_status or None,
                        lead_stage,
                        lead_source,
                        channel,
                        preferred_trip_type or None,
                        self._as_int(group_size, default=1) or 1,
                        interested_trip_ids or None,
                        suggested_trip_ids or None,
                        priority,
                        follow_up_status or None,
                        follow_up_due_date or None,
                        last_interaction_id or None,
                        interaction_count + (1 if last_interaction_id else 0),
                        notes or None,
                        booking_id or None,
                        flow_key or None,
                        current_step or None,
                        1 if handoff_required else 0,
                        handoff_reason or None,
                        language or None,
                        resolved_idempotency_key or None,
                        lead_id,
                    ),
                )
            else:
                created = True
                lead_id = self._next_prefixed_id(connection, "leads", "lead_id", "LD", 5)
                connection.execute(
                    """
                    INSERT INTO leads (
                        lead_id, created_at, updated_at, customer_name, raw_phone, integrated_whatsapp,
                        phone_lookup_key, traveler_id, traveler_status, customer_tier, match_status,
                        lead_stage, lead_source, channel, preferred_trip_type, group_size, interested_trip_ids,
                        suggested_trip_ids, priority, follow_up_status, follow_up_due_date,
                        last_interaction_id, interaction_count, notes, booking_id, flow_key,
                        current_step, handoff_required, handoff_reason, language, idempotency_key
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        lead_id,
                        now,
                        now,
                        customer_name,
                        phone["local_number"] or raw_phone,
                        phone["normalized_whatsapp"],
                        phone["lookup_key"],
                        traveler_id,
                        traveler_status or None,
                        customer_tier or None,
                        match_status or None,
                        lead_stage,
                        lead_source,
                        channel,
                        preferred_trip_type or None,
                        self._as_int(group_size, default=1) or 1,
                        interested_trip_ids or None,
                        suggested_trip_ids or None,
                        priority,
                        follow_up_status or None,
                        follow_up_due_date or None,
                        last_interaction_id or None,
                        1 if last_interaction_id else 0,
                        notes or None,
                        booking_id or None,
                        flow_key or None,
                        current_step or None,
                        1 if handoff_required else 0,
                        handoff_reason or None,
                        language or None,
                        resolved_idempotency_key or None,
                    ),
                )
            if traveler_id:
                connection.execute(
                    """
                    UPDATE travelers
                    SET last_lead_id = ?, last_contacted_at = ?
                    WHERE traveler_id = ?
                    """,
                    (lead_id, now, traveler_id),
                )
            connection.commit()
        contract = self.write_result_contract(
            status="success" if created else "reused",
            executed=created,
            reused=not created,
            record_type="lead",
            record_id=lead_id,
            idempotency_key=resolved_idempotency_key,
            customer_confirmation_allowed=True,
            safe_customer_message_key="lead.created" if created else "lead.already_recorded",
            audit={"phone_lookup_key": phone.get("lookup_key") or "", "interest_key": interest_key},
        )
        return {
            "lead_id": lead_id,
            "lead_stage": lead_stage,
            "priority": priority,
            "follow_up_status": follow_up_status,
            "follow_up_due_date": follow_up_due_date,
            "interaction_id": last_interaction_id,
            "created": created,
            "reused": not created,
            "idempotency_key": resolved_idempotency_key,
            "write_result_contract": contract,
        }

    def record_agent_outcome(
        self,
        *,
        full_name: str | None = None,
        raw_phone: str = "",
        country_code: str = "20",
        trip_type: str | None = None,
        channel: str = "web",
        source: str = "System",
        agent_notes: str = "",
        birthday: str = "",
        gender: str = "",
        nationality: str = "",
        preferred_currency: str = "",
        preferred_trip_id: str = "",
        language: str = "",
        manual_lead_stage: str = "",
        manual_priority: str = "",
        manual_follow_up_due_date: str = "",
        force_create_new_lead: bool = False,
        group_size: int | str = 1,
        session_id: str = "",
        idempotency_key: str = "",
        **legacy_kwargs: Any,
    ) -> dict[str, Any]:
        """Persist an agent qualification outcome to the system DB, then sync sheets."""
        if not full_name:
            full_name = str(legacy_kwargs.pop("customer_name", "")).strip()
        if not agent_notes:
            agent_notes = str(legacy_kwargs.pop("status_snapshot", "") or "").strip()
        if not source:
            source = str(legacy_kwargs.pop("intent", "") or "System").strip() or "System"
        if not preferred_currency:
            preferred_currency = str(legacy_kwargs.pop("currency", "") or "").strip().upper()
        if "group_size" in legacy_kwargs and group_size in (1, "1", "", None):
            group_size = legacy_kwargs.pop("group_size")
        legacy_kwargs.clear()
        self.ensure_operational_schema()
        resolution = self.resolve_identity(full_name, raw_phone, country_code)
        trip_result = self.build_trip_result(trip_type) if trip_type else {"open_trips": [], "date_tbd_trips": []}
        timestamp = _utc_now().replace(microsecond=0)

        created_traveler: dict[str, Any] | None = None
        traveler = resolution.traveler.copy() if resolution.traveler else None
        traveler_id = str(traveler.get("traveler_id")) if traveler else ""
        status_snapshot = str(traveler.get("status") or "") if traveler else ""
        actions = list(resolution.actions)

        if resolution.match_status == "not_found":
            created_traveler = self.create_traveler(
                full_name=full_name,
                raw_phone=raw_phone,
                birthday=birthday,
                gender=gender,
                nationality=nationality,
                preferred_currency=preferred_currency,
                lead_source=source,
                agent_notes=agent_notes,
                country_code=country_code,
            )
            traveler = created_traveler.copy()
            traveler_id = created_traveler["traveler_id"]
            actions.append("created_new_traveler")
        elif resolution.match_status == "single_match" and traveler_id and not resolution.handoff_required:
            profile_updates = self.update_existing_traveler_profile(
                traveler_id,
                birthday=birthday,
                gender=gender,
                nationality=nationality,
                preferred_currency=preferred_currency,
                timestamp=timestamp,
            )
            if profile_updates and traveler:
                traveler.update(profile_updates)
                actions.append("updated_missing_traveler_profile")

        suggested_trip_ids = self._collect_trip_ids(trip_result)
        if preferred_trip_id:
            interested_trip_ids = preferred_trip_id
        else:
            interested_trip_ids = ", ".join(suggested_trip_ids)

        commercial_context = self.resolve_commercial_context(
            traveler=traveler,
            trip_type=trip_type,
            requested_group_size=self._as_int(group_size, default=1) or 1,
            trip_id=preferred_trip_id or (suggested_trip_ids[0] if suggested_trip_ids else ""),
        )
        effective_handoff_required = bool(
            resolution.handoff_required or commercial_context.get("manual_handoff_recommended")
        )
        effective_handoff_reason = (
            str(resolution.handoff_reason or "").strip()
            or str(commercial_context.get("handoff_reason") or "").strip()
        )
        if commercial_context.get("group", {}).get("manual_quote_required"):
            actions.append("group_booking_quote_required")

        preview: dict[str, Any] = {
            "match_status": resolution.match_status,
            "handoff_required": effective_handoff_required,
            "handoff_reason": effective_handoff_reason,
            "traveler": traveler,
            "lookup_phone": resolution.lookup_phone,
            "name_match_status": resolution.name_match_status,
            "actions": actions,
            "trip_result": trip_result,
            "commercial_context": commercial_context,
        }

        derived_stage, derived_priority = self.derive_lead_stage(preview)
        follow_up_status, derived_due_date = self.derive_follow_up(preview, timestamp)

        if effective_handoff_required:
            lead_stage = derived_stage
            priority = derived_priority
            follow_up_due_date = derived_due_date
        else:
            lead_stage = manual_lead_stage or derived_stage
            priority = manual_priority or derived_priority
            follow_up_due_date = manual_follow_up_due_date or derived_due_date

        traveler_status = str(traveler.get("status") or "") if traveler else ""
        customer_tier = self.infer_customer_tier(preview)

        interaction = self.create_interaction(
            timestamp=timestamp,
            channel=channel,
            customer_name=full_name,
            raw_phone=raw_phone,
            integrated_whatsapp=resolution.lookup_phone.get("normalized_whatsapp", ""),
            phone_lookup_key=resolution.lookup_phone.get("lookup_key", ""),
            traveler_id=traveler_id,
            matched_row=None,
            status_snapshot=status_snapshot or ("MULTIPLE_MATCHES" if resolution.match_status == "multiple_matches" else ""),
            intent="sales_inquiry",
            trip_type=trip_type or "",
            suggested_trips=", ".join(suggested_trip_ids),
            action_taken=", ".join(actions),
            handoff_required=effective_handoff_required,
            handoff_reason=effective_handoff_reason,
            agent_notes=agent_notes,
            flow_key="booking",
            step_key="lead_qualification",
            language=language,
            outcome=lead_stage,
        )

        lead_update = self.upsert_lead(
            customer_name=full_name,
            raw_phone=raw_phone,
            traveler_id=traveler_id or None,
            lead_stage=lead_stage,
            lead_source=source,
            channel=channel,
            preferred_trip_type=trip_type or "",
            interested_trip_ids=interested_trip_ids,
            suggested_trip_ids=", ".join(suggested_trip_ids),
            priority=priority,
            follow_up_status=follow_up_status,
            follow_up_due_date=follow_up_due_date,
            notes=agent_notes,
            traveler_status=traveler_status,
            customer_tier=customer_tier,
            match_status=resolution.match_status,
            last_interaction_id=interaction["interaction_id"],
            flow_key="booking",
            current_step="lead_qualification",
            handoff_required=effective_handoff_required,
            handoff_reason=effective_handoff_reason,
            language=language,
            country_code=country_code,
            force_create_new=force_create_new_lead and not (idempotency_key or session_id),
            session_id=session_id,
            idempotency_key=idempotency_key,
        )

        events = [
            self.create_booking_event(
                event_type="inquiry_received",
                event_label="Inquiry received",
                traveler_id=traveler_id,
                lead_id=lead_update["lead_id"],
                trip_id=preferred_trip_id,
                interaction_id=interaction["interaction_id"],
                channel=channel,
                notes=agent_notes,
                metadata={"trip_type": trip_type or ""},
                occurred_at=timestamp,
            )
        ]
        if suggested_trip_ids:
            events.append(
                self.create_booking_event(
                    event_type="trip_suggested",
                    event_label="Trip suggested",
                    traveler_id=traveler_id,
                    lead_id=lead_update["lead_id"],
                    trip_id=preferred_trip_id or suggested_trip_ids[0],
                    interaction_id=interaction["interaction_id"],
                    channel=channel,
                    notes=", ".join(suggested_trip_ids),
                    metadata={"suggested_trip_ids": suggested_trip_ids},
                    occurred_at=timestamp,
                )
            )
        if lead_stage in QUALIFIED_LEAD_STAGES:
            events.append(
                self.create_booking_event(
                    event_type="lead_qualified",
                    event_label="Lead qualified",
                    traveler_id=traveler_id,
                    lead_id=lead_update["lead_id"],
                    trip_id=preferred_trip_id or (suggested_trip_ids[0] if suggested_trip_ids else ""),
                    interaction_id=interaction["interaction_id"],
                    channel=channel,
                    notes=lead_stage,
                    metadata={"priority": priority},
                    occurred_at=timestamp,
                )
            )
        if effective_handoff_required:
            events.append(
                self.create_booking_event(
                    event_type="handoff_required",
                    event_label="Handoff required",
                    traveler_id=traveler_id,
                    lead_id=lead_update["lead_id"],
                    interaction_id=interaction["interaction_id"],
                    channel=channel,
                    notes=effective_handoff_reason,
                    metadata={"handoff_reason": effective_handoff_reason, "commercial_context": commercial_context},
                    occurred_at=timestamp,
                )
            )

        self.sync_agent_write_to_sheet(
            traveler_id=traveler_id,
            lead_id=lead_update["lead_id"],
            interaction_id=interaction["interaction_id"],
            event_ids=[event["event_id"] for event in events],
        )

        preview["write_result"] = {
            "created_traveler": created_traveler,
            "interaction_log": {
                "interaction_id": interaction["interaction_id"],
                "row": None,
            },
            "lead_update": lead_update,
            "event_trail": events,
            "write_result_contract": lead_update.get("write_result_contract") if isinstance(lead_update, dict) else None,
        }
        return preview

    def create_handoff_case(
        self,
        *,
        lead_id: str = "",
        traveler_id: str = "",
        trip_id: str = "",
        flow_key: str = "",
        reason_code: str = "",
        reason_text: str = "",
        priority: str = "",
        channel: str = "",
        status: str = "Pending",
        customer_name: str = "",
        agent_summary: str = "",
        customer_summary: str = "",
        notes: str = "",
        metadata: dict[str, Any] | None = None,
        lead_stage_override: str = "",
        update_lead: bool = True,
        deduplicate_open: bool = False,
        session_id: str = "",
        idempotency_key: str = "",
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        reason_code = str(reason_code or "").strip()
        reason_text = str(reason_text or reason_code or "manual_handoff").strip()
        priority = priority or self._default_handoff_priority(reason_code)
        package = {
            "reason_code": reason_code,
            "reason_text": reason_text,
            "customer_name": customer_name,
            "customer_summary": customer_summary,
            "agent_summary": agent_summary,
            "metadata": metadata or {},
        }
        display_reason = self._handoff_display_summary(reason_code, reason_text, package)
        notes_blob = self._handoff_notes_blob(display_reason, package, notes)
        timestamp = _utc_now().replace(microsecond=0)
        scope_key = traveler_id or lead_id or flow_key or customer_name
        resolved_idempotency_key = str(idempotency_key or "").strip() or self.handoff_idempotency_key(
            traveler_id_or_phone=scope_key,
            reason_code=reason_code or reason_text,
            session_id_or_open_lead_id=session_id or lead_id or flow_key,
        )
        deduplicated = False
        existing_status = ""

        with self.connect() as connection:
            existing = None
            if resolved_idempotency_key:
                existing = connection.execute(
                    """
                    SELECT handoff_id, status, priority, reason
                    FROM handoff_queue
                    WHERE idempotency_key = ?
                      AND LOWER(COALESCE(status, '')) IN ('new', 'pending', 'in progress')
                    ORDER BY created_at DESC
                    LIMIT 1
                    """,
                    (resolved_idempotency_key,),
                ).fetchone()
            if not existing and deduplicate_open:
                scope_sql = ""
                scope_value = ""
                if traveler_id:
                    scope_sql = "traveler_id = ?"
                    scope_value = traveler_id
                elif lead_id:
                    scope_sql = "lead_id = ?"
                    scope_value = lead_id
                elif flow_key:
                    scope_sql = "flow_key = ?"
                    scope_value = flow_key
                if scope_sql:
                    existing = connection.execute(
                        f"""
                        SELECT handoff_id, status, priority, reason
                        FROM handoff_queue
                        WHERE LOWER(COALESCE(status, '')) IN ('new', 'pending', 'in progress')
                          AND reason = ?
                          AND {scope_sql}
                        ORDER BY created_at DESC
                        LIMIT 1
                        """,
                        (display_reason, scope_value),
                    ).fetchone()

            if existing:
                handoff_id = str(existing["handoff_id"] or "")
                existing_status = str(existing["status"] or "Pending")
                existing_priority = str(existing["priority"] or "")
                priority_rank = {"Low": 1, "Medium": 2, "High": 3, "Critical": 4}
                if priority_rank.get(existing_priority, 0) > priority_rank.get(priority, 0):
                    priority = existing_priority
                connection.execute(
                    """
                    UPDATE handoff_queue
                    SET lead_id = COALESCE(lead_id, ?),
                        traveler_id = COALESCE(traveler_id, ?),
                        trip_id = COALESCE(trip_id, ?),
                        priority = ?
                    WHERE handoff_id = ?
                    """,
                    (lead_id or None, traveler_id or None, trip_id or None, priority, handoff_id),
                )
                deduplicated = True
            else:
                handoff_id = self._next_prefixed_id(connection, "handoff_queue", "handoff_id", "H-", 8)
                connection.execute(
                    """
                    INSERT INTO handoff_queue (
                        handoff_id, created_at, lead_id, traveler_id, trip_id, flow_key,
                        reason, priority, channel, status, notes, idempotency_key
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        handoff_id,
                        timestamp.isoformat(timespec="seconds"),
                        lead_id or None,
                        traveler_id or None,
                        trip_id or None,
                        flow_key or None,
                        display_reason,
                        priority,
                        channel or None,
                        status or "Pending",
                        notes_blob or None,
                        resolved_idempotency_key or None,
                    ),
                )
            if update_lead and lead_id:
                stage = lead_stage_override or ("Blocked" if reason_code == "blacklisted_customer" else "Needs Review")
                connection.execute(
                    """
                    UPDATE leads
                    SET handoff_required = 1,
                        handoff_reason = ?,
                        handoff_id = ?,
                        lead_stage = ?,
                        priority = ?,
                        updated_at = ?
                    WHERE lead_id = ?
                    """,
                    (
                        reason_code or reason_text,
                        handoff_id,
                        stage,
                        priority,
                        timestamp.isoformat(timespec="seconds"),
                        lead_id,
                    ),
                )
            connection.commit()

        event = None
        if not deduplicated:
            try:
                event = self.create_booking_event(
                    event_type="handoff_required",
                    event_label="Handoff required",
                    traveler_id=traveler_id,
                    lead_id=lead_id,
                    trip_id=trip_id,
                    channel=channel,
                    actor="system",
                    notes=display_reason,
                    metadata={"handoff_id": handoff_id, "priority": priority, "reason_code": reason_code, "package": package},
                    occurred_at=timestamp,
                )
                self.sync_agent_write_to_sheet(
                    traveler_id=traveler_id,
                    lead_id=lead_id,
                    trip_id=trip_id,
                    event_ids=[event["event_id"]] if event else [],
                )
                self.sync_record_to_sheet("Handoff Queue", handoff_id)
            except Exception:
                pass

        return {
            "handoff_id": handoff_id,
            "lead_id": lead_id,
            "traveler_id": traveler_id,
            "trip_id": trip_id,
            "reason_code": reason_code,
            "reason_text": display_reason,
            "priority": priority,
            "status": existing_status or status or "Pending",
            "package": package,
            "event_id": event["event_id"] if event else "",
            "deduplicated": deduplicated,
            "reused": deduplicated,
            "idempotency_key": resolved_idempotency_key,
            "write_result_contract": self.write_result_contract(
                status="reused" if deduplicated else "success",
                executed=not deduplicated,
                reused=deduplicated,
                record_type="handoff",
                record_id=handoff_id,
                idempotency_key=resolved_idempotency_key,
                customer_confirmation_allowed=True,
                safe_customer_message_key="handoff.already_recorded" if deduplicated else "handoff.created",
                audit={"reason_code": reason_code, "lead_id": lead_id, "traveler_id": traveler_id},
            ),
        }

    def update_existing_traveler_profile(
        self,
        traveler_id: str,
        *,
        birthday: str = "",
        gender: str = "",
        nationality: str = "",
        preferred_currency: str = "",
        timestamp: datetime | None = None,
    ) -> dict[str, Any]:
        timestamp = timestamp or _utc_now().replace(microsecond=0)
        updates: dict[str, Any] = {}
        with self.connect() as connection:
            row = connection.execute(
                "SELECT birthday, gender, nationality, preferred_currency FROM travelers WHERE traveler_id = ?",
                (traveler_id,),
            ).fetchone()
            if not row:
                return updates
            fields: dict[str, Any] = {"last_contacted_at": timestamp.isoformat()}
            if birthday and not row["birthday"]:
                fields["birthday"] = birthday
                updates["birthday"] = birthday
            if gender and not row["gender"]:
                fields["gender"] = gender
                updates["gender"] = gender
            if nationality and not row["nationality"]:
                fields["nationality"] = nationality
                updates["nationality"] = nationality
            if preferred_currency and not row["preferred_currency"]:
                fields["preferred_currency"] = preferred_currency
                updates["preferred_currency"] = preferred_currency
            assignments = ", ".join(f"{key} = ?" for key in fields)
            connection.execute(
                f"UPDATE travelers SET {assignments} WHERE traveler_id = ?",
                (*fields.values(), traveler_id),
            )
            connection.commit()
        return updates

    def set_guardian_consent(
        self,
        traveler_id: str,
        *,
        is_minor: bool,
        guardian_name: str = "",
        guardian_phone: str = "",
    ) -> dict[str, Any]:
        if not traveler_id:
            return {}
        self.ensure_operational_schema()
        with self.connect() as connection:
            connection.execute(
                """
                UPDATE travelers
                SET is_minor = ?, guardian_name = ?, guardian_phone = ?
                WHERE traveler_id = ?
                """,
                (1 if is_minor else 0, guardian_name or None, guardian_phone or None, traveler_id),
            )
            connection.commit()
        return {
            "traveler_id": traveler_id,
            "is_minor": bool(is_minor),
            "guardian_name": guardian_name or None,
            "guardian_phone": guardian_phone or None,
        }

    def flag_lead_guardian_approval(self, lead_id: str, *, requires_guardian_approval: bool) -> dict[str, Any]:
        if not lead_id:
            return {}
        self.ensure_operational_schema()
        with self.connect() as connection:
            connection.execute(
                "UPDATE leads SET requires_guardian_approval = ? WHERE lead_id = ?",
                (1 if requires_guardian_approval else 0, lead_id),
            )
            connection.commit()
        return {"lead_id": lead_id, "requires_guardian_approval": bool(requires_guardian_approval)}

    def create_interaction(
        self,
        *,
        timestamp: datetime,
        channel: str,
        customer_name: str,
        raw_phone: str,
        integrated_whatsapp: str,
        phone_lookup_key: str,
        traveler_id: str,
        matched_row: int | None,
        status_snapshot: str,
        intent: str,
        trip_type: str,
        suggested_trips: str,
        action_taken: str,
        handoff_required: bool,
        handoff_reason: str,
        agent_notes: str,
        flow_key: str = "",
        step_key: str = "",
        message_key: str = "",
        language: str = "",
        outcome: str = "",
    ) -> dict[str, Any]:
        with self.connect() as connection:
            interaction_id = self._next_daily_id(connection, "interactions", "interaction_id", "INT", timestamp)
            connection.execute(
                """
                INSERT INTO interactions (
                    interaction_id, timestamp, channel, customer_name, raw_phone, integrated_whatsapp,
                    phone_lookup_key, traveler_id, matched_row, status_snapshot, intent, trip_type,
                    suggested_trips, action_taken, handoff_required, handoff_reason, agent_notes,
                    flow_key, step_key, message_key, language, outcome
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    interaction_id,
                    timestamp.isoformat(timespec="seconds"),
                    channel,
                    customer_name,
                    raw_phone,
                    integrated_whatsapp or None,
                    phone_lookup_key or None,
                    traveler_id or None,
                    matched_row,
                    status_snapshot or None,
                    intent or None,
                    trip_type or None,
                    suggested_trips or None,
                    action_taken or None,
                    1 if handoff_required else 0,
                    handoff_reason or None,
                    agent_notes or None,
                    flow_key or None,
                    step_key or None,
                    message_key or None,
                    language or None,
                    outcome or None,
                ),
            )
            connection.commit()
        return {"interaction_id": interaction_id}

    def record_inbound_channel_event(
        self,
        *,
        channel: str,
        message_key: str,
        sender_id: str,
        recipient_id: str = "",
        text: str = "",
        attachments: list[dict[str, Any]] | None = None,
        timestamp: datetime | None = None,
        customer_name: str = "",
        flow_key: str = "",
        step_key: str = "",
        outcome: str = "",
    ) -> dict[str, Any]:
        """Persist an inbound channel message as an interaction with idempotency on message_key."""
        self.ensure_operational_schema()
        safe_timestamp = (timestamp or _utc_now()).replace(microsecond=0)
        safe_channel = str(channel or "").strip() or "External"
        safe_message_key = str(message_key or "").strip()
        safe_sender_id = str(sender_id or "").strip()
        safe_recipient_id = str(recipient_id or "").strip()
        safe_text = str(text or "").strip()
        serialized_attachments = attachments or []

        with self.connect() as connection:
            if safe_message_key:
                existing = connection.execute(
                    """
                    SELECT interaction_id
                    FROM interactions
                    WHERE channel = ? AND message_key = ?
                    LIMIT 1
                    """,
                    (safe_channel, safe_message_key),
                ).fetchone()
                if existing:
                    return {"interaction_id": str(existing["interaction_id"]), "created": False}

        notes_parts = []
        if safe_text:
            notes_parts.append(f"Inbound text: {safe_text}")
        if safe_recipient_id:
            notes_parts.append(f"Recipient ID: {safe_recipient_id}")
        if serialized_attachments:
            notes_parts.append(
                "Attachments: "
                + json.dumps(serialized_attachments, ensure_ascii=False, separators=(",", ":"))
            )

        try:
            interaction = self.create_interaction(
                timestamp=safe_timestamp,
                channel=safe_channel,
                customer_name=customer_name or f"{safe_channel} sender {safe_sender_id}",
                raw_phone=safe_sender_id,
                integrated_whatsapp="",
                phone_lookup_key="",
                traveler_id="",
                matched_row=None,
                status_snapshot="WEBHOOK_RECEIVED",
                intent="inbound_message",
                trip_type="",
                suggested_trips="",
                action_taken="webhook_received",
                handoff_required=False,
                handoff_reason="",
                agent_notes="\n".join(notes_parts),
                flow_key=flow_key or safe_channel.lower(),
                step_key=step_key or "inbound_webhook",
                message_key=safe_message_key,
                language="",
                outcome=outcome or "received",
            )
        except sqlite3.IntegrityError:
            if not safe_message_key:
                raise
            with self.connect() as connection:
                existing = connection.execute(
                    """
                    SELECT interaction_id
                    FROM interactions
                    WHERE channel = ? AND message_key = ?
                    LIMIT 1
                    """,
                    (safe_channel, safe_message_key),
                ).fetchone()
            if existing:
                return {"interaction_id": str(existing["interaction_id"]), "created": False}
            raise
        interaction["created"] = True
        return interaction

    def create_booking_draft(
        self,
        *,
        trip_id: str,
        traveler_id: str,
        traveler_name: str,
        room_type: str,
        room_group: str = "",
        boys_rooms_requested: int | str = 0,
        girls_rooms_requested: int | str = 0,
        room_requirements: list[dict[str, Any]] | dict[str, Any] | str | None = None,
        channel: str = "",
        lead_id: str = "",
        flight_option: str = "",
        date_option: str = "",
        currency: str = "",
        source: str = "",
        agent_notes: str = "",
        passport_required: bool = False,
        passport_status: str = "",
        group_size: int | str = 1,
        boys_count: int | str = 0,
        girls_count: int | str = 0,
        family_units: int | str = 0,
        session_id: str = "",
        idempotency_key: str = "",
        require_explicit_confirmation: bool = False,
        customer_confirmed: bool | str = True,
    ) -> dict[str, Any]:
        if require_explicit_confirmation and not self._as_bool(customer_confirmed, default=False):
            return {
                "booking_id": "",
                "booking_status": "blocked",
                "write_result_contract": self.write_result_contract(
                    status="blocked",
                    executed=False,
                    reused=False,
                    record_type="booking",
                    idempotency_key=str(idempotency_key or ""),
                    customer_confirmation_allowed=False,
                    error_code="confirmation_required",
                    safe_customer_message_key="booking.confirmation_required",
                    audit={"source": source, "session_id": session_id},
                ),
                "write_result": {},
            }
        return self.create_booking(
            trip_id=trip_id,
            traveler_id=traveler_id,
            traveler_name=traveler_name,
            room_type=room_type,
            room_group=room_group,
            boys_rooms_requested=boys_rooms_requested,
            girls_rooms_requested=girls_rooms_requested,
            room_requirements=room_requirements,
            channel=channel,
            lead_id=lead_id,
            flight_option=flight_option,
            date_option=date_option,
            currency=currency,
            booking_status="Draft",
            booking_source=source,
            payment_status="Pending",
            booking_notes=agent_notes,
            passport_required=passport_required,
            passport_status=passport_status,
            group_size=group_size,
            boys_count=boys_count,
            girls_count=girls_count,
            family_units=family_units,
            session_id=session_id,
            idempotency_key=idempotency_key,
        )

    @staticmethod
    def _validate_booking_transition(
        old_status: str | None,
        new_status: str | None,
        *,
        allow_employee_correction: bool = False,
        correction_note: str = "",
    ) -> None:
        current = str(old_status or "").strip() or "Draft"
        target = str(new_status or "").strip()
        if not target:
            raise ValueError("Booking status is required.")
        if current == target:
            return
        if target == "Cancelled" and current in BOOKING_ACTIVE_STATUSES:
            return
        allowed = BOOKING_TRANSITIONS.get(current, set())
        if target not in allowed:
            if allow_employee_correction and target in EMPLOYEE_BOOKING_STATUSES:
                if not str(correction_note or "").strip():
                    raise ValueError(
                        "Add a reason before making a non-standard booking status change."
                    )
                return
            raise ValueError(f"Invalid booking status transition: {current} -> {target}")

    @staticmethod
    def _read_booking_history(connection: sqlite3.Connection, booking_id: str) -> list[dict[str, Any]]:
        rows = connection.execute(
            """
            SELECT old_status, new_status, changed_at, changed_by, change_source, notes
            FROM booking_status_history
            WHERE booking_id = ?
            ORDER BY changed_at ASC, history_id ASC
            """,
            (booking_id,),
        ).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def _reconcile_trip_room_holds(connection: sqlite3.Connection, trip_id: str = "") -> int:
        bookings_table = connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'trip_bookings'"
        ).fetchone()
        if not bookings_table:
            return 0

        params: tuple[Any, ...] = (trip_id,) if trip_id else ()
        where_clause = "WHERE trip_id = ?" if trip_id else ""
        trip_rows = connection.execute(
            f"""
                SELECT *
                FROM trips
                {where_clause}
                """,
            params,
        ).fetchall()
        if not trip_rows:
            return 0

        booking_columns = UnifiedCRMService._table_columns(connection, "trip_bookings")
        select_fields = ["trip_id", "room_type", "booking_status"]
        for optional in ("room_group", "boys_rooms_requested", "girls_rooms_requested", "room_requirements_json"):
            if optional in booking_columns:
                select_fields.append(optional)
        booking_where = "AND trip_id = ?" if trip_id else ""
        booking_rows = connection.execute(
            f"""
            SELECT {', '.join(select_fields)}
            FROM trip_bookings
            WHERE lower(trim(COALESCE(booking_status, ''))) <> 'cancelled'
            {booking_where}
            """,
            params,
        ).fetchall()
        booking_counts: dict[str, dict[str, int]] = {}
        for row in booking_rows:
            trip_key = str(row["trip_id"] or "")
            counts = booking_counts.setdefault(
                trip_key,
                {
                    "single": 0,
                    "double": 0,
                    "triple": 0,
                    "boys_double": 0,
                    "girls_double": 0,
                    "boys_triple": 0,
                    "girls_triple": 0,
                },
            )
            requirements = UnifiedCRMService._booking_row_room_requirements(row)
            for requirement in requirements:
                room_type = str(requirement.get("room_type") or "").strip().lower()
                group = str(requirement.get("room_group") or "").strip().lower()
                rooms = max(0, int(requirement.get("rooms") or 0))
                if room_type in {"single", "double", "triple"}:
                    counts[room_type] += rooms
                if room_type in {"double", "triple"} and group in {"boys", "girls"}:
                    counts[f"{group}_{room_type}"] += rooms

        updated = 0
        for row in trip_rows:
            expected = booking_counts.get(
                str(row["trip_id"]),
                {
                    "single": 0,
                    "double": 0,
                    "triple": 0,
                    "boys_double": 0,
                    "girls_double": 0,
                    "boys_triple": 0,
                    "girls_triple": 0,
                },
            )
            update_values = {
                "draft_holds_single": int(expected["single"]),
                "draft_holds_double": int(expected["double"]),
                "draft_holds_triple": int(expected["triple"]),
            }
            trip_columns = UnifiedCRMService._table_columns(connection, "trips")
            for key in ("boys_double", "girls_double", "boys_triple", "girls_triple"):
                hold_key = f"draft_holds_{key}"
                if hold_key in trip_columns:
                    update_values[hold_key] = int(expected[key])
            current = {key: int(row[key] or 0) for key in update_values}
            if current == update_values:
                continue
            connection.execute(
                f"""
                UPDATE trips
                SET {', '.join(f'{key} = ?' for key in update_values)}
                WHERE trip_id = ?
                """,
                (*update_values.values(), row["trip_id"]),
            )
            updated += 1
        return updated

    def reconcile_trip_room_holds(self, trip_id: str = "") -> dict[str, Any]:
        self.ensure_operational_schema()
        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            updated = self._reconcile_trip_room_holds(connection, trip_id)
            connection.commit()
        return {"trip_id": trip_id, "updated_trips": updated}

    def update_booking_status(
        self,
        booking_id: str,
        *,
        new_status: str | None = None,
        new_payment_status: str | None = None,
        changed_by: str = "system-ui",
        change_source: str = "crm-ui",
        notes: str = "",
        allow_employee_correction: bool = False,
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        if not new_status and not new_payment_status and not notes:
            raise ValueError("No booking update supplied.")
        timestamp = _utc_now().replace(microsecond=0)
        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            booking = connection.execute(
                """
                SELECT booking_id, booking_status, payment_status, traveler_id, lead_id, trip_id, interaction_id, booking_source
                FROM trip_bookings
                WHERE booking_id = ?
                """,
                (booking_id,),
            ).fetchone()
            if not booking:
                raise ValueError(f"Booking was not found: {booking_id}")

            old_status = str(booking["booking_status"] or "Draft").strip() or "Draft"
            # Validate BOTH statuses before writing either. This path used to
            # check the booking status and apply the payment status with no
            # check at all, so it could move a booking backwards from Fully
            # Paid to Deposit Paid -- something the CRM route has always
            # refused. Same rules, same override convention, one definition in
            # payment_rules.py.
            if new_payment_status:
                validate_payment_transition(
                    booking["payment_status"],
                    new_payment_status,
                    allow_employee_correction=allow_employee_correction,
                    correction_note=notes,
                )
            if new_status:
                self._validate_booking_transition(
                    old_status,
                    new_status,
                    allow_employee_correction=allow_employee_correction,
                    correction_note=notes,
                )
                connection.execute(
                    "UPDATE trip_bookings SET booking_status = ? WHERE booking_id = ?",
                    (new_status, booking_id),
                )
                connection.execute(
                    """
                    INSERT INTO booking_status_history (
                        booking_id, old_status, new_status, changed_at, changed_by, change_source, notes
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        booking_id,
                        old_status,
                        new_status,
                        timestamp.isoformat(timespec="seconds"),
                        changed_by or None,
                        change_source or None,
                        notes or None,
                    ),
                )
            if new_payment_status:
                connection.execute(
                    "UPDATE trip_bookings SET payment_status = ? WHERE booking_id = ?",
                    (new_payment_status, booking_id),
                )
            self._reconcile_trip_room_holds(connection, str(booking["trip_id"] or ""))
            connection.commit()

        event_notes = []
        if new_status:
            event_notes.append(f"Booking status: {old_status} -> {new_status}")
        if new_payment_status:
            event_notes.append(f"Payment status: {new_payment_status}")
        if notes:
            event_notes.append(notes)
        if event_notes:
            self.create_booking_event(
                event_type="booking_status_updated",
                event_label="Booking lifecycle updated",
                traveler_id=str(booking["traveler_id"] or ""),
                lead_id=str(booking["lead_id"] or ""),
                booking_id=booking_id,
                trip_id=str(booking["trip_id"] or ""),
                interaction_id=str(booking["interaction_id"] or ""),
                channel=str(booking["booking_source"] or ""),
                actor=changed_by or "system-ui",
                notes=" | ".join(event_notes),
                metadata={
                    "old_status": old_status,
                    "new_status": new_status,
                    "payment_status": new_payment_status,
                    "change_source": change_source,
                },
                occurred_at=timestamp,
            )
        traveler_id = str(booking["traveler_id"] or "").strip()
        if traveler_id:
            self.recalculate_traveler_stats(traveler_id)
            self.sync_record_to_sheet("Travelers", traveler_id)

        with self.connect() as connection:
            current = connection.execute(
                "SELECT booking_id, booking_status, payment_status FROM trip_bookings WHERE booking_id = ?",
                (booking_id,),
            ).fetchone()
            history = self._read_booking_history(connection, booking_id)
        return {
            "booking_id": booking_id,
            "booking_status": current["booking_status"] if current else None,
            "payment_status": current["payment_status"] if current else None,
            "history": history,
        }

    def _existing_booking_write_result(
        self,
        *,
        row: sqlite3.Row,
        normalized_requirements: list[dict[str, Any]],
        idempotency_key: str,
        status: str,
        safe_customer_message_key: str,
    ) -> dict[str, Any]:
        booking_id = str(row["booking_id"] or "")
        trip_id = str(row["trip_id"] or "")
        traveler_id = str(row["traveler_id"] or "")
        lead_id = str(self._row_value(row, "lead_id") or "")
        room_requirements = self._booking_row_room_requirements(row) or normalized_requirements
        contract = self.write_result_contract(
            status=status,
            executed=False,
            reused=True,
            record_type="booking",
            record_id=booking_id,
            idempotency_key=idempotency_key,
            customer_confirmation_allowed=True,
            safe_customer_message_key=safe_customer_message_key,
            audit={"traveler_id": traveler_id, "trip_id": trip_id},
        )
        result = {
            "booking_id": booking_id,
            "booking_row": None,
            "trip_id": trip_id,
            "trip_name": str(self._row_value(row, "trip_name") or ""),
            "traveler_id": traveler_id,
            "traveler_name": str(self._row_value(row, "traveler_name") or ""),
            "room_type": str(self._row_value(row, "room_type") or ""),
            "room_group": str(self._row_value(row, "room_group") or ""),
            "boys_rooms_requested": self._as_int(self._row_value(row, "boys_rooms_requested"), default=0) or 0,
            "girls_rooms_requested": self._as_int(self._row_value(row, "girls_rooms_requested"), default=0) or 0,
            "room_requirements": room_requirements,
            "flight_option": str(self._row_value(row, "flight_option") or ""),
            "date_option": str(self._row_value(row, "date_option") or ""),
            "currency": str(self._row_value(row, "currency") or ""),
            "booking_status": str(self._row_value(row, "booking_status") or "Draft"),
            "payment_status": str(self._row_value(row, "payment_status") or "Pending"),
            "refund_amount": self._row_value(row, "refund_amount"),
            "passport_required": bool(self._as_int(self._row_value(row, "passport_required"), default=0) or 0),
            "passport_status": str(self._row_value(row, "passport_status") or ""),
            "interaction": {"interaction_id": str(self._row_value(row, "interaction_id") or "")},
            "lead_update": None,
            "event_trail": [],
            "reused": True,
            "idempotency_key": idempotency_key,
            "write_result_contract": contract,
        }
        result["write_result"] = {
            "booking_draft": {
                "booking_id": booking_id,
                "trip_id": trip_id,
                "traveler_id": traveler_id,
                "lead_id": lead_id,
                "room_group": result["room_group"],
                "boys_rooms_requested": result["boys_rooms_requested"],
                "girls_rooms_requested": result["girls_rooms_requested"],
                "room_requirements": room_requirements,
            },
            "interaction_log": result["interaction"],
            "lead_update": None,
            "event_trail": [],
            "write_result_contract": contract,
        }
        return result

    def create_booking(
        self,
        *,
        trip_id: str,
        traveler_id: str,
        traveler_name: str,
        room_type: str,
        room_group: str = "",
        boys_rooms_requested: int | str = 0,
        girls_rooms_requested: int | str = 0,
        room_requirements: list[dict[str, Any]] | dict[str, Any] | str | None = None,
        channel: str = "",
        lead_id: str = "",
        flight_option: str = "",
        date_option: str = "",
        currency: str = "",
        booking_status: str = "Draft",
        booking_source: str = "",
        interaction_id: str = "",
        payment_status: str = "Pending",
        booking_notes: str = "",
        passport_required: bool = False,
        passport_status: str = "",
        group_size: int | str = 1,
        boys_count: int | str = 0,
        girls_count: int | str = 0,
        family_units: int | str = 0,
        session_id: str = "",
        idempotency_key: str = "",
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        timestamp = _utc_now().replace(microsecond=0)
        if room_type not in ROOM_HOLD_COLUMNS:
            raise ValueError(f"Unsupported room type {room_type!r}")
        normalized_requirements = self._normalize_room_requirements(
            room_type=room_type,
            room_group=room_group,
            boys_rooms_requested=boys_rooms_requested,
            girls_rooms_requested=girls_rooms_requested,
            room_requirements=room_requirements,
        )
        if not normalized_requirements:
            raise ValueError("At least one room requirement is required.")
        boys_rooms_total = sum(int(item["rooms"]) for item in normalized_requirements if item.get("room_group") == "boys")
        girls_rooms_total = sum(int(item["rooms"]) for item in normalized_requirements if item.get("room_group") == "girls")
        resolved_room_group = str(room_group or "").strip().lower()
        if boys_rooms_total and girls_rooms_total:
            resolved_room_group = "mixed"
        elif boys_rooms_total:
            resolved_room_group = "boys"
        elif girls_rooms_total:
            resolved_room_group = "girls"
        room_requirements_json = json.dumps(normalized_requirements, ensure_ascii=False, separators=(",", ":"))
        resolved_idempotency_key = str(idempotency_key or "").strip() or self.booking_idempotency_key(
            traveler_id=traveler_id,
            trip_id=trip_id,
            room_requirements=normalized_requirements,
            session_id=session_id,
        )
        if booking_status not in BOOKING_LIFECYCLE_STATUSES:
            raise ValueError(f"Unsupported booking status {booking_status!r}")
        if booking_status not in {"Draft", "Waiting Customer"}:
            raise ValueError("New bookings must start in Draft or Waiting Customer.")

        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._reconcile_trip_room_holds(connection, trip_id)
            trip = connection.execute(
                """
                SELECT *
                 FROM trips
                 WHERE trip_id = ?
                 """,
                 (trip_id,),
             ).fetchone()
            if not trip:
                raise ValueError(f"Trip was not found for booking: {trip_id}")
            sales_status = str(trip["sales_status"] or "").strip().lower()
            if sales_status and sales_status != "open":
                raise ValueError(f"Trip {trip_id} is not bookable because its status is {trip['sales_status']}")

            existing_by_key = connection.execute(
                """
                SELECT * FROM trip_bookings
                WHERE idempotency_key = ?
                  AND LOWER(TRIM(COALESCE(booking_status, ''))) NOT IN ('cancelled', 'canceled', 'completed')
                ORDER BY draft_created_at DESC
                LIMIT 1
                """,
                (resolved_idempotency_key,),
            ).fetchone() if resolved_idempotency_key else None
            if existing_by_key:
                connection.commit()
                return self._existing_booking_write_result(
                    row=existing_by_key,
                    normalized_requirements=normalized_requirements,
                    idempotency_key=resolved_idempotency_key,
                    status="reused",
                    safe_customer_message_key="booking.already_recorded",
                )

            existing_active = connection.execute(
                """
                SELECT * FROM trip_bookings
                WHERE traveler_id = ?
                  AND trip_id = ?
                  AND LOWER(TRIM(COALESCE(booking_status, ''))) NOT IN ('cancelled', 'canceled', 'completed')
                ORDER BY draft_created_at DESC
                LIMIT 1
                """,
                (traveler_id, trip_id),
            ).fetchone()
            if existing_active:
                connection.commit()
                return self._existing_booking_write_result(
                    row=existing_active,
                    normalized_requirements=normalized_requirements,
                    idempotency_key=resolved_idempotency_key,
                    status="duplicate",
                    safe_customer_message_key="booking.duplicate_active",
                )

            aggregate_availability: dict[str, int] = {}
            category_availability: dict[str, int] = {}
            for requirement in normalized_requirements:
                req_room_type = str(requirement.get("room_type") or "").strip().title()
                group = str(requirement.get("room_group") or "").strip().lower()
                rooms = max(0, int(requirement.get("rooms") or 0))
                remaining = self._as_int(trip[ROOM_REMAINING_COLUMNS[req_room_type]])
                current_hold = self._as_int(trip[ROOM_HOLD_COLUMNS[req_room_type]], default=0) or 0
                if remaining is None:
                    raise ValueError(f"Capacity is not configured for room type {req_room_type}")
                aggregate_available = remaining - current_hold
                aggregate_availability[req_room_type] = aggregate_available
                requested_same_type = sum(
                    int(item.get("rooms") or 0)
                    for item in normalized_requirements
                    if str(item.get("room_type") or "").strip().title() == req_room_type
                )
                if requested_same_type > aggregate_available:
                    raise ValueError(f"No remaining draftable capacity for {req_room_type}")
                if req_room_type in {"Double", "Triple"} and group in {"boys", "girls"}:
                    capacity_key = f"{group}_{req_room_type.lower()}"
                    hold_key = f"draft_holds_{capacity_key}"
                    category_capacity = self._as_int(self._row_value(trip, capacity_key), default=0) or 0
                    opposite_group = "girls" if group == "boys" else "boys"
                    opposite_capacity = self._as_int(
                        self._row_value(trip, f"{opposite_group}_{req_room_type.lower()}"),
                        default=0,
                    ) or 0
                    if not (category_capacity or opposite_capacity):
                        continue
                    category_hold = self._as_int(self._row_value(trip, hold_key), default=0) or 0
                    category_available = category_capacity - category_hold
                    category_availability[capacity_key] = category_available
                    if rooms > category_available:
                        raise ValueError(
                            f"{group.title()} {req_room_type.lower()} rooms are unavailable: requested {rooms}, available {max(category_available, 0)}."
                        )
            available = min(aggregate_availability.values()) if aggregate_availability else 0

            if not interaction_id:
                interaction_id = self._next_daily_id(connection, "interactions", "interaction_id", "INT", timestamp)
                connection.execute(
                    """
                    INSERT INTO interactions (
                        interaction_id, timestamp, channel, customer_name, raw_phone,
                        traveler_id, status_snapshot, intent, trip_type, suggested_trips,
                        action_taken, handoff_required, agent_notes, flow_key, step_key, outcome
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        interaction_id,
                        timestamp.isoformat(timespec="seconds"),
                        channel or None,
                        traveler_name,
                        "",
                        traveler_id,
                        "BOOKING_DRAFT",
                        "booking_request",
                        str(trip["type"] or ""),
                        trip_id,
                        "booking_draft_created",
                        0,
                        booking_notes or None,
                        "booking",
                        "booking_draft_created",
                        "Booking Draft Created",
                    ),
                )

            booking_id = self._next_booking_id(connection, trip_id, traveler_id)
            booking_columns = [
                "booking_id", "trip_id", "trip_name", "traveler_id", "traveler_name", "room_type",
                "flight_option", "date_option", "currency", "booking_status", "draft_created_at",
                "booking_source", "lead_id", "interaction_id", "alert_id", "payment_status",
                "passport_required", "passport_status", "group_size", "boys_count", "girls_count",
                "family_units", "booking_notes",
            ]
            booking_values = [
                booking_id,
                trip_id,
                str(trip["trip_name"]) if trip else "",
                traveler_id,
                traveler_name,
                room_type,
                flight_option or None,
                date_option or None,
                currency or None,
                booking_status,
                timestamp.isoformat(timespec="seconds"),
                booking_source or None,
                lead_id or None,
                interaction_id or None,
                None,
                payment_status or None,
                1 if passport_required else 0,
                passport_status or None,
                self._as_int(group_size, default=1) or 1,
                self._as_int(boys_count, default=0) or 0,
                self._as_int(girls_count, default=0) or 0,
                self._as_int(family_units, default=0) or 0,
                booking_notes or None,
            ]
            booking_table_columns = self._table_columns(connection, "trip_bookings")
            optional_booking_values = {
                "room_group": resolved_room_group or None,
                "boys_rooms_requested": boys_rooms_total,
                "girls_rooms_requested": girls_rooms_total,
                "room_requirements_json": room_requirements_json,
                "idempotency_key": resolved_idempotency_key or None,
            }
            for column, value in optional_booking_values.items():
                if column in booking_table_columns:
                    booking_columns.append(column)
                    booking_values.append(value)
            connection.execute(
                f"""
                INSERT INTO trip_bookings ({', '.join(booking_columns)})
                VALUES ({', '.join(['?'] * len(booking_columns))})
                """,
                booking_values,
            )
            self._reconcile_trip_room_holds(connection, trip_id)
            if traveler_id:
                connection.execute(
                    "UPDATE travelers SET last_booking_id = ? WHERE traveler_id = ?",
                    (booking_id, traveler_id),
                )
            lead_update = None
            if lead_id:
                lead_update = self._update_lead_for_booking(
                    connection,
                    lead_id=lead_id,
                    booking_id=booking_id,
                    trip_id=trip_id,
                    interaction_id=interaction_id,
                    timestamp=timestamp,
                )
            connection.commit()

        booking_event = self.create_booking_event(
            event_type="booking_draft_created",
            event_label="Booking draft created",
            traveler_id=traveler_id,
            lead_id=lead_id,
            booking_id=booking_id,
            trip_id=trip_id,
            interaction_id=interaction_id,
            channel=channel,
            notes=booking_notes,
            metadata={
                "room_type": room_type,
                "room_group": resolved_room_group,
                "room_requirements": normalized_requirements,
                "available_before_draft": available,
                "available_after_draft": available - sum(int(item["rooms"]) for item in normalized_requirements),
                "category_availability_before": category_availability,
            },
            occurred_at=timestamp,
        )
        follow_up_event = self.create_booking_event(
            event_type="payment_follow_up",
            event_label="Payment / follow-up pending",
            traveler_id=traveler_id,
            lead_id=lead_id,
            booking_id=booking_id,
            trip_id=trip_id,
            interaction_id=interaction_id,
            channel=channel,
            notes=payment_status,
            metadata={"payment_status": payment_status},
            occurred_at=timestamp,
        )

        if traveler_id:
            self.recalculate_traveler_stats(traveler_id)

        self.sync_agent_write_to_sheet(
            traveler_id=traveler_id,
            lead_id=lead_id,
            interaction_id=interaction_id,
            booking_id=booking_id,
            trip_id=trip_id,
            event_ids=[booking_event["event_id"], follow_up_event["event_id"]],
        )

        result = {
            "booking_id": booking_id,
            "booking_row": None,
            "trip_id": trip_id,
            "trip_name": str(trip["trip_name"] or ""),
            "traveler_id": traveler_id,
            "traveler_name": traveler_name,
            "room_type": room_type,
            "room_group": resolved_room_group,
            "boys_rooms_requested": boys_rooms_total,
            "girls_rooms_requested": girls_rooms_total,
            "room_requirements": normalized_requirements,
            "flight_option": flight_option,
            "date_option": date_option,
            "currency": currency,
            "booking_status": booking_status,
            "payment_status": payment_status or "Pending",
            "refund_amount": None,
            "passport_required": bool(passport_required),
            "passport_status": passport_status or ("pending" if passport_required else ""),
            "boys_count": self._as_int(boys_count, default=0) or 0,
            "girls_count": self._as_int(girls_count, default=0) or 0,
            "family_units": self._as_int(family_units, default=0) or 0,
            "available_before_draft": available,
            "available_after_draft": available - sum(int(item["rooms"]) for item in normalized_requirements),
            "interaction": {"interaction_id": interaction_id},
            "lead_update": lead_update,
            "event_trail": [booking_event, follow_up_event],
            "reused": False,
            "idempotency_key": resolved_idempotency_key,
        }
        contract = self.write_result_contract(
            status="success",
            executed=True,
            reused=False,
            record_type="booking",
            record_id=booking_id,
            idempotency_key=resolved_idempotency_key,
            customer_confirmation_allowed=True,
            safe_customer_message_key="booking.created",
            audit={"traveler_id": traveler_id, "trip_id": trip_id},
        )
        result["write_result_contract"] = contract
        result["write_result"] = {
            "booking_draft": {
                "booking_id": booking_id,
                "trip_id": trip_id,
                "traveler_id": traveler_id,
                "lead_id": lead_id,
                "room_group": resolved_room_group,
                "boys_rooms_requested": boys_rooms_total,
                "girls_rooms_requested": girls_rooms_total,
                "room_requirements": normalized_requirements,
            },
            "interaction_log": {"interaction_id": interaction_id},
            "lead_update": lead_update,
            "event_trail": [booking_event, follow_up_event],
            "write_result_contract": contract,
        }
        return result

    def qualify_lead(self, lead_id: str, *, channel: str = "") -> bool:
        self.ensure_operational_schema()
        timestamp = _utc_now().replace(microsecond=0)
        with self.connect() as connection:
            lead = connection.execute(
                """
                SELECT lead_id, traveler_id, customer_name, raw_phone, channel, interested_trip_ids,
                       preferred_trip_type, interaction_count
                FROM leads
                WHERE lead_id = ?
                """,
                (lead_id,),
            ).fetchone()
            if not lead:
                return False
            interaction_id = self._next_daily_id(connection, "interactions", "interaction_id", "INT", timestamp)
            resolved_channel = channel or str(lead["channel"] or "")
            connection.execute(
                """
                INSERT INTO interactions (
                    interaction_id, timestamp, channel, customer_name, raw_phone, traveler_id,
                    status_snapshot, intent, trip_type, suggested_trips, action_taken,
                    handoff_required, flow_key, step_key, outcome
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    interaction_id,
                    timestamp.isoformat(timespec="seconds"),
                    resolved_channel or None,
                    lead["customer_name"],
                    lead["raw_phone"],
                    lead["traveler_id"],
                    "QUALIFIED",
                    "lead_qualification",
                    lead["preferred_trip_type"],
                    lead["interested_trip_ids"],
                    "lead_qualified",
                    0,
                    "booking",
                    "lead_qualified",
                    "Qualified",
                ),
            )
            connection.execute(
                """
                UPDATE leads
                SET lead_stage = ?, priority = ?, follow_up_status = ?, follow_up_due_date = ?,
                    updated_at = ?, last_interaction_id = ?, interaction_count = ?
                WHERE lead_id = ?
                """,
                (
                    "Qualified",
                    "High",
                    "Follow Up Soon",
                    (timestamp.date() + timedelta(days=2)).isoformat(),
                    timestamp.isoformat(timespec="seconds"),
                    interaction_id,
                    (self._as_int(lead["interaction_count"], default=0) or 0) + 1,
                    lead_id,
                ),
            )
            connection.commit()

        event = self.create_booking_event(
            event_type="lead_qualified",
            event_label="Lead qualified",
            traveler_id=str(lead["traveler_id"] or ""),
            lead_id=lead_id,
            trip_id=str(lead["interested_trip_ids"] or "").split(",", 1)[0].strip(),
            interaction_id=interaction_id,
            channel=resolved_channel,
            notes="Manual qualification",
            occurred_at=timestamp,
        )
        self.sync_agent_write_to_sheet(
            traveler_id=str(lead["traveler_id"] or ""),
            lead_id=lead_id,
            interaction_id=interaction_id,
            event_ids=[event["event_id"]],
        )
        return True

    def update_lead_stage(
        self,
        lead_id: str,
        *,
        requested_stage: str,
        priority: str = "",
        follow_up_status: str = "",
        follow_up_due_date: str = "",
        notes: str = "",
        channel: str = "",
        flow_key: str = "",
        current_step: str = "",
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        stage = str(requested_stage or "").strip()
        if not lead_id:
            raise ValueError("Lead ID is required.")
        if not stage:
            raise ValueError("Lead stage is required.")

        timestamp = _utc_now().replace(microsecond=0)
        with self.connect() as connection:
            lead = connection.execute(
                """
                SELECT lead_id, traveler_id, customer_name, raw_phone, lead_stage, priority,
                       follow_up_status, follow_up_due_date, interaction_count
                FROM leads
                WHERE lead_id = ?
                """,
                (lead_id,),
            ).fetchone()
            if not lead:
                raise ValueError(f"Lead was not found: {lead_id}")

            resolved_priority = priority or str(lead["priority"] or "") or "Medium"
            resolved_follow_up_status = follow_up_status or str(lead["follow_up_status"] or "") or ""
            resolved_follow_up_due_date = follow_up_due_date or str(lead["follow_up_due_date"] or "") or ""
            connection.execute(
                """
                UPDATE leads
                SET lead_stage = ?, priority = ?, follow_up_status = ?, follow_up_due_date = ?,
                    updated_at = ?
                WHERE lead_id = ?
                """,
                (
                    stage,
                    resolved_priority,
                    resolved_follow_up_status or None,
                    resolved_follow_up_due_date or None,
                    timestamp.isoformat(timespec="seconds"),
                    lead_id,
                ),
            )
            connection.commit()

        event = self.create_booking_event(
            event_type="lead_stage_updated",
            event_label="Lead stage updated",
            traveler_id=str(lead["traveler_id"] or ""),
            lead_id=lead_id,
            channel=channel,
            actor="system",
            notes=notes or f"Lead stage updated to {stage}",
            metadata={"requested_stage": stage, "flow_key": flow_key, "current_step": current_step},
            occurred_at=timestamp,
        )
        try:
            self.sync_agent_write_to_sheet(
                traveler_id=str(lead["traveler_id"] or ""),
                lead_id=lead_id,
                event_ids=[event["event_id"]],
            )
        except Exception:
            pass
        try:
            self.auto_create_booking_from_lead_if_ready(lead_id, actor="agent", trigger_source="agent")
        except Exception:
            pass
        return {
            "lead_id": lead_id,
            "lead_stage": stage,
            "priority": resolved_priority,
            "follow_up_status": resolved_follow_up_status,
            "follow_up_due_date": resolved_follow_up_due_date,
            "traveler_id": str(lead["traveler_id"] or ""),
            "customer_name": str(lead["customer_name"] or ""),
            "event_id": event["event_id"],
        }

    def auto_create_booking_from_lead_if_ready(
        self,
        lead_id: str,
        *,
        actor: str = "agent",
        trigger_source: str = "agent",
    ) -> dict[str, Any] | None:
        """Create a Draft Booking the first time a Lead reaches Booking Draft.

        Fires from update_lead_stage (agent path, this SQLite backend) so a
        Booking Draft never depends on someone remembering to create the
        Booking as a separate step. Idempotent by lead_id/traveler_id: a lead
        bouncing Qualified -> Booking Draft -> Waiting Customer Reply ->
        Booking Draft again reuses the booking it already made the first
        time, matching apps/api/app/services/booking_automation.py's
        Postgres/SQLAlchemy sibling for the manual-CRM-UI and Postgres-agent
        paths (see services/crm/README.md for how the two backends relate).
        """
        self.ensure_operational_schema()
        timestamp = _utc_now().replace(microsecond=0)
        with self.connect() as connection:
            lead = connection.execute(
                """
                SELECT lead_id, traveler_id, customer_name, lead_stage, channel, language,
                       preferred_trip_type, interested_trip_ids, suggested_trip_ids,
                       group_size, booking_id, assigned_to, assigned_to_user_id
                FROM leads
                WHERE lead_id = ?
                """,
                (lead_id,),
            ).fetchone()
            if not lead or str(lead["lead_stage"] or "").strip() not in BOOKING_DRAFT_LEAD_STAGES:
                return None

            existing = connection.execute(
                "SELECT * FROM trip_bookings WHERE lead_id = ? ORDER BY draft_created_at DESC LIMIT 1",
                (lead_id,),
            ).fetchone()
            traveler_id = str(lead["traveler_id"] or "").strip()
            if existing is None and traveler_id:
                placeholders = ",".join("?" for _ in BOOKING_ACTIVE_STATUSES)
                existing = connection.execute(
                    f"""
                    SELECT * FROM trip_bookings
                    WHERE traveler_id = ?
                      AND booking_status IN ({placeholders})
                    ORDER BY draft_created_at DESC LIMIT 1
                    """,
                    (traveler_id, *BOOKING_ACTIVE_STATUSES),
                ).fetchone()
            if existing is not None:
                if not lead["booking_id"]:
                    connection.execute(
                        "UPDATE leads SET booking_id = ? WHERE lead_id = ?",
                        (existing["booking_id"], lead_id),
                    )
                    connection.commit()
                return {"booking_id": existing["booking_id"], "created": False}

            if not traveler_id:
                return None
            traveler = connection.execute(
                "SELECT traveler_id, full_name FROM travelers WHERE traveler_id = ?",
                (traveler_id,),
            ).fetchone()
            if not traveler:
                return None

            trip_id = str(lead["interested_trip_ids"] or lead["suggested_trip_ids"] or "").split(",")[0].strip()
            trip = connection.execute("SELECT * FROM trips WHERE trip_id = ?", (trip_id,)).fetchone() if trip_id else None

            missing_fields = ["room_type"]
            if not trip:
                missing_fields.append("trip")

            booking_id = self._next_booking_id(connection, trip_id or "", traveler_id)
            notes = (
                f"Auto-created from Lead {lead_id} upon reaching Booking Draft "
                f"(channel: {lead['channel'] or 'n/a'}, language: {lead['language'] or 'n/a'}, "
                f"trip type: {lead['preferred_trip_type'] or 'n/a'})."
            )
            connection.execute(
                """
                INSERT INTO trip_bookings (
                    booking_id, trip_id, trip_name, traveler_id, traveler_name, room_type,
                    booking_status, draft_created_at, booking_source, lead_id, payment_status,
                    passport_required, group_size, booking_notes, missing_info,
                    assigned_to, assigned_to_user_id, assigned_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    booking_id,
                    trip["trip_id"] if trip else None,
                    trip["trip_name"] if trip else None,
                    traveler_id,
                    str(traveler["full_name"] or ""),
                    None,
                    "Draft",
                    timestamp.isoformat(timespec="seconds"),
                    f"Auto ({trigger_source})",
                    lead_id,
                    "Pending",
                    1 if trip and str(trip["type"] or "").strip().lower() == "international" else 0,
                    self._as_int(lead["group_size"], default=1) or 1,
                    notes,
                    1,
                    str(lead["assigned_to"] or "") or None,
                    lead["assigned_to_user_id"],
                    timestamp.isoformat(timespec="seconds") if lead["assigned_to_user_id"] else None,
                ),
            )
            connection.execute("UPDATE leads SET booking_id = ? WHERE lead_id = ?", (booking_id, lead_id))
            connection.commit()

        try:
            self.create_booking_event(
                event_type="booking_auto_created",
                event_label="Booking auto-created",
                traveler_id=traveler_id,
                lead_id=lead_id,
                booking_id=booking_id,
                trip_id=trip_id,
                channel=str(lead["channel"] or ""),
                actor=actor,
                notes="Booking auto-created upon reaching Booking Draft stage",
                metadata={
                    "booking_id": booking_id,
                    "trigger_source": trigger_source,
                    "missing_info": True,
                    "missing_fields": missing_fields,
                },
                occurred_at=timestamp,
            )
        except Exception:
            pass
        return {"booking_id": booking_id, "created": True, "missing_info": True, "missing_fields": missing_fields}

    @staticmethod
    def _merge_agent_snapshot_notes(existing_notes: str, snapshot_lines: list[str]) -> str:
        marker = "[AI Agent Snapshot]"
        base = str(existing_notes or "").strip()
        if marker in base:
            base = base.split(marker, 1)[0].rstrip()
        snapshot = "\n".join(line for line in snapshot_lines if str(line).strip()).strip()
        if not snapshot:
            return base
        if base:
            return f"{base}\n\n{marker}\n{snapshot}"
        return f"{marker}\n{snapshot}"

    def sync_live_agent_lead(
        self,
        lead_id: str,
        *,
        customer_name: str = "",
        raw_phone: str = "",
        traveler_id: str = "",
        preferred_trip_type: str = "",
        interested_trip_ids: str = "",
        suggested_trip_ids: str = "",
        group_size: int | str | None = None,
        room_type: str = "",
        room_group: str = "",
        flight_option: str = "",
        preferred_date: str = "",
        current_step: str = "",
        language: str = "",
        booking_id: str = "",
        passport_attachment_ref: str = "",
        passport_status: str = "",
        channel: str = "",
        lead_source: str = "",
        handoff_required: bool | None = None,
        handoff_reason: str = "",
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        lead_key = str(lead_id or "").strip()
        if not lead_key:
            raise ValueError("lead_id is required")

        normalized_trip_type = self.normalize_trip_type(preferred_trip_type)
        trip_type_value = normalized_trip_type or str(preferred_trip_type or "").strip()
        traveler_key = str(traveler_id or "").strip()
        interested_value = str(interested_trip_ids or "").strip()
        suggested_value = str(suggested_trip_ids or "").strip()
        room_type_value = str(room_type or "").strip()
        room_group_value = str(room_group or "").strip()
        flight_option_value = str(flight_option or "").strip()
        preferred_date_value = str(preferred_date or "").strip()
        attachment_ref = str(passport_attachment_ref or "").strip()
        passport_state = str(passport_status or "").strip()
        stage_value = str(current_step or "").strip()
        language_value = str(language or "").strip()
        booking_value = str(booking_id or "").strip()
        channel_value = str(channel or "").strip()
        source_value = str(lead_source or "").strip()
        handoff_reason_value = str(handoff_reason or "").strip()

        timestamp = _utc_now().replace(microsecond=0)
        phone = self.normalize_phone(raw_phone, "20") if raw_phone else {}
        group_size_value = self._as_int(group_size, default=1) if group_size not in (None, "") else None
        uploaded_event: dict[str, Any] | None = None

        with self.connect() as connection:
            row = connection.execute(
                """
                SELECT lead_id, customer_name, raw_phone, integrated_whatsapp, phone_lookup_key,
                       traveler_id, preferred_trip_type, group_size, interested_trip_ids,
                       suggested_trip_ids, notes, current_step, language, booking_id,
                       handoff_required, handoff_reason, channel, lead_source,
                       passport_attachment_ref, passport_status
                FROM leads
                WHERE lead_id = ?
                """,
                (lead_key,),
            ).fetchone()
            if not row:
                raise ValueError(f"Lead was not found: {lead_key}")

            existing_notes = str(row["notes"] or "")
            snapshot_lines = [
                f"Current step: {stage_value}" if stage_value else "",
                f"Trip type: {trip_type_value}" if trip_type_value else "",
                f"Interested trip: {interested_value}" if interested_value else "",
                f"Suggested trips: {suggested_value}" if suggested_value else "",
                f"Room choice: {room_type_value}{f' ({room_group_value})' if room_group_value else ''}" if room_type_value else "",
                f"Travelers: {group_size_value}" if group_size_value else "",
                f"Preferred date: {preferred_date_value}" if preferred_date_value else "",
                f"Flight option: {flight_option_value}" if flight_option_value else "",
                f"Passport status: {passport_state}" if passport_state else "",
                f"Passport file: {Path(attachment_ref).name}" if attachment_ref else "",
            ]
            merged_notes = self._merge_agent_snapshot_notes(existing_notes, snapshot_lines)

            resolved_customer_name = str(customer_name or row["customer_name"] or "").strip()
            resolved_raw_phone = str(raw_phone or row["raw_phone"] or "").strip()
            resolved_traveler_id = traveler_key or str(row["traveler_id"] or "").strip()
            resolved_trip_type = trip_type_value or str(row["preferred_trip_type"] or "").strip()
            resolved_group_size = group_size_value if group_size_value is not None else self._as_int(row["group_size"], default=1) or 1
            resolved_interested = interested_value or str(row["interested_trip_ids"] or "").strip()
            resolved_suggested = suggested_value or str(row["suggested_trip_ids"] or "").strip()
            resolved_current_step = stage_value or str(row["current_step"] or "").strip()
            resolved_language = language_value or str(row["language"] or "").strip()
            resolved_booking_id = booking_value or str(row["booking_id"] or "").strip()
            resolved_handoff_required = int(bool(handoff_required)) if handoff_required is not None else int(bool(row["handoff_required"]))
            resolved_handoff_reason = handoff_reason_value or str(row["handoff_reason"] or "").strip()
            resolved_channel = channel_value or str(row["channel"] or "").strip()
            resolved_source = source_value or str(row["lead_source"] or "").strip()
            resolved_attachment_ref = attachment_ref or str(row["passport_attachment_ref"] or "").strip()
            resolved_passport_status = passport_state or str(row["passport_status"] or "").strip()

            connection.execute(
                """
                UPDATE leads
                SET updated_at = ?, customer_name = ?, raw_phone = ?, integrated_whatsapp = ?,
                    phone_lookup_key = ?, traveler_id = ?, preferred_trip_type = ?, group_size = ?,
                    interested_trip_ids = ?, suggested_trip_ids = ?, notes = ?, current_step = ?,
                    language = ?, booking_id = ?, handoff_required = ?, handoff_reason = ?,
                    channel = ?, lead_source = ?, passport_attachment_ref = ?, passport_status = ?
                WHERE lead_id = ?
                """,
                (
                    timestamp.isoformat(timespec="seconds"),
                    resolved_customer_name or None,
                    resolved_raw_phone or None,
                    phone.get("normalized_whatsapp") or row["integrated_whatsapp"],
                    phone.get("lookup_key") or row["phone_lookup_key"],
                    resolved_traveler_id or None,
                    resolved_trip_type or None,
                    resolved_group_size,
                    resolved_interested or None,
                    resolved_suggested or None,
                    merged_notes or None,
                    resolved_current_step or None,
                    resolved_language or None,
                    resolved_booking_id or None,
                    resolved_handoff_required,
                    resolved_handoff_reason or None,
                    resolved_channel or None,
                    resolved_source or None,
                    resolved_attachment_ref or None,
                    resolved_passport_status or None,
                    lead_key,
                ),
            )
            connection.commit()

            previous_attachment_ref = str(row["passport_attachment_ref"] or "").strip()
            if resolved_attachment_ref and resolved_attachment_ref != previous_attachment_ref:
                uploaded_event = self.create_booking_event(
                    event_type="passport_attachment_received",
                    event_label="Passport attachment received",
                    traveler_id=resolved_traveler_id,
                    lead_id=lead_key,
                    booking_id=resolved_booking_id,
                    trip_id=resolved_interested.split(",", 1)[0].strip(),
                    channel=resolved_channel,
                    actor="ai-agent",
                    notes=f"Passport file attached: {Path(resolved_attachment_ref).name}",
                    metadata={
                        "passport_attachment_ref": resolved_attachment_ref,
                        "current_step": resolved_current_step,
                        "passport_status": resolved_passport_status,
                    },
                    occurred_at=timestamp,
                )

        try:
            self.sync_agent_write_to_sheet(
                traveler_id=resolved_traveler_id,
                lead_id=lead_key,
                booking_id=resolved_booking_id,
                event_ids=[uploaded_event["event_id"]] if uploaded_event else None,
            )
        except Exception:
            pass

        return {
            "lead_id": lead_key,
            "traveler_id": resolved_traveler_id,
            "preferred_trip_type": resolved_trip_type,
            "interested_trip_ids": resolved_interested,
            "suggested_trip_ids": resolved_suggested,
            "group_size": resolved_group_size,
            "current_step": resolved_current_step,
            "passport_attachment_ref": resolved_attachment_ref,
            "passport_status": resolved_passport_status,
            "booking_id": resolved_booking_id,
            "event_id": uploaded_event["event_id"] if uploaded_event else "",
        }

    def create_booking_event(
        self,
        *,
        event_type: str,
        event_label: str,
        traveler_id: str = "",
        lead_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        interaction_id: str = "",
        channel: str = "",
        actor: str = "agent",
        notes: str = "",
        metadata: dict[str, Any] | None = None,
        occurred_at: datetime | None = None,
    ) -> dict[str, Any]:
        occurred_at = occurred_at or _utc_now().replace(microsecond=0)
        metadata_json = json.dumps(metadata or {}, ensure_ascii=True, sort_keys=True)
        with self.connect() as connection:
            self._ensure_booking_event_trail_table(connection)
            event_id = self._next_daily_id(connection, "booking_event_trail", "event_id", "EVT", occurred_at)
            connection.execute(
                """
                INSERT INTO booking_event_trail (
                    event_id, occurred_at, event_type, event_label, traveler_id, lead_id,
                    booking_id, trip_id, interaction_id, channel, actor, notes, metadata_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    event_id,
                    occurred_at.isoformat(timespec="seconds"),
                    event_type,
                    event_label,
                    traveler_id or None,
                    lead_id or None,
                    booking_id or None,
                    trip_id or None,
                    interaction_id or None,
                    channel or None,
                    actor or None,
                    notes or None,
                    metadata_json,
                ),
            )
            connection.commit()
        return {
            "event_id": event_id,
            "occurred_at": occurred_at.isoformat(timespec="seconds"),
            "event_type": event_type,
            "event_label": event_label,
            "traveler_id": traveler_id,
            "lead_id": lead_id,
            "booking_id": booking_id,
            "trip_id": trip_id,
            "interaction_id": interaction_id,
        }

    def sync_agent_write_to_sheet(
        self,
        *,
        traveler_id: str = "",
        lead_id: str = "",
        interaction_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        event_ids: list[str] | None = None,
    ) -> dict[str, Any]:
        if not self.settings.sheet_export_enabled:
            return {
                "status": "disabled",
                "backend": self.settings.sheet_backend,
                "reason": "CRM is the operational authority; automatic sheet mirroring is disabled.",
            }
        if self.settings.sheet_backend == "excel":
            try:
                synced = self._run_sync_agent_write_to_sheet_worker(
                    traveler_id=traveler_id,
                    lead_id=lead_id,
                    interaction_id=interaction_id,
                    booking_id=booking_id,
                    trip_id=trip_id,
                    event_ids=event_ids,
                )
                return {"status": "ok", "backend": "excel", "synced": synced}
            except Exception as exc:
                return {"status": "failed", "backend": "excel", "reason": str(exc)}

        # Spin up a daemon thread to process sheet writing asynchronously
        import threading

        def run_async():
            try:
                self._run_sync_agent_write_to_sheet_worker(
                    traveler_id=traveler_id,
                    lead_id=lead_id,
                    interaction_id=interaction_id,
                    booking_id=booking_id,
                    trip_id=trip_id,
                    event_ids=event_ids,
                )
            except Exception:
                # Daemon threads should suppress unexpected runtime worker crashes
                pass

        thread = threading.Thread(target=run_async)
        thread.daemon = True
        thread.start()
        return {"status": "queued", "message": "Background sync thread initiated."}

    def _run_sync_agent_write_to_sheet_worker(
        self,
        *,
        traveler_id: str = "",
        lead_id: str = "",
        interaction_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        event_ids: list[str] | None = None,
    ) -> dict[str, Any]:
        if self.settings.sheet_backend == "excel" and type(self).sync_record_to_sheet is UnifiedCRMService.sync_record_to_sheet:
            return self._sync_agent_write_to_excel(
                traveler_id=traveler_id,
                lead_id=lead_id,
                interaction_id=interaction_id,
                booking_id=booking_id,
                trip_id=trip_id,
                event_ids=event_ids,
            )

        synced: dict[str, Any] = {}
        if traveler_id:
            synced["traveler"] = self.sync_record_to_sheet("Travelers", traveler_id)
        if lead_id:
            synced["lead"] = self.sync_record_to_sheet("Leads", lead_id)
        if interaction_id:
            synced["interaction"] = self.sync_record_to_sheet("Interactions", interaction_id)
        if booking_id:
            synced["booking"] = self.sync_record_to_sheet("Trip Bookings", booking_id)
        for event_id in event_ids or []:
            synced.setdefault("events", []).append(self.sync_record_to_sheet("Booking Event Trail", event_id))
        return synced

    def _sync_agent_write_to_excel(
        self,
        *,
        traveler_id: str = "",
        lead_id: str = "",
        interaction_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        event_ids: list[str] | None = None,
    ) -> dict[str, Any]:
        record_jobs: list[tuple[str, dict[str, Any], str, dict[str, Any]]] = []
        if traveler_id:
            record = self._fetch_table_record("travelers", "traveler_id", traveler_id)
            if record:
                record_jobs.append(("Travelers", SHEET_TABLE_MAPPINGS["Travelers"], "traveler_id", record))
        if lead_id:
            record = self._fetch_table_record("leads", "lead_id", lead_id)
            if record:
                record_jobs.append(("Leads", SHEET_TABLE_MAPPINGS["Leads"], "lead_id", record))
        if interaction_id:
            record = self._fetch_table_record("interactions", "interaction_id", interaction_id)
            if record:
                record_jobs.append(("Interactions", SHEET_TABLE_MAPPINGS["Interactions"], "interaction_id", record))
        if booking_id:
            record = self._fetch_table_record("trip_bookings", "booking_id", booking_id)
            if record:
                record_jobs.append(("Trip Bookings", SHEET_TABLE_MAPPINGS["Trip Bookings"], "booking_id", record))
        for event_id in event_ids or []:
            record = self._fetch_table_record("booking_event_trail", "event_id", event_id)
            if record:
                record_jobs.append(("Booking Event Trail", SHEET_TABLE_MAPPINGS["Booking Event Trail"], "event_id", record))

        synced: dict[str, Any] = {}
        if not record_jobs:
            return synced

        backend = self.settings.sheet_backend
        try:
            updated = 0
            for workbook_path in self._excel_outcome_workbooks():
                self._upsert_records_in_excel(workbook_path, record_jobs)
                updated += 1
            for mapping_name, _mapping, key_column, record in record_jobs:
                record_id = str(record.get(key_column) or "").strip()
                if record_id:
                    self._record_sync_success(mapping_name, record_id)
                    if mapping_name == "Booking Event Trail":
                        synced.setdefault("events", []).append({"status": "ok", "backend": backend, "updated_workbooks": updated})
                    elif mapping_name == "Trip Bookings":
                        synced["booking"] = {"status": "ok", "backend": backend, "updated_workbooks": updated}
                    elif mapping_name == "Interactions":
                        synced["interaction"] = {"status": "ok", "backend": backend, "updated_workbooks": updated}
                    elif mapping_name == "Leads":
                        synced["lead"] = {"status": "ok", "backend": backend, "updated_workbooks": updated}
                    elif mapping_name == "Travelers":
                        synced["traveler"] = {"status": "ok", "backend": backend, "updated_workbooks": updated}
            return synced
        except Exception as exc:
            for mapping_name, _mapping, key_column, record in record_jobs:
                record_id = str(record.get(key_column) or "").strip()
                if record_id:
                    self._record_sync_failure(mapping_name, record_id, str(exc))
            return {"status": "failed", "backend": backend, "reason": str(exc)}

    def sync_record_to_sheet(self, mapping_name: str, record_id: str) -> dict[str, Any]:
        if not self.settings.sheet_export_enabled:
            return {
                "status": "disabled",
                "backend": self.settings.sheet_backend,
                "reason": "CRM is the operational authority; export must be explicitly enabled.",
            }
        if not record_id:
            return {"status": "skipped", "reason": "empty_record_id"}
        mapping = SHEET_TABLE_MAPPINGS[mapping_name]
        table_name = self._table_for_mapping(mapping_name)
        key_column = self._key_column_for_mapping(mapping_name)
        record = self._fetch_table_record(table_name, key_column, record_id)
        if not record:
            return {"status": "skipped", "reason": "record_not_found", "record_id": record_id}

        backend = self.settings.sheet_backend
        try:
            if backend == "excel":
                updated = 0
                for workbook_path in self._excel_outcome_workbooks():
                    self._upsert_record_in_excel(workbook_path, mapping, key_column, record)
                    updated += 1
                self._record_sync_success(mapping_name, record_id)
                return {"status": "ok", "backend": backend, "updated_workbooks": updated}
            if backend == "google_sheets":
                self._upsert_record_in_google_sheet(mapping, key_column, record)
                self._record_sync_success(mapping_name, record_id)
                return {"status": "ok", "backend": backend, "updated_workbooks": 1}
        except Exception as e:
            self._record_sync_failure(mapping_name, record_id, str(e))
            return {"status": "failed", "backend": backend, "reason": str(e)}
        return {"status": "skipped", "backend": backend, "reason": "Sheet sync backend not supported."}

    def remove_record_from_sheet(self, mapping_name: str, record_id: str) -> dict[str, Any]:
        if not self.settings.sheet_export_enabled:
            return {
                "status": "disabled",
                "backend": self.settings.sheet_backend,
                "reason": "CRM is the operational authority; export must be explicitly enabled.",
            }
        if not record_id:
            return {"status": "skipped", "reason": "empty_record_id"}

        mapping = SHEET_TABLE_MAPPINGS[mapping_name]
        key_column = self._key_column_for_mapping(mapping_name)
        backend = self.settings.sheet_backend
        try:
            if backend == "excel":
                updated = 0
                for workbook_path in self._excel_outcome_workbooks():
                    if self._delete_record_in_excel(workbook_path, mapping, key_column, record_id):
                        updated += 1
                self._record_sync_success(mapping_name, record_id)
                return {"status": "ok", "backend": backend, "updated_workbooks": updated}
            if backend == "google_sheets":
                deleted = self._delete_record_in_google_sheet(mapping, key_column, record_id)
                self._record_sync_success(mapping_name, record_id)
                return {"status": "ok", "backend": backend, "updated_workbooks": 1 if deleted else 0}
        except Exception as e:
            self._record_sync_failure(mapping_name, record_id, str(e))
            return {"status": "failed", "backend": backend, "reason": str(e)}
        return {"status": "skipped", "backend": backend, "reason": "Sheet sync backend not supported."}

    def _record_sync_success(self, mapping_name: str, record_id: str) -> None:
        with self.connect() as connection:
            self._ensure_sync_queue_table(connection)
            connection.execute(
                "DELETE FROM sync_queue WHERE mapping_name = ? AND record_id = ?",
                (mapping_name, record_id),
            )
            connection.commit()

    def _record_sync_failure(self, mapping_name: str, record_id: str, error_message: str) -> None:
        timestamp = _utc_now().replace(microsecond=0).isoformat(timespec="seconds")
        with self.connect() as connection:
            self._ensure_sync_queue_table(connection)
            connection.execute(
                """
                INSERT INTO sync_queue (
                    mapping_name, record_id, status, error_message, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(mapping_name, record_id) DO UPDATE SET
                    status = excluded.status,
                    error_message = excluded.error_message,
                    updated_at = excluded.updated_at
                """,
                (mapping_name, record_id, "failed", error_message, timestamp, timestamp),
            )
            connection.commit()

    def get_sync_queue(self) -> list[dict[str, Any]]:
        with self.connect() as connection:
            connection.row_factory = sqlite3.Row
            rows = connection.execute("SELECT * FROM sync_queue ORDER BY updated_at DESC").fetchall()
            return [dict(r) for r in rows]

    def ensure_operational_schema(self) -> None:
        db_key = str(Path(self.settings.db_path).resolve())
        if db_key in self._schema_ready_paths:
            return

        with _SCHEMA_READY_LOCK:
            if db_key in self._schema_ready_paths:
                return
            with self.connect() as connection:
                self._ensure_handoff_queue_table(connection)
                self._ensure_booking_event_trail_table(connection)
                self._ensure_booking_status_history_table(connection)
                self._ensure_traveler_documents_table(connection)
                self._ensure_sync_queue_table(connection)
                self._migrate_travelers_passport_columns(connection)
                self._migrate_trips_room_columns(connection)
                self._migrate_trip_booking_passport_columns(connection)
                self._migrate_lead_columns(connection)
                self._migrate_idempotency_columns(connection)
                self._migrate_interaction_message_key_unique_index(connection)
                connection.commit()
            self._schema_ready_paths.add(db_key)

    @staticmethod
    def _ensure_handoff_queue_table(connection: sqlite3.Connection) -> None:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS handoff_queue (
                handoff_id TEXT PRIMARY KEY,
                created_at TEXT,
                lead_id TEXT,
                traveler_id TEXT,
                trip_id TEXT,
                flow_key TEXT,
                reason TEXT,
                priority TEXT,
                channel TEXT,
                status TEXT DEFAULT 'Pending',
                owner TEXT,
                assigned_to TEXT,
                notes TEXT
            )
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_handoff_queue_status_created
            ON handoff_queue (status, created_at)
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_handoff_queue_traveler_status
            ON handoff_queue (traveler_id, status)
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_handoff_queue_lead_status
            ON handoff_queue (lead_id, status)
            """
        )

    @staticmethod
    def _migrate_travelers_passport_columns(connection: sqlite3.Connection) -> None:
        """Add passport and attachment metadata columns to travelers if they don't exist.

        Migration notes (Demo Phase 1):
        - passport_name TEXT: Full name as it appears on passport
        - passport_number TEXT: Passport document number
        - passport_expiry TEXT: Expiry date in ISO format (YYYY-MM-DD)
        - passport_nationality TEXT: Nationality as on passport
        - passport_attachment_ref TEXT: Local filesystem reference path (metadata only)
        - preferred_currency TEXT: Traveler payment preference such as EGP or USD
        These columns are nullable and backward-compatible with existing traveler rows.
        No data is lost on existing records. Added columns will be NULL by default.
        """
        import logging
        _log = logging.getLogger(__name__)
        existing_cols: set[str] = set()
        try:
            rows = connection.execute("PRAGMA table_info(travelers)").fetchall()
            existing_cols = {row[1] for row in rows}
        except Exception:
            return  # travelers table does not exist yet; schema will be created elsewhere

        passport_columns = [
            ("passport_name", "TEXT"),
            ("passport_number", "TEXT"),
            ("passport_expiry", "TEXT"),
            ("passport_nationality", "TEXT"),
            ("passport_attachment_ref", "TEXT"),
            ("preferred_currency", "TEXT"),
            ("is_minor", "INTEGER"),
            ("guardian_name", "TEXT"),
            ("guardian_phone", "TEXT"),
        ]
        for col_name, col_type in passport_columns:
            if col_name not in existing_cols:
                connection.execute(f"ALTER TABLE travelers ADD COLUMN {col_name} {col_type}")
                _log.info(f"[MIGRATION] Added column travelers.{col_name} ({col_type}) -- Demo Phase 1 passport support")

    @staticmethod
    def _migrate_trips_room_columns(connection: sqlite3.Connection) -> None:
        try:
            rows = connection.execute("PRAGMA table_info(trips)").fetchall()
        except Exception:
            return
        existing_cols = {row[1] for row in rows}
        if not existing_cols:
            return
        room_columns = [
            ("boys_double", "INTEGER"),
            ("girls_double", "INTEGER"),
            ("boys_triple", "INTEGER"),
            ("girls_triple", "INTEGER"),
            ("draft_holds_boys_double", "INTEGER DEFAULT 0"),
            ("draft_holds_girls_double", "INTEGER DEFAULT 0"),
            ("draft_holds_boys_triple", "INTEGER DEFAULT 0"),
            ("draft_holds_girls_triple", "INTEGER DEFAULT 0"),
            ("itinerary", "TEXT"),
            ("inclusions", "TEXT"),
            ("exclusions", "TEXT"),
            ("room_prices_json", "TEXT"),
        ]
        for col_name, col_type in room_columns:
            if col_name not in existing_cols:
                connection.execute(f"ALTER TABLE trips ADD COLUMN {col_name} {col_type}")

    @staticmethod
    def _table_columns(connection: sqlite3.Connection, table_name: str) -> set[str]:
        try:
            rows = connection.execute(f"PRAGMA table_info({table_name})").fetchall()
        except Exception:
            return set()
        return {row[1] for row in rows}

    def _trip_select_columns(self, connection: sqlite3.Connection) -> str:
        base_columns = [
            "trip_id",
            "trip_name",
            "type",
            "year",
            "trip_leader",
            "start_date",
            "end_date",
            "sales_status",
            "data_audit",
            "trip_window_status",
            "trip_availability_note",
            "next_reengage_date",
            "single_total",
            "double_total",
            "triple_total",
            "single_remaining",
            "double_remaining",
            "triple_remaining",
            "draft_holds_single",
            "draft_holds_double",
            "draft_holds_triple",
            "public_price",
            "public_description",
            "sales_notes",
        ]
        optional_columns = [
            "trip_name_ar",
            "itinerary",
            "inclusions",
            "exclusions",
            "room_prices_json",
            "boys_double",
            "girls_double",
            "boys_triple",
            "girls_triple",
            "draft_holds_boys_double",
            "draft_holds_girls_double",
            "draft_holds_boys_triple",
            "draft_holds_girls_triple",
        ]
        existing = self._table_columns(connection, "trips")
        return ", ".join([*base_columns, *[column for column in optional_columns if column in existing]])

    @staticmethod
    def _row_value(row: Any, name: str, default: Any = None) -> Any:
        try:
            return row[name]
        except (KeyError, IndexError, TypeError):
            return default

    @classmethod
    def _booking_row_room_requirements(cls, row: Any) -> list[dict[str, Any]]:
        raw_json = str(cls._row_value(row, "room_requirements_json") or "").strip()
        if raw_json:
            try:
                parsed = json.loads(raw_json)
            except (TypeError, ValueError):
                parsed = []
            if isinstance(parsed, list):
                requirements = []
                for item in parsed:
                    if not isinstance(item, dict):
                        continue
                    room_type = str(item.get("room_type") or "").strip().title()
                    group = str(item.get("room_group") or "").strip().lower()
                    rooms = cls._as_int(item.get("rooms"), default=0) or 0
                    if room_type in ROOM_HOLD_COLUMNS and rooms > 0:
                        requirements.append({"room_type": room_type, "room_group": group, "rooms": rooms})
                if requirements:
                    return requirements

        room_type = str(cls._row_value(row, "room_type") or "").strip().title()
        if room_type not in ROOM_HOLD_COLUMNS:
            return []
        boys_rooms = cls._as_int(cls._row_value(row, "boys_rooms_requested"), default=0) or 0
        girls_rooms = cls._as_int(cls._row_value(row, "girls_rooms_requested"), default=0) or 0
        requirements = []
        if boys_rooms:
            requirements.append({"room_type": room_type, "room_group": "boys", "rooms": boys_rooms})
        if girls_rooms:
            requirements.append({"room_type": room_type, "room_group": "girls", "rooms": girls_rooms})
        if requirements:
            return requirements

        group = str(cls._row_value(row, "room_group") or "").strip().lower()
        return [{"room_type": room_type, "room_group": group if group in {"boys", "girls"} else "", "rooms": 1}]

    @classmethod
    def _normalize_room_requirements(
        cls,
        *,
        room_type: str,
        room_group: str = "",
        boys_rooms_requested: int | str = 0,
        girls_rooms_requested: int | str = 0,
        room_requirements: list[dict[str, Any]] | dict[str, Any] | str | None = None,
    ) -> list[dict[str, Any]]:
        parsed: Any = room_requirements
        if isinstance(parsed, str) and parsed.strip():
            try:
                parsed = json.loads(parsed)
            except (TypeError, ValueError):
                parsed = None
        if isinstance(parsed, dict):
            parsed = parsed.get("requirements") or parsed.get("rooms") or []
        requirements: list[dict[str, Any]] = []
        if isinstance(parsed, list):
            for item in parsed:
                if not isinstance(item, dict):
                    continue
                item_room_type = str(item.get("room_type") or room_type or "").strip().title()
                group = str(item.get("room_group") or item.get("group") or "").strip().lower()
                rooms = cls._as_int(item.get("rooms") or item.get("count"), default=0) or 0
                if item_room_type in ROOM_HOLD_COLUMNS and group in {"", "boys", "girls", "family"} and rooms > 0:
                    requirements.append({"room_type": item_room_type, "room_group": group, "rooms": rooms})
        if requirements:
            return requirements

        canonical_room_type = str(room_type or "").strip().title()
        if canonical_room_type not in ROOM_HOLD_COLUMNS:
            return []
        boys_rooms = cls._as_int(boys_rooms_requested, default=0) or 0
        girls_rooms = cls._as_int(girls_rooms_requested, default=0) or 0
        if boys_rooms or girls_rooms:
            if boys_rooms:
                requirements.append({"room_type": canonical_room_type, "room_group": "boys", "rooms": boys_rooms})
            if girls_rooms:
                requirements.append({"room_type": canonical_room_type, "room_group": "girls", "rooms": girls_rooms})
            return requirements
        group = str(room_group or "").strip().lower()
        return [{"room_type": canonical_room_type, "room_group": group if group in {"boys", "girls", "family"} else "", "rooms": 1}]

    @staticmethod
    def _migrate_trip_booking_passport_columns(connection: sqlite3.Connection) -> None:
        try:
            rows = connection.execute("PRAGMA table_info(trip_bookings)").fetchall()
        except Exception:
            return
        if not rows:
            return
        existing_cols = {row[1] for row in rows}
        if 'passport_required' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN passport_required INTEGER DEFAULT 0")
        if 'passport_status' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN passport_status TEXT")
        if 'group_size' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN group_size INTEGER DEFAULT 1")
        if 'boys_count' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN boys_count INTEGER DEFAULT 0")
        if 'girls_count' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN girls_count INTEGER DEFAULT 0")
        if 'family_units' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN family_units INTEGER DEFAULT 0")
        if 'room_group' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN room_group TEXT")
        if 'boys_rooms_requested' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN boys_rooms_requested INTEGER DEFAULT 0")
        if 'girls_rooms_requested' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN girls_rooms_requested INTEGER DEFAULT 0")
        if 'room_requirements_json' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN room_requirements_json TEXT")
        if 'refund_amount' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN refund_amount REAL")
        if 'missing_info' not in existing_cols:
            connection.execute("ALTER TABLE trip_bookings ADD COLUMN missing_info INTEGER DEFAULT 0")

    @staticmethod
    def _migrate_lead_columns(connection: sqlite3.Connection) -> None:
        try:
            rows = connection.execute("PRAGMA table_info(leads)").fetchall()
        except Exception:
            return
        if not rows:
            return
        existing_cols = {row[1] for row in rows}
        if 'group_size' not in existing_cols:
            connection.execute("ALTER TABLE leads ADD COLUMN group_size INTEGER DEFAULT 1")
        if 'passport_attachment_ref' not in existing_cols:
            connection.execute("ALTER TABLE leads ADD COLUMN passport_attachment_ref TEXT")
        if 'passport_status' not in existing_cols:
            connection.execute("ALTER TABLE leads ADD COLUMN passport_status TEXT")
        if 'requires_guardian_approval' not in existing_cols:
            connection.execute("ALTER TABLE leads ADD COLUMN requires_guardian_approval INTEGER")

    @staticmethod
    def _migrate_idempotency_columns(connection: sqlite3.Connection) -> None:
        table_columns = {
            "trip_bookings": UnifiedCRMService._table_columns(connection, "trip_bookings"),
            "leads": UnifiedCRMService._table_columns(connection, "leads"),
            "handoff_queue": UnifiedCRMService._table_columns(connection, "handoff_queue"),
            "traveler_documents": UnifiedCRMService._table_columns(connection, "traveler_documents"),
        }
        for table_name, columns in table_columns.items():
            if columns and "idempotency_key" not in columns:
                connection.execute(f"ALTER TABLE {table_name} ADD COLUMN idempotency_key TEXT")
        if table_columns["trip_bookings"]:
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_trip_bookings_idempotency_key ON trip_bookings (idempotency_key)"
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_trip_bookings_active_traveler_trip ON trip_bookings (traveler_id, trip_id, booking_status)"
            )
        if table_columns["leads"]:
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_idempotency_key ON leads (idempotency_key)"
            )
        if table_columns["handoff_queue"]:
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_handoff_queue_idempotency_key ON handoff_queue (idempotency_key)"
            )
        if table_columns["traveler_documents"]:
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_traveler_documents_idempotency_key ON traveler_documents (idempotency_key)"
            )

    @staticmethod
    def _migrate_interaction_message_key_unique_index(connection: sqlite3.Connection) -> None:
        columns = UnifiedCRMService._table_columns(connection, "interactions")
        if not {"channel", "message_key"}.issubset(columns):
            return
        duplicate = connection.execute(
            """
            SELECT channel, message_key, COUNT(*) AS duplicate_count
            FROM interactions
            WHERE message_key IS NOT NULL
              AND message_key <> ''
            GROUP BY channel, message_key
            HAVING COUNT(*) > 1
            LIMIT 1
            """
        ).fetchone()
        if duplicate:
            raise RuntimeError(
                "Cannot add unique webhook message index: duplicate interaction exists "
                f"for channel={duplicate['channel']!r}, message_key={duplicate['message_key']!r}."
            )
        connection.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS ux_interactions_channel_message_key
            ON interactions (channel, message_key)
            WHERE message_key IS NOT NULL AND message_key <> ''
            """
        )

    @staticmethod
    def _ensure_traveler_documents_table(connection: sqlite3.Connection) -> None:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS traveler_documents (
                document_id INTEGER PRIMARY KEY AUTOINCREMENT,
                traveler_id TEXT NOT NULL,
                category TEXT NOT NULL DEFAULT 'document',
                file_name TEXT NOT NULL,
                original_file_name TEXT,
                mime_type TEXT,
                file_extension TEXT,
                file_size INTEGER,
                storage_path TEXT NOT NULL,
                uploaded_at TEXT NOT NULL,
                uploaded_by TEXT,
                passport_full_name TEXT,
                passport_number TEXT,
                passport_nationality TEXT,
                passport_expiry TEXT,
                verification_status TEXT DEFAULT 'pending',
                notes TEXT,
                FOREIGN KEY(traveler_id) REFERENCES travelers (traveler_id)
            )
            """
        )

    @staticmethod
    def _ensure_sync_queue_table(connection: sqlite3.Connection) -> None:
        # TODO(production): replace this lightweight sync queue with a durable retry worker
        # and dead-letter handling for sheet/API failures.
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS sync_queue (
                mapping_name TEXT NOT NULL,
                record_id TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending',
                error_message TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (mapping_name, record_id)
            )
            """
        )

    @staticmethod
    def _ensure_booking_status_history_table(connection: sqlite3.Connection) -> None:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS booking_status_history (
                history_id INTEGER PRIMARY KEY AUTOINCREMENT,
                booking_id TEXT NOT NULL,
                old_status TEXT,
                new_status TEXT,
                changed_at TEXT NOT NULL,
                changed_by TEXT,
                change_source TEXT,
                notes TEXT
            )
            """
        )

    def traveler_has_completed_trips(self, traveler_id: str) -> bool:
        """Return True if the traveler has at least one booking whose trip end_date is in the past."""
        today = date.today().isoformat()
        try:
            with self.connect() as connection:
                row = connection.execute(
                    """
                    SELECT COUNT(*) as cnt
                    FROM trip_bookings tb
                    JOIN trips t ON tb.trip_id = t.trip_id
                    WHERE tb.traveler_id = ?
                      AND t.end_date IS NOT NULL
                      AND t.end_date < ?
                    """,
                    (traveler_id, today),
                ).fetchone()
                return bool(row and row["cnt"] > 0)
        except Exception:
            return False

    def save_traveler_passport(
        self,
        traveler_id: str,
        *,
        passport_name: str = "",
        passport_number: str = "",
        passport_expiry: str = "",
        passport_nationality: str = "",
        passport_attachment_ref: str = "",
        uploaded_by: str = "ai-agent",
        attachment_file_name: str = "",
        attachment_original_name: str = "",
        attachment_mime_type: str = "",
        attachment_size: int | None = None,
        notes: str = "",
    ) -> dict[str, Any]:
        traveler_key = str(traveler_id or "").strip()
        if not traveler_key:
            raise ValueError("traveler_id is required")

        passport_name = str(passport_name or "").strip()
        passport_number = str(passport_number or "").strip()
        passport_nationality = str(passport_nationality or "").strip()
        passport_attachment_ref = str(passport_attachment_ref or "").strip()
        attachment_file_name = str(attachment_file_name or "").strip()
        attachment_original_name = str(attachment_original_name or "").strip()
        attachment_mime_type = str(attachment_mime_type or "").strip()
        notes = str(notes or "").strip()
        normalized_expiry = self._normalize_date_string(passport_expiry)
        timestamp = _utc_now().replace(microsecond=0).isoformat(timespec="seconds")

        with self.connect() as connection:
            self._ensure_traveler_documents_table(connection)
            self._migrate_travelers_passport_columns(connection)
            traveler_exists = connection.execute(
                "SELECT traveler_id FROM travelers WHERE traveler_id = ?",
                (traveler_key,),
            ).fetchone()
            if not traveler_exists:
                raise ValueError(f"traveler {traveler_key} was not found")

            connection.execute(
                """
                UPDATE travelers
                SET passport_name = ?,
                    passport_number = ?,
                    passport_expiry = ?,
                    passport_nationality = ?,
                    passport_attachment_ref = ?,
                    last_contacted_at = ?
                WHERE traveler_id = ?
                """,
                (
                    passport_name or None,
                    passport_number or None,
                    normalized_expiry,
                    passport_nationality or None,
                    passport_attachment_ref or None,
                    timestamp,
                    traveler_key,
                ),
            )

            document_id = None
            if passport_attachment_ref:
                file_name = attachment_file_name or Path(passport_attachment_ref).name or "passport"
                original_name = attachment_original_name or file_name
                file_extension = Path(file_name).suffix.lower().lstrip(".")
                existing_document = connection.execute(
                    """
                    SELECT document_id
                    FROM traveler_documents
                    WHERE traveler_id = ? AND storage_path = ?
                    ORDER BY document_id DESC
                    LIMIT 1
                    """,
                    (traveler_key, passport_attachment_ref),
                ).fetchone()
                if existing_document:
                    document_id = int(existing_document["document_id"])
                    connection.execute(
                        """
                        UPDATE traveler_documents
                        SET category = ?,
                            file_name = ?,
                            original_file_name = ?,
                            mime_type = ?,
                            file_extension = ?,
                            file_size = ?,
                            uploaded_at = ?,
                            uploaded_by = ?,
                            passport_full_name = ?,
                            passport_number = ?,
                            passport_nationality = ?,
                            passport_expiry = ?,
                            verification_status = ?,
                            notes = ?
                        WHERE document_id = ?
                        """,
                        (
                            "passport",
                            file_name,
                            original_name or None,
                            attachment_mime_type or None,
                            file_extension or None,
                            attachment_size,
                            timestamp,
                            uploaded_by or None,
                            passport_name or None,
                            passport_number or None,
                            passport_nationality or None,
                            normalized_expiry,
                            "pending",
                            notes or None,
                            document_id,
                        ),
                    )
                else:
                    cursor = connection.execute(
                        """
                        INSERT INTO traveler_documents (
                            traveler_id, category, file_name, original_file_name, mime_type,
                            file_extension, file_size, storage_path, uploaded_at, uploaded_by,
                            passport_full_name, passport_number, passport_nationality,
                            passport_expiry, verification_status, notes
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            traveler_key,
                            "passport",
                            file_name,
                            original_name or None,
                            attachment_mime_type or None,
                            file_extension or None,
                            attachment_size,
                            passport_attachment_ref,
                            timestamp,
                            uploaded_by or None,
                            passport_name or None,
                            passport_number or None,
                            passport_nationality or None,
                            normalized_expiry,
                            "pending",
                            notes or None,
                        ),
                    )
                    document_id = int(cursor.lastrowid or 0)

            connection.commit()

        return {
            "traveler_id": traveler_key,
            "passport_name": passport_name,
            "passport_number": passport_number,
            "passport_expiry": normalized_expiry,
            "passport_nationality": passport_nationality,
            "passport_attachment_ref": passport_attachment_ref,
            "document_id": document_id,
            "uploaded_at": timestamp,
        }

    @staticmethod
    def infer_customer_tier(result: dict[str, Any]) -> str:
        traveler = result.get("traveler")
        if isinstance(traveler, dict):
            status = str(traveler.get("status") or "").strip()
            if status in {"VIP", "Repeat"}:
                return status
            if status:
                return status
        return "Standard"

    @classmethod
    def resolve_commercial_context(
        cls,
        *,
        traveler: dict[str, Any] | None,
        trip_type: str | None = None,
        requested_group_size: int = 1,
        trip_id: str = "",
    ) -> dict[str, Any]:
        traveler_status = str((traveler or {}).get("status") or "").strip()
        offer_config = cls.load_trip_commercial_config(trip_id)
        normalized_group_size = max(int(requested_group_size or 1), 1)
        vip_recognized = traveler_status == "VIP"
        group_threshold = max(int(offer_config.get("group_threshold") or GROUP_DISCOUNT_THRESHOLD), 2)
        group_eligible = normalized_group_size >= group_threshold
        return {
            "flight_policy": {
                "mode": "without_flights_default",
                "summary": "Most Rahma trips are offered without flights by default.",
                "customer_can_request_flights": True,
            },
            "vip": {
                "recognized": vip_recognized,
                "discount_configured": bool(str(offer_config.get("vip_offer") or "").strip()),
                "offer_text": str(offer_config.get("vip_offer") or "").strip(),
                "manual_review_recommended": vip_recognized,
            },
            "group": {
                "size": normalized_group_size,
                "threshold": group_threshold,
                "eligible": group_eligible,
                "discount_configured": bool(str(offer_config.get("group_offer") or "").strip()),
                "offer_text": str(offer_config.get("group_offer") or "").strip(),
                "manual_quote_required": group_eligible,
            },
            "manual_handoff_recommended": group_eligible,
            "handoff_reason": "group_booking_quote" if group_eligible else "",
            "trip_type": cls.normalize_trip_type(trip_type) or (trip_type or ""),
            "trip_id": trip_id,
        }

    @staticmethod
    def _default_handoff_priority(reason_code: str) -> str:
        reason = str(reason_code or "").strip().lower()
        if reason == "blacklisted_customer":
            return "Critical"
        if reason in {"duplicate_phone_match", "phone_name_conflict", "group_booking_quote", "customer_requested_human_agent"}:
            return "High"
        return "Medium"

    @staticmethod
    def _handoff_display_summary(reason_code: str, reason_text: str, package: dict[str, Any]) -> str:
        reason = str(reason_code or "").strip().lower()
        metadata = package.get("metadata") if isinstance(package.get("metadata"), dict) else {}
        customer_name = str(package.get("customer_name") or "").strip()
        trip_id = str(metadata.get("trip_id") or "").strip()
        trip_type = str(metadata.get("trip_type") or "").strip()
        group_size = str(metadata.get("group_size") or "").strip()

        if reason == "blacklisted_customer":
            return (
                f"Blocked traveler match for {customer_name or 'this customer'}. "
                "The WhatsApp number or traveler profile is marked as blacklisted, so sales follow-up must stop until an authorized employee reviews the case."
            )
        if reason == "phone_name_conflict":
            return (
                f"Name and phone conflict for {customer_name or 'this customer'}. "
                "The WhatsApp number matches an existing traveler, but the submitted name does not match the traveler profile."
            )
        if reason == "duplicate_phone_match":
            return (
                "Duplicate traveler match detected. "
                "This WhatsApp number is linked to more than one traveler profile, so the employee must resolve the identity before continuing."
            )
        if reason == "archived_traveler":
            return (
                f"Archived or inactive traveler match for {customer_name or 'this customer'}. "
                "The employee should confirm whether the traveler should be reactivated before continuing sales work."
            )
        if reason == "group_booking_quote":
            scope = []
            if group_size:
                scope.append(f"group size {group_size}")
            if trip_id:
                scope.append(f"trip {trip_id}")
            elif trip_type:
                scope.append(f"{trip_type.lower()} trip request")
            detail = f" for {', '.join(scope)}" if scope else ""
            return (
                f"Manual commercial review needed{detail}. "
                "The request qualifies for a group booking quote and should be reviewed by an employee before confirmation."
            )
        if reason == "customer_requested_human_agent":
            return (
                f"Customer requested a human agent for {customer_name or 'this conversation'}. "
                "The employee should continue the conversation directly and confirm the next step."
            )
        if reason == "manual_admin_handoff":
            return str(reason_text or "Manual handoff requested by an employee.").strip()
        return str(reason_text or reason_code or "Manual handoff").strip()

    @classmethod
    def _handoff_action_steps(cls, reason_code: str, package: dict[str, Any]) -> list[str]:
        reason = str(reason_code or "").strip().lower()
        metadata = package.get("metadata") if isinstance(package.get("metadata"), dict) else {}
        steps: list[str] = []

        if reason == "blacklisted_customer":
            steps.extend([
                "Review the matched traveler profile and confirm why it was blacklisted.",
                "Do not continue automated sales follow-up until the blacklist decision is verified.",
                "Escalate to a supervisor if the traveler claims the blacklist is incorrect.",
            ])
        elif reason == "phone_name_conflict":
            steps.extend([
                "Verify who owns the WhatsApp number before editing the traveler profile.",
                "Compare the matched traveler name with the newly submitted customer name.",
                "Continue only after confirming whether this is the same traveler or a true mismatch.",
            ])
        elif reason == "duplicate_phone_match":
            steps.extend([
                "Open the matching traveler profiles and identify which profile is valid.",
                "Resolve or merge duplicates before attaching new bookings or leads.",
                "Document the final identity decision in the handoff notes.",
            ])
        elif reason == "archived_traveler":
            steps.extend([
                "Check why the matched traveler was archived or marked inactive.",
                "Confirm whether reactivation is allowed before continuing the sales flow.",
            ])
        elif reason == "group_booking_quote":
            group_size = metadata.get("group_size")
            trip_id = metadata.get("trip_id")
            steps.extend([
                "Review current capacity, room mix, and commercial terms before promising availability.",
                "Prepare the group pricing or discount decision for the traveler.",
            ])
            if group_size:
                steps.append(f"Confirm the requested group size: {group_size}.")
            if trip_id:
                steps.append(f"Review the requested trip: {trip_id}.")
        elif reason == "customer_requested_human_agent":
            steps.extend([
                "Continue the conversation directly with the traveler.",
                "Confirm the reason for the handoff and the traveler’s expected next step.",
            ])
        elif reason == "manual_admin_handoff":
            steps.append("Review the employee notes and continue handling the case manually.")

        return steps

    @staticmethod
    def _handoff_notes_blob(reason_text: str, package: dict[str, Any], notes: str) -> str:
        reason_code = str(package.get("reason_code") or "").strip()
        lines = [reason_text.strip()] if str(reason_text or "").strip() else []
        action_steps = UnifiedCRMService._handoff_action_steps(reason_code, package)
        if action_steps:
            lines.append("Recommended actions:\n- " + "\n- ".join(action_steps))
        if notes.strip():
            lines.append(f"Employee notes:\n{notes.strip()}")
        compact = json.dumps(package, ensure_ascii=False, separators=(",", ":"))
        lines.append(f"handoff_context={compact}")
        return "\n\n".join(line for line in lines if line)

    @classmethod
    def derive_lead_stage(cls, result: dict[str, Any]) -> tuple[str, str]:
        traveler = result.get("traveler")
        traveler_status = str(traveler.get("status") or "").strip() if isinstance(traveler, dict) else ""
        trip_result = result.get("trip_result") or {}
        open_trips = trip_result.get("open_trips", [])
        tbd_trips = trip_result.get("date_tbd_trips", [])

        if result.get("handoff_reason") == "blacklisted_customer":
            return "Blocked", "Critical"
        if result.get("handoff_required"):
            return "Needs Review", "High"
        if traveler_status == "VIP":
            return ("VIP Priority", "High") if open_trips else ("VIP Follow Up", "High")
        if traveler_status == "Repeat":
            return ("Repeat Priority", "Medium") if open_trips else ("Repeat Follow Up", "Medium")
        if open_trips:
            return "Qualified", "Medium"
        if tbd_trips:
            return "Follow Up Needed", "Medium"
        if result.get("match_status") == "not_found":
            return "New Lead", "Medium"
        if result.get("match_status") == "single_match":
            return "Existing Traveler", "Low"
        return "New Inquiry", "Low"

    @classmethod
    def derive_follow_up(cls, result: dict[str, Any], timestamp: datetime) -> tuple[str, str]:
        stage, _priority = cls.derive_lead_stage(result)
        if stage == "Blocked":
            return "Do Not Contact", ""
        if result.get("handoff_required"):
            return "Urgent", timestamp.date().isoformat()
        trip_result = result.get("trip_result") or {}
        if trip_result.get("open_trips"):
            return "Follow Up Soon", (timestamp.date() + timedelta(days=2)).isoformat()
        if trip_result.get("date_tbd_trips"):
            return "Awaiting Dates", (timestamp.date() + timedelta(days=7)).isoformat()
        if result.get("match_status") == "not_found":
            return "Qualify Lead", (timestamp.date() + timedelta(days=3)).isoformat()
        return "Monitor", (timestamp.date() + timedelta(days=14)).isoformat()

    @staticmethod
    def _collect_trip_ids(trip_result: dict[str, list[dict[str, Any]]]) -> list[str]:
        open_ids = [str(item.get("trip_id") or "") for item in trip_result.get("open_trips", []) if item.get("trip_id")]
        tbd_ids = [str(item.get("trip_id") or "") for item in trip_result.get("date_tbd_trips", []) if item.get("trip_id")]
        return open_ids or tbd_ids

    def _update_lead_for_booking(
        self,
        connection: sqlite3.Connection,
        *,
        lead_id: str,
        booking_id: str,
        trip_id: str,
        interaction_id: str,
        timestamp: datetime,
    ) -> dict[str, Any] | None:
        row = connection.execute(
            "SELECT customer_tier FROM leads WHERE lead_id = ?",
            (lead_id,),
        ).fetchone()
        if not row:
            return None
        customer_tier = str(row["customer_tier"] or "").strip()
        if customer_tier == "VIP":
            stage = "VIP Booking Draft"
        elif customer_tier == "Repeat":
            stage = "Repeat Booking Draft"
        else:
            stage = "Booking Draft Created"
        follow_up_due = (timestamp.date() + timedelta(days=1)).isoformat()
        connection.execute(
            """
            UPDATE leads
            SET lead_stage = ?, priority = ?, follow_up_status = ?, follow_up_due_date = ?,
                updated_at = ?, interested_trip_ids = ?, suggested_trip_ids = ?,
                last_interaction_id = ?, booking_id = ?
            WHERE lead_id = ?
            """,
            (
                stage,
                "High",
                "Awaiting Deposit",
                follow_up_due,
                timestamp.isoformat(timespec="seconds"),
                trip_id,
                trip_id,
                interaction_id,
                booking_id,
                lead_id,
            ),
        )
        return {
            "lead_id": lead_id,
            "lead_stage": stage,
            "follow_up_status": "Awaiting Deposit",
            "follow_up_due_date": follow_up_due,
            "booking_id": booking_id,
        }

    def get_available_trips(self, trip_type: str | None, *, today: date | None = None) -> list[dict[str, Any]]:
        normalized_type = self.normalize_trip_type(trip_type)
        if trip_type and normalized_type is None:
            raise ValueError(f"Unsupported trip type {trip_type!r}")
        today = today or date.today()
        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._reconcile_trip_room_holds(connection)
            select_columns = self._trip_select_columns(connection)
            rows = connection.execute(
                f"""
                SELECT {select_columns}
                FROM trips
                ORDER BY start_date ASC, trip_name ASC
                """
            ).fetchall()
            connection.commit()

        results: list[dict[str, Any]] = []
        for row in rows:
            record = self._trip_row_to_dict(row)
            if normalized_type and record["trip_type"] != normalized_type:
                continue
            if not self._trip_is_candidate(record, today=today):
                continue
            results.append(record)
        return results

    def build_trip_result(self, trip_type: str | None, *, today: date | None = None) -> dict[str, list[dict[str, Any]]]:
        trips = self.get_available_trips(trip_type, today=today)
        open_trips: list[dict[str, Any]] = []
        date_tbd_trips: list[dict[str, Any]] = []
        for trip in trips:
            if trip["start_date"]:
                open_trips.append(trip)
            else:
                date_tbd_trips.append(trip)
        open_trips.sort(key=lambda item: (item["start_date"] or "", item["trip_name"]))
        date_tbd_trips.sort(key=lambda item: ((item["year"] or 0), item["trip_name"]))
        return {"open_trips": open_trips, "date_tbd_trips": date_tbd_trips}

    def sync_trip_to_sheet(self, trip_id: str | dict[str, Any]) -> dict[str, Any]:
        if not self.settings.sheet_export_enabled:
            return {
                "status": "disabled",
                "backend": self.settings.sheet_backend,
                "reason": "CRM is the operational authority; export must be explicitly enabled.",
            }
        if self.settings.sheet_backend == "excel":
            return self._run_sync_trip_to_sheet_worker(trip_id)

        # Spin up a daemon thread to process trip updates to sheets asynchronously
        import threading

        def run_async():
            try:
                self._run_sync_trip_to_sheet_worker(trip_id)
            except Exception:
                pass

        thread = threading.Thread(target=run_async)
        thread.daemon = True
        thread.start()
        return {"status": "queued", "message": "Background trip sync thread initiated."}

    def _run_sync_trip_to_sheet_worker(self, trip_id: str | dict[str, Any]) -> dict[str, Any]:
        record = trip_id if isinstance(trip_id, dict) else self._fetch_trip_record(trip_id)
        if not record:
            raise ValueError(f"Trip was not found for sync: {trip_id}")
        backend = self.settings.sheet_backend
        trip_key = str(record.get("trip_id") or "")
        try:
            if backend == "excel":
                updated = 0
                workbook_paths = [self.settings.excel_runtime_workbook]
                if self.settings.allow_source_workbook_writes:
                    workbook_paths.insert(0, self.settings.excel_source_workbook)
                for workbook_path in workbook_paths:
                    if not workbook_path:
                        continue
                    self._upsert_trip_in_excel(workbook_path, record)
                    updated += 1
                self._record_sync_success("Trips", trip_key)
                return {"status": "ok", "backend": backend, "updated_workbooks": updated}
            if backend == "google_sheets":
                self._upsert_trip_in_google_sheet(record)
                self._record_sync_success("Trips", trip_key)
                return {"status": "ok", "backend": backend, "updated_workbooks": 1}
        except Exception as e:
            self._record_sync_failure("Trips", trip_key, str(e))
            return {"status": "failed", "backend": backend, "reason": str(e)}
        return {"status": "skipped", "backend": backend, "reason": "Sheet sync backend not supported in this phase."}

    def _fetch_trip_record(self, trip_id: str) -> dict[str, Any] | None:
        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._reconcile_trip_room_holds(connection, trip_id)
            select_columns = self._trip_select_columns(connection)
            row = connection.execute(
                f"""
                SELECT {select_columns}
                FROM trips
                WHERE trip_id = ?
                """,
                (trip_id,),
            ).fetchone()
            connection.commit()
        return self._trip_row_to_dict(row) if row else None

    def _trip_row_to_dict(self, row: sqlite3.Row) -> dict[str, Any]:
        def value(name: str, default: Any = None) -> Any:
            try:
                return row[name]
            except (KeyError, IndexError, TypeError):
                return default

        single_remaining = self._as_int(value("single_remaining"))
        double_remaining = self._as_int(value("double_remaining"))
        triple_remaining = self._as_int(value("triple_remaining"))
        boys_double = self._as_int(value("boys_double"), default=0)
        girls_double = self._as_int(value("girls_double"), default=0)
        boys_triple = self._as_int(value("boys_triple"), default=0)
        girls_triple = self._as_int(value("girls_triple"), default=0)
        draft_holds_boys_double = self._as_int(value("draft_holds_boys_double"), default=0) or 0
        draft_holds_girls_double = self._as_int(value("draft_holds_girls_double"), default=0) or 0
        draft_holds_boys_triple = self._as_int(value("draft_holds_boys_triple"), default=0) or 0
        draft_holds_girls_triple = self._as_int(value("draft_holds_girls_triple"), default=0) or 0
        draft_holds_single = self._as_int(value("draft_holds_single"), default=0)
        draft_holds_double = self._as_int(value("draft_holds_double"), default=0)
        draft_holds_triple = self._as_int(value("draft_holds_triple"), default=0)
        available_single = max(single_remaining - draft_holds_single, 0) if single_remaining is not None else None
        available_double = max(double_remaining - draft_holds_double, 0) if double_remaining is not None else None
        available_triple = max(triple_remaining - draft_holds_triple, 0) if triple_remaining is not None else None
        availability_note = str(value("trip_availability_note") or "").strip()
        values = [value for value in (available_single, available_double, available_triple) if value is not None]
        remaining_places = sum(values) if values else 0
        if remaining_places > 0:
            availability_status = "available"
        elif availability_note:
            availability_status = "manual_follow_up"
        else:
            availability_status = "full"
        result = {
            "row": 0,
            "trip_id": str(value("trip_id") or "").strip(),
            "trip_name": str(value("trip_name") or "").strip(),
            "trip_name_ar": str(value("trip_name_ar") or "").strip(),
            "trip_type": str(value("type") or "").strip(),
            "year": self._as_int(value("year")),
            "start_date": self._iso_date(value("start_date")),
            "end_date": self._iso_date(value("end_date")),
            "sales_status": str(value("sales_status") or "").strip() or "Open",
            "data_audit": str(value("data_audit") or "").strip(),
            "trip_window_status": str(value("trip_window_status") or "").strip(),
            "trip_availability_note": availability_note,
            "next_reengage_date": self._iso_date(value("next_reengage_date")),
            "trip_leader": str(value("trip_leader") or "").strip(),
            "single_total": self._as_int(value("single_total"), default=0),
            "double_total": self._as_int(value("double_total"), default=0),
            "triple_total": self._as_int(value("triple_total"), default=0),
            "single_remaining": single_remaining,
            "double_remaining": double_remaining,
            "triple_remaining": triple_remaining,
            "draft_holds_single": draft_holds_single,
            "draft_holds_double": draft_holds_double,
            "draft_holds_triple": draft_holds_triple,
            "boys_double": max(boys_double - draft_holds_boys_double, 0),
            "girls_double": max(girls_double - draft_holds_girls_double, 0),
            "boys_triple": max(boys_triple - draft_holds_boys_triple, 0),
            "girls_triple": max(girls_triple - draft_holds_girls_triple, 0),
            "draft_holds_boys_double": draft_holds_boys_double,
            "draft_holds_girls_double": draft_holds_girls_double,
            "draft_holds_boys_triple": draft_holds_boys_triple,
            "draft_holds_girls_triple": draft_holds_girls_triple,
            "available_single": available_single,
            "available_double": available_double,
            "available_triple": available_triple,
            "remaining_places": remaining_places,
            "availability_status": availability_status,
            "public_price": str(value("public_price") or "").strip(),
            "room_prices_json": str(value("room_prices_json") or "").strip(),
            "room_prices": parse_room_prices(value("room_prices_json")),
            "public_description": str(value("public_description") or "").strip(),
            "itinerary": str(value("itinerary") or "").strip(),
            "inclusions": str(value("inclusions") or "").strip(),
            "exclusions": str(value("exclusions") or "").strip(),
            "sales_notes": str(value("sales_notes") or "").strip(),
        }
        result["program"] = build_trip_program(result)
        return result

    def _trip_is_candidate(self, trip: dict[str, Any], *, today: date) -> bool:
        status = (trip.get("sales_status") or "").strip().lower()
        if status in INACTIVE_TRIP_STATUSES:
            return False
        start_date = trip.get("start_date")
        availability_note = (trip.get("trip_availability_note") or "").strip()
        has_capacity = bool((trip.get("remaining_places") or 0) > 0)
        if start_date:
            try:
                return date.fromisoformat(start_date) >= today and (has_capacity or bool(availability_note))
            except ValueError:
                return False
        return bool(availability_note or status == "date tbd" or has_capacity)

    def _upsert_trip_in_excel(self, workbook_path: Path, record: dict[str, Any]) -> None:
        if not workbook_path.exists():
            return
        with _SHEET_WRITE_LOCK:
            wb = load_workbook(workbook_path)
            try:
                ws = wb[TRIPS_SHEET_NAME]
                target_row = None
                for row_idx in range(TRIPS_DATA_START_ROW, ws.max_row + 1):
                    current_trip_id = str(ws.cell(row_idx, TRIPS_SHEET_COLUMNS["trip_id"]).value or "").strip()
                    if current_trip_id == record["trip_id"]:
                        target_row = row_idx
                        break
                if target_row is None:
                    target_row = ws.max_row + 1
                    if target_row < TRIPS_DATA_START_ROW:
                        target_row = TRIPS_DATA_START_ROW
                for field_name, col_idx in TRIPS_SHEET_COLUMNS.items():
                    ws.cell(target_row, col_idx).value = self._sheet_value(field_name, self._record_value(record, field_name))
                wb.save(workbook_path)
            finally:
                wb.close()

    def _upsert_trip_in_google_sheet(self, record: dict[str, Any]) -> None:
        if not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            raise RuntimeError("Google Sheets sync is not configured.")

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        ws = client.open_by_key(self.settings.google_sheet_id).worksheet(TRIPS_SHEET_NAME)
        rows = ws.get_all_values()
        target_row = None
        for row_idx in range(TRIPS_DATA_START_ROW, len(rows) + 1):
            current_trip_id = rows[row_idx - 1][TRIPS_SHEET_COLUMNS["trip_id"] - 1] if len(rows[row_idx - 1]) >= TRIPS_SHEET_COLUMNS["trip_id"] else ""
            if str(current_trip_id).strip() == record["trip_id"]:
                target_row = row_idx
                break
        if target_row is None:
            target_row = max(len(rows) + 1, TRIPS_DATA_START_ROW)
            row_values = [""] * max(TRIPS_SHEET_COLUMNS.values())
            for field_name, col_idx in TRIPS_SHEET_COLUMNS.items():
                row_values[col_idx - 1] = self._sheet_value(field_name, self._record_value(record, field_name))
            ws.append_row(row_values, value_input_option="USER_ENTERED")
            return

        cell_list = []
        for field_name, col_idx in TRIPS_SHEET_COLUMNS.items():
            cell_list.append(gspread.Cell(target_row, col_idx, self._sheet_value(field_name, self._record_value(record, field_name))))
        ws.update_cells(cell_list, value_input_option="USER_ENTERED")

    @staticmethod
    def _table_for_mapping(mapping_name: str) -> str:
        return {
            "Travelers": "travelers",
            "Leads": "leads",
            "Interactions": "interactions",
            "Trip Bookings": "trip_bookings",
            "Booking Event Trail": "booking_event_trail",
            "Handoff Queue": "handoff_queue",
        }[mapping_name]

    @staticmethod
    def _key_column_for_mapping(mapping_name: str) -> str:
        return {
            "Travelers": "traveler_id",
            "Leads": "lead_id",
            "Interactions": "interaction_id",
            "Trip Bookings": "booking_id",
            "Booking Event Trail": "event_id",
            "Handoff Queue": "handoff_id",
        }[mapping_name]

    def _fetch_table_record(self, table_name: str, key_column: str, record_id: str) -> dict[str, Any] | None:
        with self.connect() as connection:
            if table_name == "booking_event_trail":
                self._ensure_booking_event_trail_table(connection)
            row = connection.execute(
                f"SELECT * FROM {table_name} WHERE {key_column} = ?",
                (record_id,),
            ).fetchone()
        return dict(row) if row else None

    def _excel_outcome_workbooks(self) -> list[Path]:
        paths: list[Path] = []
        if self.settings.excel_runtime_workbook:
            paths.append(self.settings.excel_runtime_workbook)
        elif self.settings.excel_source_workbook:
            paths.append(self.settings.excel_source_workbook)
        return paths

    def _upsert_record_in_excel(
        self,
        workbook_path: Path,
        mapping: dict[str, Any],
        key_column: str,
        record: dict[str, Any],
    ) -> None:
        if not workbook_path or not workbook_path.exists():
            return
        with _SHEET_WRITE_LOCK:
            wb = load_workbook(workbook_path)
            try:
                ws = wb[mapping["sheet_name"]] if mapping["sheet_name"] in wb.sheetnames else wb.create_sheet(mapping["sheet_name"])
                header_row = int(mapping["header_row"])
                header_map = self._ensure_sheet_headers(ws, header_row, list(mapping["columns"].values()))
                key_header = mapping["columns"][key_column]
                key_value = str(record.get(key_column) or "").strip()
                target_row = self._find_sheet_row(ws, header_map[key_header], key_value, header_row + 1)
                if target_row is None:
                    target_row = ws.max_row + 1
                    if target_row <= header_row:
                        target_row = header_row + 1
                for field_name, header in mapping["columns"].items():
                    value = self._sheet_value(field_name, record.get(field_name))
                    ws.cell(target_row, header_map[header]).value = value
                wb.save(workbook_path)
            finally:
                wb.close()

    def _delete_record_in_excel(
        self,
        workbook_path: Path | None,
        mapping: dict[str, Any],
        key_column: str,
        record_id: str,
    ) -> bool:
        if not workbook_path or not workbook_path.exists():
            return False
        with _SHEET_WRITE_LOCK:
            wb = load_workbook(workbook_path)
            try:
                if mapping["sheet_name"] not in wb.sheetnames:
                    return False
                ws = wb[mapping["sheet_name"]]
                header_row = int(mapping["header_row"])
                header_map = self._ensure_sheet_headers(ws, header_row, list(mapping["columns"].values()))
                key_header = mapping["columns"][key_column]
                target_row = self._find_sheet_row(ws, header_map[key_header], str(record_id).strip(), header_row + 1)
                if target_row is None:
                    return False
                ws.delete_rows(target_row, 1)
                wb.save(workbook_path)
                return True
            finally:
                wb.close()

    def _upsert_records_in_excel(
        self,
        workbook_path: Path,
        record_jobs: list[tuple[str, dict[str, Any], str, dict[str, Any]]],
    ) -> None:
        if not workbook_path or not workbook_path.exists() or not record_jobs:
            return
        with _SHEET_WRITE_LOCK:
            wb = load_workbook(workbook_path)
            try:
                for _mapping_name, mapping, key_column, record in record_jobs:
                    ws = wb[mapping["sheet_name"]] if mapping["sheet_name"] in wb.sheetnames else wb.create_sheet(mapping["sheet_name"])
                    header_row = int(mapping["header_row"])
                    header_map = self._ensure_sheet_headers(ws, header_row, list(mapping["columns"].values()))
                    key_header = mapping["columns"][key_column]
                    key_value = str(record.get(key_column) or "").strip()
                    target_row = self._find_sheet_row(ws, header_map[key_header], key_value, header_row + 1)
                    if target_row is None:
                        target_row = ws.max_row + 1
                        if target_row <= header_row:
                            target_row = header_row + 1
                    for field_name, header in mapping["columns"].items():
                        value = self._sheet_value(field_name, record.get(field_name))
                        ws.cell(target_row, header_map[header]).value = value
                wb.save(workbook_path)
            finally:
                wb.close()

    def _upsert_record_in_google_sheet(
        self,
        mapping: dict[str, Any],
        key_column: str,
        record: dict[str, Any],
    ) -> None:
        if not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            raise RuntimeError("Google Sheets sync is not configured.")

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        spreadsheet = client.open_by_key(self.settings.google_sheet_id)
        try:
            ws = spreadsheet.worksheet(mapping["sheet_name"])
        except gspread.exceptions.WorksheetNotFound:
            ws = spreadsheet.add_worksheet(title=mapping["sheet_name"], rows=100, cols=max(len(mapping["columns"]), 10))

        header_row = int(mapping["header_row"])
        rows = ws.get_all_values()
        while len(rows) < header_row:
            rows.append([])
        header_map = self._ensure_google_headers(ws, rows, header_row, list(mapping["columns"].values()))
        key_header = mapping["columns"][key_column]
        key_value = str(record.get(key_column) or "").strip()
        target_row = None
        key_idx = header_map[key_header] - 1
        for row_idx in range(header_row + 1, len(rows) + 1):
            row = rows[row_idx - 1]
            current = row[key_idx] if len(row) > key_idx else ""
            if str(current).strip() == key_value:
                target_row = row_idx
                break
        values_by_col = {
            header_map[header]: self._sheet_value(field_name, record.get(field_name))
            for field_name, header in mapping["columns"].items()
        }
        if target_row is None:
            row_values = [""] * max(values_by_col)
            for col_idx, value in values_by_col.items():
                row_values[col_idx - 1] = value
            ws.append_row(row_values, value_input_option="USER_ENTERED")
            return

        cells = [gspread.Cell(target_row, col_idx, value) for col_idx, value in values_by_col.items()]
        ws.update_cells(cells, value_input_option="USER_ENTERED")

    def _delete_record_in_google_sheet(
        self,
        mapping: dict[str, Any],
        key_column: str,
        record_id: str,
    ) -> bool:
        if not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            raise RuntimeError("Google Sheets sync is not configured.")

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        ws = client.open_by_key(self.settings.google_sheet_id).worksheet(mapping["sheet_name"])
        rows = ws.get_all_values()
        header_map = self._ensure_google_headers(ws, rows, int(mapping["header_row"]), list(mapping["columns"].values()))
        key_header = mapping["columns"][key_column]
        target_row = None
        key_idx = header_map[key_header] - 1
        for row_idx in range(int(mapping["header_row"]) + 1, len(rows) + 1):
            current = rows[row_idx - 1][key_idx] if len(rows[row_idx - 1]) > key_idx else ""
            if str(current or "").strip() == str(record_id).strip():
                target_row = row_idx
                break
        if target_row is None:
            return False
        ws.delete_rows(target_row)
        return True

    @staticmethod
    def _ensure_sheet_headers(ws, header_row: int, headers: list[str]) -> dict[str, int]:
        header_map: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(header_row, col_idx).value or "").strip()
            if header and header not in header_map:
                header_map[header] = col_idx
        next_col = ws.max_column + 1
        for header in headers:
            if header in header_map:
                continue
            ws.cell(header_row, next_col).value = header
            header_map[header] = next_col
            next_col += 1
        return header_map

    @staticmethod
    def _ensure_google_headers(ws, rows: list[list[str]], header_row: int, headers: list[str]) -> dict[str, int]:
        current = rows[header_row - 1] if len(rows) >= header_row else []
        header_map: dict[str, int] = {}
        for col_idx, header in enumerate(current, start=1):
            clean = str(header or "").strip()
            if clean and clean not in header_map:
                header_map[clean] = col_idx
        updates: list[Any] = []
        next_col = max(len(current), 0) + 1
        import gspread

        for header in headers:
            if header in header_map:
                continue
            updates.append(gspread.Cell(header_row, next_col, header))
            header_map[header] = next_col
            next_col += 1
        if updates:
            ws.update_cells(updates, value_input_option="USER_ENTERED")
        return header_map

    @staticmethod
    def _find_sheet_row(ws, key_col: int, key_value: str, start_row: int) -> int | None:
        for row_idx in range(start_row, ws.max_row + 1):
            if str(ws.cell(row_idx, key_col).value or "").strip() == key_value:
                return row_idx
        return None

    @staticmethod
    def _clear_dead_local_proxy() -> None:
        dead_proxy = {"http://127.0.0.1:9", "https://127.0.0.1:9", "127.0.0.1:9"}
        for key in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy"):
            value = (os.getenv(key) or "").strip().lower()
            if value in dead_proxy:
                os.environ.pop(key, None)

    @staticmethod
    def _ensure_booking_event_trail_table(connection: sqlite3.Connection) -> None:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS booking_event_trail (
                event_id TEXT PRIMARY KEY,
                occurred_at TEXT,
                event_type TEXT,
                event_label TEXT,
                traveler_id TEXT,
                lead_id TEXT,
                booking_id TEXT,
                trip_id TEXT,
                interaction_id TEXT,
                channel TEXT,
                actor TEXT,
                notes TEXT,
                metadata_json TEXT
            )
            """
        )

    @staticmethod
    def _next_daily_id(
        connection: sqlite3.Connection,
        table_name: str,
        column_name: str,
        prefix: str,
        timestamp: datetime,
    ) -> str:
        day_prefix = timestamp.strftime(f"{prefix}-%Y%m%d-")
        rows = connection.execute(
            f"SELECT {column_name} FROM {table_name} WHERE {column_name} LIKE ?",
            (f"{day_prefix}%",),
        ).fetchall()
        max_seq = 0
        for row in rows:
            value = str(row[column_name] or "")
            tail = value[len(day_prefix) :]
            if value.startswith(day_prefix) and tail.isdigit():
                max_seq = max(max_seq, int(tail))
        return f"{day_prefix}{max_seq + 1:04d}"

    @staticmethod
    def _next_booking_id(connection: sqlite3.Connection, trip_id: str, traveler_id: str) -> str:
        rows = connection.execute(
            "SELECT booking_id FROM trip_bookings WHERE trip_id = ?",
            (trip_id,),
        ).fetchall()
        sequence = len(rows) + 1
        traveler_digits = traveler_id[2:] if traveler_id.startswith("TR") else traveler_id
        if traveler_digits.isdigit():
            traveler_part = str(int(traveler_digits))
        else:
            traveler_part = traveler_id or "NEW"
        parts = trip_id.split("-")
        trip_mode = "I" if "-INT-" in trip_id else "L"
        trip_year = parts[2] if len(parts) > 2 else _utc_now().strftime("%y")
        trip_serial = parts[3] if len(parts) > 3 else "000"
        candidate = f"{traveler_part}-{trip_mode}{trip_year}{trip_serial}-{sequence:03d}"
        existing = {str(row["booking_id"]) for row in rows}
        while candidate in existing:
            sequence += 1
            candidate = f"{traveler_part}-{trip_mode}{trip_year}{trip_serial}-{sequence:03d}"
        return candidate

    @staticmethod
    def _sheet_value(field_name: str, value: Any) -> Any:
        if value in (None, ""):
            return ""
        if field_name in {"start_date", "end_date", "next_reengage_date"}:
            return str(value)
        return value

    @staticmethod
    def _record_value(record: dict[str, Any], field_name: str) -> Any:
        if field_name == "type":
            return record.get("trip_type") or record.get("type")
        return record.get(field_name)

    @staticmethod
    def _as_bool(value: Any, *, default: bool = False) -> bool:
        if value in (None, ""):
            return default
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}

    @staticmethod
    def _as_int(value: Any, default: int | None = None) -> int | None:
        if value in (None, ""):
            return default
        try:
            return int(value)
        except (TypeError, ValueError):
            try:
                return int(float(value))
            except (TypeError, ValueError):
                return default

    @staticmethod
    def _normalize_date_string(value: Any) -> str | None:
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        text = str(value).strip()
        if not text:
            return None
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(text).date().isoformat()
        except ValueError:
            return text

    @staticmethod
    def _iso_date(value: Any) -> str | None:
        if value in (None, ""):
            return None
        text = str(value)
        if "T" in text:
            text = text.split("T", 1)[0]
        if " " in text:
            text = text.split(" ", 1)[0]
        return text

    @staticmethod
    def _split_name(full_name: str) -> tuple[str, str]:
        parts = [part for part in (full_name or "").strip().split() if part]
        if not parts:
            return "", ""
        if len(parts) == 1:
            return parts[0], ""
        return parts[0], " ".join(parts[1:])

    @staticmethod
    def _next_prefixed_id(
        connection: sqlite3.Connection,
        table_name: str,
        column_name: str,
        prefix: str,
        padding: int,
    ) -> str:
        rows = connection.execute(f"SELECT {column_name} FROM {table_name}").fetchall()
        max_number = 0
        for row in rows:
            value = str(row[column_name] or "").strip()
            match = re.fullmatch(rf"{re.escape(prefix)}(\d+)", value)
            if match:
                max_number = max(max_number, int(match.group(1)))
        return f"{prefix}{max_number + 1:0{padding}d}"
