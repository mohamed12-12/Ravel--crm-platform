from __future__ import annotations

import argparse
import json
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.phase0_cleanup import ensure_headers, normalize_text  # noqa: E402
from scripts.phase4_sales_intelligence import ensure_leads_sheet  # noqa: E402


SOURCE_WORKBOOK = PROJECT_ROOT / "RT - Travelers Database.phase5.ready.xlsx"
RUNTIME_WORKBOOK = PROJECT_ROOT / "RT - Travelers Database.phase5.demo.xlsx"

TRAVELERS_SHEET = "Travelers"
TRIPS_SHEET = "Trips"
LEADS_SHEET = "Leads"
INTERACTIONS_SHEET = "Interactions"

FLOW_STATES_SHEET = "DM Flow States"
DM_COPY_LIBRARY_SHEET = "DM Copy Library"
LANGUAGE_TEMPLATES_SHEET = "Language Templates"
WAITLIST_SHEET = "Waitlist"
FOLLOW_UPS_SHEET = "Follow Ups"
TRIP_DRIPS_SHEET = "Trip Drips"
HANDOFF_QUEUE_SHEET = "Handoff Queue"

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TRIPS_BASE_HEADERS = [
    "Trip ID",
    "Trip Name",
    "Type",
    "Year",
    "Trip Leader",
    "Start Date",
    "End Date",
    "Sales Status",
    "Data Audit",
]

TRAVELERS_BASE_HEADERS = [
    "Status",
    "Traveler ID",
    "Full Name",
    "First Name",
    "Last Name",
    "Birthday",
    "Gender",
    "Nationality",
    "Code",
    "WhatsApp",
    "Email",
    "Community Whatsapp",
    "Residence",
    "Loc. Trips",
    "Int. Trips",
    "Total trips",
    "Comm. Events",
    "Lifetime Revenue",
    "Notes",
    "Introduce yourself",
    "Emergency Contact",
    "Emergency Phone",
    "Medical Notes",
    "Room Preference",
    "⭐ Rating (1–5)",
]

LEADS_BASE_HEADERS = [
    "Lead ID",
    "Created At",
    "Updated At",
    "Customer Name",
    "Raw Phone",
    "Integrated WhatsApp",
    "Phone Lookup Key",
    "Traveler ID",
    "Traveler Status",
    "Customer Tier",
    "Match Status",
    "Lead Stage",
    "Lead Source",
    "Channel",
    "Preferred Trip Type",
    "Interested Trip IDs",
    "Suggested Trip IDs",
    "Priority",
    "Follow Up Status",
    "Follow Up Due Date",
    "Last Interaction ID",
    "Interaction Count",
    "Handoff Required",
    "Handoff Reason",
    "Notes",
]

INTERACTIONS_BASE_HEADERS = [
    "Interaction ID",
    "Timestamp",
    "Channel",
    "Customer Name",
    "Raw Phone",
    "Integrated WhatsApp",
    "Phone Lookup Key",
    "Traveler ID",
    "Matched Row",
    "Status Snapshot",
    "Intent",
    "Trip Type",
    "Suggested Trips",
    "Action Taken",
    "Handoff Required",
    "Handoff Reason",
    "Agent Notes",
]

LEADS_EXTENSION_HEADERS = [
    "Flow Key",
    "Current Step",
    "Language",
    "Trigger Keyword",
    "Waitlist ID",
    "Handoff ID",
    "Booking ID",
    "Source Sheet",
    "Source Row",
]

INTERACTIONS_EXTENSION_HEADERS = [
    "Flow Key",
    "Step Key",
    "Message Key",
    "Language",
    "Source Sheet",
    "Source Row",
    "Outcome",
]

TRAVELERS_EXTENSION_HEADERS = [
    "Primary Language",
    "Current Flow Key",
    "Current Step",
    "Last Lead ID",
    "Last Booking ID",
]

TRIPS_EXTENSION_HEADERS = [
    "Trip Window Status",
    "Trip Availability Note",
    "Next Reengage Date",
]

FLOW_STATES_HEADERS = [
    "Flow Key",
    "Flow Name",
    "Trigger",
    "Reads From",
    "Writes To",
    "Shared Module",
    "Current Step",
    "Next Step",
    "Language",
    "Owner",
    "Status",
    "Notes",
]

DM_COPY_LIBRARY_HEADERS = [
    "Message Key",
    "Flow Key",
    "Step Key",
    "Channel",
    "Arabic Copy",
    "English Copy",
    "Buttons Arabic",
    "Buttons English",
    "Variables",
    "Source Sheet",
    "Active",
    "Notes",
]

LANGUAGE_TEMPLATE_HEADERS = [
    "Template Key",
    "Language",
    "Use Case",
    "Approved Copy",
    "Variables",
    "Status",
    "Notes",
]

WAITLIST_HEADERS = [
    "Waitlist ID",
    "Created At",
    "Lead ID",
    "Traveler ID",
    "Customer Name",
    "WhatsApp",
    "Language",
    "Flow Key",
    "Trip ID",
    "Trip Name",
    "Reason",
    "Reengage At",
    "Status",
    "Owner",
    "Notes",
]

FOLLOW_UP_HEADERS = [
    "Follow Up ID",
    "Created At",
    "Lead ID",
    "Traveler ID",
    "Trip ID",
    "Flow Key",
    "Trigger",
    "Due At",
    "Priority",
    "Language",
    "Channel",
    "Status",
    "Owner",
    "Notes",
]

TRIP_DRIP_HEADERS = [
    "Drip ID",
    "Trip ID",
    "Trip Name",
    "Trip Row",
    "Language",
    "Drip Type",
    "Send At",
    "Message Key",
    "Trigger Source",
    "Status",
    "Owner",
    "Notes",
]

HANDOFF_HEADERS = [
    "Handoff ID",
    "Created At",
    "Lead ID",
    "Traveler ID",
    "Trip ID",
    "Flow Key",
    "Reason",
    "Priority",
    "Channel",
    "Status",
    "Owner",
    "Assigned To",
    "Notes",
]

FLOW_STATE_ROWS = [
    {
        "Flow Key": "FLOW-01",
        "Flow Name": "Trigger Keyword DM",
        "Trigger": 'Comment keyword like "Siwa", "Maldives", or "island hopping"',
        "Reads From": "Trips, Travelers, DM Copy Library",
        "Writes To": "Leads, Interactions, Waitlist, Handoff Queue",
        "Shared Module": "Yes",
        "Current Step": "keyword_detected",
        "Next Step": "language_choice",
        "Language": "AR/EN",
        "Owner": "Sales Agent",
        "Status": "Active",
        "Notes": "Prompts must come from sheet copy only.",
    },
    {
        "Flow Key": "FLOW-02",
        "Flow Name": "New Follower Welcome",
        "Trigger": "New follower detected",
        "Reads From": "DM Copy Library, Travelers, Leads",
        "Writes To": "Leads, Interactions",
        "Shared Module": "Yes",
        "Current Step": "welcome",
        "Next Step": "browse_book_ask",
        "Language": "AR/EN",
        "Owner": "Sales Agent",
        "Status": "Active",
        "Notes": "All branches converge to the shared booking module.",
    },
    {
        "Flow Key": "SHARED-BOOKING",
        "Flow Name": "Shared Booking Module",
        "Trigger": "Call from Flow 1 or Flow 2",
        "Reads From": "Trips, Travelers, Leads, DM Copy Library",
        "Writes To": "Trip Bookings, Booking Alerts, Interactions",
        "Shared Module": "Yes",
        "Current Step": "personal_info",
        "Next Step": "deposit_details",
        "Language": "AR/EN",
        "Owner": "Sales Agent",
        "Status": "Active",
        "Notes": "Single reusable booking path for all booking entry points.",
    },
    {
        "Flow Key": "FLOW-03",
        "Flow Name": "Booking Confirmation + Pre-Trip Drip",
        "Trigger": "Deposit confirmed by team",
        "Reads From": "Trip Bookings, Trips, Travelers, DM Copy Library",
        "Writes To": "Interactions, Follow Ups, Trip Drips, Booking Alerts",
        "Shared Module": "No",
        "Current Step": "deposit_confirmed",
        "Next Step": "trip_drips",
        "Language": "AR/EN",
        "Owner": "Operations",
        "Status": "Active",
        "Notes": "Reminder timing should come from trip dates.",
    },
    {
        "Flow Key": "FLOW-04",
        "Flow Name": "Abandoned Inquiry / No Reply",
        "Trigger": "No reply after 24h",
        "Reads From": "Leads, Interactions, DM Copy Library",
        "Writes To": "Follow Ups, Waitlist, Leads",
        "Shared Module": "No",
        "Current Step": "waiting_reply",
        "Next Step": "nudges",
        "Language": "AR/EN",
        "Owner": "Sales Agent",
        "Status": "Active",
        "Notes": "Use sheet timing, not model memory.",
    },
    {
        "Flow Key": "FLOW-05",
        "Flow Name": "Custom / Private Trip Inquiry",
        "Trigger": 'Keyword like "private trip" or "customized"',
        "Reads From": "Leads, DM Copy Library",
        "Writes To": "Leads, Handoff Queue, Interactions",
        "Shared Module": "No",
        "Current Step": "qualifying_questions",
        "Next Step": "manual_quote",
        "Language": "AR/EN",
        "Owner": "Sales Team",
        "Status": "Active",
        "Notes": "Never auto-price private trips.",
    },
    {
        "Flow Key": "FLOW-06",
        "Flow Name": "Post-Trip Review + Re-Engagement",
        "Trigger": "Trip end date reached",
        "Reads From": "Trips, Trip Bookings, Travelers, DM Copy Library",
        "Writes To": "Interactions, Follow Ups, Leads",
        "Shared Module": "No",
        "Current Step": "post_trip",
        "Next Step": "review_offer",
        "Language": "AR/EN",
        "Owner": "Operations",
        "Status": "Active",
        "Notes": "Trigger by trip end date, not by memory.",
    },
]

DM_COPY_ROWS = [
    {
        "Message Key": "session.start",
        "Flow Key": "FLOW-02",
        "Step Key": "welcome",
        "Channel": "web-demo",
        "Arabic Copy": "مرحبًا، أنا مساعد مبيعات رحمة ترافلر. للبدء، من فضلك اكتب اسمك الكامل.",
        "English Copy": "Hello. I am Rahma Traveler's sales agent. To start, please share your full name.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{customer_name}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Initial demo prompt.",
    },
    {
        "Message Key": "session.ask_phone",
        "Flow Key": "FLOW-02",
        "Step Key": "phone",
        "Channel": "web-demo",
        "Arabic Copy": "شكرًا يا {customer_name}. من فضلك أرسل رقم واتساب الخاص بالعميل.",
        "English Copy": "Thanks, {customer_name}. Please share the customer's WhatsApp number.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{customer_name}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Lead capture prompt.",
    },
    {
        "Message Key": "session.ask_trip_type",
        "Flow Key": "FLOW-02",
        "Step Key": "trip_type",
        "Channel": "web-demo",
        "Arabic Copy": "رائع. هل الاستفسار عن رحلة داخلية أم دولية؟",
        "English Copy": "Great. Is this inquiry for a local trip or an international trip?",
        "Buttons Arabic": "داخلية;دولية",
        "Buttons English": "Local;International",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Trip type routing prompt.",
    },
    {
        "Message Key": "session.ask_trip_type_retry",
        "Flow Key": "FLOW-02",
        "Step Key": "trip_type_retry",
        "Channel": "web-demo",
        "Arabic Copy": "من فضلك رد بـ داخلية أو دولية حتى أتحقق من الرحلات المتاحة.",
        "English Copy": "Please reply with either local or international so I can check upcoming trips.",
        "Buttons Arabic": "داخلية;دولية",
        "Buttons English": "Local;International",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Retry prompt.",
    },
    {
        "Message Key": "session.handoff",
        "Flow Key": "SAFETY-HANDOFF",
        "Step Key": "review",
        "Channel": "web-demo",
        "Arabic Copy": "تمت إحالة الطلب للمراجعة من فريق مختص، وسيتابع معك أحد أفراد الفريق قريبًا.",
        "English Copy": "I flagged this request for a senior team review. A team member will follow up shortly.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Human handoff message.",
    },
    {
        "Message Key": "session.result_open_trips",
        "Flow Key": "FLOW-02",
        "Step Key": "results",
        "Channel": "web-demo",
        "Arabic Copy": "هذه هي الخيارات القادمة: {trip_lines}",
        "English Copy": "Here are upcoming options: {trip_lines}",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{trip_lines}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Used when Trips contains open inventory.",
    },
    {
        "Message Key": "session.result_date_tbd",
        "Flow Key": "FLOW-02",
        "Step Key": "results",
        "Channel": "web-demo",
        "Arabic Copy": "وجدت رحلات في هذا النوع لكن التواريخ لم تتأكد بعد: {trip_names}. يمكنني إضافتك للمتابعة فور فتح التواريخ.",
        "English Copy": "I found trips in this category, but the dates are not confirmed yet: {trip_names}. I can flag you for follow-up as soon as dates open.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{trip_names}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Fallback for future trips with missing dates.",
    },
    {
        "Message Key": "session.result_no_trip",
        "Flow Key": "FLOW-02",
        "Step Key": "results",
        "Channel": "web-demo",
        "Arabic Copy": "لا توجد رحلات مؤكدة قادمة في هذا التصنيف حاليًا.",
        "English Copy": "There are no confirmed upcoming trips in that category right now.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "No inventory fallback.",
    },
    {
        "Message Key": "flow1.teaser",
        "Flow Key": "FLOW-01",
        "Step Key": "teaser",
        "Channel": "instagram",
        "Arabic Copy": "تم فتح رحلة جديدة. هل تريد التفاصيل الكاملة؟",
        "English Copy": "A new trip just opened. Want the full details?",
        "Buttons Arabic": "نعم;لاحقًا",
        "Buttons English": "Yes;Later",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Keyword trigger teaser.",
    },
    {
        "Message Key": "flow1.language_choice",
        "Flow Key": "FLOW-01",
        "Step Key": "language",
        "Channel": "instagram",
        "Arabic Copy": "اختر اللغة: العربية / English",
        "English Copy": "Choose your language: AR / EN",
        "Buttons Arabic": "العربية;English",
        "Buttons English": "AR;EN",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Language selector.",
    },
    {
        "Message Key": "flow1.full_details",
        "Flow Key": "FLOW-01",
        "Step Key": "details",
        "Channel": "instagram",
        "Arabic Copy": "التواريخ، البرنامج، نطاق السعر، وما يشمله العرض.",
        "English Copy": "Dates, itinerary, price range, inclusions.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Trip detail block.",
    },
    {
        "Message Key": "shared.personal_info",
        "Flow Key": "SHARED-BOOKING",
        "Step Key": "personal_info",
        "Channel": "instagram",
        "Arabic Copy": "من فضلك أرسل الاسم الكامل، الهاتف، البريد الإلكتروني، والنوع.",
        "English Copy": "Please share full name, phone, email, and gender.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{customer_name};{phone};{email};{gender}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Shared booking module.",
    },
    {
        "Message Key": "shared.room_type",
        "Flow Key": "SHARED-BOOKING",
        "Step Key": "room_type",
        "Channel": "instagram",
        "Arabic Copy": "اختر نوع الغرفة: مفردة أو مزدوجة.",
        "English Copy": "Choose room type: Single or Double.",
        "Buttons Arabic": "مفردة;مزدوجة",
        "Buttons English": "Single;Double",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Room selection.",
    },
    {
        "Message Key": "shared.flights",
        "Flow Key": "SHARED-BOOKING",
        "Step Key": "flights",
        "Channel": "instagram",
        "Arabic Copy": "هل تفضل مع طيران أم بدون طيران؟",
        "English Copy": "Flights with or without?",
        "Buttons Arabic": "مع طيران;بدون طيران",
        "Buttons English": "With flight;Without flight",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Flight preference.",
    },
    {
        "Message Key": "shared.date",
        "Flow Key": "SHARED-BOOKING",
        "Step Key": "date",
        "Channel": "instagram",
        "Arabic Copy": "إذا كانت هناك أكثر من موعد، اختر موعدًا واحدًا.",
        "English Copy": "If multiple dates exist, choose one.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Date selector.",
    },
    {
        "Message Key": "shared.currency",
        "Flow Key": "SHARED-BOOKING",
        "Step Key": "currency",
        "Channel": "instagram",
        "Arabic Copy": "اختر العملة: EGP أو USD.",
        "English Copy": "Choose currency: EGP or USD.",
        "Buttons Arabic": "EGP;USD",
        "Buttons English": "EGP;USD",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Currency selector.",
    },
    {
        "Message Key": "shared.deposit",
        "Flow Key": "SHARED-BOOKING",
        "Step Key": "deposit",
        "Channel": "instagram",
        "Arabic Copy": "سيتم إرسال قيمة الدفعة المقدمة وبيانات البنك الآن.",
        "English Copy": "Deposit amount and bank details will be sent next.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{amount};{bank_details}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Payment handoff.",
    },
    {
        "Message Key": "flow2.welcome",
        "Flow Key": "FLOW-02",
        "Step Key": "welcome",
        "Channel": "instagram",
        "Arabic Copy": "أهلًا بك في رحمة ترافلر.",
        "English Copy": "Welcome aboard Rahma Traveler.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "New follower welcome.",
    },
    {
        "Message Key": "flow2.actions",
        "Flow Key": "FLOW-02",
        "Step Key": "actions",
        "Channel": "instagram",
        "Arabic Copy": "ماذا تريد أن تفعل: تصفح الرحلات، الحجز الآن، أم طرح سؤال؟",
        "English Copy": "What would you like to do: browse trips, book now, or ask a question?",
        "Buttons Arabic": "تصفح الرحلات;احجز الآن;اسأل سؤالًا",
        "Buttons English": "Browse trips;Book now;Ask a question",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Branch choice.",
    },
    {
        "Message Key": "flow3.confirmation",
        "Flow Key": "FLOW-03",
        "Step Key": "confirmation",
        "Channel": "instagram",
        "Arabic Copy": "تم تأكيد الدفعة المقدمة. إليك ملخص الحجز.",
        "English Copy": "Your deposit was confirmed. Here is the full recap.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Booking confirmation.",
    },
    {
        "Message Key": "flow3.pretrip_d7",
        "Flow Key": "FLOW-03",
        "Step Key": "d7",
        "Channel": "instagram",
        "Arabic Copy": "باقي أسبوع على الرحلة. إليك الترتيبات والتذكيرات.",
        "English Copy": "Your trip is one week away. Here are the logistics and reminders.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{start_date}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Pre-trip drip T-7.",
    },
    {
        "Message Key": "flow3.pretrip_d3",
        "Flow Key": "FLOW-03",
        "Step Key": "d3",
        "Channel": "instagram",
        "Arabic Copy": "باقي ثلاثة أيام. راجع حقيبتك وتفاصيل التجمع.",
        "English Copy": "Three days to go. Please review your packing and meeting details.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{start_date}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Pre-trip drip T-3.",
    },
    {
        "Message Key": "flow3.pretrip_d1",
        "Flow Key": "FLOW-03",
        "Step Key": "d1",
        "Channel": "instagram",
        "Arabic Copy": "غدًا يوم السفر. يرجى تجهيز المستندات.",
        "English Copy": "Tomorrow is travel day. Please be ready with your documents.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{start_date}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Pre-trip drip T-1.",
    },
    {
        "Message Key": "flow3.day_of",
        "Flow Key": "FLOW-03",
        "Step Key": "day_of",
        "Channel": "instagram",
        "Arabic Copy": "اليوم موعد الرحلة. نتمنى لك سفرًا آمنًا ورجاءً تواصل مع الفريق عند الوصول.",
        "English Copy": "Today is the trip. Safe travels and please check in with the team.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{start_date}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Day-of check-in.",
    },
    {
        "Message Key": "flow4.nudge_24h",
        "Flow Key": "FLOW-04",
        "Step Key": "nudge_24h",
        "Channel": "instagram",
        "Arabic Copy": "الأماكن تمتلئ بسرعة. هل ما زلت مهتمًا؟",
        "English Copy": "Spots are filling up. Still interested?",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "24h nudge.",
    },
    {
        "Message Key": "flow4.nudge_48h",
        "Flow Key": "FLOW-04",
        "Step Key": "nudge_48h",
        "Channel": "instagram",
        "Arabic Copy": "تبقى آخر X أماكن. يرجى التأكيد قبل {date}.",
        "English Copy": "Last X spots. Please confirm by {date}.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{date}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "48h urgency nudge.",
    },
    {
        "Message Key": "flow5.private_intro",
        "Flow Key": "FLOW-05",
        "Step Key": "intro",
        "Channel": "instagram",
        "Arabic Copy": "أخبرنا باختصار عن طلب الرحلة الخاصة.",
        "English Copy": "Tell us a bit about your private trip request.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Private trip qualifier.",
    },
    {
        "Message Key": "flow5.questions",
        "Flow Key": "FLOW-05",
        "Step Key": "questions",
        "Channel": "instagram",
        "Arabic Copy": "عدد الأشخاص، الوجهة، الميزانية، والتواريخ من فضلك.",
        "English Copy": "Group size, destination, budget, and dates please.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "{group_size};{destination};{budget};{dates}",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Qualifying questions.",
    },
    {
        "Message Key": "flow6.review",
        "Flow Key": "FLOW-06",
        "Step Key": "review",
        "Channel": "instagram",
        "Arabic Copy": "نأمل أن تكون الرحلة قد نالت إعجابك. من فضلك أرسل تقييمًا سريعًا.",
        "English Copy": "We hope you had a great trip. Please leave a quick rating.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Post-trip review.",
    },
    {
        "Message Key": "flow6.post_trip",
        "Flow Key": "FLOW-06",
        "Step Key": "post_trip",
        "Channel": "instagram",
        "Arabic Copy": "نأمل أن تكون الرحلة قد أعجبتك. من فضلك شاركنا تقييمك أو صورة من الرحلة.",
        "English Copy": "We hope you had a great trip. Please share your rating or a photo from the trip.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Trip end-date trigger.",
    },
    {
        "Message Key": "flow6.ugc",
        "Flow Key": "FLOW-06",
        "Step Key": "ugc",
        "Channel": "instagram",
        "Arabic Copy": "شاركنا منشن في ستوري أو مراجعة أو صورة.",
        "English Copy": "Share a story tag, review, or photo.",
        "Buttons Arabic": "",
        "Buttons English": "",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "UGC prompt.",
    },
    {
        "Message Key": "flow6.reengage",
        "Flow Key": "FLOW-06",
        "Step Key": "reengage",
        "Channel": "instagram",
        "Arabic Copy": "هل ترغب في عرض الحجز المبكر للرحلة القادمة؟",
        "English Copy": "Would you like the early-bird offer for the next trip?",
        "Buttons Arabic": "نعم;لا",
        "Buttons English": "Yes;No",
        "Variables": "",
        "Source Sheet": "DM Copy Library",
        "Active": "Yes",
        "Notes": "Re-engagement prompt.",
    },
]

LANGUAGE_TEMPLATE_ROWS = [
    {
        "Template Key": "rules.no_hallucination",
        "Language": "EN",
        "Use Case": "System rule",
        "Approved Copy": "Never invent dates, prices, availability, or customer data. Read from the workbook or hand off.",
        "Variables": "",
        "Status": "Active",
        "Notes": "Safety boundary.",
    },
    {
        "Template Key": "rules.no_hallucination",
        "Language": "AR",
        "Use Case": "System rule",
        "Approved Copy": "لا تخترع التواريخ أو الأسعار أو التوفر أو بيانات العميل. اقرأ من الشيت أو حوّل الطلب لفريق بشري.",
        "Variables": "",
        "Status": "Active",
        "Notes": "Safety boundary.",
    },
    {
        "Template Key": "fallback.no_data",
        "Language": "EN",
        "Use Case": "Missing workbook data",
        "Approved Copy": "I could not find that in the workbook, so I am flagging it for review.",
        "Variables": "",
        "Status": "Active",
        "Notes": "Missing data fallback.",
    },
    {
        "Template Key": "fallback.no_data",
        "Language": "AR",
        "Use Case": "Missing workbook data",
        "Approved Copy": "لم أجد هذه المعلومة في الشيت، لذلك سأحوّلها للمراجعة.",
        "Variables": "",
        "Status": "Active",
        "Notes": "Missing data fallback.",
    },
    {
        "Template Key": "fallback.handoff",
        "Language": "EN",
        "Use Case": "Human handoff",
        "Approved Copy": "A team member will handle this next.",
        "Variables": "",
        "Status": "Active",
        "Notes": "Human review message.",
    },
    {
        "Template Key": "fallback.handoff",
        "Language": "AR",
        "Use Case": "Human handoff",
        "Approved Copy": "سيتولى أحد أفراد الفريق هذه الخطوة التالية.",
        "Variables": "",
        "Status": "Active",
        "Notes": "Human review message.",
    },
    {
        "Template Key": "fallback.choose_language",
        "Language": "EN",
        "Use Case": "Language choice",
        "Approved Copy": "Choose your language: AR / EN",
        "Variables": "",
        "Status": "Active",
        "Notes": "Language selector.",
    },
    {
        "Template Key": "fallback.choose_language",
        "Language": "AR",
        "Use Case": "Language choice",
        "Approved Copy": "اختر اللغة: العربية / English",
        "Variables": "",
        "Status": "Active",
        "Notes": "Language selector.",
    },
]


@dataclass
class WorkbookChangeSummary:
    output_workbook: str
    runtime_copy: str
    added_sheets: list[str]
    appended_columns: dict[str, list[str]]
    seeded_rows: dict[str, int]


def style_header_row(ws, row: int = 1) -> None:
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row, col_idx)
        if cell.value in (None, ""):
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT


def clear_sheet(ws, rows_to_keep: int = 1) -> None:
    if ws.max_row > rows_to_keep:
        ws.delete_rows(rows_to_keep + 1, ws.max_row - rows_to_keep)


def write_table(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    clear_sheet(ws, rows_to_keep=0)
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(1, col_idx).value = header
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row_idx, col_idx).value = row.get(header)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    style_header_row(ws, 1)
    auto_width(ws)


def ensure_or_append_headers(ws, header_row: int, base_headers: list[str], new_headers: list[str]) -> list[str]:
    header_map = ensure_headers(
        ws,
        header_row=header_row,
        expected_headers=base_headers,
        helper_headers=[],
    )

    appended: list[str] = []
    reference_col = ws.max_column
    for header in new_headers:
        if header in header_map:
            continue
        reference_col += 1
        ws.cell(header_row, reference_col).value = header
        if header_row == 1:
            style_header_row(ws, 1)
        else:
            # Copy a nearby header style to keep the workbook visually consistent.
            source = ws.cell(header_row, reference_col - 1)
            target = ws.cell(header_row, reference_col)
            target.fill = copy(source.fill)
            target.font = copy(source.font)
            target.alignment = copy(source.alignment)
            target.border = copy(source.border)
            target.number_format = source.number_format
        header_map[header] = reference_col
        appended.append(header)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{max(ws.max_row, header_row)}"
    auto_width(ws)
    return appended


def auto_width(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_length = 0
        for row_idx in range(1, min(ws.max_row, 200) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value in (None, ""):
                continue
            text = str(value)
            max_length = max(max_length, min(len(text), 48))
        if max_length:
            ws.column_dimensions[col_letter].width = min(max_length + 2, 52)


def ensure_sheet(workbook, name: str, headers: list[str], *, clear_existing: bool = True, header_row: int = 1):
    if name in workbook.sheetnames:
        ws = workbook[name]
        if clear_existing:
            clear_sheet(ws, rows_to_keep=0)
    else:
        ws = workbook.create_sheet(name)
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(header_row, col_idx).value = header
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(ws.max_row, header_row)}"
    style_header_row(ws, header_row)
    auto_width(ws)
    return ws


def sheet_headers(ws, row: int) -> dict[str, int]:
    return {
        normalize_text(ws.cell(row, col).value): col
        for col in range(1, ws.max_column + 1)
        if normalize_text(ws.cell(row, col).value)
    }


def meaningful_trip_rows(ws) -> list[int]:
    rows: list[int] = []
    blank_streak = 0
    for row_idx in range(3, ws.max_row + 1):
        trip_id = normalize_text(ws.cell(row_idx, 1).value)
        trip_name = normalize_text(ws.cell(row_idx, 2).value)
        if not trip_id and not trip_name:
            blank_streak += 1
            if rows and blank_streak >= 25:
                break
            continue
        blank_streak = 0
        rows.append(row_idx)
    return rows


def date_formula(sheet_name: str, row_idx: int, offset_days: int, date_col_letter: str) -> str:
    reference = f"'{sheet_name}'!{date_col_letter}{row_idx}"
    if offset_days == 0:
        return f'=IF({reference}="","",{reference})'
    sign = "+" if offset_days > 0 else "-"
    return f'=IF({reference}="","",{reference}{sign}{abs(offset_days)})'


def seed_trip_drips(wb) -> int:
    if TRIPS_SHEET not in wb.sheetnames:
        return 0

    trips_ws = wb[TRIPS_SHEET]
    trip_headers = sheet_headers(trips_ws, 2)
    start_col = get_column_letter(trip_headers["Start Date"])
    end_col = get_column_letter(trip_headers["End Date"])

    ws = ensure_sheet(wb, TRIP_DRIPS_SHEET, TRIP_DRIP_HEADERS, clear_existing=True)

    drip_events = [
        ("D-7", "flow3.pretrip_d7", -7, "Start Date"),
        ("D-3", "flow3.pretrip_d3", -3, "Start Date"),
        ("D-1", "flow3.pretrip_d1", -1, "Start Date"),
        ("Day-of", "flow3.day_of", 0, "Start Date"),
        ("Post-trip", "flow6.post_trip", 1, "End Date"),
    ]

    languages = ["AR", "EN"]
    row_idx = 2
    count = 0
    for trip_row in meaningful_trip_rows(trips_ws):
        trip_id = normalize_text(trips_ws.cell(trip_row, trip_headers["Trip ID"]).value)
        trip_name = normalize_text(trips_ws.cell(trip_row, trip_headers["Trip Name"]).value)
        if not trip_id or not trip_name:
            continue
        for language in languages:
            for drip_type, message_key, offset_days, date_source in drip_events:
                ws.cell(row_idx, 1).value = f"{trip_id}-{language}-{drip_type.replace(' ', '').replace('/', '').lower()}"
                ws.cell(row_idx, 2).value = trip_id
                ws.cell(row_idx, 3).value = trip_name
                ws.cell(row_idx, 4).value = trip_row
                ws.cell(row_idx, 5).value = language
                ws.cell(row_idx, 6).value = drip_type
                if date_source == "Start Date":
                    ws.cell(row_idx, 7).value = date_formula(TRIPS_SHEET, trip_row, offset_days, start_col)
                else:
                    ws.cell(row_idx, 7).value = date_formula(TRIPS_SHEET, trip_row, offset_days, end_col)
                ws.cell(row_idx, 8).value = message_key
                ws.cell(row_idx, 9).value = f"{TRIPS_SHEET}!{date_source}"
                ws.cell(row_idx, 10).value = "Pending"
                ws.cell(row_idx, 11).value = "Automation"
                ws.cell(row_idx, 12).value = "Auto-generated from trip dates."
                row_idx += 1
                count += 1

    style_header_row(ws, 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    auto_width(ws)
    add_language_validation(ws, 5, 2, ws.max_row)
    add_status_validation(ws, 10, 2, ws.max_row)
    return count


def add_language_validation(ws, col_idx: int, start_row: int, end_row: int) -> None:
    if end_row < start_row:
        return
    dv = DataValidation(type="list", formula1='"AR,EN,Both"', allow_blank=True)
    col = get_column_letter(col_idx)
    dv.add(f"{col}{start_row}:{col}{end_row}")
    ws.add_data_validation(dv)


def add_status_validation(ws, col_idx: int, start_row: int, end_row: int) -> None:
    if end_row < start_row:
        return
    dv = DataValidation(type="list", formula1='"Draft,Active,Paused,Done,Needs Review,Pending"', allow_blank=True)
    col = get_column_letter(col_idx)
    dv.add(f"{col}{start_row}:{col}{end_row}")
    ws.add_data_validation(dv)


def apply_trip_helpers(wb) -> list[str]:
    if TRIPS_SHEET not in wb.sheetnames:
        return []

    trips_ws = wb[TRIPS_SHEET]
    appended = ensure_or_append_headers(
        trips_ws,
        header_row=2,
        base_headers=TRIPS_BASE_HEADERS,
        new_headers=TRIPS_EXTENSION_HEADERS,
    )
    trip_headers = sheet_headers(trips_ws, 2)
    start_col = get_column_letter(trip_headers["Start Date"])
    end_col = get_column_letter(trip_headers["End Date"])
    status_col = trip_headers["Sales Status"]
    window_col = trip_headers["Trip Window Status"]
    note_col = trip_headers["Trip Availability Note"]
    reengage_col = trip_headers["Next Reengage Date"]

    for row_idx in meaningful_trip_rows(trips_ws):
        trips_ws.cell(row_idx, window_col).value = (
            f'=IF(OR({start_col}{row_idx}="",{end_col}{row_idx}=""),"Dates Missing",'
            f'IF({end_col}{row_idx}<TODAY(),"Expired",IF({start_col}{row_idx}>TODAY(),"Future","Live")))'
        )
        trips_ws.cell(row_idx, note_col).value = (
            f'=IF(OR({start_col}{row_idx}="",{end_col}{row_idx}=""),"Need dates from sheet",'
            f'IF({end_col}{row_idx}<{start_col}{row_idx},"End date before start date","OK"))'
        )
        trips_ws.cell(row_idx, reengage_col).value = (
            f'=IF({end_col}{row_idx}="","",{end_col}{row_idx}+2)'
        )
        if not normalize_text(trips_ws.cell(row_idx, status_col).value):
            trips_ws.cell(row_idx, status_col).value = "Date TBD"

    style_header_row(trips_ws, 2)
    trips_ws.freeze_panes = "A3"
    trips_ws.auto_filter.ref = f"A2:{get_column_letter(trips_ws.max_column)}{trips_ws.max_row}"
    auto_width(trips_ws)
    return appended


def apply_traveler_helpers(wb) -> list[str]:
    if TRAVELERS_SHEET not in wb.sheetnames:
        return []

    travelers_ws = wb[TRAVELERS_SHEET]
    appended = ensure_or_append_headers(
        travelers_ws,
        header_row=1,
        base_headers=TRAVELERS_BASE_HEADERS,
        new_headers=TRAVELERS_EXTENSION_HEADERS,
    )
    style_header_row(travelers_ws, 1)
    travelers_ws.freeze_panes = "A2"
    travelers_ws.auto_filter.ref = f"A1:{get_column_letter(travelers_ws.max_column)}{travelers_ws.max_row}"
    auto_width(travelers_ws)
    return appended


def apply_lead_helpers(wb) -> list[str]:
    leads_ws, _ = ensure_leads_sheet(wb)
    appended = ensure_or_append_headers(
        leads_ws,
        header_row=1,
        base_headers=LEADS_BASE_HEADERS,
        new_headers=LEADS_EXTENSION_HEADERS,
    )
    style_header_row(leads_ws, 1)
    leads_ws.freeze_panes = "A2"
    leads_ws.auto_filter.ref = f"A1:{get_column_letter(leads_ws.max_column)}{leads_ws.max_row}"
    auto_width(leads_ws)
    return appended


def apply_interaction_helpers(wb) -> list[str]:
    if INTERACTIONS_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(INTERACTIONS_SHEET)
        for col_idx, header in enumerate(INTERACTIONS_BASE_HEADERS, start=1):
            ws.cell(1, col_idx).value = header
    else:
        ws = wb[INTERACTIONS_SHEET]
    appended = ensure_or_append_headers(
        ws,
        header_row=1,
        base_headers=INTERACTIONS_BASE_HEADERS,
        new_headers=INTERACTIONS_EXTENSION_HEADERS,
    )
    style_header_row(ws, 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    auto_width(ws)
    return appended


def build_state_sheets(wb) -> dict[str, int]:
    ensure_sheet(wb, FLOW_STATES_SHEET, FLOW_STATES_HEADERS, clear_existing=True)
    ensure_sheet(wb, DM_COPY_LIBRARY_SHEET, DM_COPY_LIBRARY_HEADERS, clear_existing=True)
    ensure_sheet(wb, LANGUAGE_TEMPLATES_SHEET, LANGUAGE_TEMPLATE_HEADERS, clear_existing=True)
    ensure_sheet(wb, WAITLIST_SHEET, WAITLIST_HEADERS, clear_existing=True)
    ensure_sheet(wb, FOLLOW_UPS_SHEET, FOLLOW_UP_HEADERS, clear_existing=True)
    ensure_sheet(wb, HANDOFF_QUEUE_SHEET, HANDOFF_HEADERS, clear_existing=True)

    flow_ws = wb[FLOW_STATES_SHEET]
    for row_idx, row in enumerate(FLOW_STATE_ROWS, start=2):
        for col_idx, header in enumerate(FLOW_STATES_HEADERS, start=1):
            flow_ws.cell(row_idx, col_idx).value = row.get(header)

    copy_ws = wb[DM_COPY_LIBRARY_SHEET]
    for row_idx, row in enumerate(DM_COPY_ROWS, start=2):
        for col_idx, header in enumerate(DM_COPY_LIBRARY_HEADERS, start=1):
            copy_ws.cell(row_idx, col_idx).value = row.get(header)

    template_ws = wb[LANGUAGE_TEMPLATES_SHEET]
    for row_idx, row in enumerate(LANGUAGE_TEMPLATE_ROWS, start=2):
        for col_idx, header in enumerate(LANGUAGE_TEMPLATE_HEADERS, start=1):
            template_ws.cell(row_idx, col_idx).value = row.get(header)

    for title in [FLOW_STATES_SHEET, DM_COPY_LIBRARY_SHEET, LANGUAGE_TEMPLATES_SHEET, WAITLIST_SHEET, FOLLOW_UPS_SHEET, HANDOFF_QUEUE_SHEET]:
        ws = wb[title]
        style_header_row(ws, 1)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        auto_width(ws)

    add_language_validation(flow_ws, 9, 2, flow_ws.max_row)
    add_status_validation(flow_ws, 11, 2, flow_ws.max_row)
    add_status_validation(copy_ws, 11, 2, copy_ws.max_row)
    add_language_validation(template_ws, 2, 2, template_ws.max_row)
    add_status_validation(template_ws, 6, 2, template_ws.max_row)
    add_language_validation(wb[WAITLIST_SHEET], 7, 2, wb[WAITLIST_SHEET].max_row)
    add_status_validation(wb[WAITLIST_SHEET], 13, 2, wb[WAITLIST_SHEET].max_row)
    add_language_validation(wb[FOLLOW_UPS_SHEET], 10, 2, wb[FOLLOW_UPS_SHEET].max_row)
    add_status_validation(wb[FOLLOW_UPS_SHEET], 12, 2, wb[FOLLOW_UPS_SHEET].max_row)
    add_status_validation(wb[HANDOFF_QUEUE_SHEET], 10, 2, wb[HANDOFF_QUEUE_SHEET].max_row)

    return {
        FLOW_STATES_SHEET: len(FLOW_STATE_ROWS),
        DM_COPY_LIBRARY_SHEET: len(DM_COPY_ROWS),
        LANGUAGE_TEMPLATES_SHEET: len(LANGUAGE_TEMPLATE_ROWS),
    }


def align_workbook(input_path: Path, output_path: Path, runtime_copy_path: Path | None = None) -> WorkbookChangeSummary:
    wb = load_workbook(input_path)
    added_sheets: list[str] = []
    appended_columns: dict[str, list[str]] = {}

    for sheet_name in [
        FLOW_STATES_SHEET,
        DM_COPY_LIBRARY_SHEET,
        LANGUAGE_TEMPLATES_SHEET,
        WAITLIST_SHEET,
        FOLLOW_UPS_SHEET,
        TRIP_DRIPS_SHEET,
        HANDOFF_QUEUE_SHEET,
    ]:
        if sheet_name not in wb.sheetnames:
            added_sheets.append(sheet_name)

    seeded_rows = build_state_sheets(wb)

    traveler_appended = apply_traveler_helpers(wb)
    if traveler_appended:
        appended_columns[TRAVELERS_SHEET] = traveler_appended

    trip_appended = apply_trip_helpers(wb)
    if trip_appended:
        appended_columns[TRIPS_SHEET] = trip_appended

    lead_appended = apply_lead_helpers(wb)
    if lead_appended:
        appended_columns[LEADS_SHEET] = lead_appended

    interaction_appended = apply_interaction_helpers(wb)
    if interaction_appended:
        appended_columns[INTERACTIONS_SHEET] = interaction_appended

    seeded_rows[TRIP_DRIPS_SHEET] = seed_trip_drips(wb)

    wb.save(output_path)
    wb.close()

    if runtime_copy_path is not None:
        runtime_copy_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(output_path, runtime_copy_path)

    return WorkbookChangeSummary(
        output_workbook=str(output_path),
        runtime_copy=str(runtime_copy_path) if runtime_copy_path else "",
        added_sheets=added_sheets,
        appended_columns=appended_columns,
        seeded_rows=seeded_rows,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Align the Rahma Traveler workbook flows and control sheets.")
    parser.add_argument("--workbook", type=Path, default=SOURCE_WORKBOOK, help="Workbook to update in place.")
    parser.add_argument("--runtime-copy", type=Path, default=RUNTIME_WORKBOOK, help="Optional runtime workbook copy to refresh.")
    parser.add_argument("--no-runtime-copy", action="store_true", help="Skip copying the aligned workbook to the runtime workbook.")
    args = parser.parse_args()

    runtime_copy = None if args.no_runtime_copy else args.runtime_copy
    result = align_workbook(args.workbook, args.workbook, runtime_copy)
    print(json.dumps(result.__dict__, indent=2, ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    main()
