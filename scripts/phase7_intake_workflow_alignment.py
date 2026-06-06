from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.phase0_cleanup import ensure_headers
from scripts.phase2_controlled_agent import normalize_text
from scripts.phase4_sales_intelligence import ensure_leads_sheet


TRAVELER_HELPERS = ["Birthday", "Gender", "Nationality"]

COPY_ROWS = [
    {
        "Message Key": "session.intake_start",
        "English Copy": "Hello. I am Rahma Traveler's sales agent. Please complete the intake form so I can check the CRM safely.",
        "Arabic Copy": "اهلا، انا مساعد مبيعات رحمة ترافيل. من فضلك املأ بيانات التعارف حتى اتأكد من بيانات العميل في ال CRM بشكل صحيح.",
        "Flow Key": "shared_booking",
        "Step Key": "intake_form",
        "Active": "Yes",
    },
    {
        "Message Key": "session.preview_open_trips",
        "English Copy": "Here are upcoming options: {trip_lines}\nReply yes to save this as a lead, or no to stop.",
        "Arabic Copy": "دي الرحلات المتاحة القادمة: {trip_lines}\nرد بنعم لحفظها ك lead، او لا لايقاف المتابعة.",
        "Flow Key": "shared_booking",
        "Step Key": "trip_preview",
        "Active": "Yes",
    },
    {
        "Message Key": "session.preview_date_tbd",
        "English Copy": "I found trips in this category, but the dates are not confirmed yet: {trip_names}. Reply yes to save this lead for follow-up, or no to stop.",
        "Arabic Copy": "لقيت رحلات في نفس النوع لكن مواعيدها لسه مش مؤكدة: {trip_names}. رد بنعم لحفظ ال lead للمتابعة، او لا للايقاف.",
        "Flow Key": "shared_booking",
        "Step Key": "trip_preview",
        "Active": "Yes",
    },
]


def ensure_copy_library(wb) -> int:
    if "DM Copy Library" not in wb.sheetnames:
        ws = wb.create_sheet("DM Copy Library")
        headers = ["Message Key", "Flow Key", "Step Key", "English Copy", "Arabic Copy", "Active"]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(1, col_idx).value = header
    else:
        ws = wb["DM Copy Library"]

    headers = {
        normalize_text(ws.cell(1, col).value): col
        for col in range(1, ws.max_column + 1)
        if normalize_text(ws.cell(1, col).value)
    }
    for header in ["Message Key", "Flow Key", "Step Key", "English Copy", "Arabic Copy", "Active"]:
        if header not in headers:
            col_idx = ws.max_column + 1
            ws.cell(1, col_idx).value = header
            headers[header] = col_idx

    existing_rows = {
        normalize_text(ws.cell(row_idx, headers["Message Key"]).value): row_idx
        for row_idx in range(2, ws.max_row + 1)
        if normalize_text(ws.cell(row_idx, headers["Message Key"]).value)
    }

    writes = 0
    for row_data in COPY_ROWS:
        row_idx = existing_rows.get(row_data["Message Key"])
        if row_idx is None:
            row_idx = ws.max_row + 1
        for header, value in row_data.items():
            ws.cell(row_idx, headers[header]).value = value
        writes += 1
    return writes


def align_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path)
    if "Travelers" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{path} has no Travelers sheet")

    travelers_ws = wb["Travelers"]
    traveler_headers = ensure_headers(
        travelers_ws,
        header_row=1,
        expected_headers=["Traveler ID", "Full Name", "Code", "WhatsApp", "Integrated WhatsApp", "Phone Lookup Key"],
        helper_headers=TRAVELER_HELPERS,
    )
    _leads_ws, lead_headers = ensure_leads_sheet(wb)
    copy_rows = ensure_copy_library(wb)

    wb.save(path)
    wb.close()
    return {
        "workbook": str(path),
        "travelerProfileColumnsReady": all(header in traveler_headers for header in TRAVELER_HELPERS),
        "leadProfileColumnsReady": all(header in lead_headers for header in TRAVELER_HELPERS),
        "copyRowsWritten": copy_rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Align workbook sheets with the form-first intake workflow.")
    parser.add_argument("workbooks", nargs="+", type=Path)
    args = parser.parse_args()

    for workbook in args.workbooks:
        result = align_workbook(workbook)
        print(result)


if __name__ == "__main__":
    main()
