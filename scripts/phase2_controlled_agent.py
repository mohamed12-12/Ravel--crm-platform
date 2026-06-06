from __future__ import annotations

import argparse
import json
import sys
from copy import copy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.phase0_cleanup import (  # noqa: E402
    INTERACTIONS_SHEET,
    TODAY,
    canonicalize_status,
    ensure_headers,
    normalize_phone_fields,
)
from scripts.phase1_readonly_agent import (  # noqa: E402
    TRAVELERS_EXPECTED_HEADERS,
    TRIPS_EXPECTED_HEADERS,
    build_agent_response,
    build_agent_response_from_wb,
)
from scripts.phase4_sales_intelligence import (  # noqa: E402
    ensure_leads_sheet,
    upsert_lead,
)


TRAVELERS_SHEET = "Travelers"
TRIPS_SHEET = "Trips"

WRITEABLE_TRAVELER_FIELDS = [
    "Status",
    "Traveler ID",
    "Full Name",
    "Birthday",
    "Gender",
    "Nationality",
    "Code",
    "WhatsApp",
    "Integrated WhatsApp",
    "Normalized WhatsApp",
    "Phone Lookup Key",
    "Lead Source",
    "Created At",
    "Last Contacted At",
    "Agent Notes",
    "Data Audit",
]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clone_cell_style(source, target) -> None:
    if not all(hasattr(source, attr) and hasattr(target, attr) for attr in ("font", "fill", "border", "alignment", "number_format", "protection")):
        return
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def load_write_workbook(workbook_path: Path):
    wb = load_workbook(workbook_path, data_only=False)
    travelers_ws = wb[TRAVELERS_SHEET]
    trips_ws = wb[TRIPS_SHEET]
    interactions_ws = wb[INTERACTIONS_SHEET]
    leads_ws, lead_headers = ensure_leads_sheet(wb)
    traveler_headers = ensure_headers(
        travelers_ws,
        header_row=1,
        expected_headers=TRAVELERS_EXPECTED_HEADERS,
        helper_headers=[
            "Birthday",
            "Gender",
            "Nationality",
            "Lead Source",
            "Created At",
            "Last Contacted At",
            "Agent Notes",
            "Normalized WhatsApp",
        ],
    )
    trip_headers = ensure_headers(
        trips_ws,
        header_row=2,
        expected_headers=TRIPS_EXPECTED_HEADERS,
        helper_headers=[],
    )
    interaction_headers = {
        normalize_text(interactions_ws.cell(1, col).value): col
        for col in range(1, interactions_ws.max_column + 1)
        if normalize_text(interactions_ws.cell(1, col).value)
    }
    return wb, travelers_ws, trips_ws, interactions_ws, leads_ws, traveler_headers, trip_headers, interaction_headers, lead_headers


def find_first_blank_traveler_row(ws, header_map: dict[str, int]) -> int:
    # A row is safe to overwrite only if ALL key identity and contact fields are empty.
    # This prevents overwriting partially filled historical rows or rows with existing notes.
    key_headers = [
        "Traveler ID",
        "Full Name",
        "Code",
        "WhatsApp",
        "Integrated WhatsApp",
        "Phone Lookup Key",
        "Email",
        "Notes",
    ]
    # Filter to only headers that actually exist in the sheet
    active_keys = [h for h in key_headers if h in header_map]

    for row_idx in range(2, ws.max_row + 1):
        values = [ws.cell(row_idx, header_map[header]).value for header in active_keys]
        if not any(value not in (None, "") for value in values):
            return row_idx
    return ws.max_row + 1


def translate_formula(formula: str, origin_cell: str, target_cell: str) -> str:
    return Translator(formula, origin=origin_cell).translate_formula(dest=target_cell)


def prepare_traveler_row(ws, header_map: dict[str, int], row_idx: int) -> None:
    if row_idx <= ws.max_row:
        existing_id = ws.cell(row_idx, header_map["Traveler ID"]).value
        existing_name = ws.cell(row_idx, header_map["Full Name"]).value
        if not existing_id and not existing_name:
            return

    template_row = row_idx - 1
    if template_row < 2:
        template_row = 2

    for col_idx in range(1, ws.max_column + 1):
        source = ws.cell(template_row, col_idx)
        target = ws.cell(row_idx, col_idx)
        clone_cell_style(source, target)
        if hasattr(source, "number_format") and hasattr(target, "number_format"):
            target.number_format = source.number_format
        if isinstance(source.value, str) and source.value.startswith("="):
            target.value = translate_formula(source.value, source.coordinate, target.coordinate)
        elif col_idx not in {header_map[h] for h in WRITEABLE_TRAVELER_FIELDS if h in header_map}:
            target.value = None


def generate_next_traveler_id(ws, header_map: dict[str, int]) -> str:
    max_number = 0
    id_col = header_map["Traveler ID"]
    existing_ids: set[str] = set()
    for row_idx in range(2, ws.max_row + 1):
        traveler_id = normalize_text(ws.cell(row_idx, id_col).value)
        if traveler_id:
            existing_ids.add(traveler_id)
        if traveler_id.startswith("TR") and traveler_id[2:].isdigit():
            max_number = max(max_number, int(traveler_id[2:]))

    next_number = max_number + 1
    next_id = f"TR{next_number:05d}"
    while next_id in existing_ids:
        next_number += 1
        next_id = f"TR{next_number:05d}"
    return next_id


def derive_traveler_audit(flags: list[str]) -> str:
    keep = [flag for flag in flags if flag not in {"PHONE_INCLUDED_COUNTRY_CODE", "LEADING_ZERO_REMOVED", "CODE_DERIVED_FROM_PHONE", "CODE_DERIVED_FROM_EGYPT_LOCAL"}]
    return "; ".join(sorted(set(keep)))


def create_traveler(
    ws,
    header_map: dict[str, int],
    full_name: str,
    raw_phone: str,
    country_code: str,
    source: str,
    timestamp: str,
    birthday: str = "",
    gender: str = "",
    nationality: str = "",
    agent_notes: str = "",
) -> dict[str, Any]:
    normalized = normalize_phone_fields(country_code, raw_phone)
    row_idx = find_first_blank_traveler_row(ws, header_map)
    prepare_traveler_row(ws, header_map, row_idx)
    traveler_id = generate_next_traveler_id(ws, header_map)

    ws.cell(row_idx, header_map["Status"]).value = None
    ws.cell(row_idx, header_map["Traveler ID"]).value = traveler_id
    ws.cell(row_idx, header_map["Full Name"]).value = full_name.strip()
    if "Birthday" in header_map:
        ws.cell(row_idx, header_map["Birthday"]).value = birthday or None
    if "Gender" in header_map:
        ws.cell(row_idx, header_map["Gender"]).value = gender or None
    if "Nationality" in header_map:
        ws.cell(row_idx, header_map["Nationality"]).value = nationality or None
    ws.cell(row_idx, header_map["Code"]).value = normalized.code or None
    ws.cell(row_idx, header_map["WhatsApp"]).value = normalized.local_number or None
    ws.cell(row_idx, header_map["Integrated WhatsApp"]).value = normalized.normalized_whatsapp or None
    ws.cell(row_idx, header_map["Normalized WhatsApp"]).value = normalized.normalized_whatsapp or None
    ws.cell(row_idx, header_map["Phone Lookup Key"]).value = normalized.lookup_key or None
    ws.cell(row_idx, header_map["Lead Source"]).value = source
    ws.cell(row_idx, header_map["Created At"]).value = timestamp
    ws.cell(row_idx, header_map["Last Contacted At"]).value = timestamp
    ws.cell(row_idx, header_map["Agent Notes"]).value = agent_notes or "Created by Phase 2 controlled write"
    ws.cell(row_idx, header_map["Data Audit"]).value = derive_traveler_audit(normalized.flags) or None

    for header in ("Code", "WhatsApp", "Integrated WhatsApp", "Normalized WhatsApp", "Phone Lookup Key", "Created At", "Last Contacted At"):
        cell = ws.cell(row_idx, header_map[header])
        if hasattr(cell, "number_format"):
            cell.number_format = "@"

    return {
        "row": row_idx,
        "traveler_id": traveler_id,
        "full_name": full_name.strip(),
        "birthday": birthday,
        "gender": gender,
        "nationality": nationality,
        "code": normalized.code,
        "whatsapp": normalized.local_number,
        "integrated_whatsapp": normalized.normalized_whatsapp,
        "phone_lookup_key": normalized.lookup_key,
        "data_audit": derive_traveler_audit(normalized.flags),
    }


def update_existing_traveler_profile(
    ws,
    header_map: dict[str, int],
    row_idx: int,
    *,
    birthday: str = "",
    gender: str = "",
    nationality: str = "",
    timestamp: str,
) -> dict[str, Any]:
    updates: dict[str, Any] = {}
    for header, value in {
        "Birthday": birthday,
        "Gender": gender,
        "Nationality": nationality,
    }.items():
        if not value or header not in header_map:
            continue
        current = normalize_text(ws.cell(row_idx, header_map[header]).value)
        if current:
            continue
        ws.cell(row_idx, header_map[header]).value = value
        updates[header] = value

    if "Last Contacted At" in header_map:
        ws.cell(row_idx, header_map["Last Contacted At"]).value = timestamp
        updates["Last Contacted At"] = timestamp

    return updates


def next_interaction_row(ws) -> int:
    for row_idx in range(2, ws.max_row + 1):
        if not any(ws.cell(row_idx, col).value not in (None, "") for col in range(1, ws.max_column + 1)):
            return row_idx
    return ws.max_row + 1


def generate_interaction_id(ws, header_map: dict[str, int], timestamp: datetime) -> str:
    col = header_map["Interaction ID"]
    max_seq = 0
    prefix = timestamp.strftime("INT-%Y%m%d-")
    for row_idx in range(2, ws.max_row + 1):
        value = normalize_text(ws.cell(row_idx, col).value)
        if value.startswith(prefix):
            tail = value[len(prefix) :]
            if tail.isdigit():
                max_seq = max(max_seq, int(tail))
    return f"{prefix}{max_seq + 1:04d}"


def log_interaction(
    ws,
    header_map: dict[str, int],
    *,
    timestamp: datetime,
    channel: str,
    customer_name: str,
    raw_phone: str,
    integrated_whatsapp: str,
    phone_lookup_key: str,
    traveler_id: str,
    matched_row: int | None,
    status_snapshot: str,
    intent: str,
    trip_type: str,
    suggested_trips: list[str],
    action_taken: str,
    handoff_required: bool,
    handoff_reason: str,
    agent_notes: str,
) -> dict[str, Any]:
    row_idx = next_interaction_row(ws)
    interaction_id = generate_interaction_id(ws, header_map, timestamp)
    values = {
        "Interaction ID": interaction_id,
        "Timestamp": timestamp.isoformat(timespec="seconds"),
        "Channel": channel,
        "Customer Name": customer_name,
        "Raw Phone": raw_phone,
        "Integrated WhatsApp": integrated_whatsapp or None,
        "Phone Lookup Key": phone_lookup_key or None,
        "Traveler ID": traveler_id or None,
        "Matched Row": matched_row,
        "Status Snapshot": status_snapshot or None,
        "Intent": intent or None,
        "Trip Type": trip_type or None,
        "Suggested Trips": ", ".join(suggested_trips) or None,
        "Action Taken": action_taken,
        "Handoff Required": "Yes" if handoff_required else "No",
        "Handoff Reason": handoff_reason or None,
        "Agent Notes": agent_notes or None,
    }
    for header, value in values.items():
        ws.cell(row_idx, header_map[header]).value = value
    return {"interaction_id": interaction_id, "row": row_idx}


def run_phase2(
    workbook_path: Path,
    output_path: Path,
    *,
    full_name: str,
    raw_phone: str,
    country_code: str,
    trip_type: str | None,
    channel: str,
    source: str,
    agent_notes: str,
    birthday: str = "",
    gender: str = "",
    nationality: str = "",
    preferred_trip_id: str = "",
    trip_result_override: dict[str, list[dict[str, Any]]] | None = None,
) -> dict[str, Any]:
    wb, travelers_ws, trips_ws, interactions_ws, leads_ws, traveler_headers, trip_headers, interaction_headers, lead_headers = load_write_workbook(workbook_path)
    
    # We pass the real workbook objects to the _from_wb function
    preview, mutations = run_phase2_from_wb(
        wb=wb,
        travelers_ws=travelers_ws,
        trips_ws=trips_ws,
        interactions_ws=interactions_ws,
        leads_ws=leads_ws,
        traveler_headers=traveler_headers,
        trip_headers=trip_headers,
        interaction_headers=interaction_headers,
        lead_headers=lead_headers,
        full_name=full_name,
        raw_phone=raw_phone,
        country_code=country_code,
        trip_type=trip_type,
        channel=channel,
        source=source,
        agent_notes=agent_notes,
        birthday=birthday,
        gender=gender,
        nationality=nationality,
        preferred_trip_id=preferred_trip_id,
        trip_result_override=trip_result_override,
    )
    
    wb.save(output_path)
    wb.close()
    
    preview["write_result"]["output_workbook"] = str(output_path)
    return preview


def run_phase2_from_wb(
    wb: Any,
    travelers_ws: Any,
    trips_ws: Any,
    interactions_ws: Any,
    leads_ws: Any,
    traveler_headers: dict[str, int],
    trip_headers: dict[str, int],
    interaction_headers: dict[str, int],
    lead_headers: dict[str, int],
    *,
    full_name: str,
    raw_phone: str,
    country_code: str,
    trip_type: str | None,
    channel: str,
    source: str,
    agent_notes: str,
    birthday: str = "",
    gender: str = "",
    nationality: str = "",
    preferred_trip_id: str = "",
    trip_result_override: dict[str, list[dict[str, Any]]] | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    # Use the workbook-compatible preview function
    preview = build_agent_response_from_wb(
        wb,
        full_name,
        raw_phone,
        trip_type,
        country_code,
        trip_result_override=trip_result_override,
    )
    timestamp = datetime.now()

    created_traveler: dict[str, Any] | None = None
    traveler_id = ""
    matched_row: int | None = None
    status_snapshot = ""

    if preview["match_status"] == "not_found":
        created_traveler = create_traveler(
            travelers_ws,
            traveler_headers,
            full_name=full_name,
            raw_phone=raw_phone,
            country_code=country_code,
            source=source,
            timestamp=timestamp.isoformat(timespec="seconds"),
            birthday=birthday,
            gender=gender,
            nationality=nationality,
            agent_notes=agent_notes,
        )
        traveler_id = created_traveler["traveler_id"]
        matched_row = created_traveler["row"]
        status_snapshot = ""
        preview["actions"].append("created_new_traveler")
        preview["traveler"] = created_traveler
    elif preview["match_status"] == "single_match":
        traveler = preview["traveler"]
        traveler_id = traveler["traveler_id"]
        matched_row = traveler["row"]
        status_snapshot = traveler["status"]
        if not preview["handoff_required"]:
            profile_updates = update_existing_traveler_profile(
                travelers_ws,
                traveler_headers,
                matched_row,
                birthday=birthday,
                gender=gender,
                nationality=nationality,
                timestamp=timestamp.isoformat(timespec="seconds"),
            )
            if profile_updates:
                traveler.update({key.lower().replace(" ", "_"): value for key, value in profile_updates.items()})
                preview["actions"].append("updated_missing_traveler_profile")
    elif preview["match_status"] == "multiple_matches":
        matched_row = None
        status_snapshot = "MULTIPLE_MATCHES"

    lookup_phone = preview["lookup_phone"]
    suggested = preview["trip_result"] or {"open_trips": [], "date_tbd_trips": []}
    suggested_trip_ids = [item["trip_id"] for item in suggested.get("open_trips", [])]
    if not suggested_trip_ids:
        suggested_trip_ids = [item["trip_id"] for item in suggested.get("date_tbd_trips", [])]

    interaction_log = log_interaction(
        interactions_ws,
        interaction_headers,
        timestamp=timestamp,
        channel=channel,
        customer_name=full_name,
        raw_phone=raw_phone,
        integrated_whatsapp=lookup_phone.get("normalized_whatsapp", ""),
        phone_lookup_key=lookup_phone.get("lookup_key", ""),
        traveler_id=traveler_id,
        matched_row=matched_row,
        status_snapshot=status_snapshot,
        intent="sales_inquiry",
        trip_type=trip_type or "",
        suggested_trips=suggested_trip_ids,
        action_taken=", ".join(preview["actions"]),
        handoff_required=preview["handoff_required"],
        handoff_reason=preview["handoff_reason"],
        agent_notes=agent_notes,
    )

    lead_update = upsert_lead(
        leads_ws,
        lead_headers,
        result=preview,
        customer_name=full_name,
        raw_phone=raw_phone,
        phone_lookup_key=lookup_phone.get("lookup_key", ""),
        integrated_whatsapp=lookup_phone.get("normalized_whatsapp", ""),
        traveler_id=traveler_id,
        channel=channel,
        source=source,
        trip_type=trip_type or "",
        interaction_id=interaction_log["interaction_id"],
        timestamp=timestamp,
        notes=agent_notes,
        birthday=birthday,
        gender=gender,
        nationality=nationality,
        preferred_trip_id=preferred_trip_id,
    )

    preview["write_result"] = {
        "created_traveler": created_traveler,
        "interaction_log": interaction_log,
        "lead_update": lead_update,
    }
    
    # Return both the preview result and the mutations list tracked in the workbook wrapper
    # If it's a real openpyxl workbook, .mutations won't exist, which is fine
    mutations = getattr(wb, "mutations", [])
    return preview, mutations


def main() -> None:
    parser = argparse.ArgumentParser(description="Phase 2 controlled-write sales agent prototype.")
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--name", required=True)
    parser.add_argument("--phone", required=True)
    parser.add_argument("--country-code", default="")
    parser.add_argument("--trip-type", default=None)
    parser.add_argument("--channel", default="manual-test")
    parser.add_argument("--source", default="DM")
    parser.add_argument("--agent-notes", default="")
    args = parser.parse_args()

    result = run_phase2(
        workbook_path=args.workbook,
        output_path=args.output,
        full_name=args.name,
        raw_phone=args.phone,
        country_code=args.country_code,
        trip_type=args.trip_type,
        channel=args.channel,
        source=args.source,
        agent_notes=args.agent_notes,
    )
    print(json.dumps(result, indent=2, ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    main()
