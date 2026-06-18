from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import or_

REPO_ROOT = Path(__file__).resolve().parents[1]
API_ROOT = REPO_ROOT / "apps" / "api"
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from services.crm.system_services.phone_normalization import normalize_phone_input  # noqa: E402


SHEET_HEADER_ROWS = {
    "Travelers": 1,
    "Trips": 2,
    "Trip Bookings": 2,
    "Leads": 1,
}

BOOKING_STATUSES = {
    "Draft",
    "Waiting Customer",
    "Pending Confirmation",
    "Confirmed",
    "Payment Pending",
    "Paid",
    "Completed",
    "Cancelled",
}


@dataclass
class SheetSummary:
    scanned: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    quarantined: int = 0
    errors: int = 0


@dataclass
class MigrationSummary:
    dry_run: bool
    workbook: str
    database_uri: str
    scope: str = "full"
    sheets: dict[str, SheetSummary] = field(default_factory=dict)
    travelers_imported: int = 0
    travelers_merged: int = 0
    travelers_quarantined: int = 0
    trips_imported: int = 0
    trips_quarantined: int = 0
    bookings_matched: int = 0
    bookings_quarantined: int = 0
    booking_lifecycle_counts: dict[str, int] = field(default_factory=dict)
    validation: dict[str, Any] = field(default_factory=dict)
    quarantine_count: int = 0
    errors: list[str] = field(default_factory=list)

    def sheet(self, name: str) -> SheetSummary:
        return self.sheets.setdefault(name, SheetSummary())


@dataclass
class TravelerDecision:
    group_id: str
    classification: str
    decision: str
    primary_row: int | None
    secondary_rows: list[int]
    traveler_id_to_keep: str
    new_traveler_id_if_needed: str
    action: str
    reason: str


@dataclass
class BookingCandidate:
    booking_id: str
    trip_id: str
    trip_name: str
    traveler_id: str
    traveler_name: str
    room_type: str | None
    currency: str
    booking_status: str
    payment_status: str


class ExcelToCRMMigrator:
    def __init__(
        self,
        workbook_path: Path,
        *,
        dry_run: bool = True,
        scope: str = "full",
        decisions_path: Path | None = None,
    ) -> None:
        self.workbook_path = workbook_path
        self.dry_run = dry_run
        self.scope = scope
        self.decisions_path = decisions_path or (REPO_ROOT / "MIGRATION_DECISIONS.csv")
        self.quarantine: list[dict[str, Any]] = []
        self.summary = MigrationSummary(
            dry_run=dry_run,
            workbook=str(workbook_path),
            database_uri=os.environ.get("DATABASE_URL", ""),
            scope=scope,
        )
        self._planned_traveler_ids: set[str] = set()
        self._planned_trip_ids: set[str] = set()
        self._planned_booking_ids: set[str] = set()
        self._traveler_id_counts: Counter[str] = Counter()
        self._trip_id_counts: Counter[str] = Counter()
        self._booking_id_counts: Counter[str] = Counter()
        self._phone_counts: Counter[str] = Counter()
        self._handled_traveler_rows: set[int] = set()
        self._trip_ids_missing_split: set[str] = set()
        self._trip_end_dates: dict[str, date | None] = {}

    def run(self) -> MigrationSummary:
        apply_database_url_to_app_config()
        from app import create_app
        from app.extensions import db

        if not self.workbook_path.exists():
            raise FileNotFoundError(f"Workbook not found: {self.workbook_path}")

        self.wb = load_workbook(self.workbook_path, data_only=True, read_only=True)
        self._precompute_duplicate_keys()
        self._traveler_decisions = self._load_traveler_decisions()

        app = create_app()
        with app.app_context():
            self.db = db
            db.create_all()
            self._prime_existing_keys()
            if self.scope == "travelers":
                self._migrate_travelers_only()
            elif self.scope == "trips":
                self._migrate_trips_only()
            elif self.scope == "bookings":
                self._migrate_bookings_only()
            elif self.scope == "bookings_preview":
                self._preview_bookings_only()
            else:
                self._migrate_travelers()
                self._migrate_trips()
                self._migrate_bookings()
                self._migrate_leads()
            self.summary.quarantine_count = len(self.quarantine)
            self.summary.travelers_imported = self.summary.sheets.get("Travelers", SheetSummary()).created + self.summary.sheets.get("Travelers", SheetSummary()).updated
            self.summary.travelers_quarantined = self.summary.sheets.get("Travelers", SheetSummary()).quarantined
            self.summary.trips_imported = self.summary.sheets.get("Trips", SheetSummary()).created + self.summary.sheets.get("Trips", SheetSummary()).updated
            self.summary.trips_quarantined = self.summary.sheets.get("Trips", SheetSummary()).quarantined
            if self.dry_run:
                db.session.rollback()
            else:
                db.session.commit()
        return self.summary

    def _prime_existing_keys(self) -> None:
        from app.models import Trip, TripBooking, Traveler

        self._planned_traveler_ids = {row[0] for row in self.db.session.query(Traveler.traveler_id).all()}
        self._planned_trip_ids = {row[0] for row in self.db.session.query(Trip.trip_id).all()}
        self._planned_booking_ids = {row[0] for row in self.db.session.query(TripBooking.booking_id).all()}
        self._trip_end_dates = {trip.trip_id: trip.end_date for trip in self.db.session.query(Trip).all()}

    def _precompute_duplicate_keys(self) -> None:
        for sheet, key, counter in [
            ("Travelers", "Traveler ID", self._traveler_id_counts),
            ("Trips", "Trip ID", self._trip_id_counts),
            ("Trip Bookings", "Booking ID", self._booking_id_counts),
        ]:
            for _, row in self._iter_sheet_rows(sheet):
                value = clean(row.get(key))
                if value:
                    counter[value] += 1

        for _, row in self._iter_sheet_rows("Travelers"):
            phone = self._normalize_row_phone(row)
            if phone["lookup_key"]:
                self._phone_counts[phone["lookup_key"]] += 1

    def _load_traveler_decisions(self) -> dict[int, TravelerDecision]:
        decisions: dict[int, TravelerDecision] = {}
        if not self.decisions_path.exists():
            return decisions

        with self.decisions_path.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            for raw in reader:
                try:
                    primary_row = int(clean(raw.get("primary_row")) or 0) or None
                except ValueError:
                    primary_row = None
                secondary_rows = []
                for part in re.split(r"[,\s]+", clean(raw.get("secondary_rows"))):
                    if not part:
                        continue
                    if part.isdigit():
                        secondary_rows.append(int(part))
                decision = TravelerDecision(
                    group_id=clean(raw.get("group_id")),
                    classification=clean(raw.get("classification")),
                    decision=clean(raw.get("decision")),
                    primary_row=primary_row,
                    secondary_rows=secondary_rows,
                    traveler_id_to_keep=clean(raw.get("traveler_id_to_keep")),
                    new_traveler_id_if_needed=clean(raw.get("new_traveler_id_if_needed")),
                    action=clean(raw.get("action")),
                    reason=clean(raw.get("reason")),
                )
                if primary_row:
                    decisions[primary_row] = decision
                for secondary_row in secondary_rows:
                    decisions[secondary_row] = decision
        return decisions

    def _decision_rows(self, decision: TravelerDecision) -> set[int]:
        rows = set(decision.secondary_rows)
        if decision.primary_row:
            rows.add(decision.primary_row)
        return rows

    def _generate_new_traveler_id(self) -> str:
        from services.crm.system_services import UnifiedCRMService

        service = UnifiedCRMService()
        existing = set(self._planned_traveler_ids)
        try:
            candidate = service.next_traveler_id()
        except Exception:
            candidate = ""
        if candidate and candidate not in existing:
            return candidate
        max_number = 0
        for traveler_id in existing:
            match = re.fullmatch(r"TR(\d+)", traveler_id or "")
            if match:
                max_number = max(max_number, int(match.group(1)))
        return f"TR{max_number + 1:05d}"

    def _merge_row_values(self, base_row: dict[str, Any], extra_rows: list[dict[str, Any]]) -> dict[str, Any]:
        merged = dict(base_row)
        for extra_row in extra_rows:
            for key, value in extra_row.items():
                if merged.get(key) in ("", None) and value not in ("", None):
                    merged[key] = value
        return merged

    def _row_phone_identity(self, row: dict[str, Any]) -> str:
        phone = self._normalize_row_phone(row)
        return phone.get("lookup_key") or phone.get("normalized_e164") or ""

    def _traveler_decision_for_row(self, row_number: int) -> TravelerDecision | None:
        return getattr(self, "_traveler_decisions", {}).get(row_number)

    def _row_is_primary(self, row_number: int, decision: TravelerDecision) -> bool:
        return decision.primary_row == row_number

    def _row_is_secondary(self, row_number: int, decision: TravelerDecision) -> bool:
        return row_number in decision.secondary_rows

    def _row_by_number(self, sheet: str, target_row: int) -> dict[str, Any] | None:
        for row_number, row in self._iter_sheet_rows(sheet):
            if row_number == target_row:
                return row
        return None

    def _find_existing_traveler(self, row: dict[str, Any], *, allow_phone_match: bool = True):
        from app.models import Traveler

        traveler_id = clean(row.get("Traveler ID"))
        phone = self._normalize_row_phone(row)
        lookup_key = phone["lookup_key"]
        if traveler_id:
            traveler = self.db.session.get(Traveler, traveler_id)
            if traveler:
                return traveler
        if allow_phone_match:
            candidates: list[Any] = []
            variants = [clean(v) for v in phone.get("phone_variants_for_lookup", []) if clean(v)]
            if variants:
                candidates.extend(Traveler.query.filter(Traveler.phone_lookup_key.in_(variants)).all())
            normalized = clean(phone.get("normalized_e164"))
            if normalized:
                candidates.extend(
                    Traveler.query.filter(
                        or_(
                            Traveler.normalized_whatsapp == normalized,
                            Traveler.integrated_whatsapp == normalized,
                        )
                    ).all()
                )
            if len(candidates) == 1:
                return candidates[0]
            if len(candidates) > 1:
                return candidates[0]
        if lookup_key:
            return Traveler.query.filter_by(phone_lookup_key=lookup_key).first()
        return None

    def _apply_row_to_traveler(self, traveler: Any, row: dict[str, Any]) -> bool:
        phone = self._normalize_row_phone(row)
        changed = False
        full_name = clean(row.get("Full Name"))
        changed |= fill_blank(traveler, "full_name", full_name)
        changed |= fill_blank(traveler, "status", map_traveler_status(row.get("Status")))
        changed |= fill_blank(traveler, "first_name", clean(row.get("First Name")) or first_name(full_name))
        changed |= fill_blank(traveler, "last_name", clean(row.get("Last Name")) or last_name(full_name))
        changed |= fill_blank(traveler, "birthday", parse_date(row.get("Birthday")))
        changed |= fill_blank(traveler, "gender", clean(row.get("Gender")))
        changed |= fill_blank(traveler, "nationality", clean(row.get("Nationality")))
        changed |= fill_blank(traveler, "phone_code", phone["country_code"] or clean_code(row.get("Code")))
        changed |= fill_blank(traveler, "whatsapp_raw", clean(row.get("WhatsApp")))
        changed |= fill_blank(traveler, "email", clean(row.get("Email")))
        changed |= fill_blank(traveler, "community_whatsapp", clean(row.get("Community Whatsapp")))
        changed |= fill_blank(traveler, "residence", clean(row.get("Residence")))
        changed |= fill_blank(traveler, "local_trips_count", parse_int(row.get("Loc. Trips")))
        changed |= fill_blank(traveler, "international_trips_count", parse_int(row.get("Int. Trips")))
        changed |= fill_blank(traveler, "total_trips", parse_int(row.get("Total trips")))
        changed |= fill_blank(traveler, "community_events_count", parse_int(row.get("Comm. Events")))
        changed |= fill_blank(traveler, "lifetime_revenue", parse_float(row.get("Lifetime Revenue")))
        changed |= fill_blank(traveler, "notes", clean(row.get("Notes")))
        changed |= fill_blank(traveler, "introduce_yourself", clean(row.get("Introduce yourself")))
        changed |= fill_blank(traveler, "emergency_contact", clean(row.get("Emergency Contact")))
        changed |= fill_blank(traveler, "emergency_phone", clean(row.get("Emergency Phone")))
        changed |= fill_blank(traveler, "medical_notes", clean(row.get("Medical Notes")))
        changed |= fill_blank(traveler, "room_preference", clean(row.get("Room Preference")))
        changed |= fill_blank(traveler, "rating", parse_float(row.get("Rating") or row.get("⭐ Rating (1–5)")) or None)
        changed |= fill_blank(traveler, "integrated_whatsapp", phone["normalized_e164"])
        changed |= fill_blank(traveler, "normalized_whatsapp", phone["normalized_e164"])
        changed |= fill_blank(traveler, "phone_lookup_key", phone["lookup_key"])
        changed |= fill_blank(traveler, "lead_source", "Excel Migration")
        changed |= fill_blank(traveler, "data_audit", "Updated/fill-blanks from Excel migration")
        return changed

    def _create_traveler_from_row(self, row: dict[str, Any], traveler_id: str) -> Any:
        from app.models import Traveler

        phone = self._normalize_row_phone(row)
        full_name = clean(row.get("Full Name"))
        return Traveler(
            traveler_id=traveler_id,
            status=map_traveler_status(row.get("Status")),
            full_name=full_name,
            first_name=clean(row.get("First Name")) or first_name(full_name),
            last_name=clean(row.get("Last Name")) or last_name(full_name),
            birthday=parse_date(row.get("Birthday")),
            gender=clean(row.get("Gender")),
            nationality=clean(row.get("Nationality")),
            phone_code=phone["country_code"] or clean_code(row.get("Code")),
            whatsapp_raw=clean(row.get("WhatsApp")),
            email=clean(row.get("Email")),
            community_whatsapp=clean(row.get("Community Whatsapp")),
            residence=clean(row.get("Residence")),
            local_trips_count=parse_int(row.get("Loc. Trips")),
            international_trips_count=parse_int(row.get("Int. Trips")),
            total_trips=parse_int(row.get("Total trips")),
            community_events_count=parse_int(row.get("Comm. Events")),
            lifetime_revenue=parse_float(row.get("Lifetime Revenue")),
            notes=clean(row.get("Notes")),
            introduce_yourself=clean(row.get("Introduce yourself")),
            emergency_contact=clean(row.get("Emergency Contact")),
            emergency_phone=clean(row.get("Emergency Phone")),
            medical_notes=clean(row.get("Medical Notes")),
            room_preference=clean(row.get("Room Preference")),
            rating=parse_float(row.get("Rating") or row.get("⭐ Rating (1–5)")) or None,
            integrated_whatsapp=phone["normalized_e164"],
            normalized_whatsapp=phone["normalized_e164"],
            phone_lookup_key=phone["lookup_key"],
            lead_source="Excel Migration",
            data_audit=f"Imported from Travelers row {row.get('_row_number', '')}",
        )

    def _travelers_validation(self) -> dict[str, Any]:
        from app.models import Traveler

        travelers = Traveler.query.all()
        ids = [t.traveler_id for t in travelers if t.traveler_id]
        phones = [t.phone_lookup_key for t in travelers if t.phone_lookup_key]
        duplicate_ids = [key for key, count in Counter(ids).items() if count > 1]
        duplicate_phones = [key for key, count in Counter(phones).items() if count > 1]
        return {
            "duplicate_traveler_ids": duplicate_ids,
            "duplicate_phone_lookup_keys": duplicate_phones,
            "traveler_count": len(travelers),
        }

    def _trips_validation(self) -> dict[str, Any]:
        from app.models import Trip

        trips = Trip.query.all()
        ids = [trip.trip_id for trip in trips if trip.trip_id]
        duplicate_ids = [key for key, count in Counter(ids).items() if count > 1]
        return {
            "duplicate_trip_ids": duplicate_ids,
            "missing_boys_girls_split_trip_ids": sorted(self._trip_ids_missing_split),
            "trip_count": len(trips),
        }

    def _import_or_update_traveler(self, row: dict[str, Any], traveler_id: str, *, source_note: str) -> str:
        from app.models import Traveler

        phone = self._normalize_row_phone(row)
        existing = self.db.session.get(Traveler, traveler_id)
        if existing:
            changed = self._apply_row_to_traveler(existing, row)
            if changed:
                self.summary.sheet("Travelers").updated += 1
            else:
                self.summary.sheet("Travelers").skipped += 1
            return traveler_id

        traveler = self._create_traveler_from_row(row, traveler_id)
        traveler.data_audit = f"{traveler.data_audit}; {source_note}" if traveler.data_audit else source_note
        self.db.session.add(traveler)
        self._planned_traveler_ids.add(traveler_id)
        self.summary.sheet("Travelers").created += 1
        return traveler_id

    def _migrate_travelers_only(self) -> None:
        sheet = "Travelers"
        stats = self.summary.sheet(sheet)
        if sheet not in self.wb.sheetnames:
            stats.skipped += 1
            self.summary.validation = self._travelers_validation()
            return

        rows_by_number = {row_number: row for row_number, row in self._iter_sheet_rows(sheet)}
        for row_number in sorted(rows_by_number):
            stats.scanned += 1
            if row_number in self._handled_traveler_rows:
                continue

            row = rows_by_number[row_number]
            decision = self._traveler_decision_for_row(row_number)
            traveler_id = clean(row.get("Traveler ID"))
            full_name = clean(row.get("Full Name"))
            phone = self._normalize_row_phone(row)
            lookup_key = phone["lookup_key"]

            if decision and decision.decision.lower() == "merge" and self._row_is_secondary(row_number, decision):
                continue

            if decision and decision.decision.lower() in {"merge", "manual review"}:
                group_rows = sorted(self._decision_rows(decision))
                if decision.decision.lower() == "manual review":
                    for group_row in group_rows:
                        if group_row in self._handled_traveler_rows:
                            continue
                        row_data = rows_by_number.get(group_row)
                        if row_data is None:
                            continue
                        self._quarantine(sheet, group_row, "decision_requires_manual_review", row_data, key=decision.group_id)
                        self._handled_traveler_rows.add(group_row)
                    continue

                primary_row = rows_by_number.get(decision.primary_row or row_number)
                if primary_row is None:
                    self._quarantine(sheet, row_number, "missing_merge_primary_row", row, key=decision.group_id)
                    self._handled_traveler_rows.add(row_number)
                    continue
                merged_row = self._merge_row_values(
                    primary_row,
                    [rows_by_number[r] for r in decision.secondary_rows if r in rows_by_number],
                )
                merged_traveler_id = clean(merged_row.get("Traveler ID")) or decision.traveler_id_to_keep
                if not merged_traveler_id:
                    merged_traveler_id = self._generate_new_traveler_id()
                self._import_or_update_traveler(
                    merged_row,
                    merged_traveler_id,
                    source_note=f"merged from rows {', '.join(str(r) for r in group_rows)}",
                )
                self.summary.travelers_merged += max(len(group_rows) - 1, 0)
                self._handled_traveler_rows.update(group_rows)
                continue

            if decision and decision.decision.lower() == "keep separate":
                if self._row_is_primary(row_number, decision):
                    target_traveler_id = decision.traveler_id_to_keep or traveler_id
                elif self._row_is_secondary(row_number, decision):
                    target_traveler_id = clean(decision.new_traveler_id_if_needed)
                    if not target_traveler_id or target_traveler_id.upper() == "TBD":
                        target_traveler_id = self._generate_new_traveler_id()
                else:
                    target_traveler_id = traveler_id or self._generate_new_traveler_id()
                if not target_traveler_id:
                    self._quarantine(sheet, row_number, "missing_traveler_id_for_keep_separate", row, key=decision.group_id)
                    self._handled_traveler_rows.add(row_number)
                    continue
                self._import_or_update_traveler(
                    row,
                    target_traveler_id,
                    source_note=f"decision {decision.group_id} ({decision.action})",
                )
                self._handled_traveler_rows.add(row_number)
                continue

            if not traveler_id and not full_name and not lookup_key:
                self._quarantine(sheet, row_number, "missing_phone_and_name", row)
                self._handled_traveler_rows.add(row_number)
                continue
            if traveler_id and self._traveler_id_counts[traveler_id] > 1:
                self._quarantine(sheet, row_number, "duplicate_traveler_id_in_workbook", row, key=traveler_id)
                self._handled_traveler_rows.add(row_number)
                continue
            if lookup_key and self._phone_counts[lookup_key] > 1:
                self._quarantine(sheet, row_number, "duplicate_phone_identity_in_workbook", row, key=lookup_key)
                self._handled_traveler_rows.add(row_number)
                continue
            if not traveler_id:
                traveler_id = self._generate_new_traveler_id()
            self._import_or_update_traveler(row, traveler_id, source_note="travelers-only migration")
            self._handled_traveler_rows.add(row_number)

        self.summary.validation = self._travelers_validation()

    def _migrate_travelers(self) -> None:
        from app.models import Traveler

        sheet = "Travelers"
        stats = self.summary.sheet(sheet)
        for row_number, row in self._iter_sheet_rows(sheet):
            stats.scanned += 1
            traveler_id = clean(row.get("Traveler ID"))
            full_name = clean(row.get("Full Name"))
            phone = self._normalize_row_phone(row)
            lookup_key = phone["lookup_key"]

            if not traveler_id and not full_name and not lookup_key:
                self._quarantine(sheet, row_number, "missing_phone_and_name", row)
                continue
            if traveler_id and self._traveler_id_counts[traveler_id] > 1:
                self._quarantine(sheet, row_number, "duplicate_traveler_id_in_workbook", row, key=traveler_id)
                continue
            if lookup_key and self._phone_counts[lookup_key] > 1:
                self._quarantine(sheet, row_number, "duplicate_phone_identity_in_workbook", row, key=lookup_key)
                continue
            if not traveler_id and not lookup_key:
                self._quarantine(sheet, row_number, "missing_identity_key", row)
                continue

            matches = self._find_travelers(traveler_id=traveler_id, phone=phone)
            if len(matches) > 1:
                self._quarantine(sheet, row_number, "duplicate_uncertain_match", row, key=lookup_key or traveler_id)
                continue

            if matches:
                existing = matches[0]
                if traveler_id and existing.traveler_id != traveler_id:
                    self._quarantine(
                        sheet,
                        row_number,
                        "conflicting_phone_traveler_id",
                        row,
                        key=f"{lookup_key} maps to {existing.traveler_id}, row has {traveler_id}",
                    )
                    continue
                changed = self._fill_traveler(existing, row, phone)
                if changed:
                    stats.updated += 1
                else:
                    stats.skipped += 1
                continue

            if not traveler_id:
                self._quarantine(sheet, row_number, "missing_traveler_id_for_new_record", row, key=lookup_key)
                continue
            if not full_name:
                self._quarantine(sheet, row_number, "missing_name_for_new_record", row, key=traveler_id)
                continue

            traveler = Traveler(
                traveler_id=traveler_id,
                full_name=full_name,
                status=map_traveler_status(row.get("Status")),
                first_name=clean(row.get("First Name")) or first_name(full_name),
                last_name=clean(row.get("Last Name")) or last_name(full_name),
                birthday=parse_date(row.get("Birthday")),
                gender=clean(row.get("Gender")),
                nationality=clean(row.get("Nationality")),
                phone_code=phone["country_code"] or clean_code(row.get("Code")),
                whatsapp_raw=clean(row.get("WhatsApp")),
                email=clean(row.get("Email")),
                community_whatsapp=clean(row.get("Community Whatsapp")),
                residence=clean(row.get("Residence")),
                local_trips_count=parse_int(row.get("Loc. Trips")),
                international_trips_count=parse_int(row.get("Int. Trips")),
                total_trips=parse_int(row.get("Total trips")),
                community_events_count=parse_int(row.get("Comm. Events")),
                lifetime_revenue=parse_float(row.get("Lifetime Revenue")),
                notes=clean(row.get("Notes")),
                introduce_yourself=clean(row.get("Introduce yourself")),
                emergency_contact=clean(row.get("Emergency Contact")),
                emergency_phone=clean(row.get("Emergency Phone")),
                medical_notes=clean(row.get("Medical Notes")),
                room_preference=clean(row.get("Room Preference")),
                rating=parse_float(row.get("Rating") or row.get("⭐ Rating (1–5)")) or None,
                integrated_whatsapp=phone["normalized_e164"],
                normalized_whatsapp=phone["normalized_e164"],
                phone_lookup_key=lookup_key,
                lead_source="Excel Migration",
                data_audit=f"Imported from {sheet} row {row_number}",
            )
            self.db.session.add(traveler)
            self._planned_traveler_ids.add(traveler_id)
            stats.created += 1

    def _trip_id_and_name(self, row: dict[str, Any]) -> tuple[str, str]:
        return clean(row.get("Trip ID")), clean(row.get("Trip Name"))

    def _trip_has_missing_split(self, row: dict[str, Any]) -> bool:
        return not any(clean(row.get(field)) for field in ("Boys Double", "Girls Double", "Boys Triple", "Girls Triple"))

    def _trip_public_price(self, row: dict[str, Any]) -> str:
        return first_non_empty(row.get("Flight Price EGP"), row.get("Flight Price USD"))

    def _trip_sales_notes(self, row: dict[str, Any]) -> str:
        parts = []
        revenue_egp = clean(row.get("EGP"))
        revenue_usd = clean(row.get("USD"))
        if revenue_egp:
            parts.append(f"Revenue EGP: {revenue_egp}")
        if revenue_usd:
            parts.append(f"Revenue USD: {revenue_usd}")
        return "; ".join(parts)

    def _trip_manual_note(self, row: dict[str, Any]) -> str:
        if self._trip_has_missing_split(row):
            return "Boys/girls split missing from source workbook; set to 0 for manual completion."
        return ""

    def _build_trip(self, row: dict[str, Any], row_number: int, trip_id: str):
        from app.models import Trip

        return Trip(
            trip_id=trip_id,
            trip_name=clean(row.get("Trip Name")),
            type=clean(row.get("Type")),
            year=parse_int(row.get("Year")) or None,
            trip_leader=clean(row.get("Trip Leader")),
            start_date=parse_date(row.get("Start Date")),
            end_date=parse_date(row.get("End Date")),
            sales_status="Closed" if is_past_trip(row.get("End Date")) else "Open",
            data_audit=f"Imported from Trips row {row_number}",
            trip_availability_note=self._trip_manual_note(row) or None,
            single_total=parse_int(row.get("Single")),
            double_total=parse_int(row.get("Double")),
            triple_total=parse_int(row.get("Triple")),
            single_remaining=parse_int(row.get("Single #2")),
            double_remaining=parse_int(row.get("Double #2")),
            triple_remaining=parse_int(row.get("Triple #2")),
            boys_double=0,
            girls_double=0,
            boys_triple=0,
            girls_triple=0,
            public_price=self._trip_public_price(row),
            sales_notes=self._trip_sales_notes(row) or None,
        )

    def _apply_trip_update(self, trip: Any, row: dict[str, Any], row_number: int) -> bool:
        changed = False
        changed |= fill_blank(trip, "trip_name", clean(row.get("Trip Name")))
        changed |= fill_blank(trip, "type", clean(row.get("Type")))
        changed |= fill_blank(trip, "year", parse_int(row.get("Year")) or None)
        changed |= fill_blank(trip, "trip_leader", clean(row.get("Trip Leader")))
        changed |= fill_blank(trip, "start_date", parse_date(row.get("Start Date")))
        changed |= fill_blank(trip, "end_date", parse_date(row.get("End Date")))
        changed |= fill_blank(trip, "sales_status", "Closed" if is_past_trip(row.get("End Date")) else "Open")
        changed |= fill_blank(trip, "single_total", parse_int(row.get("Single")))
        changed |= fill_blank(trip, "double_total", parse_int(row.get("Double")))
        changed |= fill_blank(trip, "triple_total", parse_int(row.get("Triple")))
        changed |= fill_blank(trip, "single_remaining", parse_int(row.get("Single #2")))
        changed |= fill_blank(trip, "double_remaining", parse_int(row.get("Double #2")))
        changed |= fill_blank(trip, "triple_remaining", parse_int(row.get("Triple #2")))
        changed |= fill_blank(trip, "public_price", self._trip_public_price(row))
        changed |= fill_blank(trip, "sales_notes", self._trip_sales_notes(row) or None)
        manual_note = self._trip_manual_note(row)
        if manual_note:
            changed |= fill_blank(trip, "trip_availability_note", manual_note)
        if self._trip_has_missing_split(row):
            self._trip_ids_missing_split.add(clean(row.get("Trip ID")))
        return changed

    def _migrate_trips_only(self) -> None:
        from app.models import Trip

        sheet = "Trips"
        stats = self.summary.sheet(sheet)
        if sheet not in self.wb.sheetnames:
            stats.skipped += 1
            self.summary.validation = self._trips_validation()
            return

        for row_number, row in self._iter_sheet_rows(sheet):
            stats.scanned += 1
            trip_id, trip_name = self._trip_id_and_name(row)
            if not trip_id and not trip_name:
                stats.skipped += 1
                continue
            if not trip_id or not trip_name:
                self._quarantine(sheet, row_number, "missing_trip_required_field", row, key=trip_id or trip_name)
                continue
            if self._trip_id_counts[trip_id] > 1:
                self._quarantine(sheet, row_number, "duplicate_trip_id_in_workbook", row, key=trip_id)
                continue

            existing = self.db.session.get(Trip, trip_id)
            if existing:
                if self._apply_trip_update(existing, row, row_number):
                    stats.updated += 1
                else:
                    stats.skipped += 1
                if self._trip_has_missing_split(row):
                    self._trip_ids_missing_split.add(trip_id)
                continue

            trip = self._build_trip(row, row_number, trip_id)
            self.db.session.add(trip)
            self._planned_trip_ids.add(trip_id)
            stats.created += 1
            if self._trip_has_missing_split(row):
                self._trip_ids_missing_split.add(trip_id)

        self.summary.validation = self._trips_validation()

    def _migrate_trips(self) -> None:
        from app.models import Trip

        sheet = "Trips"
        stats = self.summary.sheet(sheet)
        for row_number, row in self._iter_sheet_rows(sheet):
            stats.scanned += 1
            trip_id, trip_name = self._trip_id_and_name(row)
            if not trip_id and not trip_name:
                stats.skipped += 1
                continue
            if not trip_id or not trip_name:
                self._quarantine(sheet, row_number, "missing_trip_required_field", row, key=trip_id or trip_name)
                continue
            if self._trip_id_counts[trip_id] > 1:
                self._quarantine(sheet, row_number, "duplicate_trip_id_in_workbook", row, key=trip_id)
                continue

            existing = self.db.session.get(Trip, trip_id)
            if existing:
                if self._apply_trip_update(existing, row, row_number):
                    stats.updated += 1
                else:
                    stats.skipped += 1
                continue

            trip = self._build_trip(row, row_number, trip_id)
            self.db.session.add(trip)
            self._planned_trip_ids.add(trip_id)
            stats.created += 1

    def _migrate_bookings(self) -> None:
        self._process_bookings(import_safe_matches=True)

    def _migrate_bookings_only(self) -> None:
        self._process_bookings(import_safe_matches=True)

    def _preview_bookings_only(self) -> None:
        self._process_bookings(import_safe_matches=False)

    def _process_bookings(self, *, import_safe_matches: bool) -> None:
        from app.models import BookingStatusHistory, TripBooking

        sheet = "Trip Bookings"
        stats = self.summary.sheet(sheet)
        for row_number, row in self._iter_sheet_rows(sheet):
            stats.scanned += 1
            candidate, reason, key = self._classify_booking_row(row)
            if candidate is None and reason == "skip":
                stats.skipped += 1
                continue
            if candidate is None:
                self._quarantine(sheet, row_number, reason, row, key=key)
                continue

            self.summary.bookings_matched += 1
            self.summary.booking_lifecycle_counts[candidate.booking_status] = self.summary.booking_lifecycle_counts.get(candidate.booking_status, 0) + 1
            if import_safe_matches:
                booking = TripBooking(
                    booking_id=candidate.booking_id,
                    trip_id=candidate.trip_id,
                    trip_name=candidate.trip_name,
                    traveler_id=candidate.traveler_id,
                    traveler_name=candidate.traveler_name,
                    room_type=candidate.room_type,
                    currency=candidate.currency,
                    booking_status=candidate.booking_status,
                    payment_status=candidate.payment_status,
                    booking_source="Excel Migration",
                    booking_notes=f"Imported from {sheet} row {row_number}",
                )
                self.db.session.add(booking)
                self.db.session.add(
                    BookingStatusHistory(
                        booking_id=candidate.booking_id,
                        old_status=None,
                        new_status=candidate.booking_status,
                        changed_by="excel-migration",
                        change_source="excel-import",
                        notes=f"Imported from {sheet} row {row_number}",
                    )
                )
                self._planned_booking_ids.add(candidate.booking_id)
            stats.created += 1

        self.summary.bookings_quarantined = stats.quarantined
        self.summary.validation = self._bookings_validation(import_safe_matches=import_safe_matches)

    def _classify_booking_row(self, row: dict[str, Any]) -> tuple[BookingCandidate | None, str, str]:
        booking_id = clean(row.get("Booking ID"))
        trip_id = clean(row.get("Trip ID"))
        traveler_id = clean(row.get("Traveler ID"))
        traveler_name = clean(row.get("Traveler Name"))
        if not any([booking_id, trip_id, traveler_id, traveler_name]):
            return None, "skip", ""
        if not booking_id:
            return None, "missing_booking_id", ""
        if booking_id == "0-L-000" or self._booking_id_counts[booking_id] > 1:
            return None, "duplicate_or_placeholder_booking_id", booking_id
        if not traveler_id or traveler_id not in self._planned_traveler_ids:
            return None, "invalid_booking_traveler_reference", traveler_id
        if not trip_id or trip_id not in self._planned_trip_ids:
            return None, "invalid_booking_trip_reference", trip_id
        room_type = normalize_room_type(row.get("Room Type"))
        if clean(row.get("Room Type")) and not room_type:
            return None, "invalid_room_type", clean(row.get("Room Type"))
        if booking_id in self._planned_booking_ids:
            return None, "duplicate_existing_booking_id", booking_id
        lifecycle = map_booking_lifecycle(row, self._trip_end_dates.get(trip_id))
        if not lifecycle:
            return None, "unclear_booking_lifecycle", booking_id
        return (
            BookingCandidate(
                booking_id=booking_id,
                trip_id=trip_id,
                trip_name=clean(row.get("Trip Name")),
                traveler_id=traveler_id,
                traveler_name=traveler_name,
                room_type=room_type or None,
                currency=clean(row.get("Currency")),
                booking_status=lifecycle,
                payment_status=map_payment_status(row),
            ),
            "",
            "",
        )

    def _bookings_validation(self, *, import_safe_matches: bool) -> dict[str, Any]:
        reason_counts = Counter(item["reason"] for item in self.quarantine if item["sheet"] == "Trip Bookings")
        key = "bookings_import" if import_safe_matches else "bookings_preview"
        return {
            key: {
                "matched": self.summary.bookings_matched,
                "imported": self.summary.bookings_matched if import_safe_matches else 0,
                "quarantined": self.summary.bookings_quarantined,
                "status_distribution": dict(self.summary.booking_lifecycle_counts),
                "invalid_traveler_refs": reason_counts.get("invalid_booking_traveler_reference", 0),
                "invalid_trip_refs": reason_counts.get("invalid_booking_trip_reference", 0),
                "placeholder_booking_ids": reason_counts.get("duplicate_or_placeholder_booking_id", 0),
                "lifecycle_unclear_count": reason_counts.get("unclear_booking_lifecycle", 0),
                "safe_to_execute": self.summary.bookings_quarantined == 0,
            }
        }

    def _migrate_leads(self) -> None:
        sheet = "Leads"
        if sheet not in self.wb.sheetnames:
            self.summary.sheet(sheet).skipped += 1
            return
        from app.models import Lead

        stats = self.summary.sheet(sheet)
        for row_number, row in self._iter_sheet_rows(sheet):
            stats.scanned += 1
            lead_id = clean(row.get("Lead ID"))
            traveler_id = clean(row.get("Traveler ID"))
            if not lead_id:
                self._quarantine(sheet, row_number, "missing_lead_id", row)
                continue
            if traveler_id and traveler_id not in self._planned_traveler_ids:
                self._quarantine(sheet, row_number, "invalid_lead_traveler_reference", row, key=traveler_id)
                continue
            if self.db.session.get(Lead, lead_id):
                stats.skipped += 1
                continue
            lead = Lead(
                lead_id=lead_id,
                customer_name=clean(row.get("Customer Name")),
                raw_phone=clean(row.get("Raw Phone")),
                traveler_id=traveler_id or None,
                lead_stage=clean(row.get("Lead Stage")) or "New Lead",
                lead_source=clean(row.get("Lead Source")) or "Excel Migration",
                source_sheet=sheet,
                source_row=row_number,
            )
            self.db.session.add(lead)
            stats.created += 1

    def _iter_sheet_rows(self, sheet_name: str):
        if not hasattr(self, "wb"):
            self.wb = load_workbook(self.workbook_path, data_only=True, read_only=True)
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        header_row = SHEET_HEADER_ROWS[sheet_name]
        headers = dedupe_headers([clean(cell.value) or f"Column {idx}" for idx, cell in enumerate(ws[header_row], start=1)])
        for row_number, values in enumerate(ws.iter_rows(min_row=header_row + 1, max_col=len(headers), values_only=True), start=header_row + 1):
            if not any(clean(value) for value in values):
                continue
            yield row_number, {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}

    def _normalize_row_phone(self, row: dict[str, Any]) -> dict[str, str]:
        raw_phone = clean(row.get("WhatsApp") or row.get("Raw Phone") or row.get("Phone"))
        code = clean_code(row.get("Code") or row.get("Phone Code") or row.get("Country Code"))
        result = normalize_phone_input(raw_phone, code, default_country_is_explicit=bool(code))
        return result.to_dict()

    def _find_travelers(self, *, traveler_id: str, phone: dict[str, Any]) -> list[Any]:
        from app.models import Traveler

        found = {}
        if traveler_id:
            traveler = self.db.session.get(Traveler, traveler_id)
            if traveler:
                found[traveler.traveler_id] = traveler
        variants = [clean(v) for v in phone.get("phone_variants_for_lookup", []) if clean(v)]
        if variants:
            for traveler in Traveler.query.filter(Traveler.phone_lookup_key.in_(variants)).all():
                found[traveler.traveler_id] = traveler
            normalized = clean(phone.get("normalized_e164"))
            if normalized:
                for traveler in Traveler.query.filter(Traveler.normalized_whatsapp == normalized).all():
                    found[traveler.traveler_id] = traveler
                for traveler in Traveler.query.filter(Traveler.integrated_whatsapp == normalized).all():
                    found[traveler.traveler_id] = traveler
        return list(found.values())

    def _fill_traveler(self, traveler: Any, row: dict[str, Any], phone: dict[str, Any]) -> bool:
        changed = False
        full_name = clean(row.get("Full Name"))
        changed |= fill_blank(traveler, "full_name", full_name)
        changed |= fill_blank(traveler, "status", map_traveler_status(row.get("Status")))
        changed |= fill_blank(traveler, "first_name", clean(row.get("First Name")) or first_name(full_name))
        changed |= fill_blank(traveler, "last_name", clean(row.get("Last Name")) or last_name(full_name))
        changed |= fill_blank(traveler, "birthday", parse_date(row.get("Birthday")))
        changed |= fill_blank(traveler, "gender", clean(row.get("Gender")))
        changed |= fill_blank(traveler, "nationality", clean(row.get("Nationality")))
        changed |= fill_blank(traveler, "phone_code", phone["country_code"] or clean_code(row.get("Code")))
        changed |= fill_blank(traveler, "whatsapp_raw", clean(row.get("WhatsApp")))
        changed |= fill_blank(traveler, "email", clean(row.get("Email")))
        changed |= fill_blank(traveler, "residence", clean(row.get("Residence")))
        changed |= fill_blank(traveler, "integrated_whatsapp", phone["normalized_e164"])
        changed |= fill_blank(traveler, "normalized_whatsapp", phone["normalized_e164"])
        changed |= fill_blank(traveler, "phone_lookup_key", phone["lookup_key"])
        changed |= fill_blank(traveler, "data_audit", "Updated/fill-blanks from Excel migration")
        return changed

    def _quarantine(self, sheet: str, row_number: int, reason: str, row: dict[str, Any], *, key: str = "") -> None:
        self.summary.sheet(sheet).quarantined += 1
        self.quarantine.append(
            {
                "sheet": sheet,
                "row": row_number,
                "reason": reason,
                "key": key,
                "row_data": {k: serialize_value(v) for k, v in row.items()},
            }
        )


def dedupe_headers(headers: list[str]) -> list[str]:
    counts: Counter[str] = Counter()
    result = []
    for header in headers:
        counts[header] += 1
        result.append(header if counts[header] == 1 else f"{header} #{counts[header]}")
    return result


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.lower() in {"none", "nan", "nat"} else text


def clean_code(value: Any) -> str:
    code = clean(value)
    return code[:-2] if code.endswith(".0") else re.sub(r"\D", "", code)


def parse_int(value: Any) -> int:
    text = clean(value).replace(",", "")
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def parse_float(value: Any) -> float:
    text = clean(value).replace(",", "").replace("$", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def is_past_trip(value: Any) -> bool:
    parsed = parse_date(value)
    return bool(parsed and parsed < date.today())


def first_name(full_name: str) -> str:
    return full_name.split(" ", 1)[0] if full_name else ""


def last_name(full_name: str) -> str:
    parts = full_name.split(" ", 1)
    return parts[1] if len(parts) > 1 else ""


def fill_blank(obj: Any, field_name: str, value: Any) -> bool:
    if value in ("", None):
        return False
    current = getattr(obj, field_name, None)
    if current not in ("", None):
        return False
    setattr(obj, field_name, value)
    return True


def map_traveler_status(value: Any) -> str:
    status = clean(value)
    mapping = {
        "": "Active",
        "repeat": "Repeat",
        "vip": "VIP",
        "blackllisted": "Blacklisted",
        "blacklisted": "Blacklisted",
        "cancelled": "Cancelled Review",
    }
    return mapping.get(status.lower(), status or "Active")


def normalize_room_type(value: Any) -> str:
    room = clean(value).lower()
    if not room:
        return ""
    if room.startswith("single"):
        return "Single"
    if room.startswith("double"):
        return "Double"
    if room.startswith("triple"):
        return "Triple"
    return ""


def map_payment_status(row: dict[str, Any]) -> str:
    first_amount = parse_float(row.get("Amount"))
    second_amount = parse_float(row.get("Amount #2"))
    if first_amount and second_amount:
        return "Fully Paid"
    if first_amount or second_amount:
        return "Deposit Paid"
    return ""


def map_booking_status(row: dict[str, Any]) -> str:
    raw = clean(row.get("Booking Status") or row.get("Status"))
    if raw in BOOKING_STATUSES:
        return raw
    lower = raw.lower()
    if "cancel" in lower:
        return "Cancelled"
    if map_payment_status(row) == "Fully Paid":
        return "Paid"
    if map_payment_status(row) == "Deposit Paid":
        return "Payment Pending"
    return "Confirmed"


def map_booking_lifecycle(row: dict[str, Any], trip_end_date: date | None = None) -> str:
    raw = clean(row.get("Booking Status") or row.get("Status") or row.get("Lifecycle"))
    payment_status = map_payment_status(row)
    if raw:
        lower = raw.lower()
        if "cancel" in lower:
            return "Cancelled"
        if "draft" in lower:
            return "Draft"
        if "waiting customer" in lower:
            return "Waiting Customer"
        if "pending confirmation" in lower:
            return "Pending Confirmation"
        if "confirm" in lower:
            return "Confirmed"
        if "paid" in lower:
            if payment_status == "Fully Paid" and trip_end_date and trip_end_date < date.today():
                return "Completed"
            return "Paid"
        if "complete" in lower:
            return "Completed"

    if payment_status == "Fully Paid":
        if trip_end_date and trip_end_date < date.today():
            return "Completed"
        return "Paid"
    if payment_status == "Deposit Paid":
        return "Payment Pending"
    if trip_end_date and trip_end_date < date.today():
        return "Completed"
    if trip_end_date:
        return "Confirmed"
    return ""


def first_non_empty(*values: Any) -> str:
    for value in values:
        text = clean(value)
        if text:
            return text
    return ""


def serialize_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return clean(value)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def write_report(path: Path, summary: MigrationSummary, quarantine: list[dict[str, Any]]) -> None:
    lines = [
        "# Excel to CRM Migration Run Report",
        "",
        f"- Mode: {'dry-run' if summary.dry_run else 'execute'}",
        f"- Scope: {summary.scope}",
        f"- Workbook: `{summary.workbook}`",
        f"- Database URI: `{summary.database_uri}`",
        f"- Quarantined rows: {len(quarantine)}",
        "",
    ]
    if summary.bookings_matched or summary.bookings_quarantined:
        lines.extend(
            [
                "## Booking Preview",
                "",
                f"- Matched bookings: {summary.bookings_matched}",
                f"- Imported bookings: {summary.bookings_matched if summary.scope == 'bookings' and not summary.dry_run else 0}",
                f"- Quarantined bookings: {summary.bookings_quarantined}",
                f"- Safe to execute: {summary.bookings_quarantined == 0}",
                "",
            ]
        )
    if summary.booking_lifecycle_counts:
        lines.extend(["## Booking Lifecycle Counts", ""])
        for lifecycle, count in sorted(summary.booking_lifecycle_counts.items()):
            lines.append(f"- {lifecycle}: {count}")
        lines.append("")
    lines.extend(
        [
            "## Summary",
            "",
            "| Sheet | Scanned | Created | Updated | Skipped | Quarantined | Errors |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for sheet, stats in summary.sheets.items():
        lines.append(
            f"| {sheet} | {stats.scanned} | {stats.created} | {stats.updated} | "
            f"{stats.skipped} | {stats.quarantined} | {stats.errors} |"
        )
    reason_counts = Counter(row["reason"] for row in quarantine)
    if reason_counts:
        lines.extend(["", "## Quarantine Reasons", ""])
        for reason, count in reason_counts.most_common():
            lines.append(f"- {reason}: {count}")
    if summary.validation:
        lines.extend(["", "## Validation", ""])
        for key, value in summary.validation.items():
            lines.append(f"- {key}: {value}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def apply_database_url_to_app_config() -> None:
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        return
    try:
        from app.config import DevelopmentConfig, ProductionConfig, config
    except Exception:
        return
    DevelopmentConfig.SQLALCHEMY_DATABASE_URI = database_url
    ProductionConfig.SQLALCHEMY_DATABASE_URI = database_url
    config["development"].SQLALCHEMY_DATABASE_URI = database_url
    config["default"].SQLALCHEMY_DATABASE_URI = database_url


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Safely migrate Rahma Excel data into the CRM database.")
    parser.add_argument("--workbook", default=str(REPO_ROOT / "RT - Travelers Database.xlsx"))
    parser.add_argument("--db", dest="db_path", default="")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="Validate and report without committing writes. Default.")
    mode.add_argument("--execute", action="store_true", help="Commit safe records to the CRM database.")
    parser.add_argument("--travelers-only", action="store_true", help="Import only traveler records.")
    parser.add_argument("--trips-only", action="store_true", help="Import only trip records.")
    parser.add_argument("--bookings-only", action="store_true", help="Import only safe trip bookings.")
    parser.add_argument("--bookings-preview", action="store_true", help="Analyze trip bookings without writing records.")
    parser.add_argument("--merge-decisions", default=str(REPO_ROOT / "MIGRATION_DECISIONS.csv"))
    parser.add_argument("--quarantine-output", default=str(REPO_ROOT / "migration_quarantine.json"))
    parser.add_argument("--report-output", default=str(REPO_ROOT / "migration_report.md"))
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.db_path:
        os.environ["DATABASE_URL"] = f"sqlite:///{Path(args.db_path).resolve().as_posix()}"
    dry_run = not args.execute
    scope = "travelers" if args.travelers_only else "trips" if args.trips_only else "bookings" if args.bookings_only else "bookings_preview" if args.bookings_preview else "full"
    decisions_path = Path(args.merge_decisions) if args.merge_decisions else None
    migrator = ExcelToCRMMigrator(
        Path(args.workbook),
        dry_run=dry_run,
        scope=scope,
        decisions_path=decisions_path,
    )
    summary = migrator.run()
    write_json(Path(args.quarantine_output), migrator.quarantine)
    write_report(Path(args.report_output), summary, migrator.quarantine)
    print(json.dumps(asdict(summary), indent=2, ensure_ascii=False))
    print(f"Quarantine output: {args.quarantine_output}")
    print(f"Report output: {args.report_output}")
    return 0 if not summary.errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
