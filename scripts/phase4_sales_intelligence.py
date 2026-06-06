from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


LEADS_SHEET = "Leads"

LEADS_HEADERS = [
    "Lead ID",
    "Created At",
    "Updated At",
    "Customer Name",
    "Raw Phone",
    "Integrated WhatsApp",
    "Phone Lookup Key",
    "Traveler ID",
    "Traveler Status",
    "Customer Tier",
    "Match Status",
    "Lead Stage",
    "Lead Source",
    "Channel",
    "Preferred Trip Type",
    "Interested Trip IDs",
    "Suggested Trip IDs",
    "Priority",
    "Follow Up Status",
    "Follow Up Due Date",
    "Last Interaction ID",
    "Interaction Count",
    "Handoff Required",
    "Handoff Reason",
    "Notes",
    "Birthday",
    "Gender",
    "Nationality",
]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def set_auto_filter_if_supported(ws: Worksheet, ref: str) -> None:
    auto_filter = getattr(ws, "auto_filter", None)
    if auto_filter is not None:
        auto_filter.ref = ref


def ensure_leads_sheet(wb) -> tuple[Worksheet, dict[str, int]]:
    if LEADS_SHEET in wb.sheetnames:
        ws = wb[LEADS_SHEET]
    else:
        ws = wb.create_sheet(LEADS_SHEET)

    header_map: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = normalize_text(ws.cell(1, col_idx).value)
        if header:
            header_map[header] = col_idx

    if not header_map:
        for col_idx, header in enumerate(LEADS_HEADERS, start=1):
            ws.cell(1, col_idx).value = header
            header_map[header] = col_idx
    else:
        next_col = ws.max_column + 1
        for header in LEADS_HEADERS:
            if header in header_map:
                continue
            ws.cell(1, next_col).value = header
            header_map[header] = next_col
            next_col += 1

    if hasattr(ws, "freeze_panes"):
        ws.freeze_panes = "A2"
    set_auto_filter_if_supported(ws, f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}")
    return ws, header_map


def next_blank_row(ws: Worksheet) -> int:
    for row_idx in range(2, ws.max_row + 1):
        if not any(ws.cell(row_idx, col).value not in (None, "") for col in range(1, ws.max_column + 1)):
            return row_idx
    return ws.max_row + 1


def generate_lead_id(ws: Worksheet, header_map: dict[str, int]) -> str:
    col = header_map["Lead ID"]
    max_number = 0
    for row_idx in range(2, ws.max_row + 1):
        lead_id = normalize_text(ws.cell(row_idx, col).value)
        if lead_id.startswith("LD") and lead_id[2:].isdigit():
            max_number = max(max_number, int(lead_id[2:]))
    return f"LD{max_number + 1:05d}"


def parse_iso_datetime(value: object) -> datetime | None:
    text = normalize_text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def infer_customer_tier(result: dict[str, Any]) -> str:
    traveler = result.get("traveler")
    if isinstance(traveler, dict):
        status = normalize_text(traveler.get("status"))
        if status == "VIP":
            return "VIP"
        if status == "Repeat":
            return "Repeat"
        if status:
            return status
    return "Standard"


def derive_lead_stage(result: dict[str, Any]) -> tuple[str, str]:
    traveler = result.get("traveler")
    traveler_status = normalize_text(traveler.get("status")) if isinstance(traveler, dict) else ""
    trip_result = result.get("trip_result") or {}
    open_trips = trip_result.get("open_trips", [])
    tbd_trips = trip_result.get("date_tbd_trips", [])

    if result.get("handoff_reason") == "blacklisted_customer":
        return "Blocked", "Critical"
    if result.get("handoff_required"):
        return "Needs Review", "High"
    if traveler_status == "VIP":
        if open_trips:
            return "VIP Priority", "High"
        if tbd_trips:
            return "VIP Follow Up", "High"
        return "VIP Priority", "High"
    if traveler_status == "Repeat":
        if open_trips:
            return "Repeat Priority", "Medium"
        if tbd_trips:
            return "Repeat Follow Up", "Medium"
        return "Repeat Priority", "Medium"
    if open_trips:
        return "Qualified", "Medium"
    if tbd_trips:
        return "Follow Up Needed", "Medium"
    if result.get("match_status") == "not_found":
        return "New Lead", "Medium"
    if result.get("match_status") == "single_match":
        return "Existing Traveler", "Low"
    return "New Inquiry", "Low"


def derive_follow_up(result: dict[str, Any], timestamp: datetime) -> tuple[str, str]:
    stage, _priority = derive_lead_stage(result)
    if stage == "Blocked":
        return "Do Not Contact", ""
    if result.get("handoff_required"):
        return "Urgent", timestamp.date().isoformat()

    trip_result = result.get("trip_result") or {}
    open_trips = trip_result.get("open_trips", [])
    tbd_trips = trip_result.get("date_tbd_trips", [])

    if open_trips:
        return "Follow Up Soon", (timestamp.date() + timedelta(days=2)).isoformat()
    if tbd_trips:
        return "Awaiting Dates", (timestamp.date() + timedelta(days=7)).isoformat()
    if result.get("match_status") == "not_found":
        return "Qualify Lead", (timestamp.date() + timedelta(days=3)).isoformat()
    return "Monitor", (timestamp.date() + timedelta(days=14)).isoformat()


def collect_trip_ids(result: dict[str, Any]) -> list[str]:
    trip_result = result.get("trip_result") or {}
    open_ids = [item["trip_id"] for item in trip_result.get("open_trips", [])]
    tbd_ids = [item["trip_id"] for item in trip_result.get("date_tbd_trips", [])]
    return open_ids or tbd_ids


def find_existing_lead_row(
    ws: Worksheet,
    header_map: dict[str, int],
    phone_lookup_key: str,
    traveler_id: str,
) -> int | None:
    for row_idx in range(2, ws.max_row + 1):
        existing_lookup = normalize_text(ws.cell(row_idx, header_map["Phone Lookup Key"]).value)
        existing_traveler_id = normalize_text(ws.cell(row_idx, header_map["Traveler ID"]).value)
        if phone_lookup_key and existing_lookup == phone_lookup_key:
            return row_idx
        if traveler_id and existing_traveler_id == traveler_id:
            return row_idx
    return None


def upsert_lead(
    ws: Worksheet,
    header_map: dict[str, int],
    *,
    result: dict[str, Any],
    customer_name: str,
    raw_phone: str,
    phone_lookup_key: str,
    integrated_whatsapp: str,
    traveler_id: str,
    channel: str,
    source: str,
    trip_type: str,
    interaction_id: str,
    timestamp: datetime,
    notes: str,
    birthday: str = "",
    gender: str = "",
    nationality: str = "",
    preferred_trip_id: str = "",
) -> dict[str, Any]:
    row_idx = find_existing_lead_row(ws, header_map, phone_lookup_key, traveler_id)
    is_new = row_idx is None
    if row_idx is None:
        row_idx = next_blank_row(ws)
        ws.cell(row_idx, header_map["Lead ID"]).value = generate_lead_id(ws, header_map)
        ws.cell(row_idx, header_map["Created At"]).value = timestamp.isoformat(timespec="seconds")
        ws.cell(row_idx, header_map["Interaction Count"]).value = 0

    stage, priority = derive_lead_stage(result)
    follow_up_status, follow_up_due = derive_follow_up(result, timestamp)
    trip_ids = collect_trip_ids(result)
    traveler = result.get("traveler")
    traveler_status = ""
    if isinstance(traveler, dict):
        traveler_status = normalize_text(traveler.get("status"))

    interaction_count = ws.cell(row_idx, header_map["Interaction Count"]).value
    if not isinstance(interaction_count, int):
        if isinstance(interaction_count, float) and interaction_count.is_integer():
            interaction_count = int(interaction_count)
        else:
            interaction_count = 0
            
    interested_ids = preferred_trip_id if preferred_trip_id else (", ".join(trip_ids) or None)

    updates = {
        "Updated At": timestamp.isoformat(timespec="seconds"),
        "Customer Name": customer_name,
        "Raw Phone": raw_phone,
        "Integrated WhatsApp": integrated_whatsapp or None,
        "Phone Lookup Key": phone_lookup_key or None,
        "Traveler ID": traveler_id or None,
        "Traveler Status": traveler_status or None,
        "Customer Tier": infer_customer_tier(result),
        "Match Status": result.get("match_status", ""),
        "Lead Stage": stage,
        "Lead Source": source,
        "Channel": channel,
        "Preferred Trip Type": trip_type or None,
        "Interested Trip IDs": interested_ids,
        "Suggested Trip IDs": ", ".join(trip_ids) or None,
        "Priority": priority,
        "Follow Up Status": follow_up_status,
        "Follow Up Due Date": follow_up_due or None,
        "Last Interaction ID": interaction_id,
        "Interaction Count": interaction_count + 1,
        "Handoff Required": "Yes" if result.get("handoff_required") else "No",
        "Handoff Reason": result.get("handoff_reason") or None,
        "Notes": notes or None,
        "Birthday": birthday or None,
        "Gender": gender or None,
        "Nationality": nationality or None,
    }

    for header, value in updates.items():
        ws.cell(row_idx, header_map[header]).value = value

    set_auto_filter_if_supported(ws, f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, row_idx)}")
    return {
        "row": row_idx,
        "lead_id": normalize_text(ws.cell(row_idx, header_map["Lead ID"]).value),
        "lead_stage": stage,
        "priority": priority,
        "follow_up_status": follow_up_status,
        "follow_up_due_date": follow_up_due,
        "created": is_new,
        "interaction_count": interaction_count + 1,
    }


def compute_sales_dashboard(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    result = compute_sales_dashboard_from_wb(wb)
    wb.close()
    return result


def compute_sales_dashboard_from_wb(wb: Any) -> dict[str, Any]:
    if LEADS_SHEET not in wb.sheetnames:
        return {
            "leadCount": 0,
            "leadStageCounts": {},
            "priorityCounts": {},
            "followUpSummary": {"dueToday": 0, "overdue": 0, "urgent": 0},
            "pipeline": {},
            "recentLeads": [],
        }

    ws = wb[LEADS_SHEET]
    headers = {
        normalize_text(ws.cell(1, col).value): col
        for col in range(1, ws.max_column + 1)
        if normalize_text(ws.cell(1, col).value)
    }

    lead_stage_counts: Counter[str] = Counter()
    priority_counts: Counter[str] = Counter()
    due_today = 0
    overdue = 0
    urgent = 0
    recent_rows: list[dict[str, Any]] = []
    lead_count = 0
    new_customers = 0
    existing_travelers = 0
    qualified = 0
    follow_up_needed = 0
    needs_review = 0
    blocked = 0
    booking_drafts = 0

    for row_idx in range(2, ws.max_row + 1):
        lead_id = normalize_text(ws.cell(row_idx, headers["Lead ID"]).value)
        if not lead_id:
            continue

        lead_count += 1
        stage = normalize_text(ws.cell(row_idx, headers["Lead Stage"]).value)
        priority = normalize_text(ws.cell(row_idx, headers["Priority"]).value)
        follow_up_status = normalize_text(ws.cell(row_idx, headers["Follow Up Status"]).value)
        due_date = normalize_text(ws.cell(row_idx, headers["Follow Up Due Date"]).value)
        updated_at = normalize_text(ws.cell(row_idx, headers["Updated At"]).value)
        match_status = normalize_text(ws.cell(row_idx, headers["Match Status"]).value)
        customer_name = normalize_text(ws.cell(row_idx, headers["Customer Name"]).value)
        customer_tier = normalize_text(ws.cell(row_idx, headers["Customer Tier"]).value)

        lead_stage_counts[stage] += 1
        priority_counts[priority] += 1

        if stage in {"Qualified", "VIP Priority", "Repeat Priority"}:
            qualified += 1
        if stage in {"Follow Up Needed", "VIP Follow Up", "Repeat Follow Up"}:
            follow_up_needed += 1
        if stage == "Needs Review":
            needs_review += 1
        if stage == "Blocked":
            blocked += 1
        if stage in {"Booking Draft Created", "VIP Booking Draft", "Repeat Booking Draft"}:
            booking_drafts += 1
        if match_status == "not_found":
            new_customers += 1
        if match_status == "single_match":
            existing_travelers += 1

        if follow_up_status == "Urgent":
            urgent += 1
        if due_date:
            try:
                parsed_due = date.fromisoformat(due_date)
                if parsed_due == date.today():
                    due_today += 1
                elif parsed_due < date.today():
                    overdue += 1
            except ValueError:
                pass

        recent_rows.append(
            {
                "leadId": lead_id,
                "customerName": customer_name,
                "leadStage": stage,
                "priority": priority,
                "customerTier": customer_tier,
                "updatedAt": updated_at,
                "followUpStatus": follow_up_status,
            }
        )

    recent_rows.sort(key=lambda item: item["updatedAt"], reverse=True)
    qualification_rate = round((qualified / lead_count) * 100, 1) if lead_count else 0.0

    return {
        "leadCount": lead_count,
        "leadStageCounts": dict(lead_stage_counts),
        "priorityCounts": dict(priority_counts),
        "followUpSummary": {"dueToday": due_today, "overdue": overdue, "urgent": urgent},
        "pipeline": {
            "newCustomers": new_customers,
            "existingTravelers": existing_travelers,
            "qualified": qualified,
            "followUpNeeded": follow_up_needed,
            "needsReview": needs_review,
            "blocked": blocked,
            "bookingDrafts": booking_drafts,
            "qualificationRate": qualification_rate,
        },
        "recentLeads": recent_rows[:5],
    }
