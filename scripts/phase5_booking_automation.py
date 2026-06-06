from __future__ import annotations

import json
import sys
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.phase0_cleanup import TODAY, ensure_headers  # noqa: E402
from scripts.phase2_controlled_agent import log_interaction, normalize_text  # noqa: E402
from scripts.phase4_sales_intelligence import ensure_leads_sheet, set_auto_filter_if_supported  # noqa: E402


TRIPS_SHEET = "Trips"
TRAVELERS_SHEET = "Travelers"
TRIP_BOOKINGS_SHEET = "Trip Bookings"
INTERACTIONS_SHEET = "Interactions"
BOOKING_ALERTS_SHEET = "Booking Alerts"

TRIP_BOOKINGS_HELPER_HEADERS = [
    "Booking Status",
    "Draft Created At",
    "Booking Source",
    "Lead ID",
    "Interaction ID",
    "Alert ID",
    "Payment Status",
    "Booking Notes",
]

TRIP_BOOKINGS_EXPECTED_HEADERS = [
    "Booking ID",
    "Trip ID",
    "Trip Name",
    "Traveler ID",
    "Traveler Name",
    "Room Type",
    "Flight Option",
    "Date Option",
    "Currency",
]

TRIPS_EXPECTED_HEADERS = [
    "Trip ID",
    "Trip Name",
    "Type",
    "Year",
    "Start Date",
    "End Date",
    "Sales Status",
    "Data Audit",
    "Draft Holds Single",
    "Draft Holds Double",
    "Draft Holds Triple",
]

TRIPS_PHASE5_HELPERS = [
    "Draft Holds Single",
    "Draft Holds Double",
    "Draft Holds Triple",
]

BOOKING_ALERT_HEADERS = [
    "Alert ID",
    "Created At",
    "Booking ID",
    "Trip ID",
    "Trip Name",
    "Traveler ID",
    "Traveler Name",
    "Lead ID",
    "Channel",
    "Priority",
    "Alert Type",
    "Alert Status",
    "Summary",
    "Owner",
    "Notes",
]

ROOM_TYPE_TO_REMAINING_COLUMN = {
    "Single": "Single",
    "Double": "Double",
    "Triple": "Triple",
}

ROOM_TYPE_TO_HOLD_HEADER = {
    "Single": "Draft Holds Single",
    "Double": "Draft Holds Double",
    "Triple": "Draft Holds Triple",
}

DEMO_TRIPS = [
    {
        "trip_id": "RT-LOC-26-900",
        "trip_name": "Siwa Discovery Demo",
        "trip_type": "Local",
        "year": 2026,
        "trip_leader": "Demo Team",
        "start_date": date(2026, 8, 14),
        "end_date": date(2026, 8, 17),
        "single_remaining": 3,
        "double_remaining": 4,
        "triple_remaining": 2,
        "single_total": 3,
        "double_total": 4,
        "triple_total": 2,
        "public_price": "Starts from 8,500 EGP",
        "public_description": "Demo inventory for local trip booking.",
    },
    {
        "trip_id": "RT-INT-26-900",
        "trip_name": "Cappadocia Preview Demo",
        "trip_type": "International",
        "year": 2026,
        "trip_leader": "Demo Team",
        "start_date": date(2026, 9, 11),
        "end_date": date(2026, 9, 15),
        "single_remaining": 2,
        "double_remaining": 3,
        "triple_remaining": 0,
        "single_total": 2,
        "double_total": 3,
        "triple_total": 0,
        "public_price": "Starts from 1,450 USD",
        "public_description": "Demo inventory for international trip booking.",
    },
]


def normalize_int(value: object) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return 0
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return int(text) if text.isdigit() else 0


def clone_cell_style(source, target) -> None:
    if not all(hasattr(source, attr) and hasattr(target, attr) for attr in ("font", "fill", "border", "alignment", "number_format", "protection")):
        return
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def ensure_trip_bookings_headers(ws) -> dict[str, int]:
    header_map = {}
    max_col = ws.max_column
    for col in range(1, max_col + 1):
        header = normalize_text(ws.cell(2, col).value)
        if header:
            header_map[header] = col

    reference_style_col = max_col
    for header in TRIP_BOOKINGS_EXPECTED_HEADERS:
        if header in header_map:
            continue
        max_col += 1
        ws.cell(2, max_col).value = header
        clone_cell_style(ws.cell(2, reference_style_col), ws.cell(2, max_col))
        ws.cell(1, max_col).value = None
        clone_cell_style(ws.cell(1, reference_style_col), ws.cell(1, max_col))
        header_map[header] = max_col

    for header in TRIP_BOOKINGS_HELPER_HEADERS:
        if header in header_map:
            continue
        max_col += 1
        ws.cell(2, max_col).value = header
        clone_cell_style(ws.cell(2, reference_style_col), ws.cell(2, max_col))
        ws.cell(1, max_col).value = None
        clone_cell_style(ws.cell(1, reference_style_col), ws.cell(1, max_col))
        header_map[header] = max_col

    return header_map


def ensure_booking_alerts_sheet(wb):
    if BOOKING_ALERTS_SHEET in wb.sheetnames:
        ws = wb[BOOKING_ALERTS_SHEET]
    else:
        ws = wb.create_sheet(BOOKING_ALERTS_SHEET)

    for col_idx, header in enumerate(BOOKING_ALERT_HEADERS, start=1):
        ws.cell(1, col_idx).value = header
    if hasattr(ws, "freeze_panes"):
        ws.freeze_panes = "A2"
    set_auto_filter_if_supported(ws, f"A1:{get_column_letter(len(BOOKING_ALERT_HEADERS))}{max(ws.max_row, 1)}")
    return ws, {header: idx for idx, header in enumerate(BOOKING_ALERT_HEADERS, start=1)}


def ensure_trip_phase5_helpers(ws):
    return ensure_headers(
        ws,
        header_row=2,
        expected_headers=["Trip ID", "Trip Name", "Type", "Year", "Start Date", "End Date", "Sales Status", "Data Audit"],
        helper_headers=TRIPS_PHASE5_HELPERS,
    )


def next_blank_row(ws, key_columns: list[int], start_row: int) -> int:
    for row_idx in range(start_row, ws.max_row + 1):
        if not any(ws.cell(row_idx, col).value not in (None, "") for col in key_columns):
            return row_idx
    return ws.max_row + 1


def generate_booking_id(ws, header_map: dict[str, int], trip_id: str, traveler_id: str) -> str:
    trip_col = header_map["Trip ID"]
    existing_for_trip = 0
    for row_idx in range(3, ws.max_row + 1):
        if normalize_text(ws.cell(row_idx, trip_col).value) == trip_id:
            existing_for_trip += 1

    traveler_digits = traveler_id[2:] if traveler_id.startswith("TR") else traveler_id
    trip_mode = "I" if "-INT-" in trip_id else "L"
    trip_year = trip_id.split("-")[2]
    trip_serial = trip_id.split("-")[3]
    return f"{int(traveler_digits):d}-{trip_mode}{trip_year}{trip_serial}-{existing_for_trip + 1:03d}"


def generate_alert_id(ws, header_map: dict[str, int], timestamp: datetime) -> str:
    prefix = timestamp.strftime("ALT-%Y%m%d-")
    max_seq = 0
    col = header_map["Alert ID"]
    for row_idx in range(2, ws.max_row + 1):
        alert_id = normalize_text(ws.cell(row_idx, col).value)
        if alert_id.startswith(prefix):
            tail = alert_id[len(prefix) :]
            if tail.isdigit():
                max_seq = max(max_seq, int(tail))
    return f"{prefix}{max_seq + 1:04d}"


def find_trip_row(ws, header_map: dict[str, int], trip_id: str) -> int | None:
    trip_id_col = header_map["Trip ID"]
    for row_idx in range(3, ws.max_row + 1):
        if normalize_text(ws.cell(row_idx, trip_id_col).value) == trip_id:
            return row_idx
    return None


def seed_demo_open_trips(trips_ws, header_map: dict[str, int]) -> dict[str, int]:
    created = 0
    for demo_trip in DEMO_TRIPS:
        existing_row = find_trip_row(trips_ws, header_map, demo_trip["trip_id"])
        if existing_row is None:
            row_idx = next_blank_row(trips_ws, [header_map["Trip ID"], header_map["Trip Name"]], 3)
            trips_ws.cell(row_idx, header_map["Trip ID"]).value = demo_trip["trip_id"]
            trips_ws.cell(row_idx, header_map["Trip Name"]).value = demo_trip["trip_name"]
            trips_ws.cell(row_idx, header_map["Type"]).value = demo_trip["trip_type"]
            trips_ws.cell(row_idx, header_map["Year"]).value = demo_trip["year"]
            trips_ws.cell(row_idx, 5).value = demo_trip["trip_leader"]
            trips_ws.cell(row_idx, header_map["Start Date"]).value = demo_trip["start_date"]
            trips_ws.cell(row_idx, header_map["End Date"]).value = demo_trip["end_date"]
            trips_ws.cell(row_idx, 8).value = demo_trip["single_total"]
            trips_ws.cell(row_idx, 9).value = demo_trip["double_total"]
            trips_ws.cell(row_idx, 10).value = demo_trip["triple_total"]
            trips_ws.cell(row_idx, 11).value = demo_trip["single_remaining"]
            trips_ws.cell(row_idx, 12).value = demo_trip["double_remaining"]
            trips_ws.cell(row_idx, 13).value = demo_trip["triple_remaining"]
            trips_ws.cell(row_idx, header_map["Sales Status"]).value = "Open"
            if "Public Description" in header_map:
                trips_ws.cell(row_idx, header_map["Public Description"]).value = demo_trip["public_description"]
            if "Public Price" in header_map:
                trips_ws.cell(row_idx, header_map["Public Price"]).value = demo_trip["public_price"]
            if "Sales Notes" in header_map:
                trips_ws.cell(row_idx, header_map["Sales Notes"]).value = "Demo-only open inventory seeded for Phase 5."
            trips_ws.cell(row_idx, header_map["Data Audit"]).value = None
            trips_ws.cell(row_idx, header_map["Draft Holds Single"]).value = 0
            trips_ws.cell(row_idx, header_map["Draft Holds Double"]).value = 0
            trips_ws.cell(row_idx, header_map["Draft Holds Triple"]).value = 0
            created += 1
        else:
            for key, header in {
                "single_remaining": 11,
                "double_remaining": 12,
                "triple_remaining": 13,
            }.items():
                if trips_ws.cell(existing_row, header).value is None:
                    trips_ws.cell(existing_row, header).value = demo_trip[key]
            for header in ("Draft Holds Single", "Draft Holds Double", "Draft Holds Triple"):
                if trips_ws.cell(existing_row, header_map[header]).value is None:
                    trips_ws.cell(existing_row, header_map[header]).value = 0

    set_auto_filter_if_supported(trips_ws, f"A2:{get_column_letter(trips_ws.max_column)}{trips_ws.max_row}")
    return {"demoTripsSeeded": created}


def update_traveler_booking_formulas(travelers_ws) -> int:
    updates = 0
    for row_idx in range(2, travelers_ws.max_row + 1):
        traveler_id = travelers_ws.cell(row_idx, 2).value
        full_name = travelers_ws.cell(row_idx, 3).value
        if not traveler_id and not full_name:
            continue

        loc_formula = f'=COUNTIFS(\'Trip Bookings\'!D:D,B{row_idx},\'Trip Bookings\'!B:B,"RT-LOC*",\'Trip Bookings\'!O:O,"<>Draft")'
        int_formula = f'=COUNTIFS(\'Trip Bookings\'!D:D,B{row_idx},\'Trip Bookings\'!B:B,"RT-INT*",\'Trip Bookings\'!O:O,"<>Draft")'
        total_formula = f"=N{row_idx}+O{row_idx}"
        travelers_ws.cell(row_idx, 14).value = loc_formula
        travelers_ws.cell(row_idx, 15).value = int_formula
        travelers_ws.cell(row_idx, 16).value = total_formula
        updates += 1
    return updates


def load_phase5_workbook(workbook_path: Path):
    wb = load_workbook(workbook_path, data_only=False)
    trips_ws = wb[TRIPS_SHEET]
    trip_bookings_ws = wb[TRIP_BOOKINGS_SHEET]
    travelers_ws = wb[TRAVELERS_SHEET]
    interactions_ws = wb[INTERACTIONS_SHEET]
    leads_ws, lead_headers = ensure_leads_sheet(wb)
    booking_alerts_ws, alert_headers = ensure_booking_alerts_sheet(wb)
    trip_headers = ensure_trip_phase5_helpers(trips_ws)
    booking_headers = ensure_trip_bookings_headers(trip_bookings_ws)
    interaction_headers = {
        normalize_text(interactions_ws.cell(1, col).value): col
        for col in range(1, interactions_ws.max_column + 1)
        if normalize_text(interactions_ws.cell(1, col).value)
    }
    return wb, trips_ws, trip_bookings_ws, travelers_ws, interactions_ws, leads_ws, booking_alerts_ws, trip_headers, booking_headers, interaction_headers, lead_headers, alert_headers


def _count_active_bookings(
    trip_bookings_ws,
    booking_headers: dict[str, int],
    trip_id: str,
    room_type: str,
) -> int:
    """Count all non-cancelled bookings for a trip+room directly from Trip Bookings sheet."""
    trip_col = booking_headers.get("Trip ID", 2)
    room_col = booking_headers.get("Room Type", 6)
    status_col = booking_headers.get("Booking Status")
    count = 0
    for row_idx in range(3, trip_bookings_ws.max_row + 1):
        if normalize_text(trip_bookings_ws.cell(row_idx, trip_col).value) != trip_id:
            continue
        if normalize_text(trip_bookings_ws.cell(row_idx, room_col).value).lower() != room_type.lower():
            continue
        if status_col:
            status = normalize_text(trip_bookings_ws.cell(row_idx, status_col).value).lower()
            if status == "cancelled":
                continue
        count += 1
    return count


def compute_available_rooms(
    trips_ws,
    trip_headers: dict[str, int],
    trip_row: int,
    room_type: str,
    trip_bookings_ws=None,
    booking_headers: dict[str, int] | None = None,
    trip_id: str = "",
) -> int | None:
    """Return how many rooms of this type are still bookable.

    Preferred path: count directly from Trip Bookings (works even when the
    Remaining column contains a formula string instead of a value).
    Fallback: read Remaining column minus Draft Holds counter.
    """
    if room_type not in ROOM_TYPE_TO_HOLD_HEADER:
        return None
    total_col = {"Single": 8, "Double": 9, "Triple": 10}.get(room_type)
    if total_col is None:
        return None
    total = normalize_int(trips_ws.cell(trip_row, total_col).value)

    if trip_bookings_ws is not None and booking_headers and trip_id:
        if total == 0:
            remaining_col = {"Single": 11, "Double": 12, "Triple": 13}[room_type]
            remaining_raw = trips_ws.cell(trip_row, remaining_col).value
            if remaining_raw is None or isinstance(remaining_raw, str):
                return None
            return normalize_int(remaining_raw)
        booked = _count_active_bookings(trip_bookings_ws, booking_headers, trip_id, room_type)
        return total - booked

    # Fallback: read the static Remaining value minus hold counter
    remaining_col = {"Single": 11, "Double": 12, "Triple": 13}[room_type]
    remaining_raw = trips_ws.cell(trip_row, remaining_col).value
    if isinstance(remaining_raw, str):          # It is a formula — cannot evaluate
        hold_header = ROOM_TYPE_TO_HOLD_HEADER[room_type]
        hold_count = normalize_int(trips_ws.cell(trip_row, trip_headers[hold_header]).value)
        return total - hold_count
    if remaining_raw is None:
        return None
    hold_header = ROOM_TYPE_TO_HOLD_HEADER[room_type]
    hold_count = normalize_int(trips_ws.cell(trip_row, trip_headers[hold_header]).value)
    return normalize_int(remaining_raw) - hold_count


def write_remaining_formulas(
    trips_ws,
    trip_headers: dict[str, int],
    booking_headers: dict[str, int],
) -> int:
    """Write COUNTIFS formulas to Remaining columns so the sheet self-governs.

    The formula calculates: Total - non-cancelled bookings for that room type.
    Humans opening the sheet in Google Sheets / Excel always see live numbers.
    Python code uses _count_active_bookings instead of reading these cells.
    """
    trip_id_letter = get_column_letter(trip_headers["Trip ID"])
    bk_trip_letter = get_column_letter(booking_headers["Trip ID"])
    bk_room_letter = get_column_letter(booking_headers["Room Type"])
    bk_status_letter = get_column_letter(booking_headers["Booking Status"])

    total_cols = {"Single": 8, "Double": 9, "Triple": 10}
    remaining_cols = {"Single": 11, "Double": 12, "Triple": 13}

    updated = 0
    for row_idx in range(3, trips_ws.max_row + 1):
        trip_id = trips_ws.cell(row_idx, trip_headers["Trip ID"]).value
        if not trip_id:
            continue
        for room_type, total_col in total_cols.items():
            total_letter = get_column_letter(total_col)
            remaining_col = remaining_cols[room_type]
            formula = (
                f"={total_letter}{row_idx}"
                f"-COUNTIFS("
                f"'Trip Bookings'!${bk_trip_letter}:${bk_trip_letter},"
                f"${trip_id_letter}{row_idx},"
                f"'Trip Bookings'!${bk_room_letter}:${bk_room_letter},"
                f'"{room_type}",'
                f"'Trip Bookings'!${bk_status_letter}:${bk_status_letter},"
                f'"<>Cancelled"'
                f")"
            )
            trips_ws.cell(row_idx, remaining_col).value = formula
        updated += 1
    return updated


def create_booking_alert(
    ws,
    header_map: dict[str, int],
    *,
    timestamp: datetime,
    booking_id: str,
    trip_id: str,
    trip_name: str,
    traveler_id: str,
    traveler_name: str,
    lead_id: str,
    channel: str,
    room_type: str,
) -> dict[str, Any]:
    row_idx = next_blank_row(ws, [header_map["Alert ID"], header_map["Booking ID"]], 2)
    alert_id = generate_alert_id(ws, header_map, timestamp)
    values = {
        "Alert ID": alert_id,
        "Created At": timestamp.isoformat(timespec="seconds"),
        "Booking ID": booking_id,
        "Trip ID": trip_id,
        "Trip Name": trip_name,
        "Traveler ID": traveler_id,
        "Traveler Name": traveler_name,
        "Lead ID": lead_id or None,
        "Channel": channel,
        "Priority": "High",
        "Alert Type": "Booking Draft",
        "Alert Status": "New",
        "Summary": f"Draft booking created for {traveler_name} on {trip_name} ({room_type}).",
        "Owner": "Sales Team",
        "Notes": "Awaiting deposit confirmation.",
    }
    for header, value in values.items():
        ws.cell(row_idx, header_map[header]).value = value
    set_auto_filter_if_supported(ws, f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, row_idx)}")
    return {"alert_id": alert_id, "row": row_idx}


def update_lead_for_booking(leads_ws, lead_headers: dict[str, int], lead_id: str, trip_id: str, interaction_id: str, timestamp: datetime) -> dict[str, Any] | None:
    if not lead_id:
        return None
    for row_idx in range(2, leads_ws.max_row + 1):
        if normalize_text(leads_ws.cell(row_idx, lead_headers["Lead ID"]).value) != lead_id:
            continue

        customer_tier = normalize_text(leads_ws.cell(row_idx, lead_headers["Customer Tier"]).value)
        if customer_tier == "VIP":
            stage = "VIP Booking Draft"
        elif customer_tier == "Repeat":
            stage = "Repeat Booking Draft"
        else:
            stage = "Booking Draft Created"

        leads_ws.cell(row_idx, lead_headers["Lead Stage"]).value = stage
        leads_ws.cell(row_idx, lead_headers["Priority"]).value = "High"
        leads_ws.cell(row_idx, lead_headers["Follow Up Status"]).value = "Awaiting Deposit"
        leads_ws.cell(row_idx, lead_headers["Follow Up Due Date"]).value = (timestamp.date() + timedelta(days=1)).isoformat()
        leads_ws.cell(row_idx, lead_headers["Updated At"]).value = timestamp.isoformat(timespec="seconds")
        leads_ws.cell(row_idx, lead_headers["Interested Trip IDs"]).value = trip_id
        leads_ws.cell(row_idx, lead_headers["Suggested Trip IDs"]).value = trip_id
        leads_ws.cell(row_idx, lead_headers["Last Interaction ID"]).value = interaction_id
        return {
            "row": row_idx,
            "lead_id": lead_id,
            "lead_stage": stage,
            "follow_up_status": "Awaiting Deposit",
            "follow_up_due_date": (timestamp.date() + timedelta(days=1)).isoformat(),
        }
    return None


def create_booking_draft(
    workbook_path: Path,
    output_path: Path,
    *,
    traveler_id: str,
    traveler_name: str,
    trip_id: str,
    room_type: str,
    channel: str,
    lead_id: str,
    source: str,
    agent_notes: str,
    flight_option: str = "",
    date_option: str = "",
    currency: str = "",
) -> dict[str, Any]:
    wb, trips_ws, trip_bookings_ws, travelers_ws, interactions_ws, leads_ws, booking_alerts_ws, trip_headers, booking_headers, interaction_headers, lead_headers, alert_headers = load_phase5_workbook(workbook_path)
    
    result, mutations = create_booking_draft_from_wb(
        wb=wb,
        trips_ws=trips_ws,
        trip_bookings_ws=trip_bookings_ws,
        travelers_ws=travelers_ws,
        interactions_ws=interactions_ws,
        leads_ws=leads_ws,
        booking_alerts_ws=booking_alerts_ws,
        trip_headers=trip_headers,
        booking_headers=booking_headers,
        interaction_headers=interaction_headers,
        lead_headers=lead_headers,
        alert_headers=alert_headers,
        traveler_id=traveler_id,
        traveler_name=traveler_name,
        trip_id=trip_id,
        room_type=room_type,
        channel=channel,
        lead_id=lead_id,
        source=source,
        agent_notes=agent_notes,
        flight_option=flight_option,
        date_option=date_option,
        currency=currency,
    )
    
    wb.save(output_path)
    wb.close()
    
    result["output_workbook"] = str(output_path)
    return result


def create_booking_draft_from_wb(
    wb: Any,
    trips_ws: Any,
    trip_bookings_ws: Any,
    travelers_ws: Any,
    interactions_ws: Any,
    leads_ws: Any,
    booking_alerts_ws: Any,
    trip_headers: dict[str, int],
    booking_headers: dict[str, int],
    interaction_headers: dict[str, int],
    lead_headers: dict[str, int],
    alert_headers: dict[str, int],
    *,
    traveler_id: str,
    traveler_name: str,
    trip_id: str,
    room_type: str,
    channel: str,
    lead_id: str,
    source: str,
    agent_notes: str,
    flight_option: str = "",
    date_option: str = "",
    currency: str = "",
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    timestamp = datetime.now()

    trip_row = find_trip_row(trips_ws, trip_headers, trip_id)
    if trip_row is None:
        raise ValueError(f"Trip {trip_id} not found")

    sales_status = normalize_text(trips_ws.cell(trip_row, trip_headers["Sales Status"]).value)
    trip_name = normalize_text(trips_ws.cell(trip_row, trip_headers["Trip Name"]).value)
    if sales_status != "Open":
        raise ValueError(f"Trip {trip_id} is not bookable because its status is {sales_status or 'blank'}")

    if room_type not in ROOM_TYPE_TO_HOLD_HEADER:
        raise ValueError(f"Unsupported room type {room_type!r}")

    # Count directly from Trip Bookings — reliable even when Remaining cells contain formulas
    available = compute_available_rooms(
        trips_ws, trip_headers, trip_row, room_type,
        trip_bookings_ws=trip_bookings_ws,
        booking_headers=booking_headers,
        trip_id=trip_id,
    )
    if available is None:
        raise ValueError(f"Capacity is not configured for room type {room_type}")
    if available <= 0:
        raise ValueError(f"No remaining draftable capacity for {room_type}")

    booking_row = next_blank_row(trip_bookings_ws, [booking_headers["Trip ID"], booking_headers["Traveler ID"]], 3)
    booking_id = generate_booking_id(trip_bookings_ws, booking_headers, trip_id, traveler_id)

    trip_bookings_ws.cell(booking_row, booking_headers["Booking ID"]).value = booking_id
    trip_bookings_ws.cell(booking_row, booking_headers["Trip ID"]).value = trip_id
    trip_bookings_ws.cell(booking_row, booking_headers["Trip Name"]).value = trip_name
    trip_bookings_ws.cell(booking_row, booking_headers["Traveler ID"]).value = traveler_id
    trip_bookings_ws.cell(booking_row, booking_headers["Traveler Name"]).value = traveler_name
    trip_bookings_ws.cell(booking_row, booking_headers["Room Type"]).value = room_type
    trip_bookings_ws.cell(booking_row, booking_headers["Flight Option"]).value = flight_option
    trip_bookings_ws.cell(booking_row, booking_headers["Date Option"]).value = date_option
    trip_bookings_ws.cell(booking_row, booking_headers["Currency"]).value = currency
    trip_bookings_ws.cell(booking_row, booking_headers["Booking Status"]).value = "Draft"
    trip_bookings_ws.cell(booking_row, booking_headers["Draft Created At"]).value = timestamp.isoformat(timespec="seconds")
    trip_bookings_ws.cell(booking_row, booking_headers["Booking Source"]).value = source
    trip_bookings_ws.cell(booking_row, booking_headers["Lead ID"]).value = lead_id or None
    trip_bookings_ws.cell(booking_row, booking_headers["Payment Status"]).value = "Awaiting Deposit"
    trip_bookings_ws.cell(booking_row, booking_headers["Booking Notes"]).value = agent_notes or "Draft created by Phase 5 automation"

    alert = create_booking_alert(
        booking_alerts_ws,
        alert_headers,
        timestamp=timestamp,
        booking_id=booking_id,
        trip_id=trip_id,
        trip_name=trip_name,
        traveler_id=traveler_id,
        traveler_name=traveler_name,
        lead_id=lead_id,
        channel=channel,
        room_type=room_type,
    )
    trip_bookings_ws.cell(booking_row, booking_headers["Alert ID"]).value = alert["alert_id"]

    hold_header = ROOM_TYPE_TO_HOLD_HEADER[room_type]
    current_hold = normalize_int(trips_ws.cell(trip_row, trip_headers[hold_header]).value)
    trips_ws.cell(trip_row, trip_headers[hold_header]).value = current_hold + 1

    interaction = log_interaction(
        interactions_ws,
        interaction_headers,
        timestamp=timestamp,
        channel=channel,
        customer_name=traveler_name,
        raw_phone="",
        integrated_whatsapp="",
        phone_lookup_key="",
        traveler_id=traveler_id,
        matched_row=None,
        status_snapshot="BOOKING_DRAFT",
        intent="booking_request",
        trip_type=normalize_text(trips_ws.cell(trip_row, trip_headers["Type"]).value),
        suggested_trips=[trip_id],
        action_taken=f"create_booking_draft:{booking_id}",
        handoff_required=False,
        handoff_reason="",
        agent_notes=agent_notes or "Booking draft created",
    )
    trip_bookings_ws.cell(booking_row, booking_headers["Interaction ID"]).value = interaction["interaction_id"]

    lead_update = update_lead_for_booking(leads_ws, lead_headers, lead_id, trip_id, interaction["interaction_id"], timestamp)

    # Rewrite Remaining formulas so the sheet reflects the new booking count
    write_remaining_formulas(trips_ws, trip_headers, booking_headers)

    result = {
        "booking_id": booking_id,
        "booking_row": booking_row,
        "trip_id": trip_id,
        "trip_name": trip_name,
        "room_type": room_type,
        "flight_option": flight_option,
        "date_option": date_option,
        "currency": currency,
        "available_before_draft": available,
        "available_after_draft": available - 1,
        "alert": alert,
        "interaction": interaction,
        "lead_update": lead_update,
    }
    
    mutations = getattr(wb, "mutations", [])
    return result, mutations


def prepare_phase5_workbook(input_path: Path, output_path: Path) -> dict[str, Any]:
    wb, trips_ws, trip_bookings_ws, travelers_ws, interactions_ws, leads_ws, booking_alerts_ws, trip_headers, booking_headers, interaction_headers, lead_headers, alert_headers = load_phase5_workbook(input_path)
    demo_trip_result = seed_demo_open_trips(trips_ws, trip_headers)
    traveler_formula_updates = update_traveler_booking_formulas(travelers_ws)
    # Write self-governing formulas to Remaining columns
    trips_formula_updates = write_remaining_formulas(trips_ws, trip_headers, booking_headers)
    wb.save(output_path)
    wb.close()
    return {
        "output_workbook": str(output_path),
        "demoTripsSeeded": demo_trip_result["demoTripsSeeded"],
        "travelerFormulaUpdates": traveler_formula_updates,
        "tripRemainingFormulasWritten": trips_formula_updates,
        "bookingAlertsSheetReady": True,
        "tripBookingsHelpersReady": True,
    }


def compute_booking_stats(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    result = compute_booking_stats_from_wb(wb)
    wb.close()
    return result


def compute_booking_stats_from_wb(wb: Any) -> dict[str, Any]:
    booking_draft_count = 0
    payment_pending_count = 0
    alert_count = 0

    if TRIP_BOOKINGS_SHEET in wb.sheetnames:
        ws = wb[TRIP_BOOKINGS_SHEET]
        booking_headers = {
            normalize_text(ws.cell(2, col).value): col
            for col in range(1, ws.max_column + 1)
            if normalize_text(ws.cell(2, col).value)
        }
        status_col = booking_headers.get("Booking Status")
        payment_col = booking_headers.get("Payment Status")
        for row_idx in range(3, ws.max_row + 1):
            trip_id = ws.cell(row_idx, 2).value
            if not trip_id:
                continue
            if status_col and normalize_text(ws.cell(row_idx, status_col).value) == "Draft":
                booking_draft_count += 1
            if payment_col and normalize_text(ws.cell(row_idx, payment_col).value) == "Awaiting Deposit":
                payment_pending_count += 1

    if BOOKING_ALERTS_SHEET in wb.sheetnames:
        ws = wb[BOOKING_ALERTS_SHEET]
        for row_idx in range(2, ws.max_row + 1):
            if ws.cell(row_idx, 1).value:
                alert_count += 1

    return {
        "bookingDraftCount": booking_draft_count,
        "paymentPendingCount": payment_pending_count,
        "bookingAlertCount": alert_count,
    }


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Phase 5 booking automation utilities.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare_parser = subparsers.add_parser("prepare")
    prepare_parser.add_argument("--input", required=True, type=Path)
    prepare_parser.add_argument("--output", required=True, type=Path)

    draft_parser = subparsers.add_parser("draft")
    draft_parser.add_argument("--workbook", required=True, type=Path)
    draft_parser.add_argument("--output", required=True, type=Path)
    draft_parser.add_argument("--traveler-id", required=True)
    draft_parser.add_argument("--traveler-name", required=True)
    draft_parser.add_argument("--trip-id", required=True)
    draft_parser.add_argument("--room-type", required=True)
    draft_parser.add_argument("--channel", default="web-demo")
    draft_parser.add_argument("--lead-id", default="")
    draft_parser.add_argument("--source", default="Booking Demo")
    draft_parser.add_argument("--agent-notes", default="")
    draft_parser.add_argument("--flight-option", default="")
    draft_parser.add_argument("--date-option", default="")
    draft_parser.add_argument("--currency", default="")

    stats_parser = subparsers.add_parser("stats")
    stats_parser.add_argument("--workbook", required=True, type=Path)

    args = parser.parse_args()

    if args.command == "prepare":
        result = prepare_phase5_workbook(args.input, args.output)
    elif args.command == "draft":
        result = create_booking_draft(
            workbook_path=args.workbook,
            output_path=args.output,
            traveler_id=args.traveler_id,
            traveler_name=args.traveler_name,
            trip_id=args.trip_id,
            room_type=args.room_type,
            channel=args.channel,
            lead_id=args.lead_id,
            source=args.source,
            agent_notes=args.agent_notes,
            flight_option=args.flight_option,
            date_option=args.date_option,
            currency=args.currency,
        )
    else:
        result = compute_booking_stats(args.workbook)

    print(json.dumps(result, indent=2, ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    main()
