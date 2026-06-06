from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.phase0_cleanup import normalize_phone_fields, normalize_text  # noqa: E402


TRAVELERS_SHEET = "Travelers"
CRM_AUDIT_SHEET = "CRM Integrity Audit"
VALID_TRAVELER_ID_RE = re.compile(r"^TR(\d+)$", re.IGNORECASE)


@dataclass
class TravelerIdStats:
    true_max_traveler_id: str
    true_max_traveler_id_number: int
    true_max_traveler_id_row: int
    next_traveler_id: str
    valid_traveler_id_count: int
    invalid_traveler_id_count: int


def header_map(ws: Worksheet, header_row: int = 1) -> dict[str, int]:
    return {
        normalize_text(ws.cell(header_row, col_idx).value): col_idx
        for col_idx in range(1, ws.max_column + 1)
        if normalize_text(ws.cell(header_row, col_idx).value)
    }


def required_headers(headers: dict[str, int], required: list[str]) -> None:
    missing = [header for header in required if header not in headers]
    if missing:
        raise ValueError(f"Missing required Travelers headers: {', '.join(missing)}")


def normalize_traveler_id(value: object) -> str:
    return normalize_text(value).upper()


def parse_traveler_id(value: object) -> int | None:
    match = VALID_TRAVELER_ID_RE.match(normalize_text(value))
    if not match:
        return None
    return int(match.group(1))


def format_traveler_id(number: int) -> str:
    return f"TR{number:05d}"


def row_has_any_data(ws: Worksheet, row_idx: int) -> bool:
    return any(normalize_text(ws.cell(row_idx, col_idx).value) for col_idx in range(1, ws.max_column + 1))


def cell_value(ws: Worksheet, headers: dict[str, int], row_idx: int, header: str) -> str:
    col_idx = headers.get(header)
    if not col_idx:
        return ""
    return normalize_text(ws.cell(row_idx, col_idx).value)


def effective_phone_lookup_key(ws: Worksheet, headers: dict[str, int], row_idx: int) -> tuple[str, bool, bool]:
    stored_lookup_key = cell_value(ws, headers, row_idx, "Phone Lookup Key")
    integrated_whatsapp = cell_value(ws, headers, row_idx, "Integrated WhatsApp")
    raw_code = cell_value(ws, headers, row_idx, "Code")
    raw_phone = cell_value(ws, headers, row_idx, "WhatsApp")

    phone_exists = bool(stored_lookup_key or integrated_whatsapp or raw_phone)
    if stored_lookup_key:
        return stored_lookup_key, phone_exists, False

    normalized = normalize_phone_fields(raw_code, raw_phone)
    if normalized.lookup_key:
        return normalized.lookup_key, phone_exists, True

    return "", phone_exists, phone_exists


def collect_traveler_rows(ws: Worksheet, headers: dict[str, int]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        if not row_has_any_data(ws, row_idx):
            continue

        traveler_id = normalize_traveler_id(cell_value(ws, headers, row_idx, "Traveler ID"))
        full_name = cell_value(ws, headers, row_idx, "Full Name")
        lookup_key, phone_exists, lookup_key_missing = effective_phone_lookup_key(ws, headers, row_idx)
        id_number = parse_traveler_id(traveler_id)

        rows.append(
            {
                "row": row_idx,
                "traveler_id": traveler_id,
                "traveler_id_number": id_number,
                "full_name": full_name,
                "lookup_key": lookup_key,
                "phone_exists": phone_exists,
                "lookup_key_missing": lookup_key_missing,
                "hidden": bool(ws.row_dimensions[row_idx].hidden),
            }
        )
    return rows


def compute_traveler_id_stats(rows: list[dict[str, Any]]) -> TravelerIdStats:
    valid_rows = [row for row in rows if row["traveler_id_number"] is not None]
    invalid_count = len([row for row in rows if row["traveler_id"] and row["traveler_id_number"] is None])

    if not valid_rows:
        return TravelerIdStats(
            true_max_traveler_id="",
            true_max_traveler_id_number=0,
            true_max_traveler_id_row=0,
            next_traveler_id="TR00001",
            valid_traveler_id_count=0,
            invalid_traveler_id_count=invalid_count,
        )

    max_row = max(valid_rows, key=lambda row: row["traveler_id_number"])
    next_number = int(max_row["traveler_id_number"]) + 1
    existing_ids = {row["traveler_id"] for row in valid_rows}
    next_id = format_traveler_id(next_number)
    while next_id in existing_ids:
        next_number += 1
        next_id = format_traveler_id(next_number)

    return TravelerIdStats(
        true_max_traveler_id=format_traveler_id(int(max_row["traveler_id_number"])),
        true_max_traveler_id_number=int(max_row["traveler_id_number"]),
        true_max_traveler_id_row=int(max_row["row"]),
        next_traveler_id=next_id,
        valid_traveler_id_count=len(valid_rows),
        invalid_traveler_id_count=invalid_count,
    )


def add_issue(
    issues: list[list[object]],
    severity: str,
    row: dict[str, Any] | None,
    issue: str,
    details: str,
) -> None:
    issues.append(
        [
            severity,
            TRAVELERS_SHEET,
            row["row"] if row else "",
            row["traveler_id"] if row else "",
            row["full_name"] if row else "",
            row["lookup_key"] if row else "",
            issue,
            details,
        ]
    )


def build_crm_integrity_audit(wb) -> tuple[dict[str, Any], list[list[object]]]:
    ws = wb[TRAVELERS_SHEET]
    headers = header_map(ws)
    required_headers(headers, ["Traveler ID", "Full Name", "Code", "WhatsApp"])

    rows = collect_traveler_rows(ws, headers)
    id_stats = compute_traveler_id_stats(rows)

    id_counts = Counter(row["traveler_id"] for row in rows if row["traveler_id"] and row["traveler_id_number"] is not None)
    phone_counts = Counter(row["lookup_key"] for row in rows if row["lookup_key"])

    rows_by_id: dict[str, list[dict[str, Any]]] = defaultdict(list)
    rows_by_phone: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row["traveler_id"] and row["traveler_id_number"] is not None:
            rows_by_id[row["traveler_id"]].append(row)
        if row["lookup_key"]:
            rows_by_phone[row["lookup_key"]].append(row)

    issues: list[list[object]] = []
    duplicate_id_keys = [traveler_id for traveler_id, count in id_counts.items() if count > 1]
    duplicate_phone_keys = [lookup_key for lookup_key, count in phone_counts.items() if count > 1]

    for traveler_id in sorted(duplicate_id_keys):
        duplicate_rows = rows_by_id[traveler_id]
        row_numbers = ", ".join(str(row["row"]) for row in duplicate_rows)
        for row in duplicate_rows:
            add_issue(issues, "Critical", row, "DUPLICATE_TRAVELER_ID", f"{traveler_id} appears in rows {row_numbers}")

    for lookup_key in sorted(duplicate_phone_keys):
        duplicate_rows = rows_by_phone[lookup_key]
        row_numbers = ", ".join(str(row["row"]) for row in duplicate_rows)
        for row in duplicate_rows:
            add_issue(issues, "Critical", row, "DUPLICATE_PHONE_LOOKUP_KEY", f"{lookup_key} appears in rows {row_numbers}")

    for row in rows:
        if row["traveler_id"] and row["traveler_id_number"] is None:
            add_issue(issues, "High", row, "INVALID_TRAVELER_ID", "Traveler ID does not match TR plus digits.")
        if row["traveler_id"] and (not row["full_name"] or not row["phone_exists"]):
            missing = []
            if not row["full_name"]:
                missing.append("Full Name")
            if not row["phone_exists"]:
                missing.append("phone")
            add_issue(issues, "High", row, "TRAVELER_ID_WITH_MISSING_IDENTITY_DATA", f"Missing: {', '.join(missing)}")
        if row["phone_exists"] and row["lookup_key_missing"]:
            add_issue(issues, "High", row, "PHONE_EXISTS_BUT_LOOKUP_KEY_MISSING", "Phone data exists but Phone Lookup Key is blank.")
        if row["hidden"]:
            add_issue(issues, "Info", row, "HIDDEN_ROW_WITH_DATA", "This data row is hidden in the workbook.")

    auto_filter_ref = str(ws.auto_filter.ref or "")
    hidden_rows_with_data = len([row for row in rows if row["hidden"]])
    has_filter_warning = bool(auto_filter_ref or hidden_rows_with_data)

    summary: dict[str, Any] = {
        "audit_timestamp": datetime.now().replace(microsecond=0).isoformat(),
        "travelers_rows_scanned": len(rows),
        **asdict(id_stats),
        "duplicate_traveler_id_keys": len(duplicate_id_keys),
        "duplicate_traveler_id_rows": sum(1 for row in rows if row["traveler_id"] in duplicate_id_keys),
        "duplicate_phone_lookup_keys": len(duplicate_phone_keys),
        "duplicate_phone_lookup_rows": sum(1 for row in rows if row["lookup_key"] in duplicate_phone_keys),
        "rows_with_id_missing_name_or_phone": sum(
            1 for row in rows if row["traveler_id"] and (not row["full_name"] or not row["phone_exists"])
        ),
        "rows_phone_exists_lookup_key_missing": sum(1 for row in rows if row["phone_exists"] and row["lookup_key_missing"]),
        "hidden_rows_with_data": hidden_rows_with_data,
        "auto_filter_ref": auto_filter_ref,
        "visible_rows_are_authoritative": "No" if has_filter_warning else "Unknown; still scan all rows",
        "visible_rows_warning": (
            "Do not use the last visible Google Sheet row for IDs. Scan the full Travelers sheet."
            if has_filter_warning
            else "No filter/hidden row metadata found, but the agent must still scan all Travelers rows."
        ),
        "issue_count": len(issues),
    }
    return summary, issues


def write_audit_sheet(wb, summary: dict[str, Any], issues: list[list[object]]) -> None:
    if CRM_AUDIT_SHEET in wb.sheetnames:
        del wb[CRM_AUDIT_SHEET]
    ws = wb.create_sheet(CRM_AUDIT_SHEET)

    title_fill = PatternFill("solid", fgColor="0F766E")
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    critical_fill = PatternFill("solid", fgColor="F4CCCC")
    high_fill = PatternFill("solid", fgColor="FCE5CD")
    info_fill = PatternFill("solid", fgColor="D9EAF7")

    ws["A1"] = "CRM Integrity Audit"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ws["A1"].fill = title_fill
    ws["A2"] = "This sheet proves the true next Traveler ID and flags CRM identity risks."

    row_idx = 4
    ws.cell(row_idx, 1).value = "Summary Key"
    ws.cell(row_idx, 2).value = "Value"
    for col_idx in range(1, 3):
        ws.cell(row_idx, col_idx).font = Font(bold=True)
        ws.cell(row_idx, col_idx).fill = header_fill

    for key, value in summary.items():
        row_idx += 1
        ws.cell(row_idx, 1).value = key
        if isinstance(value, (dict, list)):
            ws.cell(row_idx, 2).value = json.dumps(value, ensure_ascii=True, sort_keys=True)
        else:
            ws.cell(row_idx, 2).value = value

    row_idx += 2
    issue_header_row = row_idx
    issue_headers = ["Severity", "Sheet", "Row", "Traveler ID", "Full Name", "Phone Lookup Key", "Issue", "Details"]
    for col_idx, header in enumerate(issue_headers, start=1):
        cell = ws.cell(issue_header_row, col_idx)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for issue in issues:
        row_idx += 1
        for col_idx, value in enumerate(issue, start=1):
            ws.cell(row_idx, col_idx).value = value
        severity = issue[0]
        fill = critical_fill if severity == "Critical" else high_fill if severity == "High" else info_fill
        ws.cell(row_idx, 1).fill = fill

    for col_idx in range(1, len(issue_headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22
    ws.column_dimensions["H"].width = 60
    ws.freeze_panes = f"A{issue_header_row + 1}"
    ws.auto_filter.ref = f"A{issue_header_row}:H{max(issue_header_row, row_idx)}"


def write_markdown_report(report_path: Path, summary: dict[str, Any], issues: list[list[object]]) -> None:
    lines = [
        "# CRM Integrity Audit",
        "",
        f"Generated: {summary['audit_timestamp']}",
        "",
        "## Identity Summary",
        "",
        f"- Travelers rows scanned: {summary['travelers_rows_scanned']}",
        f"- Valid Traveler IDs: {summary['valid_traveler_id_count']}",
        f"- Invalid Traveler IDs: {summary['invalid_traveler_id_count']}",
        f"- True max Traveler ID: `{summary['true_max_traveler_id']}` at row `{summary['true_max_traveler_id_row']}`",
        f"- Next Traveler ID: `{summary['next_traveler_id']}`",
        f"- Duplicate Traveler ID keys: {summary['duplicate_traveler_id_keys']}",
        f"- Duplicate phone lookup keys: {summary['duplicate_phone_lookup_keys']}",
        f"- Rows with ID but missing name or phone: {summary['rows_with_id_missing_name_or_phone']}",
        f"- Rows with phone but missing lookup key: {summary['rows_phone_exists_lookup_key_missing']}",
        f"- Hidden rows with data: {summary['hidden_rows_with_data']}",
        f"- Auto filter range: `{summary['auto_filter_ref']}`",
        "",
        "## Rule",
        "",
        summary["visible_rows_warning"],
        "",
        "## Issue Sample",
        "",
    ]

    if not issues:
        lines.append("- No CRM identity issues found by this audit.")
    else:
        for issue in issues[:50]:
            severity, sheet, row, traveler_id, name, lookup_key, issue_code, details = issue
            lines.append(
                f"- `{severity}` `{sheet}` row `{row}` `{issue_code}`: "
                f"`{traveler_id}` `{name}` `{lookup_key}` - {details}"
            )

    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run_crm_integrity_audit(input_path: Path, output_path: Path, report_path: Path | None = None) -> dict[str, Any]:
    wb = load_workbook(input_path)
    summary, issues = build_crm_integrity_audit(wb)
    write_audit_sheet(wb, summary, issues)
    wb.save(output_path)
    if report_path:
        write_markdown_report(report_path, summary, issues)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Run Phase A CRM integrity audit.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    summary = run_crm_integrity_audit(args.input, args.output, args.report)
    print(json.dumps(summary, indent=2, ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    main()
