from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.phase4_sales_intelligence import ensure_leads_sheet  # noqa: E402


def prepare_phase4_workbook(input_path: Path, output_path: Path) -> dict[str, object]:
    wb = load_workbook(input_path, data_only=False)
    leads_ws, _headers = ensure_leads_sheet(wb)
    wb.save(output_path)
    wb.close()
    return {
        "output_workbook": str(output_path),
        "leads_sheet_ready": True,
        "lead_rows": max(leads_ws.max_row - 1, 0),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare a workbook copy with Phase 4 sales intelligence sheets.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    result = prepare_phase4_workbook(args.input, args.output)
    print(json.dumps(result, indent=2, ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    main()
