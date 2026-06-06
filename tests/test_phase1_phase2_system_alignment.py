from __future__ import annotations

import os
import shutil
import sqlite3
import sys
import unittest
import uuid
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook

from demo_web.app import create_app


SYSTEM_ROOT = Path(__file__).resolve().parent.parent / "rahma-traveler"
if str(SYSTEM_ROOT) not in sys.path:
    sys.path.insert(0, str(SYSTEM_ROOT))

from system_services.config import SystemServiceSettings  # noqa: E402
from system_services.unified_service import UnifiedCRMService  # noqa: E402


def create_trip_table(connection: sqlite3.Connection) -> None:
    connection.execute(
        """
        CREATE TABLE trips (
            trip_id TEXT PRIMARY KEY,
            trip_name TEXT,
            type TEXT,
            year INTEGER,
            trip_leader TEXT,
            start_date TEXT,
            end_date TEXT,
            sales_status TEXT,
            data_audit TEXT,
            trip_window_status TEXT,
            trip_availability_note TEXT,
            next_reengage_date TEXT,
            single_total INTEGER,
            double_total INTEGER,
            triple_total INTEGER,
            single_remaining INTEGER,
            double_remaining INTEGER,
            triple_remaining INTEGER,
            draft_holds_single INTEGER,
            draft_holds_double INTEGER,
            draft_holds_triple INTEGER,
            public_price TEXT,
            public_description TEXT,
            sales_notes TEXT
        )
        """
    )


class Phase1Phase2SystemAlignmentTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = Path(".tmp-test-workdirs")
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"system-align-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self._original_env = dict(os.environ)

    def tearDown(self) -> None:
        os.environ.clear()
        os.environ.update(self._original_env)
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def test_unified_service_filters_system_trip_availability(self) -> None:
        db_path = self.tmp_path / "system.db"
        with closing(sqlite3.connect(db_path)) as connection:
            create_trip_table(connection)
            connection.executemany(
                """
                INSERT INTO trips (
                    trip_id, trip_name, type, year, start_date, end_date, sales_status,
                    trip_availability_note, single_remaining, double_remaining, triple_remaining,
                    draft_holds_single, draft_holds_double, draft_holds_triple
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    ("TR-PAST", "Past Trip", "Local", 2026, "2026-01-01", "2026-01-03", "Open", "", 5, 0, 0, 0, 0, 0),
                    ("TR-CANCEL", "Cancelled Trip", "Local", 2026, "2026-09-01", "2026-09-03", "Cancelled", "", 5, 0, 0, 0, 0, 0),
                    ("TR-FULL", "Full Trip", "Local", 2026, "2026-09-01", "2026-09-03", "Open", "", 0, 0, 0, 0, 0, 0),
                    ("TR-OPEN", "System Fresh Trip", "Local", 2026, "2026-09-10", "2026-09-12", "Open", "", 2, 3, 1, 0, 0, 0),
                    ("TR-TBD", "Date TBD Trip", "Local", 2026, None, None, "Date TBD", "Waiting list open", 0, 0, 0, 0, 0, 0),
                ],
            )
            connection.commit()

        settings = SystemServiceSettings(
            repo_root=self.tmp_path,
            system_root=self.tmp_path,
            db_path=db_path,
            sheet_backend="excel",
            excel_source_workbook=None,
            excel_runtime_workbook=None,
            google_sheet_id="",
            google_application_credentials=None,
        )
        service = UnifiedCRMService(settings)

        result = service.build_trip_result("local", today=__import__("datetime").date(2026, 5, 18))
        self.assertEqual([trip["trip_id"] for trip in result["open_trips"]], ["TR-OPEN"])
        self.assertEqual([trip["trip_id"] for trip in result["date_tbd_trips"]], ["TR-TBD"])

    def test_trip_sync_updates_excel_workbooks(self) -> None:
        db_path = self.tmp_path / "system.db"
        with closing(sqlite3.connect(db_path)) as connection:
            create_trip_table(connection)
            connection.execute(
                """
                INSERT INTO trips (
                    trip_id, trip_name, type, year, trip_leader, start_date, end_date,
                    sales_status, data_audit, single_total, double_total, triple_total,
                    single_remaining, double_remaining, triple_remaining
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "TR-SYNC-1",
                    "Sheet Sync Trip",
                    "International",
                    2026,
                    "Nada",
                    "2026-11-01",
                    "2026-11-05",
                    "Open",
                    "SYNCED",
                    1,
                    2,
                    3,
                    1,
                    2,
                    3,
                ),
            )
            connection.commit()

        source = self.tmp_path / "source.xlsx"
        runtime = self.tmp_path / "runtime.xlsx"
        wb = Workbook()
        trips = wb.active
        trips.title = "Trips"
        trips["A2"] = "Trip ID"
        trips["B2"] = "Trip Name"
        trips["C2"] = "Type"
        trips["D2"] = "Year"
        trips["E2"] = "Trip Leader"
        trips["F2"] = "Start Date"
        trips["G2"] = "End Date"
        trips["Z2"] = "Sales Status"
        trips["AA2"] = "Data Audit"
        wb.save(source)
        wb.save(runtime)
        wb.close()

        settings = SystemServiceSettings(
            repo_root=self.tmp_path,
            system_root=self.tmp_path,
            db_path=db_path,
            sheet_backend="excel",
            excel_source_workbook=source,
            excel_runtime_workbook=runtime,
            google_sheet_id="",
            google_application_credentials=None,
        )
        service = UnifiedCRMService(settings)
        sync_result = service.sync_trip_to_sheet("TR-SYNC-1")
        self.assertEqual(sync_result["status"], "ok")

        for workbook_path in (source, runtime):
            synced = load_workbook(workbook_path, data_only=True)
            ws = synced["Trips"]
            self.assertEqual(ws["A3"].value, "TR-SYNC-1")
            self.assertEqual(ws["B3"].value, "Sheet Sync Trip")
            self.assertEqual(ws["C3"].value, "International")
            self.assertEqual(ws["F3"].value, "2026-11-01")
            self.assertEqual(ws["Z3"].value, "Open")
            synced.close()

    def test_demo_agent_uses_system_trip_catalog_over_stale_workbook_trips(self) -> None:
        db_path = self.tmp_path / "system.db"
        with closing(sqlite3.connect(db_path)) as connection:
            create_trip_table(connection)
            connection.execute(
                """
                INSERT INTO trips (
                    trip_id, trip_name, type, year, start_date, end_date, sales_status,
                    single_remaining, double_remaining, triple_remaining,
                    draft_holds_single, draft_holds_double, draft_holds_triple
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "TR-SYS-NEW",
                    "System Fresh Trip",
                    "Local",
                    2026,
                    "2026-09-20",
                    "2026-09-24",
                    "Open",
                    2,
                    2,
                    1,
                    0,
                    0,
                    0,
                ),
            )
            connection.commit()

        os.environ["RAHMA_SYSTEM_DB_PATH"] = str(db_path)

        source_workbook = self.tmp_path / "source.xlsx"
        runtime_workbook = self.tmp_path / "runtime.xlsx"
        wb = Workbook()
        travelers = wb.active
        travelers.title = "Travelers"
        travelers.append(["Status", "Traveler ID", "Full Name", "Code", "WhatsApp", "Loc. Trips", "Int. Trips", "Total trips", "Integrated WhatsApp", "Phone Lookup Key", "Data Audit"])

        trips = wb.create_sheet("Trips")
        trips["A2"] = "Trip ID"
        trips["B2"] = "Trip Name"
        trips["C2"] = "Type"
        trips["D2"] = "Year"
        trips["F2"] = "Start Date"
        trips["G2"] = "End Date"
        trips["Z2"] = "Sales Status"
        trips["AA2"] = "Data Audit"
        trips["A3"] = "TR-WB-OLD"
        trips["B3"] = "Workbook Old Trip"
        trips["C3"] = "Local"
        trips["D3"] = 2025
        trips["F3"] = "2025-01-01"
        trips["G3"] = "2025-01-03"
        trips["Z3"] = "Open"

        interactions = wb.create_sheet("Interactions")
        interactions.append(["Interaction ID", "Timestamp", "Channel", "Customer Name", "Raw Phone", "Integrated WhatsApp", "Phone Lookup Key", "Traveler ID", "Matched Row", "Status Snapshot", "Intent", "Trip Type", "Suggested Trips", "Action Taken", "Handoff Required", "Handoff Reason", "Agent Notes"])
        wb.save(source_workbook)
        wb.close()

        app = create_app(source_workbook=source_workbook, runtime_workbook=runtime_workbook)
        client = app.test_client()
        session = client.post("/api/session", json={}).get_json()["session"]
        session = client.post(
            f"/api/session/{session['id']}/intake",
            json={
                "fullName": "Sara Demo",
                "birthday": "1998-02-20",
                "gender": "Female",
                "nationality": "Egypt",
                "countryCode": "20",
                "rawPhone": "01012345678",
            },
        ).get_json()["session"]
        session = client.post(
            f"/api/session/{session['id']}/message",
            json={"text": "local"},
        ).get_json()["session"]

        assistant_text = session["messages"][-1]["text"]
        self.assertIn("System Fresh Trip", assistant_text)
        self.assertNotIn("Workbook Old Trip", assistant_text)


if __name__ == "__main__":
    unittest.main()
