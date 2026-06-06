from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.phase2_controlled_agent import run_phase2


class Phase2ControlledAgentTests(unittest.TestCase):
    def test_run_phase2_uses_highest_existing_traveler_id_plus_one(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase2-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "Code",
                    "WhatsApp",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Lead Source",
                    "Created At",
                    "Last Contacted At",
                    "Agent Notes",
                    "Data Audit",
                    "Normalized WhatsApp",
                ]
            )
            travelers.append(["", "TR00517", "Existing One", "20", "1000000001", None, None, None, "+201000000001", "20:1000000001", None, None, None, None, None, "+201000000001"])
            travelers.append([None] * 16)
            travelers.append(["", "TR00520", "Existing Later", "20", "1000000020", None, None, None, "+201000000020", "20:1000000020", None, None, None, None, None, "+201000000020"])

            trips = wb.create_sheet("Trips")
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Data Audit"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )
            wb.save(workbook_path)

            result = run_phase2(
                workbook_path=workbook_path,
                output_path=output_path,
                full_name="Newest Traveler",
                raw_phone="01099998888",
                country_code="20",
                trip_type="local",
                channel="web-demo",
                source="Web Demo",
                agent_notes="id regression",
            )

            self.assertEqual(result["write_result"]["created_traveler"]["traveler_id"], "TR00521")
            saved = load_workbook(output_path, data_only=False)
            travelers_ws = saved["Travelers"]
            ids = [travelers_ws.cell(row_idx, 2).value for row_idx in range(2, travelers_ws.max_row + 1)]
            self.assertEqual(ids.count("TR00521"), 1)
            self.assertEqual(ids.count("TR00520"), 1)
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)

    def test_run_phase2_creates_new_traveler_and_logs_interaction(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase2-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            traveler_headers = [
                "Status",
                "Traveler ID",
                "Full Name",
                "First Name",
                "Last Name",
                "Birthday",
                "Gender",
                "Nationality",
                "Code",
                "WhatsApp",
                "Email",
                "Community Whatsapp",
                "Residence",
                "Loc. Trips",
                "Int. Trips",
                "Total trips",
                "Comm. Events",
                "Lifetime Revenue",
                "Notes",
                "Introduce yourself",
                "Emergency Contact",
                "Emergency Phone",
                "Medical Notes",
                "Room Preference",
                "⭐ Rating (1–5)",
                "Integrated WhatsApp",
                "Normalized WhatsApp",
                "Phone Lookup Key",
                "Lead Source",
                "Created At",
                "Last Contacted At",
                "Agent Notes",
                "Data Audit",
            ]
            travelers.append(traveler_headers)
            travelers.append(
                [
                    "",
                    "TR00517",
                    "Allyson Jannel Fischer",
                    '=IFERROR(__xludf.DUMMYFUNCTION("REGEXEXTRACT(C2,""[A-Za-z]+"")"),"Allyson")',
                    '=RIGHT(C2,LEN(C2) - (FIND(" ",C2)))',
                    None,
                    "Female",
                    "America",
                    "1",
                    "5551234567",
                    None,
                    None,
                    None,
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B2,\'Trip Bookings\'!B:B,"RT-LOC*")',
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B2,\'Trip Bookings\'!B:B,"RT-INT*")',
                    "=N2+O2",
                    '=COUNTIFS(\'CE Bookings\'!D:D,B2,\'CE Bookings\'!D:D,"<>")',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "+15551234567",
                    "+15551234567",
                    "1:5551234567",
                    "DM",
                    "2026-05-11T18:00:00",
                    "2026-05-11T18:00:00",
                    None,
                    None,
                ]
            )
            travelers.append(
                [
                    None,
                    None,
                    None,
                    '=IFERROR(__xludf.DUMMYFUNCTION("REGEXEXTRACT(C3,""[A-Za-z]+"")"),"#N/A")',
                    '=RIGHT(C3,LEN(C3) - (FIND(" ",C3)))',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B3,\'Trip Bookings\'!B:B,"RT-LOC*")',
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B3,\'Trip Bookings\'!B:B,"RT-INT*")',
                    "=N3+O3",
                    '=COUNTIFS(\'CE Bookings\'!D:D,B3,\'CE Bookings\'!D:D,"<>")',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )

            trips = wb.create_sheet("Trips")
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Data Audit"
            trips["A3"] = "RT-LOC-26-001"
            trips["B3"] = "Siwa 3"
            trips["C3"] = "Local"
            trips["D3"] = 2026
            trips["Z3"] = "Date TBD"
            trips["AA3"] = "MISSING_END_DATE; MISSING_START_DATE"

            interactions = wb.create_sheet("Interactions")
            interaction_headers = [
                "Interaction ID",
                "Timestamp",
                "Channel",
                "Customer Name",
                "Raw Phone",
                "Integrated WhatsApp",
                "Phone Lookup Key",
                "Traveler ID",
                "Matched Row",
                "Status Snapshot",
                "Intent",
                "Trip Type",
                "Suggested Trips",
                "Action Taken",
                "Handoff Required",
                "Handoff Reason",
                "Agent Notes",
            ]
            interactions.append(interaction_headers)

            wb.save(workbook_path)

            result = run_phase2(
                workbook_path=workbook_path,
                output_path=output_path,
                full_name="Sara New",
                raw_phone="01012345678",
                country_code="",
                trip_type="local",
                channel="instagram",
                source="DM",
                agent_notes="phase2 test",
            )

            self.assertEqual(result["match_status"], "not_found")
            self.assertIsNotNone(result["write_result"]["created_traveler"])
            self.assertEqual(result["write_result"]["lead_update"]["lead_stage"], "Follow Up Needed")

            saved = load_workbook(output_path, data_only=False)
            travelers_ws = saved["Travelers"]
            interactions_ws = saved["Interactions"]
            leads_ws = saved["Leads"]

            self.assertEqual(travelers_ws["B3"].value, "TR00518")
            self.assertEqual(travelers_ws["C3"].value, "Sara New")
            self.assertEqual(travelers_ws["I3"].value, "20")
            self.assertEqual(travelers_ws["J3"].value, "1012345678")
            self.assertEqual(travelers_ws["Z3"].value, "+201012345678")
            self.assertEqual(travelers_ws["AB3"].value, "20:1012345678")
            self.assertTrue(str(travelers_ws["D3"].value).startswith("="))
            self.assertTrue(str(travelers_ws["N3"].value).startswith("="))

            self.assertTrue(str(interactions_ws["A2"].value).startswith("INT-"))
            self.assertEqual(interactions_ws["D2"].value, "Sara New")
            self.assertEqual(interactions_ws["H2"].value, "TR00518")
            self.assertEqual(leads_ws["A2"].value, "LD00001")
            self.assertEqual(leads_ws["D2"].value, "Sara New")
            self.assertEqual(leads_ws["L2"].value, "Follow Up Needed")
            self.assertEqual(leads_ws["R2"].value, "Medium")
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)

    def test_run_phase2_duplicate_phone_logs_without_creating(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase2-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "Code",
                    "WhatsApp",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Lead Source",
                    "Created At",
                    "Last Contacted At",
                    "Agent Notes",
                    "Data Audit",
                    "Normalized WhatsApp",
                ]
            )
            travelers.append(["Blacklisted", "TR00040", "Ahmed naim", "20", "1277445335", None, None, None, "+201277445335", "20:1277445335", None, None, None, None, "DUPLICATE_PHONE", "+201277445335"])
            travelers.append(["", "TR00354", "Ahmed Naim", "20", "1277445335", None, None, None, "+201277445335", "20:1277445335", None, None, None, None, "DUPLICATE_PHONE", "+201277445335"])

            trips = wb.create_sheet("Trips")
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Data Audit"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )
            wb.save(workbook_path)

            result = run_phase2(
                workbook_path=workbook_path,
                output_path=output_path,
                full_name="Ahmed Naim",
                raw_phone="1277445335",
                country_code="20",
                trip_type="international",
                channel="facebook",
                source="DM",
                agent_notes="duplicate test",
            )

            self.assertEqual(result["match_status"], "multiple_matches")
            self.assertTrue(result["handoff_required"])
            self.assertIsNone(result["write_result"]["created_traveler"])
            self.assertEqual(result["write_result"]["lead_update"]["lead_stage"], "Needs Review")

            saved = load_workbook(output_path, data_only=False)
            travelers_ws = saved["Travelers"]
            interactions_ws = saved["Interactions"]
            leads_ws = saved["Leads"]
            self.assertEqual(travelers_ws.max_row, 3)
            self.assertEqual(interactions_ws["O2"].value, "Yes")
            self.assertEqual(interactions_ws["P2"].value, "duplicate_phone_match")
            self.assertEqual(leads_ws["L2"].value, "Needs Review")
            self.assertEqual(leads_ws["S2"].value, "Urgent")
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)

    def test_run_phase2_blacklisted_customer_is_blocked_not_qualified(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase2-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "Code",
                    "WhatsApp",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Lead Source",
                    "Created At",
                    "Last Contacted At",
                    "Agent Notes",
                    "Data Audit",
                    "Normalized WhatsApp",
                ]
            )
            travelers.append(
                [
                    "Blacklisted",
                    "TR00040",
                    "Jane Doe",
                    "20",
                    "1005828000",
                    None,
                    None,
                    None,
                    "+201005828000",
                    "20:1005828000",
                    None,
                    None,
                    None,
                    None,
                    "",
                    "+201005828000",
                ]
            )

            trips = wb.create_sheet("Trips")
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Data Audit"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )
            wb.save(workbook_path)

            result = run_phase2(
                workbook_path=workbook_path,
                output_path=output_path,
                full_name="Different Name",
                raw_phone="1005828000",
                country_code="20",
                trip_type="local",
                channel="web-demo",
                source="Web Demo",
                agent_notes="blacklist test",
            )

            self.assertEqual(result["match_status"], "single_match")
            self.assertTrue(result["handoff_required"])
            self.assertEqual(result["handoff_reason"], "blacklisted_customer")
            self.assertIsNone(result["write_result"]["created_traveler"])
            self.assertEqual(result["write_result"]["lead_update"]["lead_stage"], "Blocked")
            self.assertEqual(result["write_result"]["lead_update"]["priority"], "Critical")

            saved = load_workbook(output_path, data_only=False)
            travelers_ws = saved["Travelers"]
            interactions_ws = saved["Interactions"]
            leads_ws = saved["Leads"]
            self.assertEqual(travelers_ws.max_row, 2)
            self.assertEqual(interactions_ws["H2"].value, "TR00040")
            self.assertEqual(interactions_ws["P2"].value, "blacklisted_customer")
            self.assertEqual(leads_ws["L2"].value, "Blocked")
            self.assertEqual(leads_ws["S2"].value, "Do Not Contact")
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)

    def test_run_phase2_same_phone_different_name_requires_identity_review(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase2-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "Birthday",
                    "Gender",
                    "Nationality",
                    "Code",
                    "WhatsApp",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Lead Source",
                    "Created At",
                    "Last Contacted At",
                    "Agent Notes",
                    "Data Audit",
                    "Normalized WhatsApp",
                ]
            )
            travelers.append(
                [
                    "",
                    "TR00212",
                    "Ahmed Selim",
                    "",
                    "",
                    "",
                    "20",
                    "1278727374",
                    None,
                    None,
                    None,
                    "+201278727374",
                    "20:1278727374",
                    "DM",
                    "2026-05-12T10:00:00",
                    "2026-05-12T10:00:00",
                    "",
                    "",
                    "+201278727374",
                ]
            )

            trips = wb.create_sheet("Trips")
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Data Audit"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )
            wb.save(workbook_path)

            result = run_phase2(
                workbook_path=workbook_path,
                output_path=output_path,
                full_name="Malak Yasser Mohamed",
                raw_phone="+201278727374",
                country_code="20",
                trip_type="local",
                channel="web-demo",
                source="Web Demo",
                agent_notes="identity conflict",
                birthday="2003-04-28",
                gender="Male",
                nationality="Egypt",
            )

            self.assertEqual(result["match_status"], "single_match")
            self.assertEqual(result["name_match_status"], "conflict")
            self.assertTrue(result["handoff_required"])
            self.assertEqual(result["handoff_reason"], "phone_name_conflict")
            self.assertIsNone(result["write_result"]["created_traveler"])
            self.assertEqual(result["write_result"]["lead_update"]["lead_stage"], "Needs Review")

            saved = load_workbook(output_path, data_only=False)
            travelers_ws = saved["Travelers"]
            interactions_ws = saved["Interactions"]
            leads_ws = saved["Leads"]
            lead_headers = {str(leads_ws.cell(1, col).value): col for col in range(1, leads_ws.max_column + 1) if leads_ws.cell(1, col).value}
            self.assertEqual(travelers_ws.max_row, 2)
            self.assertEqual(travelers_ws["C2"].value, "Ahmed Selim")
            self.assertIsNone(travelers_ws["D2"].value)
            self.assertEqual(interactions_ws["H2"].value, "TR00212")
            self.assertEqual(interactions_ws["P2"].value, "phone_name_conflict")
            self.assertEqual(leads_ws["L2"].value, "Needs Review")
            self.assertEqual(leads_ws.cell(2, lead_headers["Birthday"]).value, "2003-04-28")
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
