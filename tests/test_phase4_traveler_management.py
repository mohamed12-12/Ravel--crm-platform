from __future__ import annotations

import os
import shutil
import sqlite3
import sys
import unittest
import uuid
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook

SYSTEM_ROOT = Path(__file__).resolve().parent.parent / "apps" / "api"
if str(SYSTEM_ROOT) not in sys.path:
    sys.path.insert(0, str(SYSTEM_ROOT))

from services.crm.system_services.field_mapping import SHEET_TABLE_MAPPINGS  # noqa: E402
from services.crm.system_services.config import SystemServiceSettings  # noqa: E402
from services.crm.system_services.unified_service import UnifiedCRMService  # noqa: E402


def create_app_db_schema(app) -> None:
    from app.extensions import db

    with app.app_context():
        db.create_all()


def seed_traveler_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Travelers"
    headers = list(SHEET_TABLE_MAPPINGS["Travelers"]["columns"].values())
    ws.append(headers)
    wb.save(path)
    wb.close()


class Phase4TravelerManagementTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = Path(".tmp-test-workdirs")
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"phase4-travelers-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self._original_env = dict(os.environ)

    def tearDown(self) -> None:
        os.environ.clear()
        os.environ.update(self._original_env)
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def _build_app(self):
        db_path = self.tmp_path / "system.db"
        workbook_path = self.tmp_path / "runtime.xlsx"
        seed_traveler_workbook(workbook_path)

        os.environ["DATABASE_URL"] = f"sqlite:///{db_path.resolve().as_posix()}"
        os.environ["RAHMA_SYSTEM_DB_PATH"] = str(db_path)
        os.environ["SHEET_BACKEND"] = "excel"
        os.environ["EXCEL_RUNTIME_WORKBOOK"] = str(workbook_path)
        os.environ["EXCEL_SOURCE_WORKBOOK"] = str(workbook_path)
        os.environ["CRM_SHEET_MIRROR_ENABLED"] = "true"
        os.environ["EXCEL_EXPORT_ENABLED"] = "true"

        for module_name in list(sys.modules):
            if module_name == "app" or module_name.startswith("app."):
                sys.modules.pop(module_name, None)

        from app import create_app
        from app.extensions import db
        from app.models.booking import CEBooking, TripBooking
        from app.models.booking_event import BookingEventTrail
        from app.models.handoff import HandoffQueue
        from app.models.interaction import Interaction
        from app.models.lead import Lead
        from app.models.traveler import Traveler
        from app.models.traveler_document import TravelerDocument

        self.service = UnifiedCRMService()
        self.db = db
        self.Traveler = Traveler
        self.Lead = Lead
        self.TripBooking = TripBooking
        self.CEBooking = CEBooking
        self.Interaction = Interaction
        self.HandoffQueue = HandoffQueue
        self.BookingEventTrail = BookingEventTrail
        self.TravelerDocument = TravelerDocument

        app = create_app()
        create_app_db_schema(app)
        uri = app.config["SQLALCHEMY_DATABASE_URI"]
        resolved_db_path = Path(uri.replace("sqlite:////", "").replace("sqlite:///", "", 1))
        return app, resolved_db_path, workbook_path

    def test_create_uses_true_max_traveler_id_and_syncs_sheet(self) -> None:
        app, db_path, workbook_path = self._build_app()
        client = app.test_client()

        with app.app_context():
            self.db.session.add_all(
                [
                    self.Traveler(traveler_id="TR00009", full_name="Old Nine", status="New"),
                    self.Traveler(traveler_id="TR00150", full_name="Old Max", status="VIP"),
                    self.Traveler(traveler_id="BAD-ID", full_name="Broken Id", status="New"),
                ]
            )
            self.db.session.commit()

        response = client.post(
            "/travelers/",
            data={
                "full_name": "Sara Traveler",
                "whatsapp_raw": "01012345678",
                "email": "sara@example.com",
                "nationality": "Egyptian",
                "residence": "Cairo",
                "lead_source": "Referral",
                "status": "New",
            },
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 302)

        with app.app_context():
            traveler = self.Traveler.query.filter_by(full_name="Sara Traveler").first()
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler.traveler_id, "TR00151")
            self.assertEqual(traveler.integrated_whatsapp, "+201012345678")

        wb = load_workbook(workbook_path, data_only=True)
        ws = wb["Travelers"]
        headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
        self.assertEqual(ws.cell(2, headers["Traveler ID"]).value, "TR00151")
        self.assertEqual(ws.cell(2, headers["Full Name"]).value, "Sara Traveler")
        self.assertEqual(ws.cell(2, headers["WhatsApp"]).value, "1012345678")
        self.assertEqual(ws.cell(2, headers["Integrated WhatsApp"]).value, "+201012345678")
        wb.close()

        self.assertTrue(db_path.exists())

    def test_database_max_id_wins_over_higher_workbook_ids(self) -> None:
        app, db_path, workbook_path = self._build_app()

        wb = load_workbook(workbook_path)
        try:
            ws = wb["Travelers"]
            ws.append(["TR00909", "Active", "Workbook Higher"] + [None] * 24)
            wb.save(workbook_path)
        finally:
            wb.close()

        with app.app_context():
            self.db.session.add(
                self.Traveler(traveler_id="TR00585", full_name="DB Max", status="Active")
            )
            self.db.session.commit()

        self.assertEqual(self.service.next_traveler_id(), "TR00586")

        created = self.service.create_traveler(
            full_name="Database Priority",
            raw_phone="01022223333",
            country_code="20",
        )
        self.assertEqual(created["traveler_id"], "TR00586")

        with app.app_context():
            traveler = self.Traveler.query.filter_by(traveler_id="TR00586").first()
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler.full_name, "Database Priority")

    def test_detail_shows_related_records_and_update_syncs_back(self) -> None:
        app, db_path, workbook_path = self._build_app()
        client = app.test_client()

        with app.app_context():
            traveler = self.Traveler(
                traveler_id="TR00200",
                full_name="Profile Traveler",
                status="New",
                whatsapp_raw="1000000000",
                integrated_whatsapp="+201000000000",
                normalized_whatsapp="+201000000000",
                phone_lookup_key="2010000000",
                nationality="Egyptian",
                residence="Giza",
                lead_source="Web Demo",
                preferred_currency="EGP",
                lifetime_revenue=100.0,
            )
            lead = self.Lead(
                lead_id="LD20001",
                customer_name="Profile Traveler",
                raw_phone="1000000000",
                integrated_whatsapp="+201000000000",
                phone_lookup_key="2010000000",
                traveler_id="TR00200",
                lead_stage="Qualified",
                lead_source="Web Demo",
                channel="WhatsApp",
                priority="High",
                follow_up_status="Awaiting reply",
            )
            booking = self.TripBooking(
                booking_id="B20001",
                trip_id="RT-LOC-26-900",
                trip_name="Desert Escape",
                traveler_id="TR00200",
                traveler_name="Profile Traveler",
                room_type="Double",
                flight_option="Without Flights",
                booking_status="Draft",
                payment_status="Pending",
                booking_source="System UI",
            )
            interaction = self.Interaction(
                interaction_id="INT20001",
                timestamp=__import__("datetime").datetime.now(__import__("datetime").timezone.utc),
                channel="WhatsApp",
                customer_name="Profile Traveler",
                raw_phone="1000000000",
                integrated_whatsapp="+201000000000",
                phone_lookup_key="2010000000",
                traveler_id="TR00200",
                status_snapshot="Qualified",
                intent="Trip follow-up",
                trip_type="Local",
                suggested_trips="RT-LOC-26-900",
                action_taken="Follow-up sent",
                handoff_required=False,
                handoff_reason="",
                agent_notes="Ready for booking",
            )
            handoff = self.HandoffQueue(
                handoff_id="H20001",
                lead_id="LD20001",
                traveler_id="TR00200",
                trip_id="RT-LOC-26-900",
                flow_key="booking_flow",
                reason="Needs manager review",
                priority="Medium",
                channel="WhatsApp",
                status="Pending",
            )
            event = self.BookingEventTrail(
                event_id="EV20001",
                event_type="lead_qualified",
                event_label="Lead qualified",
                traveler_id="TR00200",
                lead_id="LD20001",
                booking_id="B20001",
                trip_id="RT-LOC-26-900",
                interaction_id="INT20001",
                channel="WhatsApp",
                actor="system-ui",
                notes="Qualified for booking",
                metadata_json='{"source":"test"}',
            )
            self.db.session.add_all([traveler, lead, booking, interaction, handoff, event])
            self.db.session.commit()

        detail_response = client.get("/travelers/TR00200")
        self.assertEqual(detail_response.status_code, 200)
        body = detail_response.get_data(as_text=True)
        self.assertIn("LD20001", body)
        self.assertIn("B20001", body)
        self.assertIn("INT20001", body)
        self.assertIn("Needs manager review", body)
        self.assertIn("Lifetime Revenue (EGP)", body)
        self.assertIn("EGP 0.00", body)
        self.assertIn("Lifetime Revenue (USD): $0.00", body)
        self.assertIn("1 USD = 50.00 EGP", body)
        self.assertIn("Preferred Payment Currency", body)

        update_response = client.put(
            "/travelers/TR00200",
            json={
                "full_name": "Profile Traveler Updated",
                "status": "VIP",
                "whatsapp_raw": "01099998888",
                "nationality": "Egyptian",
                "residence": "Cairo",
                "lead_source": "Referral",
                "preferred_currency": "egp",
                "rating": 4.8,
            },
        )
        self.assertEqual(update_response.status_code, 200)

        explicit_service = UnifiedCRMService(
            SystemServiceSettings(
                repo_root=self.tmp_path,
                system_root=self.tmp_path,
                db_path=db_path,
                sheet_backend="excel",
                excel_source_workbook=workbook_path,
                excel_runtime_workbook=workbook_path,
                google_sheet_id="",
                google_application_credentials=None,
            )
        )
        explicit_service.sync_record_to_sheet("Travelers", "TR00200")

        with app.app_context():
            traveler = self.db.session.get(self.Traveler, "TR00200")
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler.full_name, "Profile Traveler Updated")
            self.assertEqual(traveler.status, "VIP")
            self.assertEqual(traveler.integrated_whatsapp, "+201099998888")
            self.assertEqual(traveler.preferred_currency, "EGP")

        wb = load_workbook(workbook_path, data_only=True)
        ws = wb["Travelers"]
        headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
        self.assertEqual(ws.cell(2, headers["Full Name"]).value, "Profile Traveler Updated")
        self.assertEqual(ws.cell(2, headers["Status"]).value, "VIP")
        self.assertEqual(ws.cell(2, headers["Integrated WhatsApp"]).value, "+201099998888")
        self.assertEqual(ws.cell(2, headers["Preferred Currency"]).value, "EGP")
        wb.close()

    def test_index_hides_blank_travelers_from_active_list(self) -> None:
        app, db, workbook_path = self._build_app()
        client = app.test_client()

        with app.app_context():
            from app.models import Traveler

            self.db.session.add(
                Traveler(
                    traveler_id="TR00999",
                    full_name="",
                    first_name="#N/A",
                    last_name="#VALUE!",
                    whatsapp_raw="",
                    phone_code="",
                    integrated_whatsapp="",
                    normalized_whatsapp="",
                    phone_lookup_key="",
                    email="",
                    nationality="",
                    residence="",
                    status="Active",
                )
            )
            self.db.session.commit()

        response = client.get("/travelers/")
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("TR00999", response.get_data(as_text=True))


    def test_index_hides_archived_and_inactive_travelers_by_default_and_shows_filter(self) -> None:
        app, db, workbook_path = self._build_app()
        client = app.test_client()

        with app.app_context():
            from app.models import Traveler

            self.db.session.add_all([
                Traveler(
                    traveler_id="TR00880",
                    full_name="Archived Traveler",
                    whatsapp_raw="1000000880",
                    integrated_whatsapp="+201000000880",
                    normalized_whatsapp="+201000000880",
                    phone_lookup_key="20:1000000880",
                    status="Archived",
                ),
                Traveler(
                    traveler_id="TR00881",
                    full_name="Inactive Traveler",
                    whatsapp_raw="1000000881",
                    integrated_whatsapp="+201000000881",
                    normalized_whatsapp="+201000000881",
                    phone_lookup_key="20:1000000881",
                    status="Inactive",
                ),
                Traveler(
                    traveler_id="TR00882",
                    full_name="Blacklisted Traveler",
                    whatsapp_raw="1000000882",
                    integrated_whatsapp="+201000000882",
                    normalized_whatsapp="+201000000882",
                    phone_lookup_key="20:1000000882",
                    status="Blacklisted",
                ),
                Traveler(
                    traveler_id="TR00883",
                    full_name="Blocked Traveler",
                    whatsapp_raw="1000000883",
                    integrated_whatsapp="+201000000883",
                    normalized_whatsapp="+201000000883",
                    phone_lookup_key="20:1000000883",
                    status="Blocked",
                ),
            ])
            self.db.session.commit()

        default_response = client.get("/travelers/")
        default_body = default_response.get_data(as_text=True)
        self.assertNotIn("TR00880", default_body)
        self.assertNotIn("TR00881", default_body)
        self.assertNotIn("TR00882", default_body)
        self.assertNotIn("TR00883", default_body)

        archived_response = client.get("/travelers/?status=Archived")
        archived_body = archived_response.get_data(as_text=True)
        self.assertIn("TR00880", archived_body)
        self.assertIn("TR00881", archived_body)
        self.assertNotIn("TR00882", archived_body)
        self.assertNotIn("TR00883", archived_body)
        self.assertIn("Inactive", archived_body)
        self.assertIn("Blacklist / Blocked", archived_body)
        self.assertNotIn(">Blocked<", archived_body)

        inactive_response = client.get("/travelers/?status=Inactive")
        inactive_body = inactive_response.get_data(as_text=True)
        self.assertIn("TR00881", inactive_body)
        self.assertNotIn("TR00880", inactive_body)

        blacklist_response = client.get("/travelers/?status=Blacklisted")
        blacklist_body = blacklist_response.get_data(as_text=True)
        self.assertIn("TR00882", blacklist_body)
        self.assertIn("TR00883", blacklist_body)
        self.assertNotIn("TR00880", blacklist_body)
        self.assertNotIn("TR00881", blacklist_body)

    def test_edit_traveler_preserves_id_and_blocks_phone_conflict(self) -> None:
        app, db, workbook_path = self._build_app()
        client = app.test_client()

        with app.app_context():
            self.db.session.add_all([
                self.Traveler(
                    traveler_id="TR01000",
                    full_name="Primary Traveler",
                    status="Active",
                    whatsapp_raw="1000001000",
                    integrated_whatsapp="+201000001000",
                    normalized_whatsapp="+201000001000",
                    phone_lookup_key="20:1000001000",
                    email="primary@example.com",
                ),
                self.Traveler(
                    traveler_id="TR01001",
                    full_name="Conflicting Traveler",
                    status="Active",
                    whatsapp_raw="1000001001",
                    integrated_whatsapp="+201000001001",
                    normalized_whatsapp="+201000001001",
                    phone_lookup_key="20:1000001001",
                    email="conflict@example.com",
                ),
            ])
            self.db.session.commit()

        response = client.put(
            "/travelers/TR01000",
            json={
                "full_name": "Primary Traveler Updated",
                "status": "VIP",
                "whatsapp_raw": "01000001001",
                "email": "primary.updated@example.com",
                "nationality": "Egyptian",
                "residence": "Cairo",
                "notes": "Updated safely",
            },
        )
        self.assertEqual(response.status_code, 400)
        payload = response.get_json()
        self.assertIn("already belongs to traveler TR01001", payload["error"])

        with app.app_context():
            traveler = self.db.session.get(self.Traveler, "TR01000")
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler.traveler_id, "TR01000")
            self.assertEqual(traveler.full_name, "Primary Traveler")
            self.assertEqual(traveler.status, "Active")

    def test_delete_traveler_hard_deletes_related_records_and_reuses_highest_id(self) -> None:
        app, db_path, workbook_path = self._build_app()
        client = app.test_client()
        uploads_root = self.tmp_path / "uploads"
        traveler_dir = uploads_root / "TR00151"
        traveler_dir.mkdir(parents=True, exist_ok=True)
        passport_file = traveler_dir / "passport.pdf"
        passport_file.write_bytes(b"passport")
        app.config["TRAVELER_UPLOAD_ROOT"] = str(uploads_root)

        with app.app_context():
            self.db.session.add(self.Traveler(traveler_id="TR00150", full_name="Keep Traveler", status="Active"))
            traveler = self.Traveler(
                traveler_id="TR00151",
                full_name="Delete Traveler",
                status="Active",
                last_lead_id="LD00151",
            )
            lead = self.Lead(
                lead_id="LD00151",
                customer_name="Delete Traveler",
                traveler_id="TR00151",
                lead_stage="Qualified",
            )
            booking = self.TripBooking(
                booking_id="B00151",
                trip_id="RT-LOC-26-900",
                trip_name="Delete Trip",
                traveler_id="TR00151",
                traveler_name="Delete Traveler",
                booking_status="Draft",
                lead_id="LD00151",
            )
            ce_booking = self.CEBooking(
                booking_id="CE00151",
                event_id="EVT-1",
                event_name="Camp",
                traveler_id="TR00151",
                traveler_name="Delete Traveler",
                status="Draft",
            )
            interaction = self.Interaction(
                interaction_id="INT00151",
                traveler_id="TR00151",
                customer_name="Delete Traveler",
                raw_phone="01011111111",
            )
            handoff = self.HandoffQueue(
                handoff_id="H00151",
                traveler_id="TR00151",
                lead_id="LD00151",
                reason="delete-test",
            )
            event = self.BookingEventTrail(
                event_id="BE00151",
                traveler_id="TR00151",
                lead_id="LD00151",
                booking_id="B00151",
                interaction_id="INT00151",
                event_type="test",
                event_label="Test Event",
            )
            document = self.TravelerDocument(
                traveler_id="TR00151",
                category="passport",
                file_name="passport.pdf",
                original_file_name="passport.pdf",
                mime_type="application/pdf",
                file_extension="pdf",
                file_size=8,
                storage_path=str(passport_file),
            )
            self.db.session.add_all([traveler, lead, booking, ce_booking, interaction, handoff, event, document])
            self.db.session.commit()
            self.service.sync_record_to_sheet("Travelers", "TR00150")
            self.service.sync_record_to_sheet("Travelers", "TR00151")
            self.service.sync_record_to_sheet("Leads", "LD00151")

        response = client.delete("/travelers/TR00151")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["message"], "Traveler deleted permanently")

        with app.app_context():
            self.assertIsNone(self.db.session.get(self.Traveler, "TR00151"))
            self.assertIsNone(self.db.session.get(self.Lead, "LD00151"))
            self.assertIsNone(self.db.session.get(self.TripBooking, "B00151"))
            self.assertIsNone(self.db.session.get(self.CEBooking, "CE00151"))
            self.assertIsNone(self.db.session.get(self.Interaction, "INT00151"))
            self.assertIsNone(self.db.session.get(self.HandoffQueue, "H00151"))
            self.assertIsNone(self.db.session.get(self.BookingEventTrail, "BE00151"))
            self.assertEqual(self.service.next_traveler_id(), "TR00151")

        self.assertFalse(passport_file.exists())

        wb = load_workbook(workbook_path, data_only=True)
        ws = wb["Travelers"]
        headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
        traveler_ids = [
            str(ws.cell(row, headers["Traveler ID"]).value or "").strip()
            for row in range(2, ws.max_row + 1)
        ]
        wb.close()
        self.assertIn("TR00150", traveler_ids)
        self.assertNotIn("TR00151", traveler_ids)

        create_response = client.post(
            "/travelers/",
            data={
                "full_name": "Replacement Traveler",
                "whatsapp_raw": "01022221111",
                "nationality": "Egyptian",
                "status": "New",
            },
            follow_redirects=False,
        )
        self.assertEqual(create_response.status_code, 302)

        with app.app_context():
            replacement = self.db.session.get(self.Traveler, "TR00151")
            self.assertIsNotNone(replacement)
            self.assertEqual(replacement.full_name, "Replacement Traveler")

    def test_index_page_offers_a_delete_action_per_traveler(self) -> None:
        """The list/dashboard page previously had no delete affordance at
        all -- only the detail page did -- so a customer-facing traveler
        could only be removed by first navigating into their profile. This
        confirms the list page itself now offers delete, wired to the same
        DELETE /travelers/<id> endpoint the detail page already used.
        """
        app, db, workbook_path = self._build_app()
        client = app.test_client()

        with app.app_context():
            self.db.session.add(self.Traveler(traveler_id="TR00160", full_name="List Delete Traveler", status="Active"))
            self.db.session.commit()

        response = client.get("/travelers/")
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("deleteTravelerFromList", html)
        self.assertIn("TR00160", html)

        delete_response = client.delete("/travelers/TR00160")
        self.assertEqual(delete_response.status_code, 200)
        with app.app_context():
            self.assertIsNone(self.db.session.get(self.Traveler, "TR00160"))


if __name__ == "__main__":
    unittest.main()
