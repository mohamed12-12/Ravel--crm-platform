from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.phaseA_crm_integrity_audit import CRM_AUDIT_SHEET, run_crm_integrity_audit


class PhaseACrmIntegrityAuditTests(unittest.TestCase):
    def test_audit_reports_true_max_duplicates_and_missing_lookup_keys(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phaseA-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            input_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"
            report_path = tmp_path / "report.md"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "Code",
                    "WhatsApp",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                ]
            )
            travelers.append(["", "TR00001", "Alice One", "20", "1000000001", "+201000000001", "20:1000000001"])
            travelers.append(["", "TR00003", "Bob Missing Phone", "", "", "", ""])
            travelers.append(["", "TR00003", "Bob Duplicate", "20", "1000000001", "+201000000001", "20:1000000001"])
            travelers.append(["", "BAD-ID", "Bad Id", "20", "1000000002", "+201000000002", "20:1000000002"])
            travelers.append(["", "", "Phone No Lookup", "20", "1000000003", "", ""])
            travelers.append(["", "TR00010", "Hidden Max", "20", "1000000010", "+201000000010", "20:1000000010"])
            travelers.row_dimensions[7].hidden = True
            travelers.auto_filter.ref = "A1:G7"
            wb.save(input_path)

            summary = run_crm_integrity_audit(input_path, output_path, report_path)

            self.assertEqual(summary["true_max_traveler_id"], "TR00010")
            self.assertEqual(summary["true_max_traveler_id_row"], 7)
            self.assertEqual(summary["next_traveler_id"], "TR00011")
            self.assertEqual(summary["duplicate_traveler_id_keys"], 1)
            self.assertEqual(summary["duplicate_phone_lookup_keys"], 1)
            self.assertEqual(summary["invalid_traveler_id_count"], 1)
            self.assertEqual(summary["rows_with_id_missing_name_or_phone"], 1)
            self.assertEqual(summary["rows_phone_exists_lookup_key_missing"], 1)
            self.assertEqual(summary["hidden_rows_with_data"], 1)
            self.assertEqual(summary["visible_rows_are_authoritative"], "No")

            audited = load_workbook(output_path)
            self.assertIn(CRM_AUDIT_SHEET, audited.sheetnames)
            audit_ws = audited[CRM_AUDIT_SHEET]
            values = [
                audit_ws.cell(row_idx, col_idx).value
                for row_idx in range(1, audit_ws.max_row + 1)
                for col_idx in range(1, audit_ws.max_column + 1)
            ]
            self.assertIn("DUPLICATE_TRAVELER_ID", values)
            self.assertIn("DUPLICATE_PHONE_LOOKUP_KEY", values)
            self.assertIn("PHONE_EXISTS_BUT_LOOKUP_KEY_MISSING", values)
            self.assertTrue(report_path.exists())
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
