from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook

from app.sheets.sheets_adapter import FakeWorkbook, SheetRowAdapter
from scripts.phase4_sales_intelligence import compute_sales_dashboard, ensure_leads_sheet, upsert_lead


class Phase4SalesIntelligenceTests(unittest.TestCase):
    def test_ensure_leads_sheet_supports_google_sheet_adapter(self) -> None:
        wb = FakeWorkbook({"Leads": SheetRowAdapter([[]], title="Leads")})
        leads_ws, headers = ensure_leads_sheet(wb)

        self.assertEqual(leads_ws.cell(1, headers["Lead ID"]).value, "Lead ID")
        self.assertNotIn("auto_filter", dir(leads_ws))

    def test_dashboard_counts_lead_stages_and_follow_ups(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase4-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "phase4.xlsx"
            wb = Workbook()
            leads_ws = wb.active
            leads_ws.title = "Sheet"
            wb.remove(leads_ws)
            leads_ws, headers = ensure_leads_sheet(wb)

            result_qualified = {
                "match_status": "not_found",
                "handoff_required": False,
                "handoff_reason": "",
                "trip_result": {"open_trips": [{"trip_id": "RT-LOC-26-001"}], "date_tbd_trips": []},
                "traveler": None,
            }
            result_review = {
                "match_status": "multiple_matches",
                "handoff_required": True,
                "handoff_reason": "duplicate_phone_match",
                "trip_result": {"open_trips": [], "date_tbd_trips": []},
                "traveler": None,
            }

            from datetime import datetime

            upsert_lead(
                leads_ws,
                headers,
                result=result_qualified,
                customer_name="Sara",
                raw_phone="01012345678",
                phone_lookup_key="20:1012345678",
                integrated_whatsapp="+201012345678",
                traveler_id="TR00518",
                channel="demo",
                source="Web Demo",
                trip_type="local",
                interaction_id="INT-20260511-0001",
                timestamp=datetime.now(),
                notes="qualified",
            )
            upsert_lead(
                leads_ws,
                headers,
                result=result_review,
                customer_name="Ahmed",
                raw_phone="1277445335",
                phone_lookup_key="20:1277445335",
                integrated_whatsapp="+201277445335",
                traveler_id="",
                channel="demo",
                source="Web Demo",
                trip_type="international",
                interaction_id="INT-20260511-0002",
                timestamp=datetime.now(),
                notes="review",
            )

            wb.save(workbook_path)
            stats = compute_sales_dashboard(workbook_path)

            self.assertEqual(stats["leadCount"], 2)
            self.assertEqual(stats["leadStageCounts"]["Qualified"], 1)
            self.assertEqual(stats["leadStageCounts"]["Needs Review"], 1)
            self.assertEqual(stats["pipeline"]["qualified"], 1)
            self.assertEqual(stats["pipeline"]["needsReview"], 1)
            self.assertEqual(stats["followUpSummary"]["urgent"], 1)
            self.assertEqual(len(stats["recentLeads"]), 2)
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
