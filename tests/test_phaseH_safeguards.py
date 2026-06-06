from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path
from datetime import date

from openpyxl import Workbook, load_workbook

from scripts.phase2_controlled_agent import run_phase2, generate_next_traveler_id
from scripts.phase1_readonly_agent import build_agent_response, recommend_trips, TripRecord

class PhaseHSafeguardsTests(unittest.TestCase):
    def setUp(self):
        self.tmp_root = Path(".tmp-test-workdirs")
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"phaseH-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self.workbook_path = self.tmp_path / "input.xlsx"
        self.output_path = self.tmp_path / "output.xlsx"

    def tearDown(self):
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def create_base_workbook(self):
        wb = Workbook()
        travelers = wb.active
        travelers.title = "Travelers"
        traveler_headers = [
            "Status", "Traveler ID", "Full Name", "Code", "WhatsApp", 
            "Loc. Trips", "Int. Trips", "Total trips", 
            "Integrated WhatsApp", "Phone Lookup Key", "Lead Source", 
            "Created At", "Last Contacted At", "Agent Notes", "Data Audit", 
            "Normalized WhatsApp"
        ]
        travelers.append(traveler_headers)
        
        trips = wb.create_sheet("Trips")
        # Header row 1 is empty, row 2 has headers
        trips.append([])
        trips.append([
            "Trip ID", "Trip Name", "Type", "Year", "Start Date", "End Date", 
            "Sales Status", "Data Audit", "Remaining Single", "Remaining Double", 
            "Remaining Triple", "Draft Holds Single", "Draft Holds Double", "Draft Holds Triple"
        ])
        
        interactions = wb.create_sheet("Interactions")
        interactions.append([
            "Interaction ID", "Timestamp", "Channel", "Customer Name", "Raw Phone",
            "Integrated WhatsApp", "Phone Lookup Key", "Traveler ID", "Matched Row",
            "Status Snapshot", "Intent", "Trip Type", "Suggested Trips", 
            "Action Taken", "Handoff Required", "Handoff Reason", "Agent Notes"
        ])
        return wb

    def test_true_max_traveler_id_from_gaps(self):
        wb = self.create_base_workbook()
        ws = wb["Travelers"]
        # Add some IDs with gaps and "hidden" style (just non-sequential rows)
        ws.cell(row=5, column=2).value = "TR00100"
        ws.cell(row=10, column=2).value = "TR00050"
        ws.cell(row=15, column=2).value = "TR00150"
        
        # Manually verify max calculation
        header_map = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
        next_id = generate_next_traveler_id(ws, header_map)
        self.assertEqual(next_id, "TR00151")

    def test_duplicate_traveler_ids_check(self):
        # The current system doesn't explicitly 'fail' on duplicates in the sheet, 
        # but generate_next_traveler_id handles them by avoiding any existing ID.
        wb = self.create_base_workbook()
        ws = wb["Travelers"]
        ws.cell(row=2, column=2).value = "TR00001"
        ws.cell(row=3, column=2).value = "TR00001" # Duplicate
        
        header_map = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
        next_id = generate_next_traveler_id(ws, header_map)
        self.assertEqual(next_id, "TR00002")

    def test_new_customer_gets_max_plus_one(self):
        wb = self.create_base_workbook()
        ws = wb["Travelers"]
        ws.cell(row=2, column=2).value = "TR00010"
        wb.save(self.workbook_path)
        
        result = run_phase2(
            workbook_path=self.workbook_path,
            output_path=self.output_path,
            full_name="New Guy",
            raw_phone="0123456789",
            country_code="20",
            trip_type="Local",
            channel="test",
            source="test",
            agent_notes="test"
        )
        self.assertEqual(result["write_result"]["created_traveler"]["traveler_id"], "TR00011")

    def test_existing_customer_reuses_id(self):
        wb = self.create_base_workbook()
        ws = wb["Travelers"]
        ws.cell(row=2, column=2).value = "TR00010"
        ws.cell(row=2, column=3).value = "Old Guy"
        ws.cell(row=2, column=10).value = "20:0123456789"
        wb.save(self.workbook_path)
        
        # This will match the existing traveler
        result = build_agent_response(
            workbook_path=self.workbook_path,
            full_name="Old Guy",
            raw_phone="0123456789",
            country_code="20",
            trip_type="Local"
        )
        self.assertEqual(result["match_status"], "single_match")
        self.assertEqual(result["traveler"]["traveler_id"], "TR00010")

    def test_missing_trip_dates_not_confirmed(self):
        # Create a trip without dates
        trip = TripRecord(
            row=3,
            trip_id="T1",
            trip_name="T1 Name",
            trip_type="Local",
            year=2026,
            start_date=None,
            end_date=None,
            sales_status="Open",
            data_audit="",
            remaining_places=10,
            availability_status="available"
        )
        
        recs = recommend_trips([trip], "Local", today=date(2026, 1, 1))
        # Should be in tbd_trips, NOT open_trips
        self.assertEqual(len(recs["open_trips"]), 0)
        self.assertEqual(len(recs["date_tbd_trips"]), 1)
        self.assertEqual(recs["date_tbd_trips"][0]["trip_id"], "T1")

    def test_blacklisted_traveler_stopped(self):
        wb = self.create_base_workbook()
        ws = wb["Travelers"]
        ws.cell(row=2, column=1).value = "Blacklisted"
        ws.cell(row=2, column=2).value = "TR00001"
        ws.cell(row=2, column=3).value = "Bad Guy"
        ws.cell(row=2, column=10).value = "20:0123456789"
        wb.save(self.workbook_path)
        
        result = build_agent_response(
            workbook_path=self.workbook_path,
            full_name="Bad Guy",
            raw_phone="0123456789",
            country_code="20",
            trip_type="Local"
        )
        self.assertTrue(result["handoff_required"])
        self.assertEqual(result["handoff_reason"], "blacklisted_customer")

if __name__ == "__main__":
    unittest.main()
