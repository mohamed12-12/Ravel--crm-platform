from __future__ import annotations

import os
import shutil
import sys
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook

SYSTEM_ROOT = Path(__file__).resolve().parent.parent / "apps" / "api"
if str(SYSTEM_ROOT) not in sys.path:
    sys.path.insert(0, str(SYSTEM_ROOT))

from services.crm.system_services.field_mapping import SHEET_TABLE_MAPPINGS
from services.crm.system_services.config import resolve_system_db_path
from services.crm.system_services.unified_service import UnifiedCRMService

def create_app_db_schema(app) -> None:
    from app.extensions import db
    with app.app_context():
        db.create_all()

def seed_workbook(path: Path) -> None:
    wb = Workbook()
    
    # Create Travelers
    ws_travelers = wb.active
    ws_travelers.title = "Travelers"
    ws_travelers.append(list(SHEET_TABLE_MAPPINGS["Travelers"]["columns"].values()))
    
    # Create Leads
    ws_leads = wb.create_sheet("Leads")
    ws_leads.append(list(SHEET_TABLE_MAPPINGS["Leads"]["columns"].values()))
    
    # Create Interactions
    ws_interactions = wb.create_sheet("Interactions")
    ws_interactions.append(list(SHEET_TABLE_MAPPINGS["Interactions"]["columns"].values()))

    # Create Handoff Queue
    ws_handoff = wb.create_sheet("Handoff Queue")
    ws_handoff.append(list(SHEET_TABLE_MAPPINGS["Handoff Queue"]["columns"].values()))
    
    # Create Event Trail
    ws_events = wb.create_sheet("Booking Event Trail")
    ws_events.append(list(SHEET_TABLE_MAPPINGS["Booking Event Trail"]["columns"].values()))

    wb.save(path)
    wb.close()


class Phase5ManualLeadAgentLogicTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = Path(".tmp-test-workdirs")
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"phase5-manual-lead-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self._original_env = dict(os.environ)

    def tearDown(self) -> None:
        os.environ.clear()
        os.environ.update(self._original_env)
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def _build_app(self):
        db_path = self.tmp_path / "system.db"
        workbook_path = self.tmp_path / "runtime.xlsx"
        seed_workbook(workbook_path)

        os.environ["DATABASE_URL"] = f"sqlite:///{db_path.resolve().as_posix()}"
        os.environ["RAHMA_SYSTEM_DB_PATH"] = str(db_path)
        os.environ["SHEET_BACKEND"] = "excel"
        os.environ["EXCEL_RUNTIME_WORKBOOK"] = str(workbook_path)
        os.environ["EXCEL_SOURCE_WORKBOOK"] = str(workbook_path)

        for module_name in list(sys.modules):
            if module_name == "app" or module_name.startswith("app."):
                sys.modules.pop(module_name, None)

        from app import create_app
        from app.extensions import db
        from app.models.lead import Lead
        from app.models.traveler import Traveler
        from app.models.handoff import HandoffQueue

        self.db = db
        self.Traveler = Traveler
        self.Lead = Lead
        self.HandoffQueue = HandoffQueue

        app = create_app()
        app.config["TESTING"] = True
        create_app_db_schema(app)
        return app

    def test_manual_lead_creates_new_traveler_when_not_found(self):
        app = self._build_app()
        client = app.test_client()

        response = client.post(
            "/leads/",
            data={
                "customer_name": "New Manual Lead",
                "raw_phone": "201012345678",
                "lead_stage": "Interested",
                "priority": "Medium",
                "lead_source": "WhatsApp",
            },
        )
        self.assertEqual(response.status_code, 302)

        with app.app_context():
            # Verify Traveler was created
            travelers = self.Traveler.query.all()
            self.assertEqual(len(travelers), 1)
            self.assertEqual(travelers[0].full_name, "New Manual Lead")
            
            # Verify Lead was linked
            leads = self.Lead.query.all()
            self.assertEqual(len(leads), 1)
            self.assertEqual(leads[0].traveler_id, travelers[0].traveler_id)
            self.assertEqual(leads[0].lead_stage, "Interested")
            self.assertFalse(leads[0].handoff_required)
            
            # Verify Handoff Queue is empty
            handoffs = self.HandoffQueue.query.all()
            self.assertEqual(len(handoffs), 0)

    def test_manual_lead_creates_distinct_row_for_existing_traveler(self):
        app = self._build_app()
        client = app.test_client()

        with app.app_context():
            traveler = self.Traveler(
                traveler_id="TR00010",
                full_name="Existing Traveler",
                whatsapp_raw="1012345678",
                integrated_whatsapp="+201012345678",
                normalized_whatsapp="+201012345678",
                phone_lookup_key="20:1012345678",
            )
            existing_lead = self.Lead(
                lead_id="LD00001",
                customer_name="Existing Traveler",
                raw_phone="1012345678",
                integrated_whatsapp="+201012345678",
                phone_lookup_key="20:1012345678",
                traveler_id="TR00010",
                lead_stage="Contacted",
                lead_source="Manual Web UI",
                channel="System UI",
            )
            self.db.session.add_all([traveler, existing_lead])
            self.db.session.commit()

        response = client.post(
            "/leads/",
            data={
                "customer_name": "Existing Traveler",
                "raw_phone": "201012345678",
                "lead_stage": "Qualified",
                "priority": "High",
                "lead_source": "WhatsApp",
            },
        )
        self.assertEqual(response.status_code, 302)

        with app.app_context():
            leads = self.Lead.query.order_by(self.Lead.created_at.asc()).all()
            self.assertEqual(len(leads), 2)
            self.assertEqual(leads[0].lead_id, "LD00001")
            self.assertEqual(leads[1].traveler_id, "TR00010")
            self.assertEqual(leads[1].lead_stage, "Qualified")

    def test_manual_lead_conflict_creates_handoff(self):
        app = self._build_app()
        client = app.test_client()

        with app.app_context():
            # Seed existing traveler with conflicting name
            t1 = self.Traveler(traveler_id="TR00001", full_name="Different Person", whatsapp_raw="1012345678", phone_lookup_key="20:1012345678")
            self.db.session.add(t1)
            self.db.session.commit()

        # Submit manual lead with same phone but different name
        response = client.post(
            "/leads/",
            data={
                "customer_name": "Conflicting Manual Lead",
                "raw_phone": "201012345678",
                "lead_stage": "Interested",
                "priority": "Medium",
            },
        )
        self.assertEqual(response.status_code, 302)

        with app.app_context():
            # Verify no new traveler created (do not guess)
            travelers = self.Traveler.query.all()
            self.assertEqual(len(travelers), 1)
            
            leads = self.Lead.query.filter_by(customer_name="Conflicting Manual Lead").all()
            self.assertEqual(len(leads), 1)
            lead = leads[0]
            
            self.assertTrue(lead.handoff_required)
            self.assertEqual(lead.lead_stage, "Needs Review")
            self.assertEqual(lead.priority, "High")
            self.assertIsNotNone(lead.handoff_id)
            self.assertEqual(lead.handoff_reason, "phone_name_conflict")
            
            # Verify HandoffQueue entry exists
            handoff = self.HandoffQueue.query.filter_by(handoff_id=lead.handoff_id).first()
            self.assertIsNotNone(handoff)
            self.assertIn("Name and phone conflict", handoff.reason)
            self.assertIn("Recommended actions:", handoff.notes or "")

    def test_manual_lead_blacklist_creates_handoff_and_blocks(self):
        app = self._build_app()
        client = app.test_client()

        with app.app_context():
            # Seed blacklisted traveler
            t1 = self.Traveler(traveler_id="TR00002", full_name="Bad Guy", whatsapp_raw="1098765432", phone_lookup_key="20:1098765432", status="Blacklisted")
            self.db.session.add(t1)
            self.db.session.commit()

        response = client.post(
            "/leads/",
            data={
                "customer_name": "Bad Guy",
                "raw_phone": "201098765432",
                "lead_stage": "Interested",
            },
        )
        self.assertEqual(response.status_code, 302)

        with app.app_context():
            leads = self.Lead.query.all()
            self.assertEqual(len(leads), 1)
            lead = leads[0]
            
            self.assertTrue(lead.handoff_required)
            self.assertEqual(lead.lead_stage, "Blocked")
            self.assertEqual(lead.priority, "Critical")
            self.assertEqual(lead.handoff_reason, "blacklisted_customer")
            
            handoff = self.HandoffQueue.query.filter_by(handoff_id=lead.handoff_id).first()
            self.assertIsNotNone(handoff)
            self.assertIn("Blocked traveler match", handoff.reason)
            self.assertIn("Recommended actions:", handoff.notes or "")

    def test_commercial_context_flags_group_quote_review(self):
        context = UnifiedCRMService.resolve_commercial_context(
            traveler={"status": "VIP"},
            trip_type="International",
            requested_group_size=5,
            trip_id="RT-INT-26-001",
        )
        self.assertEqual(context["flight_policy"]["mode"], "without_flights_default")
        self.assertTrue(context["vip"]["recognized"])
        self.assertTrue(context["group"]["eligible"])
        self.assertTrue(context["manual_handoff_recommended"])
        self.assertEqual(context["handoff_reason"], "group_booking_quote")

    def test_default_app_database_matches_system_service_database(self):
        os.environ.pop("DATABASE_URL", None)
        os.environ.pop("RAHMA_SYSTEM_DB_PATH", None)

        for module_name in list(sys.modules):
            if module_name == "app" or module_name.startswith("app."):
                sys.modules.pop(module_name, None)

        from app.config import DevelopmentConfig

        expected_uri = f"sqlite:///{resolve_system_db_path().resolve().as_posix()}"
        self.assertEqual(DevelopmentConfig.get_sqlalchemy_uri(), expected_uri)

if __name__ == "__main__":
    unittest.main()
