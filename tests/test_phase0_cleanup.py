from __future__ import annotations

import unittest
from pathlib import Path
import shutil
import uuid

from openpyxl import Workbook, load_workbook

from scripts.phase0_cleanup import (
    AUDIT_SHEET,
    INTERACTIONS_SHEET,
    TRAVELERS_SHEET,
    TRIPS_SHEET,
    TODAY,
    canonicalize_status,
    normalize_phone_fields,
    run_phase0_cleanup,
)


class Phase0CleanupTests(unittest.TestCase):
    def test_canonicalize_status(self) -> None:
        self.assertEqual(canonicalize_status("Blackllisted"), "Blacklisted")
        self.assertEqual(canonicalize_status(" active traveler "), "Active")
        self.assertEqual(canonicalize_status("VIP"), "VIP")
        self.assertEqual(canonicalize_status(""), "")

    def test_normalize_phone_fields(self) -> None:
        normalized = normalize_phone_fields("20", 1005828000)
        self.assertEqual(normalized.code, "20")
        self.assertEqual(normalized.local_number, "1005828000")
        self.assertEqual(normalized.normalized_whatsapp, "+201005828000")
        self.assertEqual(normalized.lookup_key, "20:1005828000")
        self.assertEqual(normalized.flags, [])

        derived = normalize_phone_fields("", "201005828000")
        self.assertEqual(derived.code, "20")
        self.assertEqual(derived.local_number, "1005828000")
        self.assertIn("CODE_DERIVED_FROM_PHONE", derived.flags)

        missing_code = normalize_phone_fields("", "01005828000")
        self.assertEqual(missing_code.code, "20")
        self.assertEqual(missing_code.local_number, "1005828000")
        self.assertIn("CODE_DERIVED_FROM_EGYPT_LOCAL", missing_code.flags)

    def test_run_phase0_cleanup_updates_workbook_and_audit(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase0-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            input_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"
            report_path = tmp_path / "report.md"

            wb = Workbook()
            travelers = wb.active
            travelers.title = TRAVELERS_SHEET
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "First Name",
                    "Last Name",
                    "Birthday",
                    "Gender",
                    "Nationality",
                    "Code",
                    "WhatsApp",
                    "Email",
                    "Community Whatsapp",
                    "Residence",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                    "Comm. Events",
                    "Lifetime Revenue",
                    "Notes",
                    "Introduce yourself",
                    "Emergency Contact",
                    "Emergency Phone",
                    "Medical Notes",
                    "Room Preference",
                    "⭐ Rating (1–5)",
                ]
            )
            travelers.append(["Blackllisted", "TR00001", "Jane Doe", None, None, None, None, None, "20", 1005828000])
            travelers.append(["", "TR00002", "Janet Doe", None, None, None, None, None, "", "201005828000"])
            travelers.append(["", "TR00003", "Janet Doe 2", None, None, None, None, None, "", "201005828000"])

            trips = wb.create_sheet(TRIPS_SHEET)
            trips["A2"] = "Trip ID "
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["E2"] = "Trip Leader"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["A3"] = "RT-LOC-26-001"
            trips["B3"] = "Siwa 3"
            trips["C3"] = "Local"
            trips["D3"] = 2026
            trips["F3"] = TODAY
            trips["G3"] = TODAY
            trips["A4"] = "RT-INT-26-002"
            trips["B4"] = "Georgia"
            trips["C4"] = "International"
            trips["D4"] = 2026
            trips["G4"] = "Apri 5"
            trips["A5"] = "RT-INT-26-003"
            trips["B5"] = "Maldives 2"
            trips["C5"] = "International"
            trips["D5"] = 2026
            trips["A6"] = "RT-LOC-25-010"
            trips["B6"] = "Do Nothing 2"
            trips["C6"] = "Local"
            trips["D6"] = 2025

            wb.save(input_path)

            summary = run_phase0_cleanup(input_path, output_path, report_path)
            cleaned = load_workbook(output_path)

            travelers_ws = cleaned[TRAVELERS_SHEET]
            trips_ws = cleaned[TRIPS_SHEET]

            self.assertEqual(travelers_ws["A2"].value, "Blacklisted")
            self.assertEqual(travelers_ws["I2"].value, "20")
            self.assertEqual(travelers_ws["J2"].value, "1005828000")

            traveler_headers = [travelers_ws.cell(1, col).value for col in range(1, travelers_ws.max_column + 1)]
            self.assertIn("Integrated WhatsApp", traveler_headers)
            self.assertIn("Normalized WhatsApp", traveler_headers)
            self.assertIn("Data Audit", traveler_headers)

            header_to_col = {travelers_ws.cell(1, col).value: col for col in range(1, travelers_ws.max_column + 1)}
            self.assertEqual(
                travelers_ws.cell(2, header_to_col["Integrated WhatsApp"]).value,
                "+201005828000",
            )
            self.assertEqual(
                travelers_ws.cell(2, header_to_col["Normalized WhatsApp"]).value,
                "+201005828000",
            )
            self.assertIn(
                "DUPLICATE_PHONE",
                travelers_ws.cell(3, header_to_col["Data Audit"]).value,
            )

            trip_headers = [trips_ws.cell(2, col).value for col in range(1, trips_ws.max_column + 1)]
            self.assertIn("Sales Status", trip_headers)
            trip_header_to_col = {trips_ws.cell(2, col).value: col for col in range(1, trips_ws.max_column + 1)}
            self.assertEqual(trips_ws.cell(3, trip_header_to_col["Sales Status"]).value, "Open")
            self.assertIn(
                "INVALID_END_DATE",
                trips_ws.cell(4, trip_header_to_col["Data Audit"]).value,
            )
            self.assertEqual(trips_ws.cell(4, trip_header_to_col["Sales Status"]).value, "Date Fix Needed")
            self.assertEqual(trips_ws.cell(5, trip_header_to_col["Sales Status"]).value, "Date TBD")
            self.assertEqual(trips_ws.cell(6, trip_header_to_col["Sales Status"]).value, "Archived")

            self.assertIn(AUDIT_SHEET, cleaned.sheetnames)
            self.assertIn(INTERACTIONS_SHEET, cleaned.sheetnames)
            self.assertEqual(cleaned[INTERACTIONS_SHEET]["A1"].value, "Interaction ID")
            status_validations = [dv for dv in travelers_ws.data_validations.dataValidation if dv.type == "list"]
            self.assertTrue(any("Blacklisted" in (dv.formula1 or "") for dv in status_validations))
            self.assertGreater(summary["traveler_duplicate_phone_keys"], 0)
            self.assertEqual(summary["interactions_sheet_ready"], 1)
            self.assertTrue(report_path.exists())
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
