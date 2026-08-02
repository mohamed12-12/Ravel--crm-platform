from __future__ import annotations

import os
import shutil
import sqlite3
import unittest
import uuid
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook

from demo_web.app import create_app


def create_operational_tables(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE travelers (
            traveler_id TEXT PRIMARY KEY,
            status TEXT,
            full_name TEXT,
            first_name TEXT,
            last_name TEXT,
            birthday TEXT,
            gender TEXT,
            nationality TEXT,
            phone_code TEXT,
            whatsapp_raw TEXT,
            integrated_whatsapp TEXT,
            normalized_whatsapp TEXT,
            phone_lookup_key TEXT,
            local_trips_count INTEGER,
            international_trips_count INTEGER,
            total_trips INTEGER,
            lead_source TEXT,
            created_at TEXT,
            last_contacted_at TEXT,
            agent_notes TEXT,
            data_audit TEXT,
            last_lead_id TEXT,
            last_booking_id TEXT
        );
        CREATE TABLE trips (
            trip_id TEXT PRIMARY KEY,
            trip_name TEXT,
            type TEXT,
            year INTEGER,
            trip_leader TEXT,
            start_date TEXT,
            end_date TEXT,
            sales_status TEXT,
            data_audit TEXT,
            trip_window_status TEXT,
            trip_availability_note TEXT,
            next_reengage_date TEXT,
            single_total INTEGER,
            double_total INTEGER,
            triple_total INTEGER,
            single_remaining INTEGER,
            double_remaining INTEGER,
            triple_remaining INTEGER,
            draft_holds_single INTEGER,
            draft_holds_double INTEGER,
            draft_holds_triple INTEGER,
            boys_double INTEGER,
            girls_double INTEGER,
            boys_triple INTEGER,
            girls_triple INTEGER,
            public_price TEXT,
            public_description TEXT,
            sales_notes TEXT
        );
        CREATE TABLE leads (
            lead_id TEXT PRIMARY KEY,
            created_at TEXT,
            updated_at TEXT,
            customer_name TEXT,
            raw_phone TEXT,
            integrated_whatsapp TEXT,
            phone_lookup_key TEXT,
            traveler_id TEXT,
            traveler_status TEXT,
            customer_tier TEXT,
            match_status TEXT,
            lead_stage TEXT,
            lead_source TEXT,
            channel TEXT,
            preferred_trip_type TEXT,
            group_size INTEGER DEFAULT 1,
            interested_trip_ids TEXT,
            suggested_trip_ids TEXT,
            priority TEXT,
            follow_up_status TEXT,
            follow_up_due_date TEXT,
            last_interaction_id TEXT,
            interaction_count INTEGER,
            handoff_required INTEGER,
            handoff_reason TEXT,
            notes TEXT,
            booking_id TEXT,
            flow_key TEXT,
            current_step TEXT,
            language TEXT
        );
        CREATE TABLE interactions (
            interaction_id TEXT PRIMARY KEY,
            timestamp TEXT,
            channel TEXT,
            customer_name TEXT,
            raw_phone TEXT,
            integrated_whatsapp TEXT,
            phone_lookup_key TEXT,
            traveler_id TEXT,
            matched_row INTEGER,
            status_snapshot TEXT,
            intent TEXT,
            trip_type TEXT,
            suggested_trips TEXT,
            action_taken TEXT,
            handoff_required INTEGER,
            handoff_reason TEXT,
            agent_notes TEXT,
            flow_key TEXT,
            step_key TEXT,
            message_key TEXT,
            language TEXT,
            outcome TEXT
        );
        CREATE TABLE trip_bookings (
            booking_id TEXT PRIMARY KEY,
            trip_id TEXT,
            trip_name TEXT,
            traveler_id TEXT,
            traveler_name TEXT,
            room_type TEXT,
            flight_option TEXT,
            date_option TEXT,
            currency TEXT,
            booking_status TEXT,
            draft_created_at TEXT,
            booking_source TEXT,
            lead_id TEXT,
            interaction_id TEXT,
            alert_id TEXT,
            payment_status TEXT,
            group_size INTEGER DEFAULT 1,
            booking_notes TEXT
        );
        """
    )


def create_workbook(path: Path) -> None:
    wb = Workbook()
    travelers = wb.active
    travelers.title = "Travelers"
    travelers.append([
        "Status", "Traveler ID", "Full Name", "Code", "WhatsApp",
        "Loc. Trips", "Int. Trips", "Total trips", "Integrated WhatsApp",
        "Phone Lookup Key", "Data Audit",
    ])

    trips = wb.create_sheet("Trips")
    for col, header in {
        "A": "Trip ID",
        "B": "Trip Name",
        "C": "Type",
        "D": "Year",
        "E": "Trip Leader",
        "F": "Start Date",
        "G": "End Date",
        "H": "Single",
        "I": "Double",
        "J": "Triple",
        "K": "Single",
        "L": "Double",
        "M": "Triple",
        "Z": "Sales Status",
        "AA": "Data Audit",
    }.items():
        trips[f"{col}2"] = header

    interactions = wb.create_sheet("Interactions")
    interactions.append([
        "Interaction ID", "Timestamp", "Channel", "Customer Name", "Raw Phone",
        "Integrated WhatsApp", "Phone Lookup Key", "Traveler ID", "Matched Row",
        "Status Snapshot", "Intent", "Trip Type", "Suggested Trips",
        "Action Taken", "Handoff Required", "Handoff Reason", "Agent Notes",
    ])
    wb.save(path)
    wb.close()


class Phase3BookingWriteThroughTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = Path(".tmp-test-workdirs")
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"phase3-write-through-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self._original_env = dict(os.environ)

    def tearDown(self) -> None:
        os.environ.clear()
        os.environ.update(self._original_env)
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def test_agent_booking_writes_system_db_ui_tables_and_sheet(self) -> None:
        db_path = self.tmp_path / "system.db"
        with closing(sqlite3.connect(str(db_path))) as connection:
            create_operational_tables(connection)
            connection.execute(
                """
                INSERT INTO trips (
                    trip_id, trip_name, type, year, start_date, end_date, sales_status,
                    single_total, double_total, triple_total, single_remaining, double_remaining,
                    triple_remaining, draft_holds_single, draft_holds_double, draft_holds_triple
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "RT-LOC-26-777",
                    "System Booking Trip",
                    "Local",
                    2026,
                    "2026-10-10",
                    "2026-10-12",
                    "Open",
                    2,
                    2,
                    1,
                    2,
                    2,
                    1,
                    0,
                    0,
                    0,
                ),
            )
            connection.commit()

        os.environ["RAHMA_SYSTEM_DB_PATH"] = str(db_path)
        os.environ["AI_AGENT_MODE"] = "deterministic"
        os.environ["CRM_SHEET_MIRROR_ENABLED"] = "true"
        os.environ["EXCEL_EXPORT_ENABLED"] = "true"
        source_workbook = self.tmp_path / "source.xlsx"
        runtime_workbook = self.tmp_path / "runtime.xlsx"
        create_workbook(source_workbook)

        app = create_app(source_workbook=source_workbook, runtime_workbook=runtime_workbook)
        client = app.test_client()
        session = client.post("/api/session", json={}).get_json()["session"]
        session_id = session["id"]

        client.post(
            f"/api/session/{session_id}/intake",
            json={
                "fullName": "Sara Phase Three",
                "birthday": "1998-02-20",
                "gender": "Female",
                "nationality": "Egypt",
                "countryCode": "20",
                "rawPhone": "01012345678",
            },
        )
        client.post(f"/api/session/{session_id}/message", json={"text": "local"})
        client.post(f"/api/session/{session_id}/message", json={"text": "1"})
        client.post(f"/api/session/{session_id}/message", json={"text": "Single"})
        client.post(f"/api/session/{session_id}/message", json={"text": "no"})
        session = client.post(f"/api/session/{session_id}/message", json={"text": "EGP"}).get_json()["session"]

        self.assertEqual(session["stage"], "post_booking_support")
        self.assertTrue(session["bookingCompleted"])
        self.assertIsNotNone(session["bookingResult"])

        with closing(sqlite3.connect(str(db_path))) as connection:
            connection.row_factory = sqlite3.Row
            lead = connection.execute("SELECT * FROM leads").fetchone()
            booking = connection.execute("SELECT * FROM trip_bookings").fetchone()
            interactions = connection.execute("SELECT COUNT(*) AS c FROM interactions").fetchone()["c"]
            events = [
                row["event_type"]
                for row in connection.execute("SELECT event_type FROM booking_event_trail ORDER BY occurred_at, event_id")
            ]

        self.assertEqual(lead["lead_stage"], "Booking Draft Created")
        self.assertEqual(booking["booking_status"], "Draft")
        self.assertEqual(booking["payment_status"], "Pending")
        self.assertGreaterEqual(interactions, 2)
        self.assertIn("inquiry_received", events)
        self.assertIn("trip_suggested", events)
        self.assertIn("lead_qualified", events)
        self.assertIn("booking_draft_created", events)
        self.assertIn("payment_follow_up", events)

        wb = load_workbook(runtime_workbook, data_only=True)
        try:
            leads = wb["Leads"]
            lead_headers = {leads.cell(1, col).value: col for col in range(1, leads.max_column + 1)}
            self.assertEqual(leads.cell(2, lead_headers["Lead Stage"]).value, "Booking Draft Created")

            bookings = wb["Trip Bookings"]
            booking_headers = {bookings.cell(2, col).value: col for col in range(1, bookings.max_column + 1)}
            self.assertEqual(bookings.cell(3, booking_headers["Booking Status"]).value, "Draft")
            self.assertEqual(bookings.cell(3, booking_headers["Payment Status"]).value, "Pending")

            event_sheet = wb["Booking Event Trail"]
            event_headers = {event_sheet.cell(1, col).value: col for col in range(1, event_sheet.max_column + 1)}
            sheet_events = [
                event_sheet.cell(row, event_headers["Event Type"]).value
                for row in range(2, event_sheet.max_row + 1)
                if event_sheet.cell(row, event_headers["Event ID"]).value
            ]
            self.assertIn("booking_draft_created", sheet_events)
            self.assertIn("payment_follow_up", sheet_events)
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
