from __future__ import annotations

import os
import shutil
import sqlite3
import sys
import unittest
import uuid
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook

from demo_web.app import create_app
from services.crm.system_services.field_mapping import SHEET_TABLE_MAPPINGS

CURRENT_DIR = Path(__file__).resolve().parent
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

from test_phase3_booking_write_through import create_operational_tables


def seed_workbook(path: Path) -> None:
    wb = Workbook()
    
    # Travelers
    ws_travelers = wb.active
    ws_travelers.title = "Travelers"
    ws_travelers.append(list(SHEET_TABLE_MAPPINGS["Travelers"]["columns"].values()))
    
    # Trips
    ws_trips = wb.create_sheet("Trips")
    ws_trips["A2"] = "Trip ID"
    ws_trips["B2"] = "Trip Name"
    ws_trips["C2"] = "Type"
    ws_trips["D2"] = "Year"
    ws_trips["F2"] = "Start Date"
    ws_trips["G2"] = "End Date"
    ws_trips["Z2"] = "Sales Status"
    ws_trips["AA2"] = "Data Audit"
    ws_trips["A3"] = "RT-LOC-26-900"
    ws_trips["B3"] = "Siwa Discovery Demo"
    ws_trips["C3"] = "Local"
    ws_trips["D3"] = 2026
    ws_trips["F3"] = "2026-08-14"
    ws_trips["G3"] = "2026-08-17"
    ws_trips["Z3"] = "Open"

    # Leads
    ws_leads = wb.create_sheet("Leads")
    ws_leads.append(list(SHEET_TABLE_MAPPINGS["Leads"]["columns"].values()))
    
    # Interactions
    ws_interactions = wb.create_sheet("Interactions")
    ws_interactions.append(list(SHEET_TABLE_MAPPINGS["Interactions"]["columns"].values()))

    # Booking Event Trail
    ws_events = wb.create_sheet("Booking Event Trail")
    ws_events.append(list(SHEET_TABLE_MAPPINGS["Booking Event Trail"]["columns"].values()))

    # Trip Bookings
    ws_bookings = wb.create_sheet("Trip Bookings")
    ws_bookings.append(list(SHEET_TABLE_MAPPINGS["Trip Bookings"]["columns"].values()))

    # DM Copy Library
    ws_copy = wb.create_sheet("DM Copy Library")
    ws_copy.append(["Message Key", "Arabic Copy", "English Copy", "Active"])
    ws_copy.append(["session.intake_start", "أهلاً بك", "Hello. I am Rahma Traveler's sales agent. Please complete the intake form so I can check the CRM safely.", "Yes"])

    wb.save(path)
    wb.close()


class Phase8DemoAlignmentTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = Path(".tmp-test-workdirs")
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"phase8-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self._original_env = dict(os.environ)

    def tearDown(self) -> None:
        os.environ.clear()
        os.environ.update(self._original_env)
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def test_demo_web_calls_crm_services_safely_and_creates_records(self) -> None:
        source_workbook = self.tmp_path / "source.xlsx"
        runtime_workbook = self.tmp_path / "runtime.xlsx"
        db_path = self.tmp_path / "system.db"
        
        seed_workbook(source_workbook)

        with closing(sqlite3.connect(db_path)) as connection:
            create_operational_tables(connection)
            connection.execute(
                """
                INSERT INTO trips (
                    trip_id, trip_name, type, year, start_date, end_date, sales_status,
                    single_total, double_total, triple_total, single_remaining, double_remaining,
                    triple_remaining, draft_holds_single, draft_holds_double, draft_holds_triple
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "RT-LOC-26-900",
                    "Siwa Discovery Demo",
                    "Local",
                    2026,
                    "2026-08-14",
                    "2026-08-17",
                    "Open",
                    5,
                    5,
                    5,
                    5,
                    5,
                    5,
                    0,
                    0,
                    0,
                ),
            )
            connection.commit()

        os.environ["DATABASE_URL"] = f"sqlite:///{db_path.resolve().as_posix()}"
        os.environ["RAHMA_SYSTEM_DB_PATH"] = str(db_path)
        os.environ["SHEET_BACKEND"] = "excel"
        os.environ["EXCEL_RUNTIME_WORKBOOK"] = str(runtime_workbook)
        os.environ["EXCEL_SOURCE_WORKBOOK"] = str(source_workbook)

        # Build clean test client
        app = create_app(source_workbook=source_workbook, runtime_workbook=runtime_workbook)
        client = app.test_client()

        # 1. Bootstrap check
        resp = client.get("/api/bootstrap")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("stats", resp.get_json())

        # 2. Create intake session
        resp = client.post("/api/session", json={})
        session = resp.get_json()["session"]
        self.assertEqual(session["stage"], "awaiting_phone")

        # 3. Share WhatsApp number so the demo moves into intake capture
        resp = client.post(f"/api/session/{session['id']}/message", json={"text": "01023456789"})
        session = resp.get_json()["session"]
        self.assertEqual(session["stage"], "awaiting_intake")
        
        # 4. Submit intake form
        resp = client.post(
            f"/api/session/{session['id']}/intake",
            json={
                "fullName": "Test Alignment traveler",
                "birthday": "1995-10-10",
                "gender": "Male",
                "nationality": "Egypt",
                "countryCode": "20",
                "rawPhone": "1023456789",
                "language": "en",
            }
        )
        session = resp.get_json()["session"]
        self.assertEqual(session["stage"], "awaiting_trip_type")

        # 5. Message trip type
        resp = client.post(f"/api/session/{session['id']}/message", json={"text": "local"})
        session = resp.get_json()["session"]
        self.assertEqual(session["stage"], "awaiting_confirmation")

        # 6. Confirm lead creation
        resp = client.post(f"/api/session/{session['id']}/message", json={"text": "yes"})
        session = resp.get_json()["session"]
        self.assertEqual(session["stage"], "awaiting_room_type")
        
        self.assertIsNotNone(session["finalResult"])
        self.assertIn("write_result", session["finalResult"])
        self.assertIsNotNone(session["finalResult"]["write_result"]["created_traveler"])
        
        # Verify db has traveler, lead and interaction records
        with closing(sqlite3.connect(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            traveler = conn.execute("SELECT * FROM travelers").fetchone()
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler["full_name"], "Test Alignment traveler")
            self.assertEqual(traveler["phone_lookup_key"], "20:1023456789")

            lead = conn.execute("SELECT * FROM leads").fetchone()
            self.assertIsNotNone(lead)
            self.assertEqual(lead["traveler_id"], traveler["traveler_id"])

            interaction = conn.execute("SELECT * FROM interactions").fetchone()
            self.assertIsNotNone(interaction)
            self.assertEqual(interaction["traveler_id"], traveler["traveler_id"])


if __name__ == "__main__":
    unittest.main()
