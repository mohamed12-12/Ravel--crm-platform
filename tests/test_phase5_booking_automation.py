from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.phase5_booking_automation import create_booking_draft, prepare_phase5_workbook


class Phase5BookingAutomationTests(unittest.TestCase):
    def test_prepare_phase5_workbook_seeds_demo_trips_and_helpers(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase5-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            input_path = tmp_path / "input.xlsx"
            output_path = tmp_path / "output.xlsx"

            wb = Workbook()
            trips = wb.active
            trips.title = "Trips"
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["E2"] = "Trip Leader"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Public Description"
            trips["AB2"] = "Public Price"
            trips["AC2"] = "Sales Notes"
            trips["AD2"] = "Data Audit"

            travelers = wb.create_sheet("Travelers")
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "First Name",
                    "Last Name",
                    "Birthday",
                    "Gender",
                    "Nationality",
                    "Code",
                    "WhatsApp",
                    "Email",
                    "Community Whatsapp",
                    "Residence",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                ]
            )
            travelers.append(["", "TR00001", "Demo User", "=1", "=1", None, None, None, "20", "1000000000", None, None, None, "=1", "=1", "=1"])

            trip_bookings = wb.create_sheet("Trip Bookings")
            trip_bookings["A2"] = "Booking ID"
            trip_bookings["B2"] = "Trip ID"
            trip_bookings["C2"] = "Trip Name"
            trip_bookings["D2"] = "Traveler ID"
            trip_bookings["E2"] = "Traveler Name"
            trip_bookings["F2"] = "Room Type"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )

            leads = wb.create_sheet("Leads")
            booking_alerts = wb.create_sheet("Booking Alerts")
            wb.save(input_path)

            result = prepare_phase5_workbook(input_path, output_path)
            self.assertGreaterEqual(result["demoTripsSeeded"], 2)

            saved = load_workbook(output_path, data_only=False)
            trips_ws = saved["Trips"]
            self.assertEqual(trips_ws["A3"].value, "RT-LOC-26-900")
            self.assertEqual(trips_ws["Z3"].value, "Open")
            self.assertEqual(trips_ws["AE2"].value, "Draft Holds Single")
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)

    def test_create_booking_draft_updates_workbook(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase5-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        try:
            workbook_path = tmp_path / "phase5.xlsx"

            wb = Workbook()
            trips = wb.active
            trips.title = "Trips"
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["E2"] = "Trip Leader"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Public Description"
            trips["AB2"] = "Public Price"
            trips["AC2"] = "Sales Notes"
            trips["AD2"] = "Data Audit"
            trips["AE2"] = "Draft Holds Single"
            trips["AF2"] = "Draft Holds Double"
            trips["AG2"] = "Draft Holds Triple"
            trips["A3"] = "RT-LOC-26-900"
            trips["B3"] = "Siwa Discovery Demo"
            trips["C3"] = "Local"
            trips["D3"] = 2026
            trips["F3"] = "2026-08-14"
            trips["G3"] = "2026-08-17"
            trips["K3"] = 3
            trips["L3"] = 4
            trips["M3"] = 2
            trips["Z3"] = "Open"
            trips["AE3"] = 0
            trips["AF3"] = 0
            trips["AG3"] = 0

            travelers = wb.create_sheet("Travelers")
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "First Name",
                    "Last Name",
                    "Birthday",
                    "Gender",
                    "Nationality",
                    "Code",
                    "WhatsApp",
                    "Email",
                    "Community Whatsapp",
                    "Residence",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                ]
            )
            travelers.append(["", "TR00518", "Sara New", "=1", "=1", None, None, None, "20", "1012345678", None, None, None, "=1", "=1", "=1"])

            trip_bookings = wb.create_sheet("Trip Bookings")
            trip_bookings["A2"] = "Booking ID"
            trip_bookings["B2"] = "Trip ID"
            trip_bookings["C2"] = "Trip Name"
            trip_bookings["D2"] = "Traveler ID"
            trip_bookings["E2"] = "Traveler Name"
            trip_bookings["F2"] = "Room Type"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )

            leads = wb.create_sheet("Leads")
            leads.append(
                [
                    "Lead ID",
                    "Created At",
                    "Updated At",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Traveler Status",
                    "Customer Tier",
                    "Match Status",
                    "Lead Stage",
                    "Lead Source",
                    "Channel",
                    "Preferred Trip Type",
                    "Interested Trip IDs",
                    "Suggested Trip IDs",
                    "Priority",
                    "Follow Up Status",
                    "Follow Up Due Date",
                    "Last Interaction ID",
                    "Interaction Count",
                    "Handoff Required",
                    "Handoff Reason",
                    "Notes",
                ]
            )
            leads.append(["LD00001", None, None, "Sara New", "01012345678", "+201012345678", "20:1012345678", "TR00518", "", "Standard", "not_found", "Follow Up Needed", "Web Demo", "web-demo", "local", "RT-LOC-26-900", "RT-LOC-26-900", "Medium", "Awaiting Dates", "2026-05-18", "INT-20260511-0001", 1, "No", None, None])

            booking_alerts = wb.create_sheet("Booking Alerts")
            booking_alerts.append(
                [
                    "Alert ID",
                    "Created At",
                    "Booking ID",
                    "Trip ID",
                    "Trip Name",
                    "Traveler ID",
                    "Traveler Name",
                    "Lead ID",
                    "Channel",
                    "Priority",
                    "Alert Type",
                    "Alert Status",
                    "Summary",
                    "Owner",
                    "Notes",
                ]
            )
            wb.save(workbook_path)

            result = create_booking_draft(
                workbook_path=workbook_path,
                output_path=workbook_path,
                traveler_id="TR00518",
                traveler_name="Sara New",
                trip_id="RT-LOC-26-900",
                room_type="Double",
                channel="web-demo",
                lead_id="LD00001",
                source="Web Demo Booking",
                agent_notes="phase5 test",
            )

            self.assertEqual(result["trip_id"], "RT-LOC-26-900")
            self.assertEqual(result["room_type"], "Double")
            self.assertEqual(result["available_before_draft"], 4)
            self.assertEqual(result["available_after_draft"], 3)

            saved = load_workbook(workbook_path, data_only=False)
            bookings_ws = saved["Trip Bookings"]
            alerts_ws = saved["Booking Alerts"]
            leads_ws = saved["Leads"]
            trips_ws = saved["Trips"]
            booking_headers = {bookings_ws.cell(2, c).value: c for c in range(1, bookings_ws.max_column + 1) if bookings_ws.cell(2, c).value}

            self.assertEqual(bookings_ws["B3"].value, "RT-LOC-26-900")
            self.assertEqual(bookings_ws["F3"].value, "Double")
            self.assertEqual(bookings_ws.cell(3, booking_headers["Booking Status"]).value, "Draft")
            self.assertEqual(bookings_ws.cell(3, booking_headers["Payment Status"]).value, "Awaiting Deposit")
            self.assertTrue(str(alerts_ws["A2"].value).startswith("ALT-"))
            self.assertEqual(leads_ws["L2"].value, "Booking Draft Created")
            self.assertEqual(leads_ws["S2"].value, "Awaiting Deposit")
            self.assertEqual(trips_ws["AF3"].value, 1)
        finally:
            shutil.rmtree(tmp_path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
