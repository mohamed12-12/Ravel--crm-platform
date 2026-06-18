from __future__ import annotations

import os
import shutil
import sys
import unittest
import uuid
from pathlib import Path
from datetime import date

from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
API_ROOT = REPO_ROOT / "apps" / "api"
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from scripts.migrate_excel_to_crm import ExcelToCRMMigrator, apply_database_url_to_app_config, map_booking_lifecycle  # noqa: E402


TRAVELER_HEADERS = [
    "Status",
    "Traveler ID",
    "Full Name",
    "First Name",
    "Last Name",
    "Birthday",
    "Gender",
    "Nationality",
    "Code",
    "WhatsApp",
    "Email",
    "Community Whatsapp",
    "Residence",
    "Loc. Trips",
    "Int. Trips",
    "Total trips",
    "Comm. Events",
    "Lifetime Revenue",
    "Notes",
    "Introduce yourself",
    "Emergency Contact",
    "Emergency Phone",
    "Medical Notes",
    "Room Preference",
    "Rating",
]

TRIP_HEADERS = [
    "Trip ID",
    "Trip Name",
    "Type",
    "Year",
    "Trip Leader",
    "Start Date",
    "End Date",
    "Single",
    "Double",
    "Triple",
    "Single",
    "Double",
    "Triple",
]

BOOKING_HEADERS = [
    "Booking ID",
    "Trip ID",
    "Trip Name",
    "Traveler ID",
    "Traveler Name",
    "Room Type",
    "Currency",
    "Amount",
    "Date",
    "Payment Method",
    "Currency",
    "Amount",
    "Date",
    "Payment Method",
]


class ExcelToCRMMigrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = REPO_ROOT / ".tmp-test-workdirs"
        self.tmp_root.mkdir(exist_ok=True)
        self.tmp_path = self.tmp_root / f"excel-migration-{uuid.uuid4().hex}"
        self.tmp_path.mkdir()
        self.db_path = self.tmp_path / "crm.db"
        self.workbook_path = self.tmp_path / "source.xlsx"
        self._original_env = dict(os.environ)
        self._original_app_config = None

    def tearDown(self) -> None:
        if self._original_app_config and "app.config" in sys.modules:
            from app.config import DevelopmentConfig, ProductionConfig, config

            dev_uri, prod_uri = self._original_app_config
            DevelopmentConfig.SQLALCHEMY_DATABASE_URI = dev_uri
            ProductionConfig.SQLALCHEMY_DATABASE_URI = prod_uri
            config["development"].SQLALCHEMY_DATABASE_URI = dev_uri
            config["default"].SQLALCHEMY_DATABASE_URI = dev_uri
            config["production"].SQLALCHEMY_DATABASE_URI = prod_uri
        os.environ.clear()
        os.environ.update(self._original_env)
        shutil.rmtree(self.tmp_path, ignore_errors=True)

    def _build_schema(self):
        os.environ["DATABASE_URL"] = f"sqlite:///{self.db_path.resolve().as_posix()}"
        from app.config import DevelopmentConfig, ProductionConfig

        if self._original_app_config is None:
            self._original_app_config = (
                DevelopmentConfig.SQLALCHEMY_DATABASE_URI,
                ProductionConfig.SQLALCHEMY_DATABASE_URI,
            )
        apply_database_url_to_app_config()
        from app import create_app
        from app.extensions import db

        app = create_app()
        with app.app_context():
            db.create_all()
        return app, db

    def _write_workbook(
        self,
        *,
        travelers: list[list[object]] | None = None,
        trips: list[list[object]] | None = None,
        bookings: list[list[object]] | None = None,
    ) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Travelers"
        ws.append(TRAVELER_HEADERS)
        for row in travelers or []:
            ws.append(row)

        ws = wb.create_sheet("Trips")
        ws.append(["", "", "", "", "", "", "", "Accommodation Total", "", "", "Accommodation Remaining"])
        ws.append(TRIP_HEADERS)
        for row in trips or []:
            ws.append(row)

        ws = wb.create_sheet("Trip Bookings")
        ws.append(["", "", "", "", "", "", "First Deposit", "", "", "", "Second Deposit"])
        ws.append(BOOKING_HEADERS)
        for row in bookings or []:
            ws.append(row)

        wb.save(self.workbook_path)
        wb.close()

    def _write_sparse_workbook(
        self,
        *,
        travelers_by_row: dict[int, list[object]] | None = None,
        trips: list[list[object]] | None = None,
        bookings: list[list[object]] | None = None,
    ) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Travelers"
        for index, header in enumerate(TRAVELER_HEADERS, start=1):
            ws.cell(row=1, column=index, value=header)
        for row_number, row in (travelers_by_row or {}).items():
            for col_index, value in enumerate(row, start=1):
                ws.cell(row=row_number, column=col_index, value=value)

        ws = wb.create_sheet("Trips")
        ws.append(["", "", "", "", "", "", "", "Accommodation Total", "", "", "Accommodation Remaining"])
        ws.append(TRIP_HEADERS)
        for row in trips or []:
            ws.append(row)

        ws = wb.create_sheet("Trip Bookings")
        ws.append(["", "", "", "", "", "", "First Deposit", "", "", "", "Second Deposit"])
        ws.append(BOOKING_HEADERS)
        for row in bookings or []:
            ws.append(row)

        wb.save(self.workbook_path)
        wb.close()

    def _write_sparse_trip_workbook(
        self,
        *,
        trips_by_row: dict[int, list[object]] | None = None,
        bookings: list[list[object]] | None = None,
    ) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Travelers"
        ws.append(TRAVELER_HEADERS)

        ws = wb.create_sheet("Trips")
        ws.append(["", "", "", "", "", "", "", "Accommodation Total", "", "", "Accommodation Remaining", "", "", "Accomodation Price EGP", "", "", "", "Accomodation Price USD", "", "", "", "Total Revenue", ""])
        ws.append(TRIP_HEADERS)
        for row_number, row in (trips_by_row or {}).items():
            for col_index, value in enumerate(row, start=1):
                ws.cell(row=row_number, column=col_index, value=value)

        ws = wb.create_sheet("Trip Bookings")
        ws.append(["", "", "", "", "", "", "First Deposit", "", "", "", "Second Deposit"])
        ws.append(BOOKING_HEADERS)
        for row in bookings or []:
            ws.append(row)

        wb.save(self.workbook_path)
        wb.close()

    def _run(
        self,
        *,
        dry_run: bool,
        travelers_only: bool = False,
        trips_only: bool = False,
        bookings_only: bool = False,
        bookings_preview: bool = False,
    ):
        scope = "travelers" if travelers_only else "trips" if trips_only else "bookings" if bookings_only else "bookings_preview" if bookings_preview else "full"
        migrator = ExcelToCRMMigrator(self.workbook_path, dry_run=dry_run, scope=scope)
        summary = migrator.run()
        return migrator, summary

    def _run_sequence_on_same_db(self):
        travelers_result = self._run(dry_run=False, travelers_only=True)
        trips_result = self._run(dry_run=False, trips_only=True)
        bookings_result = self._run(dry_run=True, bookings_preview=True)
        return travelers_result, trips_result, bookings_result

    def test_dry_run_does_not_write_records(self) -> None:
        app, db = self._build_schema()
        self._write_workbook(
            travelers=[["", "TR00001", "Nada Adel", "", "", "", "Female", "Yemen", "44", "7493723281"]]
        )

        _, summary = self._run(dry_run=True)

        with app.app_context():
            from app.models import Traveler

            self.assertEqual(Traveler.query.count(), 0)
        self.assertEqual(summary.sheets["Travelers"].created, 1)

    def test_execute_preserves_traveler_id_and_non_egyptian_phone(self) -> None:
        app, db = self._build_schema()
        self._write_workbook(
            travelers=[["Repeat", "TR00001", "Nada Adel", "", "", "", "Female", "Yemen", "44", "7493723281"]]
        )

        self._run(dry_run=False)

        with app.app_context():
            from app.models import Traveler

            traveler = db.session.get(Traveler, "TR00001")
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler.traveler_id, "TR00001")
            self.assertEqual(traveler.normalized_whatsapp, "+447493723281")
            self.assertEqual(traveler.phone_lookup_key, "44:7493723281")

    def test_travelers_only_merges_approved_duplicate_group_and_skips_other_sheets(self) -> None:
        app, db = self._build_schema()
        self._write_sparse_workbook(
            travelers_by_row={
                356: ["Repeat", "TR00355", "Alpha One", "Alpha", "One", "", "Female", "Egypt", "20", "01011111111", "", "", "", "", "", "", "", "", "primary note"],
                377: ["Repeat", "TR00355", "Alpha One", "", "", "", "", "", "", "01011111111", "alpha@example.com", "", "", "", "", "", "", "", "secondary note"],
            },
            trips=[["RT-IGNORE-1", "Ignored Trip", "Local", 2026, "Leader", "2026-08-01", "2026-08-03", 1, 1, 1, 1, 1, 1]],
            bookings=[["B-IGNORE-1", "RT-IGNORE-1", "Ignored Trip", "TR00355", "Alpha One", "Single", "USD", 100]],
        )

        _, summary = self._run(dry_run=False, travelers_only=True)

        with app.app_context():
            from app.models import Traveler, Trip, TripBooking

            traveler = db.session.get(Traveler, "TR00355")
            self.assertIsNotNone(traveler)
            self.assertEqual(traveler.email, "alpha@example.com")
            self.assertEqual(Trip.query.count(), 0)
            self.assertEqual(TripBooking.query.count(), 0)
        self.assertEqual(summary.travelers_merged, 1)
        self.assertEqual(summary.sheets["Travelers"].created, 1)

    def test_existing_traveler_update_preserves_existing_id(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Traveler

            db.session.add(Traveler(traveler_id="TR00002", full_name="Existing", phone_lookup_key="966:512345678"))
            db.session.commit()
        self._write_workbook(
            travelers=[["VIP", "TR00002", "Existing", "", "", "", "", "KSA", "966", "512345678", "e@example.com"]]
        )

        _, summary = self._run(dry_run=False)

        with app.app_context():
            from app.models import Traveler

            self.assertEqual(Traveler.query.count(), 1)
            traveler = db.session.get(Traveler, "TR00002")
            self.assertEqual(traveler.email, "e@example.com")
            self.assertEqual(traveler.traveler_id, "TR00002")
        self.assertEqual(summary.sheets["Travelers"].updated, 1)

    def test_travelers_only_assigns_new_id_for_keep_separate_secondary_row(self) -> None:
        app, db = self._build_schema()
        self._write_sparse_workbook(
            travelers_by_row={
                398: ["", "TR00397", "Primary Traveler", "Primary", "Traveler", "", "", "Egypt", "20", "01022222221"],
                519: ["", "TR00397", "Secondary Traveler", "Secondary", "Traveler", "", "", "Egypt", "20", "01022222222"],
            }
        )

        _, summary = self._run(dry_run=False, travelers_only=True)

        with app.app_context():
            from app.models import Traveler

            travelers = Traveler.query.order_by(Traveler.traveler_id).all()
            self.assertEqual(len(travelers), 2)
            ids = {traveler.traveler_id for traveler in travelers}
            self.assertIn("TR00397", ids)
            self.assertGreater(len(ids), 1)
        self.assertEqual(summary.sheets["Travelers"].created, 2)
        self.assertEqual(summary.travelers_merged, 0)

    def test_duplicate_phone_rows_are_quarantined(self) -> None:
        self._build_schema()
        self._write_workbook(
            travelers=[
                ["", "TR00003", "Person One", "", "", "", "", "Egypt", "20", "01012345678"],
                ["", "TR00004", "Person Two", "", "", "", "", "Egypt", "20", "01012345678"],
            ]
        )

        migrator, summary = self._run(dry_run=True)

        self.assertEqual(summary.sheets["Travelers"].quarantined, 2)
        self.assertTrue(
            all(item["reason"] == "duplicate_phone_identity_in_workbook" for item in migrator.quarantine)
        )

    def test_booking_preview_matches_valid_traveler_and_trip_without_writing(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Traveler, Trip

            db.session.add(Traveler(traveler_id="TR00005", full_name="Safe Traveler"))
            db.session.add(Trip(trip_id="RT-LOC-26-001", trip_name="Siwa", end_date=date(2099, 1, 1)))
            db.session.commit()
        self._write_workbook(
            bookings=[["B-005", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Double"]]
        )

        migrator, summary = self._run(dry_run=True, bookings_preview=True)

        with app.app_context():
            from app.models import TripBooking

            self.assertEqual(TripBooking.query.count(), 0)
        self.assertEqual(summary.bookings_matched, 1)
        self.assertEqual(summary.bookings_quarantined, 0)
        self.assertEqual(summary.booking_lifecycle_counts["Confirmed"], 1)
        self.assertEqual(migrator.quarantine, [])

    def test_booking_preview_after_sequential_traveler_and_trip_imports_uses_same_db(self) -> None:
        app, db = self._build_schema()
        self._write_workbook(
            travelers=[["", "TR00005", "Safe Traveler", "", "", "", "", "Egypt", "20", "01012345678"]],
            trips=[["RT-LOC-26-001", "Siwa", "Local", 2026, "Mariam", "2026-08-01", "2026-08-05", 2, 4, 6, 2, 4, 6]],
            bookings=[["B-005", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Double", "EGP", 1000, "", "", "", "", "", ""]],
        )

        _, _, (migrator, summary) = self._run_sequence_on_same_db()

        with app.app_context():
            from app.models import TripBooking, Traveler, Trip

            self.assertIsNotNone(db.session.get(Traveler, "TR00005"))
            self.assertIsNotNone(db.session.get(Trip, "RT-LOC-26-001"))
            self.assertEqual(TripBooking.query.count(), 0)
        self.assertEqual(summary.bookings_matched, 1)
        self.assertEqual(summary.bookings_quarantined, 0)
        self.assertFalse(migrator.quarantine)

    def test_bookings_only_execute_imports_safe_rows_and_creates_status_history(self) -> None:
        app, db = self._build_schema()
        self._write_workbook(
            travelers=[["", "TR00005", "Safe Traveler", "", "", "", "", "Egypt", "20", "01012345678"]],
            trips=[["RT-LOC-26-001", "Siwa", "Local", 2026, "Mariam", "2026-08-01", "2026-08-05", 2, 4, 6, 2, 4, 6]],
            bookings=[["B-005", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Double", "EGP", 1000, "", "", "", "", "", ""]],
        )

        self._run(dry_run=False, travelers_only=True)
        self._run(dry_run=False, trips_only=True)
        migrator, summary = self._run(dry_run=False, bookings_only=True)

        with app.app_context():
            from app.models import BookingStatusHistory, TripBooking

            booking = db.session.get(TripBooking, "B-005")
            self.assertIsNotNone(booking)
            self.assertEqual(booking.booking_status, "Payment Pending")
            history = BookingStatusHistory.query.filter_by(booking_id="B-005").all()
            self.assertEqual(len(history), 1)
            self.assertIsNone(history[0].old_status)
            self.assertEqual(history[0].new_status, "Payment Pending")
        self.assertEqual(summary.bookings_matched, 1)
        self.assertEqual(summary.bookings_quarantined, 0)
        self.assertFalse(migrator.quarantine)

    def test_bookings_only_execute_skips_quarantined_rows(self) -> None:
        app, db = self._build_schema()
        self._write_workbook(
            travelers=[["", "TR00005", "Safe Traveler", "", "", "", "", "Egypt", "20", "01012345678"]],
            trips=[["RT-LOC-26-001", "Siwa", "Local", 2026, "Mariam", "2026-08-01", "2026-08-05", 2, 4, 6, 2, 4, 6]],
            bookings=[
                ["B-005", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Double", "EGP", 1000, "", "", "", "", "", ""],
                ["0-L-000", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Double", "EGP", 1000, "", "", "", "", "", ""],
                ["B-006", "MISSING-TRIP", "Bad Trip", "TR00005", "Safe Traveler", "Double", "EGP", 1000, "", "", "", "", "", ""],
            ],
        )

        self._run(dry_run=False, travelers_only=True)
        self._run(dry_run=False, trips_only=True)
        migrator, summary = self._run(dry_run=False, bookings_only=True)

        with app.app_context():
            from app.models import TripBooking

            self.assertIsNotNone(db.session.get(TripBooking, "B-005"))
            self.assertIsNone(db.session.get(TripBooking, "B-006"))
            self.assertIsNone(db.session.get(TripBooking, "0-L-000"))
            self.assertEqual(TripBooking.query.count(), 1)
        self.assertEqual(summary.bookings_matched, 1)
        self.assertEqual(summary.bookings_quarantined, 2)
        self.assertEqual(sorted(item["reason"] for item in migrator.quarantine), [
            "duplicate_or_placeholder_booking_id",
            "invalid_booking_trip_reference",
        ])

    def test_invalid_booking_references_are_quarantined(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Trip

            db.session.add(Trip(trip_id="RT-LOC-26-001", trip_name="Siwa", end_date=date(2099, 1, 1)))
            db.session.commit()
        self._write_workbook(
            bookings=[["B-001", "RT-LOC-26-001", "Bad Trip", "TR404", "Ghost Traveler", "Single"]]
        )

        migrator, summary = self._run(dry_run=True, bookings_preview=True)

        self.assertEqual(summary.sheets["Trip Bookings"].quarantined, 1)
        self.assertEqual(migrator.quarantine[0]["reason"], "invalid_booking_traveler_reference")

    def test_invalid_trip_reference_is_quarantined(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Traveler

            db.session.add(Traveler(traveler_id="TR00005", full_name="Safe Traveler"))
            db.session.commit()
        self._write_workbook(
            bookings=[["B-002", "MISSING-TRIP", "Bad Trip", "TR00005", "Safe Traveler", "Single"]]
        )

        migrator, summary = self._run(dry_run=True, bookings_preview=True)

        self.assertEqual(summary.sheets["Trip Bookings"].quarantined, 1)
        self.assertEqual(migrator.quarantine[0]["reason"], "invalid_booking_trip_reference")

    def test_placeholder_booking_ids_are_quarantined(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Traveler, Trip

            db.session.add(Traveler(traveler_id="TR00005", full_name="Safe Traveler"))
            db.session.add(Trip(trip_id="RT-LOC-26-001", trip_name="Siwa", end_date=date(2099, 1, 1)))
            db.session.commit()
        self._write_workbook(
            bookings=[["0-L-000", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Single"]]
        )

        migrator, summary = self._run(dry_run=True, bookings_preview=True)

        self.assertEqual(summary.sheets["Trip Bookings"].quarantined, 1)
        self.assertEqual(migrator.quarantine[0]["reason"], "duplicate_or_placeholder_booking_id")

    def test_invalid_room_type_is_quarantined(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Traveler, Trip

            db.session.add(Traveler(traveler_id="TR00005", full_name="Safe Traveler"))
            db.session.add(Trip(trip_id="RT-LOC-26-001", trip_name="Siwa", end_date=None))
            db.session.commit()
        self._write_workbook(
            bookings=[["B-003", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "USD"]]
        )

        migrator, summary = self._run(dry_run=True, bookings_preview=True)

        self.assertEqual(summary.sheets["Trip Bookings"].quarantined, 1)
        self.assertEqual(migrator.quarantine[0]["reason"], "invalid_room_type")

    def test_booking_lifecycle_mapping_prefers_payment_and_historical_completion(self) -> None:
        future_trip_end_date = date(2099, 1, 1)
        past_trip_end_date = date(2020, 1, 1)

        self.assertEqual(
            map_booking_lifecycle({"Amount": 1000, "Amount #2": 1000}, future_trip_end_date),
            "Paid",
        )
        self.assertEqual(
            map_booking_lifecycle({"Amount": 1000, "Amount #2": 1000}, past_trip_end_date),
            "Completed",
        )
        self.assertEqual(map_booking_lifecycle({"Amount": 1000}, future_trip_end_date), "Payment Pending")
        self.assertEqual(map_booking_lifecycle({"Booking Status": "Confirmed"}, future_trip_end_date), "Confirmed")
        self.assertEqual(map_booking_lifecycle({"Booking Status": "Cancelled"}, future_trip_end_date), "Cancelled")

    def test_travelers_only_quarantines_manual_review_group(self) -> None:
        self._build_schema()
        self._write_sparse_workbook(
            travelers_by_row={
                74: ["", "TR00073", "Manual One", "Manual", "One", "", "", "Egypt", "20", "01033333331"],
                349: ["", "TR00073", "Manual Two", "Manual", "Two", "", "", "Egypt", "20", "01033333332"],
            }
        )

        migrator, summary = self._run(dry_run=True, travelers_only=True)

        self.assertEqual(summary.sheets["Travelers"].quarantined, 2)
        self.assertTrue(all(item["reason"] == "decision_requires_manual_review" for item in migrator.quarantine))

    def test_trips_only_imports_unique_trips_and_flags_missing_split(self) -> None:
        app, db = self._build_schema()
        self._write_sparse_trip_workbook(
            trips_by_row={
                3: ["RT-LOC-26-001", "Siwa", "Local", 2026, "Mariam", "2026-08-01", "2026-08-05", 2, 4, 6, 2, 4, 6, "", "", "", "1200", "", "", "", "1500", "40000", "900"],
                4: ["RT-INT-25-002", "Serbia & Bosnia", "International", 2024, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                5: ["RT-INT-25-002", "Zanzibar 2", "International", 2025, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                6: ["RT-INT-25-003", "Oman", "International", 2024, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                7: ["RT-INT-25-003", "Lebanon 2", "International", 2025, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                8: ["RT-INT-25-004", "Morocco", "International", 2024, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                9: ["RT-INT-25-004", "Spain", "International", 2025, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            },
            bookings=[["B-IGNORE", "RT-LOC-26-001", "Siwa", "TR00005", "Safe Traveler", "Double", "EGP", 1000]],
        )

        migrator, summary = self._run(dry_run=False, trips_only=True)

        with app.app_context():
            from app.models import Trip, TripBooking

            self.assertEqual(Trip.query.count(), 1)
            trip = db.session.get(Trip, "RT-LOC-26-001")
            self.assertIsNotNone(trip)
            self.assertEqual(trip.trip_name, "Siwa")
            self.assertEqual(trip.boys_double, 0)
            self.assertEqual(trip.girls_double, 0)
            self.assertEqual(trip.boys_triple, 0)
            self.assertEqual(trip.girls_triple, 0)
            self.assertEqual(TripBooking.query.count(), 0)
        self.assertEqual(summary.sheets["Trips"].created, 1)
        self.assertEqual(summary.sheets["Trips"].quarantined, 6)
        self.assertIn("RT-LOC-26-001", summary.validation["missing_boys_girls_split_trip_ids"])
        self.assertTrue(all(item["reason"] == "duplicate_trip_id_in_workbook" for item in migrator.quarantine))

    def test_trips_only_preserves_existing_trip_id(self) -> None:
        app, db = self._build_schema()
        with app.app_context():
            from app.models import Trip

            db.session.add(Trip(trip_id="RT-LOC-26-001", trip_name="Existing Trip"))
            db.session.commit()
        self._write_sparse_trip_workbook(
            trips_by_row={
                3: ["RT-LOC-26-001", "Existing Trip Updated", "Local", 2026, "Leader", "2026-08-01", "2026-08-05", 2, 4, 6, 2, 4, 6, "", "", "", "1200", "", "", "", "1500", "40000", "900"],
            }
        )

        _, summary = self._run(dry_run=False, trips_only=True)

        with app.app_context():
            from app.models import Trip

            self.assertEqual(Trip.query.count(), 1)
            trip = db.session.get(Trip, "RT-LOC-26-001")
            self.assertEqual(trip.trip_id, "RT-LOC-26-001")
            self.assertEqual(trip.trip_name, "Existing Trip")
        self.assertGreaterEqual(summary.sheets["Trips"].updated, 0)

if __name__ == "__main__":
    unittest.main()
