from __future__ import annotations

import shutil
import os
import sqlite3
import sys
import unittest
import uuid
from contextlib import closing
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

from demo_web.app import create_app

CURRENT_DIR = Path(__file__).resolve().parent
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

from test_phase3_booking_write_through import create_operational_tables

# Fixed relative to today (not a hardcoded calendar date) -- a static future
# date silently becomes "in the past" as real time passes, which made this
# seeded trip fall out of the "confirmed upcoming trips" window and start
# failing the moment the wall clock caught up to it.
_TRIP_START = (date.today() + timedelta(days=7)).isoformat()
_TRIP_END = (date.today() + timedelta(days=10)).isoformat()


class Phase3DemoWebTests(unittest.TestCase):
    def test_demo_session_creates_new_traveler_and_logs_interaction(self) -> None:
        tmp_root = Path(".tmp-test-workdirs")
        tmp_root.mkdir(exist_ok=True)
        tmp_path = tmp_root / f"phase3-{uuid.uuid4().hex}"
        tmp_path.mkdir()
        original_env = dict(os.environ)
        try:
            source_workbook = tmp_path / "source.xlsx"
            runtime_workbook = tmp_path / "runtime.xlsx"
            db_path = tmp_path / "system.db"
            with closing(sqlite3.connect(db_path)) as connection:
                create_operational_tables(connection)
                connection.execute(
                    """
                    INSERT INTO trips (
                        trip_id, trip_name, type, year, start_date, end_date, sales_status,
                        single_total, double_total, triple_total, single_remaining, double_remaining,
                        triple_remaining, draft_holds_single, draft_holds_double, draft_holds_triple
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        "RT-LOC-26-900",
                        "Siwa Discovery Demo",
                        "Local",
                        2026,
                        _TRIP_START,
                        _TRIP_END,
                        "Open",
                        3,
                        4,
                        2,
                        0,
                        4,
                        2,
                        0,
                        0,
                        0,
                    ),
                )
                connection.commit()
            os.environ["RAHMA_SYSTEM_DB_PATH"] = str(db_path)
            os.environ["AI_AGENT_MODE"] = "deterministic"

            wb = Workbook()
            travelers = wb.active
            travelers.title = "Travelers"
            travelers.append(
                [
                    "Status",
                    "Traveler ID",
                    "Full Name",
                    "First Name",
                    "Last Name",
                    "Birthday",
                    "Gender",
                    "Nationality",
                    "Code",
                    "WhatsApp",
                    "Email",
                    "Community Whatsapp",
                    "Residence",
                    "Loc. Trips",
                    "Int. Trips",
                    "Total trips",
                    "Comm. Events",
                    "Lifetime Revenue",
                    "Notes",
                    "Introduce yourself",
                    "Emergency Contact",
                    "Emergency Phone",
                    "Medical Notes",
                    "Room Preference",
                    "⭐ Rating (1–5)",
                    "Integrated WhatsApp",
                    "Normalized WhatsApp",
                    "Phone Lookup Key",
                    "Lead Source",
                    "Created At",
                    "Last Contacted At",
                    "Agent Notes",
                    "Data Audit",
                ]
            )
            travelers.append(
                [
                    "",
                    "TR00517",
                    "Existing Traveler",
                    '=IFERROR(__xludf.DUMMYFUNCTION("REGEXEXTRACT(C2,""[A-Za-z]+"")"),"Existing")',
                    '=RIGHT(C2,LEN(C2) - (FIND(" ",C2)))',
                    None,
                    None,
                    None,
                    "20",
                    "1000000000",
                    None,
                    None,
                    None,
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B2,\'Trip Bookings\'!B:B,"RT-LOC*")',
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B2,\'Trip Bookings\'!B:B,"RT-INT*")',
                    "=N2+O2",
                    '=COUNTIFS(\'CE Bookings\'!D:D,B2,\'CE Bookings\'!D:D,"<>")',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "+201000000000",
                    "+201000000000",
                    "20:1000000000",
                    "DM",
                    "2026-05-11T18:00:00",
                    "2026-05-11T18:00:00",
                    None,
                    None,
                ]
            )
            travelers.append(
                [
                    None,
                    None,
                    None,
                    '=IFERROR(__xludf.DUMMYFUNCTION("REGEXEXTRACT(C3,""[A-Za-z]+"")"),"#N/A")',
                    '=RIGHT(C3,LEN(C3) - (FIND(" ",C3)))',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B3,\'Trip Bookings\'!B:B,"RT-LOC*")',
                    '=COUNTIFS(\'Trip Bookings\'!D:D,B3,\'Trip Bookings\'!B:B,"RT-INT*")',
                    "=N3+O3",
                    '=COUNTIFS(\'CE Bookings\'!D:D,B3,\'CE Bookings\'!D:D,"<>")',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )

            trips = wb.create_sheet("Trips")
            trips["A2"] = "Trip ID"
            trips["B2"] = "Trip Name"
            trips["C2"] = "Type"
            trips["D2"] = "Year"
            trips["F2"] = "Start Date"
            trips["G2"] = "End Date"
            trips["Z2"] = "Sales Status"
            trips["AA2"] = "Data Audit"
            trips["A3"] = "RT-LOC-26-001"
            trips["B3"] = "Siwa 3"
            trips["C3"] = "Local"
            trips["D3"] = 2026
            trips["Z3"] = "Date TBD"
            trips["AA3"] = "MISSING_END_DATE; MISSING_START_DATE"
            trips["A4"] = "RT-LOC-26-900"
            trips["B4"] = "Siwa Discovery Demo"
            trips["C4"] = "Local"
            trips["D4"] = 2026
            trips["F4"] = _TRIP_START
            trips["G4"] = _TRIP_END
            trips["K4"] = 3
            trips["L4"] = 4
            trips["M4"] = 2
            trips["Z4"] = "Open"

            interactions = wb.create_sheet("Interactions")
            interactions.append(
                [
                    "Interaction ID",
                    "Timestamp",
                    "Channel",
                    "Customer Name",
                    "Raw Phone",
                    "Integrated WhatsApp",
                    "Phone Lookup Key",
                    "Traveler ID",
                    "Matched Row",
                    "Status Snapshot",
                    "Intent",
                    "Trip Type",
                    "Suggested Trips",
                    "Action Taken",
                    "Handoff Required",
                    "Handoff Reason",
                    "Agent Notes",
                ]
            )

            trip_bookings = wb.create_sheet("Trip Bookings")
            trip_bookings["A2"] = "Booking ID"
            trip_bookings["B2"] = "Trip ID"
            trip_bookings["C2"] = "Trip Name"
            trip_bookings["D2"] = "Traveler ID"
            trip_bookings["E2"] = "Traveler Name"
            trip_bookings["F2"] = "Room Type"

            wb.save(source_workbook)

            app = create_app(source_workbook=source_workbook, runtime_workbook=runtime_workbook)
            client = app.test_client()

            bootstrap = client.get("/api/bootstrap")
            self.assertEqual(bootstrap.status_code, 200)

            session_resp = client.post("/api/session", json={})
            session = session_resp.get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_phone")

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "01012345678"}).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_intake")

            session = client.post(
                f"/api/session/{session['id']}/intake",
                json={
                    "fullName": "Sara Demo",
                    "birthday": "1998-02-20",
                    "gender": "Female",
                    "nationality": "Egypt",
                    "countryCode": "20",
                    "rawPhone": "01012345678",
                },
            ).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_trip_type")

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "local"}).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_confirmation")
            self.assertIsNone(session["finalResult"])

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "yes"}).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_room_type")
            self.assertIsNotNone(session["finalResult"]["write_result"]["created_traveler"])
            self.assertIsNotNone(session["finalResult"]["write_result"]["lead_update"])
            self.assertEqual(session["stats"]["interactionCount"], 1)
            self.assertEqual(session["stats"]["leadCount"], 1)
            self.assertEqual(session["selectedTripId"], "RT-LOC-26-900")

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "Single"}).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_room_type")
            self.assertIn("Single is not available", session["messages"][-1]["text"])
            self.assertEqual(session["roomType"], "")

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "Double"}).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_flight")

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "no"}).get_json()["session"]
            self.assertEqual(session["stage"], "awaiting_currency")

            session = client.post(f"/api/session/{session['id']}/message", json={"text": "EGP"}).get_json()["session"]
            self.assertEqual(session["stage"], "post_booking_support")
            self.assertTrue(session["bookingCompleted"])
            self.assertIsNotNone(session["bookingResult"])
            self.assertEqual(session["bookingResult"]["trip_id"], "RT-LOC-26-900")
            self.assertEqual(session["stats"]["bookingDraftCount"], 1)
        finally:
            os.environ.clear()
            os.environ.update(original_env)
            shutil.rmtree(tmp_path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
