from __future__ import annotations

import re
import sqlite3
import uuid
import json
import os
import threading
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import SystemServiceSettings, load_system_settings
from .field_mapping import (
    SHEET_TABLE_MAPPINGS,
    TRIPS_DATA_START_ROW,
    TRIPS_SHEET_COLUMNS,
    TRIPS_SHEET_NAME,
)


BLOCKED_STATUSES = {"blacklisted", "blacklist"}
REVIEW_STATUSES = {"payment risk", "high maintenance"}
INACTIVE_TRIP_STATUSES = {"cancelled", "closed", "archived"}
QUALIFIED_LEAD_STAGES = {"Qualified", "VIP Priority", "Repeat Priority"}
BOOKING_DRAFT_LEAD_STAGES = {"Booking Draft Created", "VIP Booking Draft", "Repeat Booking Draft"}
ROOM_HOLD_COLUMNS = {
    "Single": "draft_holds_single",
    "Double": "draft_holds_double",
    "Triple": "draft_holds_triple",
}
ROOM_REMAINING_COLUMNS = {
    "Single": "single_remaining",
    "Double": "double_remaining",
    "Triple": "triple_remaining",
}
_SHEET_WRITE_LOCK = threading.RLock()


@dataclass
class IdentityResolution:
    match_status: str
    handoff_required: bool
    handoff_reason: str
    traveler: dict[str, Any] | None
    lookup_phone: dict[str, str]
    name_match_status: str
    actions: list[str]


class UnifiedCRMService:
    def __init__(self, settings: SystemServiceSettings | None = None) -> None:
        self.settings = settings or load_system_settings()

    @contextmanager
    def connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.settings.db_path)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
        finally:
            connection.close()

    @staticmethod
    def normalize_phone(raw_phone: str, country_code: str = "20") -> dict[str, str]:
        text = (raw_phone or "").split("@", 1)[0]
        text = re.sub(r"[^\d+]", "", text)
        if not text:
            return {
                "country_code": country_code,
                "local_number": "",
                "normalized_whatsapp": "",
                "lookup_key": "",
            }
        if text.startswith("00"):
            text = "+" + text[2:]
        if text.startswith("+"):
            digits = re.sub(r"\D", "", text)
            local_number = digits[2:] if digits.startswith(country_code) else digits
            return {
                "country_code": country_code,
                "local_number": local_number,
                "normalized_whatsapp": text,
                "lookup_key": digits[-10:],
            }
        local_number = text[1:] if text.startswith("0") else text
        normalized = f"+{country_code}{local_number}"
        return {
            "country_code": country_code,
            "local_number": local_number,
            "normalized_whatsapp": normalized,
            "lookup_key": re.sub(r"\D", "", normalized)[-10:],
        }

    @staticmethod
    def normalize_trip_type(raw_type: str | None) -> str | None:
        lowered = (raw_type or "").strip().lower()
        if lowered in {"local", "inside egypt", "egypt", "egypt trips"}:
            return "Local"
        if lowered in {"international", "abroad", "outside egypt"}:
            return "International"
        return None

    @staticmethod
    def normalize_name(value: str) -> str:
        lowered = (value or "").casefold().strip()
        lowered = re.sub(r"[^0-9a-z\u0600-\u06ff\s]+", " ", lowered)
        return " ".join(lowered.split())

    @classmethod
    def names_are_compatible(cls, submitted_name: str, existing_name: str) -> bool:
        submitted = cls.normalize_name(submitted_name)
        existing = cls.normalize_name(existing_name)
        if not submitted or not existing:
            return False
        if submitted == existing:
            return True
        submitted_tokens = submitted.split()
        existing_tokens = existing.split()
        if len(submitted_tokens) < 2 or len(existing_tokens) < 2:
            return False
        if submitted_tokens[0] == existing_tokens[0] and submitted_tokens[-1] == existing_tokens[-1]:
            return True
        shared = set(submitted_tokens) & set(existing_tokens)
        shorter = min(len(set(submitted_tokens)), len(set(existing_tokens)))
        return bool(shorter and len(shared) >= 2 and (len(shared) / shorter) >= 0.80)

    def resolve_identity(self, full_name: str, raw_phone: str, country_code: str = "20") -> IdentityResolution:
        phone = self.normalize_phone(raw_phone, country_code)
        with self.connect() as connection:
            matches = connection.execute(
                """
                SELECT traveler_id, status, full_name, birthday, gender, nationality,
                       phone_code, whatsapp_raw, integrated_whatsapp, normalized_whatsapp,
                       phone_lookup_key, local_trips_count, international_trips_count, total_trips,
                       lead_source, created_at, last_contacted_at, agent_notes, data_audit
                FROM travelers
                WHERE phone_lookup_key = ?
                   OR integrated_whatsapp = ?
                   OR normalized_whatsapp = ?
                   OR whatsapp_raw = ?
                """,
                (
                    phone["lookup_key"],
                    phone["normalized_whatsapp"],
                    phone["normalized_whatsapp"],
                    phone["local_number"],
                ),
            ).fetchall()

        if not matches:
            return IdentityResolution(
                match_status="not_found",
                handoff_required=False,
                handoff_reason="",
                traveler=None,
                lookup_phone=phone,
                name_match_status="",
                actions=["collect_new_traveler_data"],
            )

        if len(matches) > 1:
            return IdentityResolution(
                match_status="multiple_matches",
                handoff_required=True,
                handoff_reason="duplicate_phone_match",
                traveler=None,
                lookup_phone=phone,
                name_match_status="",
                actions=["human_review_duplicate_phone"],
            )

        traveler = dict(matches[0])
        status = (traveler.get("status") or "").strip().lower()
        if status in BLOCKED_STATUSES:
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason="blacklisted_customer",
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="blocked_status",
                actions=["block_sales_flow"],
            )
        if not self.names_are_compatible(full_name, traveler.get("full_name") or ""):
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason="phone_name_conflict",
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="conflict",
                actions=["human_review_identity_conflict"],
            )
        if status in REVIEW_STATUSES:
            return IdentityResolution(
                match_status="single_match",
                handoff_required=True,
                handoff_reason=status.replace(" ", "_"),
                traveler=traveler,
                lookup_phone=phone,
                name_match_status="matched",
                actions=["allow_conversation_but_require_human_review"],
            )
        return IdentityResolution(
            match_status="single_match",
            handoff_required=False,
            handoff_reason="",
            traveler=traveler,
            lookup_phone=phone,
            name_match_status="matched",
            actions=["continue_sales_flow"],
        )

    def recalculate_traveler_stats(self, traveler_id: str) -> None:
        if not traveler_id:
            return
        with self.connect() as connection:
            try:
                # 1. Count local trips
                local_trips = connection.execute(
                    """
                    SELECT COUNT(*) FROM trip_bookings b
                    JOIN trips t ON TRIM(b.trip_id) = TRIM(t.trip_id)
                    WHERE TRIM(b.traveler_id) = ? 
                      AND IFNULL(b.booking_status, '') != 'Cancelled' 
                      AND TRIM(t.type) = 'Local'
                    """,
                    (traveler_id.strip(),),
                ).fetchone()[0] or 0

                # 2. Count international trips
                intl_trips = connection.execute(
                    """
                    SELECT COUNT(*) FROM trip_bookings b
                    JOIN trips t ON TRIM(b.trip_id) = TRIM(t.trip_id)
                    WHERE TRIM(b.traveler_id) = ? 
                      AND IFNULL(b.booking_status, '') != 'Cancelled' 
                      AND TRIM(t.type) = 'International'
                    """,
                    (traveler_id.strip(),),
                ).fetchone()[0] or 0

                # 3. Count community events
                comm_events = connection.execute(
                    """
                    SELECT COUNT(*) FROM ce_bookings
                    WHERE TRIM(traveler_id) = ? AND IFNULL(status, '') != 'Cancelled'
                    """,
                    (traveler_id.strip(),),
                ).fetchone()[0] or 0

                # Update the database
                connection.execute(
                    """
                    UPDATE travelers
                    SET local_trips_count = ?,
                        international_trips_count = ?,
                        total_trips = ?,
                        community_events_count = ?
                    WHERE traveler_id = ?
                    """,
                    (local_trips, intl_trips, local_trips + intl_trips, comm_events, traveler_id),
                )
                connection.commit()
            except sqlite3.OperationalError:
                pass

    def refresh_traveler_sheet_stats(self, traveler_id: str) -> bool:
        """Pull profile counters from the Travelers sheet for one traveler."""
        if not traveler_id:
            return False
        return bool(self.refresh_travelers_sheet_stats([traveler_id.strip()]))

    def refresh_travelers_sheet_stats(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        """Pull profile counters from the Travelers sheet for multiple travelers."""
        clean_ids = sorted({str(traveler_id or "").strip() for traveler_id in traveler_ids if str(traveler_id or "").strip()})
        if not clean_ids:
            return {}

        stats_by_id = self._read_traveler_stats_by_id_from_sheet(clean_ids)
        if not stats_by_id:
            return {}

        with self.connect() as connection:
            for traveler_id, values in stats_by_id.items():
                if not values:
                    continue
                assignments = []
                params: list[Any] = []
                for field_name, value in values.items():
                    assignments.append(f"{field_name} = ?")
                    params.append(value)
                params.append(traveler_id)
                connection.execute(
                    f"UPDATE travelers SET {', '.join(assignments)} WHERE traveler_id = ?",
                    params,
                )
            connection.commit()
        return stats_by_id

    def _read_traveler_stats_from_sheet(self, traveler_id: str) -> dict[str, Any]:
        stats_by_id = self._read_traveler_stats_by_id_from_sheet([traveler_id])
        return stats_by_id.get(traveler_id, {})

    def _read_traveler_stats_by_id_from_sheet(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        if self.settings.sheet_backend == "google_sheets":
            return self._read_traveler_stats_by_id_from_google_sheet(traveler_ids)
        return self._read_traveler_stats_by_id_from_excel(traveler_ids)

    def _read_traveler_stats_by_id_from_excel(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        wanted = {str(traveler_id or "").strip() for traveler_id in traveler_ids if str(traveler_id or "").strip()}
        if not wanted:
            return {}

        workbook_path = self.settings.excel_runtime_workbook or self.settings.excel_source_workbook
        if not workbook_path or not workbook_path.exists():
            return {}

        stats_by_id: dict[str, dict[str, Any]] = {}
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if "Travelers" not in wb.sheetnames:
                return {}
            ws = wb["Travelers"]
            headers = {
                str(ws.cell(1, col_idx).value or "").strip().lower(): col_idx
                for col_idx in range(1, ws.max_column + 1)
                if ws.cell(1, col_idx).value
            }
            id_col = headers.get("traveler id")
            if not id_col:
                return {}
            for row_idx in range(2, ws.max_row + 1):
                current = str(ws.cell(row_idx, id_col).value or "").strip()
                if current not in wanted:
                    continue
                stats_by_id[current] = self._extract_traveler_stats_from_cells(
                    lambda header, row_idx=row_idx: ws.cell(row_idx, headers[header]).value if header in headers else None
                )
                if len(stats_by_id) == len(wanted):
                    break
        finally:
            wb.close()
        return stats_by_id

    def _read_traveler_stats_by_id_from_google_sheet(self, traveler_ids: list[str]) -> dict[str, dict[str, Any]]:
        wanted = {str(traveler_id or "").strip() for traveler_id in traveler_ids if str(traveler_id or "").strip()}
        if not wanted or not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            return {}

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets.readonly",
            "https://www.googleapis.com/auth/drive.readonly",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        ws = client.open_by_key(self.settings.google_sheet_id).worksheet("Travelers")
        rows = ws.get_all_values()
        if not rows:
            return {}

        headers = {
            str(header or "").strip().lower(): col_idx
            for col_idx, header in enumerate(rows[0])
            if str(header or "").strip()
        }
        id_idx = headers.get("traveler id")
        if id_idx is None:
            return {}

        stats_by_id: dict[str, dict[str, Any]] = {}
        for row in rows[1:]:
            current = row[id_idx] if len(row) > id_idx else ""
            current = str(current).strip()
            if current not in wanted:
                continue
            stats_by_id[current] = self._extract_traveler_stats_from_cells(
                lambda header, row=row: row[headers[header]] if header in headers and len(row) > headers[header] else None
            )
            if len(stats_by_id) == len(wanted):
                break
        return stats_by_id

    def _read_traveler_stats_from_excel(self, traveler_id: str) -> dict[str, Any]:
        stats_by_id = self._read_traveler_stats_by_id_from_excel([traveler_id])
        return stats_by_id.get(traveler_id, {})

    def _read_traveler_stats_from_google_sheet(self, traveler_id: str) -> dict[str, Any]:
        stats_by_id = self._read_traveler_stats_by_id_from_google_sheet([traveler_id])
        return stats_by_id.get(traveler_id, {})

    def _extract_traveler_stats_from_cells(self, value_for_header) -> dict[str, Any]:
        field_headers = {
            "local_trips_count": ("loc. trips", "local trips count"),
            "international_trips_count": ("int. trips", "international trips count"),
            "total_trips": ("total trips",),
            "community_events_count": ("comm. events", "community events count"),
            "lifetime_revenue": ("lifetime revenue",),
        }
        values: dict[str, Any] = {}
        for field_name, headers in field_headers.items():
            raw_value = None
            for header in headers:
                raw_value = value_for_header(header)
                if raw_value not in (None, ""):
                    break
            if raw_value in (None, ""):
                continue
            if field_name == "lifetime_revenue":
                values[field_name] = self._parse_sheet_float(raw_value)
            else:
                values[field_name] = self._parse_sheet_int(raw_value)
        return values

    @staticmethod
    def _parse_sheet_int(value: Any) -> int:
        try:
            return int(float(str(value).replace(",", "").strip()))
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _parse_sheet_float(value: Any) -> float:
        try:
            return float(str(value).replace(",", "").replace("$", "").strip())
        except (TypeError, ValueError):
            return 0.0

    def next_traveler_id(self) -> str:
        with self.connect() as connection:
            row = connection.execute(
                """
                SELECT traveler_id 
                FROM travelers 
                WHERE traveler_id LIKE 'TR%' 
                ORDER BY CAST(SUBSTR(traveler_id, 3) AS INTEGER) DESC 
                LIMIT 1
                """
            ).fetchone()

        max_number = 0
        if row and row["traveler_id"]:
            traveler_id = str(row["traveler_id"]).strip()
            match = re.fullmatch(r"TR(\d+)", traveler_id)
            if match:
                max_number = int(match.group(1))

        next_number = max_number + 1
        return f"TR{next_number:05d}"

    def create_traveler(
        self,
        *,
        full_name: str,
        raw_phone: str,
        birthday: str = "",
        gender: str = "",
        nationality: str = "",
        lead_source: str = "",
        agent_notes: str = "",
        country_code: str = "20",
    ) -> dict[str, Any]:
        phone = self.normalize_phone(raw_phone, country_code)
        first_name, last_name = self._split_name(full_name)
        now = datetime.utcnow().replace(microsecond=0).isoformat()
        
        max_retries = 3
        for attempt in range(max_retries):
            traveler_id = self.next_traveler_id()
            try:
                with self.connect() as connection:
                    connection.execute(
                        """
                        INSERT INTO travelers (
                            traveler_id, status, full_name, first_name, last_name, birthday, gender,
                            nationality, phone_code, whatsapp_raw, integrated_whatsapp, normalized_whatsapp,
                            phone_lookup_key, lead_source, created_at, last_contacted_at, agent_notes
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            traveler_id,
                            "Active",
                            full_name.strip(),
                            first_name,
                            last_name,
                            birthday or None,
                            gender or None,
                            nationality or None,
                            phone["country_code"] or None,
                            phone["local_number"] or None,
                            phone["normalized_whatsapp"] or None,
                            phone["normalized_whatsapp"] or None,
                            phone["lookup_key"] or None,
                            lead_source or None,
                            now,
                            now,
                            agent_notes or None,
                        ),
                    )
                    connection.commit()
                break
            except sqlite3.IntegrityError:
                if attempt == max_retries - 1:
                    raise
        return {
            "traveler_id": traveler_id,
            "full_name": full_name.strip(),
            "birthday": birthday or None,
            "gender": gender or None,
            "nationality": nationality or None,
            "phone_lookup_key": phone["lookup_key"],
            "integrated_whatsapp": phone["normalized_whatsapp"],
        }

    def upsert_lead(
        self,
        *,
        customer_name: str,
        raw_phone: str,
        traveler_id: str | None,
        lead_stage: str,
        lead_source: str,
        channel: str,
        preferred_trip_type: str = "",
        interested_trip_ids: str = "",
        suggested_trip_ids: str = "",
        priority: str = "Medium",
        follow_up_status: str = "",
        follow_up_due_date: str = "",
        notes: str = "",
        booking_id: str = "",
        traveler_status: str = "",
        customer_tier: str = "",
        match_status: str = "",
        last_interaction_id: str = "",
        flow_key: str = "",
        current_step: str = "",
        handoff_required: bool = False,
        handoff_reason: str = "",
        language: str = "",
        country_code: str = "20",
    ) -> dict[str, Any]:
        phone = self.normalize_phone(raw_phone, country_code)
        now = datetime.utcnow().replace(microsecond=0).isoformat()
        with self.connect() as connection:
            existing = connection.execute(
                """
                SELECT lead_id, interaction_count FROM leads
                WHERE (? <> '' AND phone_lookup_key = ?)
                   OR (? <> '' AND traveler_id = ?)
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (phone["lookup_key"], phone["lookup_key"], traveler_id or "", traveler_id or ""),
            ).fetchone()
            if existing:
                lead_id = str(existing["lead_id"])
                interaction_count = self._as_int(existing["interaction_count"], default=0) or 0
                connection.execute(
                    """
                    UPDATE leads
                    SET updated_at = ?, customer_name = ?, raw_phone = ?, integrated_whatsapp = ?,
                        phone_lookup_key = ?, traveler_id = ?, traveler_status = ?, customer_tier = ?,
                        match_status = ?, lead_stage = ?, lead_source = ?, channel = ?,
                        preferred_trip_type = ?, interested_trip_ids = ?, suggested_trip_ids = ?,
                        priority = ?, follow_up_status = ?, follow_up_due_date = ?, last_interaction_id = ?,
                        interaction_count = ?, notes = ?, booking_id = ?, flow_key = ?, current_step = ?,
                        handoff_required = ?, handoff_reason = ?, language = ?
                    WHERE lead_id = ?
                    """,
                    (
                        now,
                        customer_name,
                        phone["local_number"] or raw_phone,
                        phone["normalized_whatsapp"],
                        phone["lookup_key"],
                        traveler_id,
                        traveler_status or None,
                        customer_tier or None,
                        match_status or None,
                        lead_stage,
                        lead_source,
                        channel,
                        preferred_trip_type or None,
                        interested_trip_ids or None,
                        suggested_trip_ids or None,
                        priority,
                        follow_up_status or None,
                        follow_up_due_date or None,
                        last_interaction_id or None,
                        interaction_count + (1 if last_interaction_id else 0),
                        notes or None,
                        booking_id or None,
                        flow_key or None,
                        current_step or None,
                        1 if handoff_required else 0,
                        handoff_reason or None,
                        language or None,
                        lead_id,
                    ),
                )
            else:
                lead_id = self._next_prefixed_id(connection, "leads", "lead_id", "LD", 5)
                connection.execute(
                    """
                    INSERT INTO leads (
                        lead_id, created_at, updated_at, customer_name, raw_phone, integrated_whatsapp,
                        phone_lookup_key, traveler_id, traveler_status, customer_tier, match_status,
                        lead_stage, lead_source, channel, preferred_trip_type, interested_trip_ids,
                        suggested_trip_ids, priority, follow_up_status, follow_up_due_date,
                        last_interaction_id, interaction_count, notes, booking_id, flow_key,
                        current_step, handoff_required, handoff_reason, language
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        lead_id,
                        now,
                        now,
                        customer_name,
                        phone["local_number"] or raw_phone,
                        phone["normalized_whatsapp"],
                        phone["lookup_key"],
                        traveler_id,
                        traveler_status or None,
                        customer_tier or None,
                        match_status or None,
                        lead_stage,
                        lead_source,
                        channel,
                        preferred_trip_type or None,
                        interested_trip_ids or None,
                        suggested_trip_ids or None,
                        priority,
                        follow_up_status or None,
                        follow_up_due_date or None,
                        last_interaction_id or None,
                        1 if last_interaction_id else 0,
                        notes or None,
                        booking_id or None,
                        flow_key or None,
                        current_step or None,
                        1 if handoff_required else 0,
                        handoff_reason or None,
                        language or None,
                    ),
                )
            if traveler_id:
                connection.execute(
                    """
                    UPDATE travelers
                    SET last_lead_id = ?, last_contacted_at = ?
                    WHERE traveler_id = ?
                    """,
                    (lead_id, now, traveler_id),
                )
            connection.commit()
        return {
            "lead_id": lead_id,
            "lead_stage": lead_stage,
            "priority": priority,
            "follow_up_status": follow_up_status,
            "follow_up_due_date": follow_up_due_date,
            "interaction_id": last_interaction_id,
        }

    def record_agent_outcome(
        self,
        *,
        full_name: str,
        raw_phone: str,
        country_code: str,
        trip_type: str | None,
        channel: str,
        source: str,
        agent_notes: str,
        birthday: str = "",
        gender: str = "",
        nationality: str = "",
        preferred_trip_id: str = "",
        language: str = "",
        manual_lead_stage: str = "",
        manual_priority: str = "",
        manual_follow_up_due_date: str = "",
    ) -> dict[str, Any]:
        """Persist an agent qualification outcome to the system DB, then sync sheets."""
        self.ensure_operational_schema()
        resolution = self.resolve_identity(full_name, raw_phone, country_code)
        trip_result = self.build_trip_result(trip_type) if trip_type else {"open_trips": [], "date_tbd_trips": []}
        timestamp = datetime.utcnow().replace(microsecond=0)

        created_traveler: dict[str, Any] | None = None
        traveler = resolution.traveler.copy() if resolution.traveler else None
        traveler_id = str(traveler.get("traveler_id")) if traveler else ""
        status_snapshot = str(traveler.get("status") or "") if traveler else ""
        actions = list(resolution.actions)

        if resolution.match_status == "not_found":
            created_traveler = self.create_traveler(
                full_name=full_name,
                raw_phone=raw_phone,
                birthday=birthday,
                gender=gender,
                nationality=nationality,
                lead_source=source,
                agent_notes=agent_notes,
                country_code=country_code,
            )
            traveler = created_traveler.copy()
            traveler_id = created_traveler["traveler_id"]
            actions.append("created_new_traveler")
        elif resolution.match_status == "single_match" and traveler_id and not resolution.handoff_required:
            profile_updates = self.update_existing_traveler_profile(
                traveler_id,
                birthday=birthday,
                gender=gender,
                nationality=nationality,
                timestamp=timestamp,
            )
            if profile_updates and traveler:
                traveler.update(profile_updates)
                actions.append("updated_missing_traveler_profile")

        suggested_trip_ids = self._collect_trip_ids(trip_result)
        if preferred_trip_id:
            interested_trip_ids = preferred_trip_id
        else:
            interested_trip_ids = ", ".join(suggested_trip_ids)

        preview: dict[str, Any] = {
            "match_status": resolution.match_status,
            "handoff_required": resolution.handoff_required,
            "handoff_reason": resolution.handoff_reason,
            "traveler": traveler,
            "lookup_phone": resolution.lookup_phone,
            "name_match_status": resolution.name_match_status,
            "actions": actions,
            "trip_result": trip_result,
        }

        derived_stage, derived_priority = self.derive_lead_stage(preview)
        follow_up_status, derived_due_date = self.derive_follow_up(preview, timestamp)

        if resolution.handoff_required:
            lead_stage = derived_stage
            priority = derived_priority
            follow_up_due_date = derived_due_date
        else:
            lead_stage = manual_lead_stage or derived_stage
            priority = manual_priority or derived_priority
            follow_up_due_date = manual_follow_up_due_date or derived_due_date

        traveler_status = str(traveler.get("status") or "") if traveler else ""
        customer_tier = self.infer_customer_tier(preview)

        interaction = self.create_interaction(
            timestamp=timestamp,
            channel=channel,
            customer_name=full_name,
            raw_phone=raw_phone,
            integrated_whatsapp=resolution.lookup_phone.get("normalized_whatsapp", ""),
            phone_lookup_key=resolution.lookup_phone.get("lookup_key", ""),
            traveler_id=traveler_id,
            matched_row=None,
            status_snapshot=status_snapshot or ("MULTIPLE_MATCHES" if resolution.match_status == "multiple_matches" else ""),
            intent="sales_inquiry",
            trip_type=trip_type or "",
            suggested_trips=", ".join(suggested_trip_ids),
            action_taken=", ".join(actions),
            handoff_required=resolution.handoff_required,
            handoff_reason=resolution.handoff_reason,
            agent_notes=agent_notes,
            flow_key="booking",
            step_key="lead_qualification",
            language=language,
            outcome=lead_stage,
        )

        lead_update = self.upsert_lead(
            customer_name=full_name,
            raw_phone=raw_phone,
            traveler_id=traveler_id or None,
            lead_stage=lead_stage,
            lead_source=source,
            channel=channel,
            preferred_trip_type=trip_type or "",
            interested_trip_ids=interested_trip_ids,
            suggested_trip_ids=", ".join(suggested_trip_ids),
            priority=priority,
            follow_up_status=follow_up_status,
            follow_up_due_date=follow_up_due_date,
            notes=agent_notes,
            traveler_status=traveler_status,
            customer_tier=customer_tier,
            match_status=resolution.match_status,
            last_interaction_id=interaction["interaction_id"],
            flow_key="booking",
            current_step="lead_qualification",
            handoff_required=resolution.handoff_required,
            handoff_reason=resolution.handoff_reason,
            language=language,
            country_code=country_code,
        )

        events = [
            self.create_booking_event(
                event_type="inquiry_received",
                event_label="Inquiry received",
                traveler_id=traveler_id,
                lead_id=lead_update["lead_id"],
                trip_id=preferred_trip_id,
                interaction_id=interaction["interaction_id"],
                channel=channel,
                notes=agent_notes,
                metadata={"trip_type": trip_type or ""},
                occurred_at=timestamp,
            )
        ]
        if suggested_trip_ids:
            events.append(
                self.create_booking_event(
                    event_type="trip_suggested",
                    event_label="Trip suggested",
                    traveler_id=traveler_id,
                    lead_id=lead_update["lead_id"],
                    trip_id=preferred_trip_id or suggested_trip_ids[0],
                    interaction_id=interaction["interaction_id"],
                    channel=channel,
                    notes=", ".join(suggested_trip_ids),
                    metadata={"suggested_trip_ids": suggested_trip_ids},
                    occurred_at=timestamp,
                )
            )
        if lead_stage in QUALIFIED_LEAD_STAGES:
            events.append(
                self.create_booking_event(
                    event_type="lead_qualified",
                    event_label="Lead qualified",
                    traveler_id=traveler_id,
                    lead_id=lead_update["lead_id"],
                    trip_id=preferred_trip_id or (suggested_trip_ids[0] if suggested_trip_ids else ""),
                    interaction_id=interaction["interaction_id"],
                    channel=channel,
                    notes=lead_stage,
                    metadata={"priority": priority},
                    occurred_at=timestamp,
                )
            )
        if resolution.handoff_required:
            events.append(
                self.create_booking_event(
                    event_type="handoff_required",
                    event_label="Handoff required",
                    traveler_id=traveler_id,
                    lead_id=lead_update["lead_id"],
                    interaction_id=interaction["interaction_id"],
                    channel=channel,
                    notes=resolution.handoff_reason,
                    metadata={"handoff_reason": resolution.handoff_reason},
                    occurred_at=timestamp,
                )
            )

        self.sync_agent_write_to_sheet(
            traveler_id=traveler_id,
            lead_id=lead_update["lead_id"],
            interaction_id=interaction["interaction_id"],
            event_ids=[event["event_id"] for event in events],
        )

        preview["write_result"] = {
            "created_traveler": created_traveler,
            "interaction_log": {
                "interaction_id": interaction["interaction_id"],
                "row": None,
            },
            "lead_update": lead_update,
            "event_trail": events,
        }
        return preview

    def update_existing_traveler_profile(
        self,
        traveler_id: str,
        *,
        birthday: str = "",
        gender: str = "",
        nationality: str = "",
        timestamp: datetime | None = None,
    ) -> dict[str, Any]:
        timestamp = timestamp or datetime.utcnow().replace(microsecond=0)
        updates: dict[str, Any] = {}
        with self.connect() as connection:
            row = connection.execute(
                "SELECT birthday, gender, nationality FROM travelers WHERE traveler_id = ?",
                (traveler_id,),
            ).fetchone()
            if not row:
                return updates
            fields: dict[str, Any] = {"last_contacted_at": timestamp.isoformat()}
            if birthday and not row["birthday"]:
                fields["birthday"] = birthday
                updates["birthday"] = birthday
            if gender and not row["gender"]:
                fields["gender"] = gender
                updates["gender"] = gender
            if nationality and not row["nationality"]:
                fields["nationality"] = nationality
                updates["nationality"] = nationality
            assignments = ", ".join(f"{key} = ?" for key in fields)
            connection.execute(
                f"UPDATE travelers SET {assignments} WHERE traveler_id = ?",
                (*fields.values(), traveler_id),
            )
            connection.commit()
        return updates

    def create_interaction(
        self,
        *,
        timestamp: datetime,
        channel: str,
        customer_name: str,
        raw_phone: str,
        integrated_whatsapp: str,
        phone_lookup_key: str,
        traveler_id: str,
        matched_row: int | None,
        status_snapshot: str,
        intent: str,
        trip_type: str,
        suggested_trips: str,
        action_taken: str,
        handoff_required: bool,
        handoff_reason: str,
        agent_notes: str,
        flow_key: str = "",
        step_key: str = "",
        message_key: str = "",
        language: str = "",
        outcome: str = "",
    ) -> dict[str, Any]:
        with self.connect() as connection:
            interaction_id = self._next_daily_id(connection, "interactions", "interaction_id", "INT", timestamp)
            connection.execute(
                """
                INSERT INTO interactions (
                    interaction_id, timestamp, channel, customer_name, raw_phone, integrated_whatsapp,
                    phone_lookup_key, traveler_id, matched_row, status_snapshot, intent, trip_type,
                    suggested_trips, action_taken, handoff_required, handoff_reason, agent_notes,
                    flow_key, step_key, message_key, language, outcome
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    interaction_id,
                    timestamp.isoformat(timespec="seconds"),
                    channel,
                    customer_name,
                    raw_phone,
                    integrated_whatsapp or None,
                    phone_lookup_key or None,
                    traveler_id or None,
                    matched_row,
                    status_snapshot or None,
                    intent or None,
                    trip_type or None,
                    suggested_trips or None,
                    action_taken or None,
                    1 if handoff_required else 0,
                    handoff_reason or None,
                    agent_notes or None,
                    flow_key or None,
                    step_key or None,
                    message_key or None,
                    language or None,
                    outcome or None,
                ),
            )
            connection.commit()
        return {"interaction_id": interaction_id}

    def create_booking_draft(
        self,
        *,
        trip_id: str,
        traveler_id: str,
        traveler_name: str,
        room_type: str,
        channel: str = "",
        lead_id: str = "",
        flight_option: str = "",
        date_option: str = "",
        currency: str = "",
        source: str = "",
        agent_notes: str = "",
    ) -> dict[str, Any]:
        return self.create_booking(
            trip_id=trip_id,
            traveler_id=traveler_id,
            traveler_name=traveler_name,
            room_type=room_type,
            channel=channel,
            lead_id=lead_id,
            flight_option=flight_option,
            date_option=date_option,
            currency=currency,
            booking_status="Draft",
            booking_source=source,
            payment_status="Awaiting Deposit",
            booking_notes=agent_notes,
        )

    def create_booking(
        self,
        *,
        trip_id: str,
        traveler_id: str,
        traveler_name: str,
        room_type: str,
        channel: str = "",
        lead_id: str = "",
        flight_option: str = "",
        date_option: str = "",
        currency: str = "",
        booking_status: str = "Draft",
        booking_source: str = "",
        interaction_id: str = "",
        payment_status: str = "Pending",
        booking_notes: str = "",
    ) -> dict[str, Any]:
        self.ensure_operational_schema()
        timestamp = datetime.utcnow().replace(microsecond=0)
        if room_type not in ROOM_HOLD_COLUMNS:
            raise ValueError(f"Unsupported room type {room_type!r}")

        with self.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            trip = connection.execute(
                """
                SELECT trip_id, trip_name, type, sales_status, single_remaining, double_remaining,
                       triple_remaining, draft_holds_single, draft_holds_double, draft_holds_triple
                 FROM trips
                 WHERE trip_id = ?
                 """,
                 (trip_id,),
             ).fetchone()
            if not trip:
                raise ValueError(f"Trip was not found for booking: {trip_id}")
            sales_status = str(trip["sales_status"] or "").strip().lower()
            if sales_status and sales_status != "open":
                raise ValueError(f"Trip {trip_id} is not bookable because its status is {trip['sales_status']}")

            remaining = self._as_int(trip[ROOM_REMAINING_COLUMNS[room_type]])
            current_hold = self._as_int(trip[ROOM_HOLD_COLUMNS[room_type]], default=0) or 0
            if remaining is None:
                raise ValueError(f"Capacity is not configured for room type {room_type}")
            available = remaining - current_hold
            if available <= 0:
                raise ValueError(f"No remaining draftable capacity for {room_type}")

            if not interaction_id:
                interaction_id = self._next_daily_id(connection, "interactions", "interaction_id", "INT", timestamp)
                connection.execute(
                    """
                    INSERT INTO interactions (
                        interaction_id, timestamp, channel, customer_name, raw_phone,
                        traveler_id, status_snapshot, intent, trip_type, suggested_trips,
                        action_taken, handoff_required, agent_notes, flow_key, step_key, outcome
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        interaction_id,
                        timestamp.isoformat(timespec="seconds"),
                        channel or None,
                        traveler_name,
                        "",
                        traveler_id,
                        "BOOKING_DRAFT",
                        "booking_request",
                        str(trip["type"] or ""),
                        trip_id,
                        "booking_draft_created",
                        0,
                        booking_notes or None,
                        "booking",
                        "booking_draft_created",
                        "Booking Draft Created",
                    ),
                )

            booking_id = self._next_booking_id(connection, trip_id, traveler_id)
            connection.execute(
                """
                INSERT INTO trip_bookings (
                    booking_id, trip_id, trip_name, traveler_id, traveler_name, room_type,
                    flight_option, date_option, currency, booking_status, draft_created_at,
                    booking_source, lead_id, interaction_id, payment_status, booking_notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    booking_id,
                    trip_id,
                    str(trip["trip_name"]) if trip else "",
                    traveler_id,
                    traveler_name,
                    room_type,
                    flight_option or None,
                    date_option or None,
                    currency or None,
                    booking_status,
                    timestamp.isoformat(timespec="seconds"),
                    booking_source or None,
                    lead_id or None,
                    interaction_id or None,
                    payment_status or None,
                    booking_notes or None,
                ),
            )
            connection.execute(
                f"UPDATE trips SET {ROOM_HOLD_COLUMNS[room_type]} = ? WHERE trip_id = ?",
                (current_hold + 1, trip_id),
            )
            if traveler_id:
                connection.execute(
                    "UPDATE travelers SET last_booking_id = ? WHERE traveler_id = ?",
                    (booking_id, traveler_id),
                )
            lead_update = None
            if lead_id:
                lead_update = self._update_lead_for_booking(
                    connection,
                    lead_id=lead_id,
                    booking_id=booking_id,
                    trip_id=trip_id,
                    interaction_id=interaction_id,
                    timestamp=timestamp,
                )
            connection.commit()

        booking_event = self.create_booking_event(
            event_type="booking_draft_created",
            event_label="Booking draft created",
            traveler_id=traveler_id,
            lead_id=lead_id,
            booking_id=booking_id,
            trip_id=trip_id,
            interaction_id=interaction_id,
            channel=channel,
            notes=booking_notes,
            metadata={"room_type": room_type, "available_before_draft": available, "available_after_draft": available - 1},
            occurred_at=timestamp,
        )
        follow_up_event = self.create_booking_event(
            event_type="payment_follow_up",
            event_label="Payment / follow-up pending",
            traveler_id=traveler_id,
            lead_id=lead_id,
            booking_id=booking_id,
            trip_id=trip_id,
            interaction_id=interaction_id,
            channel=channel,
            notes=payment_status,
            metadata={"payment_status": payment_status},
            occurred_at=timestamp,
        )

        self.sync_agent_write_to_sheet(
            traveler_id=traveler_id,
            lead_id=lead_id,
            interaction_id=interaction_id,
            booking_id=booking_id,
            trip_id=trip_id,
            event_ids=[booking_event["event_id"], follow_up_event["event_id"]],
        )

        result = {
            "booking_id": booking_id,
            "booking_row": None,
            "trip_id": trip_id,
            "trip_name": str(trip["trip_name"] or ""),
            "traveler_id": traveler_id,
            "traveler_name": traveler_name,
            "room_type": room_type,
            "flight_option": flight_option,
            "date_option": date_option,
            "currency": currency,
            "available_before_draft": available,
            "available_after_draft": available - 1,
            "interaction": {"interaction_id": interaction_id},
            "lead_update": lead_update,
            "event_trail": [booking_event, follow_up_event],
        }
        result["write_result"] = {
            "booking_draft": {
                "booking_id": booking_id,
                "trip_id": trip_id,
                "traveler_id": traveler_id,
                "lead_id": lead_id,
            },
            "interaction_log": {"interaction_id": interaction_id},
            "lead_update": lead_update,
            "event_trail": [booking_event, follow_up_event],
        }
        return result

    def qualify_lead(self, lead_id: str, *, channel: str = "") -> bool:
        self.ensure_operational_schema()
        timestamp = datetime.utcnow().replace(microsecond=0)
        with self.connect() as connection:
            lead = connection.execute(
                """
                SELECT lead_id, traveler_id, customer_name, raw_phone, channel, interested_trip_ids,
                       preferred_trip_type, interaction_count
                FROM leads
                WHERE lead_id = ?
                """,
                (lead_id,),
            ).fetchone()
            if not lead:
                return False
            interaction_id = self._next_daily_id(connection, "interactions", "interaction_id", "INT", timestamp)
            resolved_channel = channel or str(lead["channel"] or "")
            connection.execute(
                """
                INSERT INTO interactions (
                    interaction_id, timestamp, channel, customer_name, raw_phone, traveler_id,
                    status_snapshot, intent, trip_type, suggested_trips, action_taken,
                    handoff_required, flow_key, step_key, outcome
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    interaction_id,
                    timestamp.isoformat(timespec="seconds"),
                    resolved_channel or None,
                    lead["customer_name"],
                    lead["raw_phone"],
                    lead["traveler_id"],
                    "QUALIFIED",
                    "lead_qualification",
                    lead["preferred_trip_type"],
                    lead["interested_trip_ids"],
                    "lead_qualified",
                    0,
                    "booking",
                    "lead_qualified",
                    "Qualified",
                ),
            )
            connection.execute(
                """
                UPDATE leads
                SET lead_stage = ?, priority = ?, follow_up_status = ?, follow_up_due_date = ?,
                    updated_at = ?, last_interaction_id = ?, interaction_count = ?
                WHERE lead_id = ?
                """,
                (
                    "Qualified",
                    "High",
                    "Follow Up Soon",
                    (timestamp.date() + timedelta(days=2)).isoformat(),
                    timestamp.isoformat(timespec="seconds"),
                    interaction_id,
                    (self._as_int(lead["interaction_count"], default=0) or 0) + 1,
                    lead_id,
                ),
            )
            connection.commit()

        event = self.create_booking_event(
            event_type="lead_qualified",
            event_label="Lead qualified",
            traveler_id=str(lead["traveler_id"] or ""),
            lead_id=lead_id,
            trip_id=str(lead["interested_trip_ids"] or "").split(",", 1)[0].strip(),
            interaction_id=interaction_id,
            channel=resolved_channel,
            notes="Manual qualification",
            occurred_at=timestamp,
        )
        self.sync_agent_write_to_sheet(
            traveler_id=str(lead["traveler_id"] or ""),
            lead_id=lead_id,
            interaction_id=interaction_id,
            event_ids=[event["event_id"]],
        )
        return True

    def create_booking_event(
        self,
        *,
        event_type: str,
        event_label: str,
        traveler_id: str = "",
        lead_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        interaction_id: str = "",
        channel: str = "",
        actor: str = "agent",
        notes: str = "",
        metadata: dict[str, Any] | None = None,
        occurred_at: datetime | None = None,
    ) -> dict[str, Any]:
        occurred_at = occurred_at or datetime.utcnow().replace(microsecond=0)
        metadata_json = json.dumps(metadata or {}, ensure_ascii=True, sort_keys=True)
        with self.connect() as connection:
            self._ensure_booking_event_trail_table(connection)
            event_id = self._next_daily_id(connection, "booking_event_trail", "event_id", "EVT", occurred_at)
            connection.execute(
                """
                INSERT INTO booking_event_trail (
                    event_id, occurred_at, event_type, event_label, traveler_id, lead_id,
                    booking_id, trip_id, interaction_id, channel, actor, notes, metadata_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    event_id,
                    occurred_at.isoformat(timespec="seconds"),
                    event_type,
                    event_label,
                    traveler_id or None,
                    lead_id or None,
                    booking_id or None,
                    trip_id or None,
                    interaction_id or None,
                    channel or None,
                    actor or None,
                    notes or None,
                    metadata_json,
                ),
            )
            connection.commit()
        return {
            "event_id": event_id,
            "occurred_at": occurred_at.isoformat(timespec="seconds"),
            "event_type": event_type,
            "event_label": event_label,
            "traveler_id": traveler_id,
            "lead_id": lead_id,
            "booking_id": booking_id,
            "trip_id": trip_id,
            "interaction_id": interaction_id,
        }

    def sync_agent_write_to_sheet(
        self,
        *,
        traveler_id: str = "",
        lead_id: str = "",
        interaction_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        event_ids: list[str] | None = None,
    ) -> dict[str, Any]:
        if self.settings.sheet_backend == "excel":
            try:
                synced = self._run_sync_agent_write_to_sheet_worker(
                    traveler_id=traveler_id,
                    lead_id=lead_id,
                    interaction_id=interaction_id,
                    booking_id=booking_id,
                    trip_id=trip_id,
                    event_ids=event_ids,
                )
                return {"status": "ok", "backend": "excel", "synced": synced}
            except Exception as exc:
                return {"status": "failed", "backend": "excel", "reason": str(exc)}

        # Spin up a daemon thread to process sheet writing asynchronously
        import threading
        
        def run_async():
            try:
                self._run_sync_agent_write_to_sheet_worker(
                    traveler_id=traveler_id,
                    lead_id=lead_id,
                    interaction_id=interaction_id,
                    booking_id=booking_id,
                    trip_id=trip_id,
                    event_ids=event_ids,
                )
            except Exception:
                # Daemon threads should suppress unexpected runtime worker crashes
                pass

        thread = threading.Thread(target=run_async)
        thread.daemon = True
        thread.start()
        return {"status": "queued", "message": "Background sync thread initiated."}

    def _run_sync_agent_write_to_sheet_worker(
        self,
        *,
        traveler_id: str = "",
        lead_id: str = "",
        interaction_id: str = "",
        booking_id: str = "",
        trip_id: str = "",
        event_ids: list[str] | None = None,
    ) -> dict[str, Any]:
        synced: dict[str, Any] = {}
        if traveler_id:
            synced["traveler"] = self.sync_record_to_sheet("Travelers", traveler_id)
        if lead_id:
            synced["lead"] = self.sync_record_to_sheet("Leads", lead_id)
        if interaction_id:
            synced["interaction"] = self.sync_record_to_sheet("Interactions", interaction_id)
        if booking_id:
            synced["booking"] = self.sync_record_to_sheet("Trip Bookings", booking_id)
        for event_id in event_ids or []:
            synced.setdefault("events", []).append(self.sync_record_to_sheet("Booking Event Trail", event_id))
        return synced

    def sync_record_to_sheet(self, mapping_name: str, record_id: str) -> dict[str, Any]:
        if not record_id:
            return {"status": "skipped", "reason": "empty_record_id"}
        mapping = SHEET_TABLE_MAPPINGS[mapping_name]
        table_name = self._table_for_mapping(mapping_name)
        key_column = self._key_column_for_mapping(mapping_name)
        record = self._fetch_table_record(table_name, key_column, record_id)
        if not record:
            return {"status": "skipped", "reason": "record_not_found", "record_id": record_id}

        backend = self.settings.sheet_backend
        try:
            if backend == "excel":
                updated = 0
                for workbook_path in self._excel_outcome_workbooks():
                    self._upsert_record_in_excel(workbook_path, mapping, key_column, record)
                    updated += 1
                self._record_sync_success(mapping_name, record_id)
                return {"status": "ok", "backend": backend, "updated_workbooks": updated}
            if backend == "google_sheets":
                self._upsert_record_in_google_sheet(mapping, key_column, record)
                self._record_sync_success(mapping_name, record_id)
                return {"status": "ok", "backend": backend, "updated_workbooks": 1}
        except Exception as e:
            self._record_sync_failure(mapping_name, record_id, str(e))
            return {"status": "failed", "backend": backend, "reason": str(e)}
        return {"status": "skipped", "backend": backend, "reason": "Sheet sync backend not supported."}

    def _record_sync_success(self, mapping_name: str, record_id: str) -> None:
        with self.connect() as connection:
            self._ensure_sync_queue_table(connection)
            connection.execute(
                "DELETE FROM sync_queue WHERE mapping_name = ? AND record_id = ?",
                (mapping_name, record_id),
            )
            connection.commit()

    def _record_sync_failure(self, mapping_name: str, record_id: str, error_message: str) -> None:
        timestamp = datetime.utcnow().replace(microsecond=0).isoformat(timespec="seconds")
        with self.connect() as connection:
            self._ensure_sync_queue_table(connection)
            connection.execute(
                """
                INSERT INTO sync_queue (
                    mapping_name, record_id, status, error_message, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(mapping_name, record_id) DO UPDATE SET
                    status = excluded.status,
                    error_message = excluded.error_message,
                    updated_at = excluded.updated_at
                """,
                (mapping_name, record_id, "failed", error_message, timestamp, timestamp),
            )
            connection.commit()

    def get_sync_queue(self) -> list[dict[str, Any]]:
        with self.connect() as connection:
            connection.row_factory = sqlite3.Row
            rows = connection.execute("SELECT * FROM sync_queue ORDER BY updated_at DESC").fetchall()
            return [dict(r) for r in rows]

    def ensure_operational_schema(self) -> None:
        with self.connect() as connection:
            self._ensure_booking_event_trail_table(connection)
            self._ensure_sync_queue_table(connection)
            connection.commit()

    @staticmethod
    def _ensure_sync_queue_table(connection: sqlite3.Connection) -> None:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS sync_queue (
                mapping_name TEXT NOT NULL,
                record_id TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending',
                error_message TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (mapping_name, record_id)
            )
            """
        )

    @staticmethod
    def infer_customer_tier(result: dict[str, Any]) -> str:
        traveler = result.get("traveler")
        if isinstance(traveler, dict):
            status = str(traveler.get("status") or "").strip()
            if status in {"VIP", "Repeat"}:
                return status
            if status:
                return status
        return "Standard"

    @classmethod
    def derive_lead_stage(cls, result: dict[str, Any]) -> tuple[str, str]:
        traveler = result.get("traveler")
        traveler_status = str(traveler.get("status") or "").strip() if isinstance(traveler, dict) else ""
        trip_result = result.get("trip_result") or {}
        open_trips = trip_result.get("open_trips", [])
        tbd_trips = trip_result.get("date_tbd_trips", [])

        if result.get("handoff_reason") == "blacklisted_customer":
            return "Blocked", "Critical"
        if result.get("handoff_required"):
            return "Needs Review", "High"
        if traveler_status == "VIP":
            return ("VIP Priority", "High") if open_trips else ("VIP Follow Up", "High")
        if traveler_status == "Repeat":
            return ("Repeat Priority", "Medium") if open_trips else ("Repeat Follow Up", "Medium")
        if open_trips:
            return "Qualified", "Medium"
        if tbd_trips:
            return "Follow Up Needed", "Medium"
        if result.get("match_status") == "not_found":
            return "New Lead", "Medium"
        if result.get("match_status") == "single_match":
            return "Existing Traveler", "Low"
        return "New Inquiry", "Low"

    @classmethod
    def derive_follow_up(cls, result: dict[str, Any], timestamp: datetime) -> tuple[str, str]:
        stage, _priority = cls.derive_lead_stage(result)
        if stage == "Blocked":
            return "Do Not Contact", ""
        if result.get("handoff_required"):
            return "Urgent", timestamp.date().isoformat()
        trip_result = result.get("trip_result") or {}
        if trip_result.get("open_trips"):
            return "Follow Up Soon", (timestamp.date() + timedelta(days=2)).isoformat()
        if trip_result.get("date_tbd_trips"):
            return "Awaiting Dates", (timestamp.date() + timedelta(days=7)).isoformat()
        if result.get("match_status") == "not_found":
            return "Qualify Lead", (timestamp.date() + timedelta(days=3)).isoformat()
        return "Monitor", (timestamp.date() + timedelta(days=14)).isoformat()

    @staticmethod
    def _collect_trip_ids(trip_result: dict[str, list[dict[str, Any]]]) -> list[str]:
        open_ids = [str(item.get("trip_id") or "") for item in trip_result.get("open_trips", []) if item.get("trip_id")]
        tbd_ids = [str(item.get("trip_id") or "") for item in trip_result.get("date_tbd_trips", []) if item.get("trip_id")]
        return open_ids or tbd_ids

    def _update_lead_for_booking(
        self,
        connection: sqlite3.Connection,
        *,
        lead_id: str,
        booking_id: str,
        trip_id: str,
        interaction_id: str,
        timestamp: datetime,
    ) -> dict[str, Any] | None:
        row = connection.execute(
            "SELECT customer_tier FROM leads WHERE lead_id = ?",
            (lead_id,),
        ).fetchone()
        if not row:
            return None
        customer_tier = str(row["customer_tier"] or "").strip()
        if customer_tier == "VIP":
            stage = "VIP Booking Draft"
        elif customer_tier == "Repeat":
            stage = "Repeat Booking Draft"
        else:
            stage = "Booking Draft Created"
        follow_up_due = (timestamp.date() + timedelta(days=1)).isoformat()
        connection.execute(
            """
            UPDATE leads
            SET lead_stage = ?, priority = ?, follow_up_status = ?, follow_up_due_date = ?,
                updated_at = ?, interested_trip_ids = ?, suggested_trip_ids = ?,
                last_interaction_id = ?, booking_id = ?
            WHERE lead_id = ?
            """,
            (
                stage,
                "High",
                "Awaiting Deposit",
                follow_up_due,
                timestamp.isoformat(timespec="seconds"),
                trip_id,
                trip_id,
                interaction_id,
                booking_id,
                lead_id,
            ),
        )
        return {
            "lead_id": lead_id,
            "lead_stage": stage,
            "follow_up_status": "Awaiting Deposit",
            "follow_up_due_date": follow_up_due,
            "booking_id": booking_id,
        }

    def get_available_trips(self, trip_type: str | None, *, today: date | None = None) -> list[dict[str, Any]]:
        normalized_type = self.normalize_trip_type(trip_type)
        if trip_type and normalized_type is None:
            raise ValueError(f"Unsupported trip type {trip_type!r}")
        today = today or date.today()
        with self.connect() as connection:
            rows = connection.execute(
                """
                SELECT trip_id, trip_name, type, year, trip_leader, start_date, end_date, sales_status,
                       data_audit, trip_window_status, trip_availability_note, next_reengage_date,
                       single_total, double_total, triple_total, single_remaining, double_remaining,
                       triple_remaining, draft_holds_single, draft_holds_double, draft_holds_triple,
                       public_price, public_description, sales_notes
                FROM trips
                ORDER BY start_date ASC, trip_name ASC
                """
            ).fetchall()

        results: list[dict[str, Any]] = []
        for row in rows:
            record = self._trip_row_to_dict(row)
            if normalized_type and record["trip_type"] != normalized_type:
                continue
            if not self._trip_is_candidate(record, today=today):
                continue
            results.append(record)
        return results

    def build_trip_result(self, trip_type: str | None, *, today: date | None = None) -> dict[str, list[dict[str, Any]]]:
        trips = self.get_available_trips(trip_type, today=today)
        open_trips: list[dict[str, Any]] = []
        date_tbd_trips: list[dict[str, Any]] = []
        for trip in trips:
            if trip["start_date"]:
                open_trips.append(trip)
            else:
                date_tbd_trips.append(trip)
        open_trips.sort(key=lambda item: (item["start_date"] or "", item["trip_name"]))
        date_tbd_trips.sort(key=lambda item: ((item["year"] or 0), item["trip_name"]))
        return {"open_trips": open_trips, "date_tbd_trips": date_tbd_trips}

    def sync_trip_to_sheet(self, trip_id: str | dict[str, Any]) -> dict[str, Any]:
        if self.settings.sheet_backend == "excel":
            return self._run_sync_trip_to_sheet_worker(trip_id)

        # Spin up a daemon thread to process trip updates to sheets asynchronously
        import threading
        
        def run_async():
            try:
                self._run_sync_trip_to_sheet_worker(trip_id)
            except Exception:
                pass

        thread = threading.Thread(target=run_async)
        thread.daemon = True
        thread.start()
        return {"status": "queued", "message": "Background trip sync thread initiated."}

    def _run_sync_trip_to_sheet_worker(self, trip_id: str | dict[str, Any]) -> dict[str, Any]:
        record = trip_id if isinstance(trip_id, dict) else self._fetch_trip_record(trip_id)
        if not record:
            raise ValueError(f"Trip was not found for sync: {trip_id}")
        backend = self.settings.sheet_backend
        trip_key = str(record.get("trip_id") or "")
        try:
            if backend == "excel":
                updated = 0
                for workbook_path in (self.settings.excel_source_workbook, self.settings.excel_runtime_workbook):
                    if not workbook_path:
                        continue
                    self._upsert_trip_in_excel(workbook_path, record)
                    updated += 1
                self._record_sync_success("Trips", trip_key)
                return {"status": "ok", "backend": backend, "updated_workbooks": updated}
            if backend == "google_sheets":
                self._upsert_trip_in_google_sheet(record)
                self._record_sync_success("Trips", trip_key)
                return {"status": "ok", "backend": backend, "updated_workbooks": 1}
        except Exception as e:
            self._record_sync_failure("Trips", trip_key, str(e))
            return {"status": "failed", "backend": backend, "reason": str(e)}
        return {"status": "skipped", "backend": backend, "reason": "Sheet sync backend not supported in this phase."}

    def _fetch_trip_record(self, trip_id: str) -> dict[str, Any] | None:
        with self.connect() as connection:
            row = connection.execute(
                """
                SELECT trip_id, trip_name, type, year, trip_leader, start_date, end_date, sales_status,
                       data_audit, trip_window_status, trip_availability_note, next_reengage_date,
                       single_total, double_total, triple_total, single_remaining, double_remaining,
                       triple_remaining, draft_holds_single, draft_holds_double, draft_holds_triple,
                       public_price, public_description, sales_notes
                FROM trips
                WHERE trip_id = ?
                """,
                (trip_id,),
            ).fetchone()
        return self._trip_row_to_dict(row) if row else None

    def _trip_row_to_dict(self, row: sqlite3.Row) -> dict[str, Any]:
        single_remaining = self._as_int(row["single_remaining"])
        double_remaining = self._as_int(row["double_remaining"])
        triple_remaining = self._as_int(row["triple_remaining"])
        draft_holds_single = self._as_int(row["draft_holds_single"], default=0)
        draft_holds_double = self._as_int(row["draft_holds_double"], default=0)
        draft_holds_triple = self._as_int(row["draft_holds_triple"], default=0)
        available_single = max(single_remaining - draft_holds_single, 0) if single_remaining is not None else None
        available_double = max(double_remaining - draft_holds_double, 0) if double_remaining is not None else None
        available_triple = max(triple_remaining - draft_holds_triple, 0) if triple_remaining is not None else None
        availability_note = str(row["trip_availability_note"] or "").strip()
        values = [value for value in (available_single, available_double, available_triple) if value is not None]
        remaining_places = sum(values) if values else 0
        if remaining_places > 0:
            availability_status = "available"
        elif availability_note:
            availability_status = "manual_follow_up"
        else:
            availability_status = "full"
        return {
            "row": 0,
            "trip_id": str(row["trip_id"] or "").strip(),
            "trip_name": str(row["trip_name"] or "").strip(),
            "trip_type": str(row["type"] or "").strip(),
            "year": self._as_int(row["year"]),
            "start_date": self._iso_date(row["start_date"]),
            "end_date": self._iso_date(row["end_date"]),
            "sales_status": str(row["sales_status"] or "").strip() or "Open",
            "data_audit": str(row["data_audit"] or "").strip(),
            "trip_window_status": str(row["trip_window_status"] or "").strip(),
            "trip_availability_note": availability_note,
            "next_reengage_date": self._iso_date(row["next_reengage_date"]),
            "trip_leader": str(row["trip_leader"] or "").strip(),
            "single_total": self._as_int(row["single_total"], default=0),
            "double_total": self._as_int(row["double_total"], default=0),
            "triple_total": self._as_int(row["triple_total"], default=0),
            "single_remaining": single_remaining,
            "double_remaining": double_remaining,
            "triple_remaining": triple_remaining,
            "draft_holds_single": draft_holds_single,
            "draft_holds_double": draft_holds_double,
            "draft_holds_triple": draft_holds_triple,
            "available_single": available_single,
            "available_double": available_double,
            "available_triple": available_triple,
            "remaining_places": remaining_places,
            "availability_status": availability_status,
            "public_price": str(row["public_price"] or "").strip(),
            "public_description": str(row["public_description"] or "").strip(),
            "sales_notes": str(row["sales_notes"] or "").strip(),
        }

    def _trip_is_candidate(self, trip: dict[str, Any], *, today: date) -> bool:
        status = (trip.get("sales_status") or "").strip().lower()
        if status in INACTIVE_TRIP_STATUSES:
            return False
        start_date = trip.get("start_date")
        availability_note = (trip.get("trip_availability_note") or "").strip()
        has_capacity = bool((trip.get("remaining_places") or 0) > 0)
        if start_date:
            try:
                return date.fromisoformat(start_date) >= today and (has_capacity or bool(availability_note))
            except ValueError:
                return False
        return bool(availability_note or status == "date tbd" or has_capacity)

    def _upsert_trip_in_excel(self, workbook_path: Path, record: dict[str, Any]) -> None:
        if not workbook_path.exists():
            return
        with _SHEET_WRITE_LOCK:
            wb = load_workbook(workbook_path)
            try:
                ws = wb[TRIPS_SHEET_NAME]
                target_row = None
                for row_idx in range(TRIPS_DATA_START_ROW, ws.max_row + 1):
                    current_trip_id = str(ws.cell(row_idx, TRIPS_SHEET_COLUMNS["trip_id"]).value or "").strip()
                    if current_trip_id == record["trip_id"]:
                        target_row = row_idx
                        break
                if target_row is None:
                    target_row = ws.max_row + 1
                    if target_row < TRIPS_DATA_START_ROW:
                        target_row = TRIPS_DATA_START_ROW
                for field_name, col_idx in TRIPS_SHEET_COLUMNS.items():
                    ws.cell(target_row, col_idx).value = self._sheet_value(field_name, self._record_value(record, field_name))
                wb.save(workbook_path)
            finally:
                wb.close()

    def _upsert_trip_in_google_sheet(self, record: dict[str, Any]) -> None:
        if not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            raise RuntimeError("Google Sheets sync is not configured.")

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        ws = client.open_by_key(self.settings.google_sheet_id).worksheet(TRIPS_SHEET_NAME)
        rows = ws.get_all_values()
        target_row = None
        for row_idx in range(TRIPS_DATA_START_ROW, len(rows) + 1):
            current_trip_id = rows[row_idx - 1][TRIPS_SHEET_COLUMNS["trip_id"] - 1] if len(rows[row_idx - 1]) >= TRIPS_SHEET_COLUMNS["trip_id"] else ""
            if str(current_trip_id).strip() == record["trip_id"]:
                target_row = row_idx
                break
        if target_row is None:
            target_row = max(len(rows) + 1, TRIPS_DATA_START_ROW)
            row_values = [""] * max(TRIPS_SHEET_COLUMNS.values())
            for field_name, col_idx in TRIPS_SHEET_COLUMNS.items():
                row_values[col_idx - 1] = self._sheet_value(field_name, self._record_value(record, field_name))
            ws.append_row(row_values, value_input_option="USER_ENTERED")
            return

        cell_list = []
        for field_name, col_idx in TRIPS_SHEET_COLUMNS.items():
            cell_list.append(gspread.Cell(target_row, col_idx, self._sheet_value(field_name, self._record_value(record, field_name))))
        ws.update_cells(cell_list, value_input_option="USER_ENTERED")

    @staticmethod
    def _table_for_mapping(mapping_name: str) -> str:
        return {
            "Travelers": "travelers",
            "Leads": "leads",
            "Interactions": "interactions",
            "Trip Bookings": "trip_bookings",
            "Booking Event Trail": "booking_event_trail",
            "Handoff Queue": "handoff_queue",
        }[mapping_name]

    @staticmethod
    def _key_column_for_mapping(mapping_name: str) -> str:
        return {
            "Travelers": "traveler_id",
            "Leads": "lead_id",
            "Interactions": "interaction_id",
            "Trip Bookings": "booking_id",
            "Booking Event Trail": "event_id",
            "Handoff Queue": "handoff_id",
        }[mapping_name]

    def _fetch_table_record(self, table_name: str, key_column: str, record_id: str) -> dict[str, Any] | None:
        with self.connect() as connection:
            if table_name == "booking_event_trail":
                self._ensure_booking_event_trail_table(connection)
            row = connection.execute(
                f"SELECT * FROM {table_name} WHERE {key_column} = ?",
                (record_id,),
            ).fetchone()
        return dict(row) if row else None

    def _excel_outcome_workbooks(self) -> list[Path]:
        paths: list[Path] = []
        if self.settings.excel_runtime_workbook:
            paths.append(self.settings.excel_runtime_workbook)
        elif self.settings.excel_source_workbook:
            paths.append(self.settings.excel_source_workbook)
        return paths

    def _upsert_record_in_excel(
        self,
        workbook_path: Path,
        mapping: dict[str, Any],
        key_column: str,
        record: dict[str, Any],
    ) -> None:
        if not workbook_path or not workbook_path.exists():
            return
        with _SHEET_WRITE_LOCK:
            wb = load_workbook(workbook_path)
            try:
                ws = wb[mapping["sheet_name"]] if mapping["sheet_name"] in wb.sheetnames else wb.create_sheet(mapping["sheet_name"])
                header_row = int(mapping["header_row"])
                header_map = self._ensure_sheet_headers(ws, header_row, list(mapping["columns"].values()))
                key_header = mapping["columns"][key_column]
                key_value = str(record.get(key_column) or "").strip()
                target_row = self._find_sheet_row(ws, header_map[key_header], key_value, header_row + 1)
                if target_row is None:
                    target_row = ws.max_row + 1
                    if target_row <= header_row:
                        target_row = header_row + 1
                for field_name, header in mapping["columns"].items():
                    value = self._sheet_value(field_name, record.get(field_name))
                    ws.cell(target_row, header_map[header]).value = value
                wb.save(workbook_path)
            finally:
                wb.close()

    def _upsert_record_in_google_sheet(
        self,
        mapping: dict[str, Any],
        key_column: str,
        record: dict[str, Any],
    ) -> None:
        if not self.settings.google_sheet_id or not self.settings.google_application_credentials:
            raise RuntimeError("Google Sheets sync is not configured.")

        import gspread
        from google.oauth2.service_account import Credentials

        self._clear_dead_local_proxy()
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        creds = Credentials.from_service_account_file(
            str(self.settings.google_application_credentials),
            scopes=scopes,
        )
        client = gspread.authorize(creds)
        spreadsheet = client.open_by_key(self.settings.google_sheet_id)
        try:
            ws = spreadsheet.worksheet(mapping["sheet_name"])
        except gspread.exceptions.WorksheetNotFound:
            ws = spreadsheet.add_worksheet(title=mapping["sheet_name"], rows=100, cols=max(len(mapping["columns"]), 10))

        header_row = int(mapping["header_row"])
        rows = ws.get_all_values()
        while len(rows) < header_row:
            rows.append([])
        header_map = self._ensure_google_headers(ws, rows, header_row, list(mapping["columns"].values()))
        key_header = mapping["columns"][key_column]
        key_value = str(record.get(key_column) or "").strip()
        target_row = None
        key_idx = header_map[key_header] - 1
        for row_idx in range(header_row + 1, len(rows) + 1):
            row = rows[row_idx - 1]
            current = row[key_idx] if len(row) > key_idx else ""
            if str(current).strip() == key_value:
                target_row = row_idx
                break
        values_by_col = {
            header_map[header]: self._sheet_value(field_name, record.get(field_name))
            for field_name, header in mapping["columns"].items()
        }
        if target_row is None:
            row_values = [""] * max(values_by_col)
            for col_idx, value in values_by_col.items():
                row_values[col_idx - 1] = value
            ws.append_row(row_values, value_input_option="USER_ENTERED")
            return

        cells = [gspread.Cell(target_row, col_idx, value) for col_idx, value in values_by_col.items()]
        ws.update_cells(cells, value_input_option="USER_ENTERED")

    @staticmethod
    def _ensure_sheet_headers(ws, header_row: int, headers: list[str]) -> dict[str, int]:
        header_map: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(header_row, col_idx).value or "").strip()
            if header and header not in header_map:
                header_map[header] = col_idx
        next_col = ws.max_column + 1
        for header in headers:
            if header in header_map:
                continue
            ws.cell(header_row, next_col).value = header
            header_map[header] = next_col
            next_col += 1
        return header_map

    @staticmethod
    def _ensure_google_headers(ws, rows: list[list[str]], header_row: int, headers: list[str]) -> dict[str, int]:
        current = rows[header_row - 1] if len(rows) >= header_row else []
        header_map: dict[str, int] = {}
        for col_idx, header in enumerate(current, start=1):
            clean = str(header or "").strip()
            if clean and clean not in header_map:
                header_map[clean] = col_idx
        updates: list[Any] = []
        next_col = max(len(current), 0) + 1
        import gspread

        for header in headers:
            if header in header_map:
                continue
            updates.append(gspread.Cell(header_row, next_col, header))
            header_map[header] = next_col
            next_col += 1
        if updates:
            ws.update_cells(updates, value_input_option="USER_ENTERED")
        return header_map

    @staticmethod
    def _find_sheet_row(ws, key_col: int, key_value: str, start_row: int) -> int | None:
        for row_idx in range(start_row, ws.max_row + 1):
            if str(ws.cell(row_idx, key_col).value or "").strip() == key_value:
                return row_idx
        return None

    @staticmethod
    def _clear_dead_local_proxy() -> None:
        dead_proxy = {"http://127.0.0.1:9", "https://127.0.0.1:9", "127.0.0.1:9"}
        for key in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy"):
            value = (os.getenv(key) or "").strip().lower()
            if value in dead_proxy:
                os.environ.pop(key, None)

    @staticmethod
    def _ensure_booking_event_trail_table(connection: sqlite3.Connection) -> None:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS booking_event_trail (
                event_id TEXT PRIMARY KEY,
                occurred_at TEXT,
                event_type TEXT,
                event_label TEXT,
                traveler_id TEXT,
                lead_id TEXT,
                booking_id TEXT,
                trip_id TEXT,
                interaction_id TEXT,
                channel TEXT,
                actor TEXT,
                notes TEXT,
                metadata_json TEXT
            )
            """
        )

    @staticmethod
    def _next_daily_id(
        connection: sqlite3.Connection,
        table_name: str,
        column_name: str,
        prefix: str,
        timestamp: datetime,
    ) -> str:
        day_prefix = timestamp.strftime(f"{prefix}-%Y%m%d-")
        rows = connection.execute(
            f"SELECT {column_name} FROM {table_name} WHERE {column_name} LIKE ?",
            (f"{day_prefix}%",),
        ).fetchall()
        max_seq = 0
        for row in rows:
            value = str(row[column_name] or "")
            tail = value[len(day_prefix) :]
            if value.startswith(day_prefix) and tail.isdigit():
                max_seq = max(max_seq, int(tail))
        return f"{day_prefix}{max_seq + 1:04d}"

    @staticmethod
    def _next_booking_id(connection: sqlite3.Connection, trip_id: str, traveler_id: str) -> str:
        rows = connection.execute(
            "SELECT booking_id FROM trip_bookings WHERE trip_id = ?",
            (trip_id,),
        ).fetchall()
        sequence = len(rows) + 1
        traveler_digits = traveler_id[2:] if traveler_id.startswith("TR") else traveler_id
        if traveler_digits.isdigit():
            traveler_part = str(int(traveler_digits))
        else:
            traveler_part = traveler_id or "NEW"
        parts = trip_id.split("-")
        trip_mode = "I" if "-INT-" in trip_id else "L"
        trip_year = parts[2] if len(parts) > 2 else datetime.utcnow().strftime("%y")
        trip_serial = parts[3] if len(parts) > 3 else "000"
        candidate = f"{traveler_part}-{trip_mode}{trip_year}{trip_serial}-{sequence:03d}"
        existing = {str(row["booking_id"]) for row in rows}
        while candidate in existing:
            sequence += 1
            candidate = f"{traveler_part}-{trip_mode}{trip_year}{trip_serial}-{sequence:03d}"
        return candidate

    @staticmethod
    def _sheet_value(field_name: str, value: Any) -> Any:
        if value in (None, ""):
            return ""
        if field_name in {"start_date", "end_date", "next_reengage_date"}:
            return str(value)
        return value

    @staticmethod
    def _record_value(record: dict[str, Any], field_name: str) -> Any:
        if field_name == "type":
            return record.get("trip_type") or record.get("type")
        return record.get(field_name)

    @staticmethod
    def _as_int(value: Any, default: int | None = None) -> int | None:
        if value in (None, ""):
            return default
        try:
            return int(value)
        except (TypeError, ValueError):
            try:
                return int(float(value))
            except (TypeError, ValueError):
                return default

    @staticmethod
    def _iso_date(value: Any) -> str | None:
        if value in (None, ""):
            return None
        text = str(value)
        if "T" in text:
            text = text.split("T", 1)[0]
        if " " in text:
            text = text.split(" ", 1)[0]
        return text

    @staticmethod
    def _split_name(full_name: str) -> tuple[str, str]:
        parts = [part for part in (full_name or "").strip().split() if part]
        if not parts:
            return "", ""
        if len(parts) == 1:
            return parts[0], ""
        return parts[0], " ".join(parts[1:])

    @staticmethod
    def _next_prefixed_id(
        connection: sqlite3.Connection,
        table_name: str,
        column_name: str,
        prefix: str,
        padding: int,
    ) -> str:
        rows = connection.execute(f"SELECT {column_name} FROM {table_name}").fetchall()
        max_number = 0
        for row in rows:
            value = str(row[column_name] or "").strip()
            match = re.fullmatch(rf"{re.escape(prefix)}(\d+)", value)
            if match:
                max_number = max(max_number, int(match.group(1)))
        return f"{prefix}{max_number + 1:0{padding}d}"
