import openpyxl
from pathlib import Path

workbook_path = Path("RT - Travelers Database.phase5.ready.xlsx")

def setup_sheets():
    if not workbook_path.exists():
        print(f"Error: {workbook_path} not found")
        return

    wb = openpyxl.load_workbook(workbook_path)
    existing_sheets = wb.sheetnames
    print(f"Existing sheets: {existing_sheets}")

    required_sheets = {
        "DM Copy Library": ["Flow Name", "Step Name", "Message Key", "Arabic Copy", "English Copy", "Active"],
        "Waitlist": ["Contact Reference", "Trip Interest", "Language", "Last Touch", "Re-engage Time", "Re-engagement Reason"],
        "Follow Ups": ["Lead ID", "Reason", "Priority", "Owner", "Status"],
        "Trip Drips": ["Trip ID", "D-7 Message", "D-3 Message", "D-1 Message", "Day-of Message", "Post-trip Message"],
        "Handoff Queue": ["Lead ID", "Reason", "Priority", "Owner", "Status"],
        "Trip Bookings": ["Booking ID", "Trip ID", "Trip Name", "Traveler ID", "Traveler Name", "Room Type", "Flight Option", "Date Option", "Currency", "Booking Status", "Draft Created At", "Booking Source", "Lead ID", "Interaction ID", "Alert ID", "Payment Status", "Booking Notes"]
    }

    for sheet_name, headers in required_sheets.items():
        if sheet_name not in existing_sheets:
            print(f"Creating sheet: {sheet_name}")
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
        else:
            print(f"Sheet already exists: {sheet_name}")

    # Add some initial copy to DM Copy Library if it's empty
    ws_copy = wb["DM Copy Library"]
    if ws_copy.max_row == 1:
        initial_copy = [
            ["Flow 1", "Intake", "session.intake_start", "مرحباً. أنا وكيل المبيعات لـ رحمة ترافلر. يرجى إكمال نموذج البيانات حتى أتمكن من التحقق من سجلاتنا بأمان.", "Hello. I am Rahma Traveler's sales agent. Please complete the intake form so I can check the CRM safely.", "Yes"],
            ["Flow 1", "Intake", "session.completed", "هذه الجلسة مكتملة بالفعل. ابدأ جلسة جديدة للمتابعة.", "This session is already completed. Start a new session to continue.", "Yes"],
            ["Flow 1", "Trip Choice", "session.ask_trip_type", "رائع. هل هذا الاستفسار لرحلة محلية أم دولية؟", "Great. Is this inquiry for a local trip or an international trip?", "Yes"],
            ["Flow 1", "Trip Choice", "session.ask_trip_type_retry", "يرجى الرد بكلمة (محلية) أو (دولية) حتى أتمكن من التحقق من الرحلات القادمة.", "Please reply with either local or international so I can check upcoming trips.", "Yes"],
            ["Flow 1", "Preview", "session.preview_open_trips", "إليك الخيارات المتاحة قريباً: {trip_lines}\nرد بـ (نعم) لحفظ هذا الطلب، أو (لا) للتوقف.", "Here are upcoming options: {trip_lines}\nReply yes to save this as a lead, or no to stop.", "Yes"],
            ["Flow 1", "Preview", "session.preview_date_tbd", "وجدت رحلات في هذه الفئة، لكن المواعيد لم تتأكد بعد: {trip_names}. رد بـ (نعم) لحفظ هذا الطلب للمتابعة، أو (لا) للتوقف.", "I found trips in this category, but the dates are not confirmed yet: {trip_names}. Reply yes to save this lead for follow-up, or no to stop.", "Yes"],
            ["Flow 1", "Preview", "session.result_no_trip", "لا توجد رحلات مؤكدة قادمة في هذه الفئة حالياً.", "There are no confirmed upcoming trips in that category right now.", "Yes"],
            ["Flow 1", "Handoff", "session.handoff", "{reason_text}", "{reason_text}", "Yes"],
            ["Flow 1", "Booking", "session.ask_room_type", "أي نوع من الغرف تفضل؟ (فردي، ثنائي، أو ثلاثي)", "What type of room do you prefer? (Single, Double, or Triple)", "Yes"],
            ["Flow 1", "Booking", "session.ask_flight", "هل ترغب في حجز الطيران معنا؟ (نعم/لا)", "Would you like to book flights with us? (Yes/No)", "Yes"],
            ["Flow 1", "Booking", "session.ask_currency", "ما هي العملة المفضلة للدفع؟ (EGP/USD)", "What is your preferred currency for payment? (EGP/USD)", "Yes"],
            ["Flow 1", "Booking", "session.confirm_booking", "شكراً لك. لقد قمت بإنشاء مسودة حجز برقم {booking_id}. هل ترغب في تأكيدها بدفع العربون؟", "Thank you. I have created a booking draft {booking_id}. Would you like to confirm it by paying the deposit?", "Yes"],
            ["Flow 1", "Handoff", "handoff.phone_name_conflict", "رقم الواتساب هذا موجود بالفعل لاسم مسافر آخر في نظامنا. لقد أوقفت الأتمتة حتى يتمكن فريق المبيعات من مراجعتها بأمان.", "This WhatsApp number already exists for another traveler name in the CRM. I paused automation so the sales team can review it safely.", "Yes"],
            ["Flow 1", "Handoff", "handoff.blacklisted_customer", "لقد أوقفت هذا الاستفسار للمراجعة اليدوية من قبل فريق الإدارة لدينا. سيتصل بك أحد أعضاء الفريق إذا لزم الأمر.", "I have paused this inquiry for manual review by our management team. A team member will contact you if further information is needed.", "Yes"],
            ["Flow 1", "Handoff", "handoff.duplicate_phone_match", "هذا الرقم مرتبط بملفات تعريف متعددة للمسافرين في نظامنا. لقد أوقفت الأتمتة لضمان تحديث سجلك بشكل صحيح.", "This phone number is linked to multiple traveler profiles in our CRM. I have paused automation to ensure your record is updated correctly.", "Yes"],
            ["Flow 1", "Handoff", "handoff.generic", "لقد أرسلت هذا الطلب للمراجعة من قبل الفريق. برجاء الاتصال بمكتبنا مباشرة على {phone} للحصول على الدعم.", "I flagged this request for a senior team review. Please contact our office directly at {phone} for support.", "Yes"],
        ]
        for row in initial_copy:
            ws_copy.append(row)

    wb.save(workbook_path)
    print("Workbook updated successfully.")

if __name__ == "__main__":
    setup_sheets()
